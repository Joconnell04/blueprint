#!/usr/bin/env python3
"""import_scan.py: scan an .xlsx into the Workbook Planner block model.

Usage: python3 builder/import_scan.py <input.xlsx> <out.json>

Writes a plan fragment (tabs, blocks, unclassified regions) to out.json and
prints exactly one JSON line on stdout. Exit 0 on success, 1 with
{"ok": false, "errors": [...]} on failure. Python 3.9, openpyxl only.
Design: docs/research/import-scan.md. Sections follow its build order.
"""

import json
import os
import re
import sys
from collections import Counter
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.units import EMU_to_cm
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

GREY = "595959"
BLUE = "0000FF"
HEAVY = ("medium", "thick", "double")
TOTAL_RE = re.compile(r"^\s*((sub)?total|sum\b|net\s|grand)", re.I)
SUM_RE = re.compile(r"^=\s*(SUM|SUBTOTAL)\s*\(", re.I)
CHART_KINDS = {"BarChart": "bar", "BarChart3D": "bar", "LineChart": "line",
               "LineChart3D": "line", "PieChart": "pie", "PieChart3D": "pie",
               "DoughnutChart": "pie"}
SAMPLE_CAP = 200
COL_CM, ROW_CM = 1.7, 0.5  # builder chart geometry: width = cols * 1.7, height = rows * 0.5


def a1(r, c):
    return "%s%d" % (get_column_letter(c), r)


def rng(r1, c1, r2, c2):
    return "%s:%s" % (a1(r1, c1), a1(r2, c2))


# ------------------------------------------------------------ 1. cell features


class Feat(object):
    """One occupied cell: value kind, style signals and the derived flags the rules read."""
    __slots__ = ("r", "c", "v", "kind", "res", "bold", "italic", "size", "color", "fill", "wrap", "halign",
                 "b", "merged", "copy", "note_like", "blue", "nf", "is_text", "is_num", "numeric_text")

    def __init__(self, r, c):
        self.r, self.c, self.v, self.kind, self.res, self.nf = r, c, None, "empty", "number", "General"
        self.bold = self.italic = self.fill = self.wrap = self.copy = self.note_like = self.blue = False
        self.is_text = self.is_num = self.numeric_text = False
        self.size, self.color, self.halign, self.merged = 11.0, None, None, None
        self.b = {"l": None, "r": None, "t": None, "b": None}

    @property
    def has_value(self):
        return self.v is not None

    @property
    def is_f(self):
        return self.kind == "formula"

    @property
    def heavy_bottom(self):
        return self.b["b"] in HEAVY

    @property
    def boxed(self):
        return all(self.b[s] for s in "lrtb")

    def text_len(self):
        return len(str(self.v)) if self.v is not None else 0


def cell_kind(cell, v):
    if v is None:
        return "empty"
    if cell.data_type == "f" or isinstance(v, (ArrayFormula, DataTableFormula)) or (isinstance(v, str) and v.startswith("=")):
        return "formula"
    if cell.data_type == "s" or isinstance(v, str):
        return "text"
    if cell.data_type == "e":
        return "error"
    if isinstance(v, bool) or cell.data_type == "b":
        return "bool"
    return "date" if isinstance(v, (datetime, date)) or cell.data_type == "d" else "number"


def feat_from_cell(cell, cached):
    f = Feat(cell.row, cell.column)
    v = cell.value
    if isinstance(v, ArrayFormula):
        v = v.text
    elif isinstance(v, DataTableFormula):
        v = "=" + str(v.ref)
    f.v, f.kind, f.nf = v, cell_kind(cell, v), cell.number_format or "General"
    if f.kind == "formula":
        f.res = "text" if (isinstance(cached, str) if cached is not None else f.nf == "@") else "number"
    f.is_text = f.kind == "text" or (f.kind == "formula" and f.res == "text")
    f.is_num = f.kind in ("number", "date", "bool") or (f.kind == "formula" and f.res != "text")
    if f.kind == "text":
        try:
            float(str(v).replace(",", ""))
            f.numeric_text = True
        except ValueError:
            pass
    if not cell.has_style:
        return f
    font, al, bd = cell.font, cell.alignment, cell.border
    f.bold, f.italic, f.size = bool(font.bold), bool(font.italic), float(font.size or 11)
    col = font.color
    f.color = col.rgb[-6:].upper() if col is not None and col.type == "rgb" and isinstance(col.rgb, str) else None
    f.fill = bool(cell.fill is not None and cell.fill.fill_type)
    f.wrap = bool(al is not None and al.wrap_text)
    f.halign = al.horizontal if al is not None else None
    if bd is not None:
        for key, side in (("l", bd.left), ("r", bd.right), ("t", bd.top), ("b", bd.bottom)):
            f.b[key] = side.style if side is not None else None
    f.note_like = f.italic and f.kind == "text" and (f.color == GREY or not any(f.b.values()))
    f.blue = f.color == BLUE and f.kind not in ("formula", "empty")
    return f


class Grid(object):
    """Sparse feature grid of occupied cells for one worksheet."""

    def __init__(self, ws, cached_ws):
        self.ws = ws
        self.cells = {}
        self.hidden_rows = set(r for r, d in ws.row_dimensions.items() if d.hidden)
        self.hidden_cols = set()  # Excel groups <col min max>; expand every dimension over its span
        for d in ws.column_dimensions.values():
            if d.hidden and d.min is not None and d.max is not None:
                self.hidden_cols.update(range(d.min, d.max + 1))
        self.freeze_rows = range_boundaries(ws.freeze_panes)[1] - 1 if ws.freeze_panes else 0
        anchors = {}
        for m in ws.merged_cells.ranges:
            for r in range(m.min_row, m.max_row + 1):
                for c in range(m.min_col, m.max_col + 1):
                    anchors[(r, c)] = (m.min_row, m.min_col, m.max_row, m.max_col)
        sizes = Counter()
        for row in ws.iter_rows():
            for cell in row:
                key = (cell.row, cell.column)
                if key in anchors and (cell.row, cell.column) != anchors[key][:2]:
                    continue  # merged copy, filled in below
                if cell.value is None and not cell.has_style:
                    continue
                cached = cached_ws.cell(cell.row, cell.column).value if cached_ws is not None and cell.data_type == "f" else None
                f = feat_from_cell(cell, cached)
                if f.v is None and not any(f.b.values()) and not f.fill and key not in anchors:
                    continue
                f.merged = anchors.get(key)
                self.cells[key] = f
                if f.v is not None:
                    sizes[f.size] += 1
        for key, m in anchors.items():  # merged copies carry the anchor's value and style
            src = self.cells.get(m[:2])
            if key != m[:2] and src is not None:
                f = Feat(key[0], key[1])
                for name in Feat.__slots__[2:]:
                    setattr(f, name, getattr(src, name))
                f.b, f.copy = dict(src.b), True
                self.cells[key] = f
        self.modal_size = sizes.most_common(1)[0][0] if sizes else 11.0

    def row_cells(self, r, c1, c2, valued=True):
        out = [self.cells[(r, c)] for c in range(c1, c2 + 1) if (r, c) in self.cells]
        return [f for f in out if f.has_value] if valued else out

    def col_cells(self, c, r1, r2, valued=True):
        out = [self.cells[(r, c)] for r in range(r1, r2 + 1) if (r, c) in self.cells]
        return [f for f in out if f.has_value] if valued else out

    def blank_row(self, r, c1, c2):
        return r < 1 or not self.row_cells(r, c1, c2, valued=False)


# ------------------------------------------------------- 2. regions and boxes


class Box(object):
    def __init__(self, r1, c1, r2, c2):
        self.r1, self.c1, self.r2, self.c2 = r1, c1, r2, c2
        self.flags = []
        self.notes = []       # peeled note cells
        self.subtitle = []    # note-like rows directly under the title
        self.forced = None    # ListObject overrides

    @property
    def width(self):
        return self.c2 - self.c1 + 1

    def touches(self, o):
        return not (o.c1 > self.c2 + 1 or o.c2 < self.c1 - 1 or o.r1 > self.r2 + 1 or o.r2 < self.r1 - 1)

    def absorb(self, o):
        self.r1, self.c1 = min(self.r1, o.r1), min(self.c1, o.c1)
        self.r2, self.c2 = max(self.r2, o.r2), max(self.c2, o.c2)
        self.flags += o.flags

    def range(self):
        return rng(self.r1, self.c1, self.r2, self.c2)


def components(grid, skip):
    """8-connected components over occupied cells, as bounding boxes."""
    seen = set(skip)
    boxes = []
    for start in sorted(grid.cells):
        if start in seen:
            continue
        seen.add(start)
        stack, member = [start], []
        while stack:
            r, c = stack.pop()
            member.append((r, c))
            for dr in (-1, 0, 1):
                for dc in (-1, 0, 1):
                    n = (r + dr, c + dc)
                    if n in grid.cells and n not in seen:
                        seen.add(n)
                        stack.append(n)
        rs, cs = [m[0] for m in member], [m[1] for m in member]
        boxes.append(Box(min(rs), min(cs), max(rs), max(cs)))
    changed = True
    while changed:  # an L-shaped component's box may swallow or touch a neighbour: merge those
        changed, out = False, []
        for b in boxes:
            hit = next((o for o in out if o.touches(b)), None)
            if hit is None:
                out.append(b)
            else:
                hit.absorb(b)
                changed = True
        boxes = out
    return boxes


def bridge(grid, boxes, det):
    """Merge boxes separated by exactly one blank row (or column), with guards."""
    def row_ok(a, b):
        c1, c2 = min(a.c1, b.c1), max(a.c2, b.c2)
        top = grid.row_cells(b.r1, b.c1, b.c2)
        return (b.r1 == a.r2 + 2 and abs(a.c1 - b.c1) <= 1 and abs(a.c2 - b.c2) <= 1
                and not (len(top) == 1 and top[0].kind == "text")  # a lone text row is never a line item
                and grid.blank_row(a.r2 + 1, c1, c2) and not det.title_like(b.r1, b.c1, b.c2)[0]
                and det.header_score(b.r1, b.c1, b.c2) == 0 and not det.total_like(a.r2, a.c1, a.c2)
                and not any(f.heavy_bottom for f in grid.row_cells(a.r2, a.c1, a.c2, False)))

    def col_ok(a, b):
        return (b.c1 == a.c2 + 2 and a.r1 == b.r1 and a.r2 == b.r2 and not det.label_like(b.c1, b.r1, b.r2)
                and not det.title_like(b.r1, b.c1, b.c2)[0]
                and not any(f.b["r"] in HEAVY for f in grid.col_cells(a.c2, a.r1, a.r2, False)))

    changed = True
    while changed:
        changed = False
        boxes.sort(key=lambda b: (b.r1, b.c1))
        for i, a in enumerate(boxes):
            hit = next((b for b in boxes[i + 1:] if row_ok(a, b) or col_ok(a, b)), None)
            if hit is not None:
                a.absorb(hit)
                a.flags.append("bridged")
                boxes.remove(hit)
                changed = True
                break
    return boxes


def peel_notes(grid, box):
    """Drop note-only columns from the right edge; tighten the box to occupied cells."""
    while box.c2 > box.c1:
        cells = grid.col_cells(box.c2, box.r1, box.r2)
        notes = [f for f in cells if f.note_like]
        others = [f for f in cells if not f.note_like and not f.copy]
        if not notes or others:
            break
        box.notes.extend(f for f in notes if not f.copy)
        box.c2 -= 1
    tighten(grid, box)


def tighten(grid, box):
    """Shrink the box until every edge row and column holds an occupied cell."""
    col = lambda c: grid.col_cells(c, box.r1, box.r2, False)
    row = lambda r: grid.row_cells(r, box.c1, box.c2, False)
    while box.c2 > box.c1 and not col(box.c2):
        box.c2 -= 1
    while box.c1 < box.c2 and not col(box.c1):
        box.c1 += 1
    while box.r2 > box.r1 and not row(box.r2):
        box.r2 -= 1
    while box.r1 < box.r2 and not row(box.r1):
        box.r1 += 1


# ------------------------------------------------------------ 3. splitting


def piece(box, r1, c1, r2, c2, flag):
    p = Box(r1, c1, r2, c2)
    p.flags = list(box.flags) + ([flag] if flag else [])
    return p


def split_box(grid, det, box):
    """Cut a box into pieces before rows where the split score reaches 2, then on label columns."""
    pieces, start, flag = [], box.r1, None
    prev_title = det.title_like(box.r1, box.c1, box.c2)[0]
    headed = prev_title and det.header_score(box.r1 + 1, box.c1, box.c2) > 0
    for r in range(box.r1 + 1, box.r2 + 1):
        cells = grid.row_cells(r, box.c1, box.c2)
        above = grid.row_cells(r - 1, box.c1, box.c2)
        is_title, tsig = det.title_like(r, box.c1, box.c2)[:2]
        if not cells or prev_title:  # never cut right after a title row
            prev_title = is_title
            continue
        score = 0
        same_merge = (above and above[0].merged and cells[0].merged and not cells[0].bold
                      and above[0].merged[1::2] == cells[0].merged[1::2])  # another merged paragraph
        if is_title and not same_merge and not (headed and det.header_score(r + 1, box.c1, box.c2) == 0):
            # house titles are unambiguous (3). A title-like row inside a headed table with no header
            # row after it is a section divider, not a cut.
            score += 3 if "title_house" in tsig else 2
        all_text = all(f.is_text for f in cells)
        bold_now, bold_above = [f.bold for f in cells], [f.bold for f in above]
        span = lambda fs: (max(f.c for f in fs) - min(f.c for f in fs) + 1) if fs else 0
        score += all(f.b["t"] in HEAVY for f in cells) and not any(f.b["t"] in HEAVY for f in above)
        score += all_text and det.total_like(r - 1, box.c1, box.c2)
        score += bool(above) and all_text and all(bold_now) != all(bold_above) and (all(bold_now) or not any(bold_now))
        score += bool(above) and all(f.b["b"] for f in above) and not any(f.b["t"] or f.b["b"] for f in cells)
        score += bool(above) and abs(span(cells) - span(above)) > box.width / 2.0
        if score >= 2:
            pieces.append(piece(box, start, box.c1, r - 1, box.c2, flag or ("split_weak" if score == 2 else "split_strong")))
            start, flag = r, "split_weak" if score == 2 else "split_strong"
        prev_title = is_title
    pieces.append(piece(box, start, box.c1, box.r2, box.c2, flag))
    out = []
    for p in pieces:
        p.notes = [f for f in box.notes if p.r1 <= f.r <= p.r2]
        out.extend(split_cols(grid, det, p))
    return out


def split_cols(grid, det, box):
    """Cut before a label column that follows a value (or note) column, unless a shared header spans both."""
    pieces, start = [], box.c1
    for c in range(box.c1 + 2, box.c2 + 1):
        prev = grid.col_cells(c - 1, box.r1, box.r2)
        prev_ok = any(f.is_num for f in prev) or (prev and all(f.note_like for f in prev))
        if not prev_ok or not det.label_like(c, box.r1, box.r2):
            continue
        pair = lambda r: (grid.cells.get((r, c - 1)), grid.cells.get((r, c)))
        shared = [r for r in range(box.r1, box.r2 + 1) if all(pair(r)) and all(f.is_text and not f.note_like for f in pair(r))
                  and pair(r)[0].bold == pair(r)[1].bold]
        if shared or any(f.merged and f.merged[1] < c <= f.merged[3] for f in grid.col_cells(c, box.r1, box.r2, False)):
            continue
        pieces.append(piece(box, box.r1, start, box.r2, c - 1, "split_strong"))
        start = c
    pieces.append(piece(box, box.r1, start, box.r2, box.c2, "split_strong" if pieces else None))
    for p in pieces:
        tighten(grid, p)
    pieces[-1].notes = list(box.notes)
    return pieces


# ------------------------------------- 4. title, header, label column, totals


class Detector(object):
    def __init__(self, grid):
        self.g = grid

    def title_like(self, r, c1, c2):
        """(is_title, signals, merge_width, text) for row r within columns c1..c2."""
        cells = [f for f in self.g.row_cells(r, c1, c2) if not f.note_like and not f.copy]
        if len(cells) != 1 or cells[0].kind != "text":
            return False, [], 1, ""
        f = cells[0]
        if f.text_len() > 60:
            return False, [], 1, ""
        mw = (f.merged[3] - f.merged[1] + 1) if f.merged else 1
        if f.c > c1 and mw >= 2 and any(x.kind == "text" for x in self.g.row_cells(r + 1, c1, f.c - 1)):
            return False, [], 1, ""  # a group header over value columns, not a title
        sig = [name for name, on in (("bold title", f.bold), ("large title", f.size > self.g.modal_size),
                                     ("heavy bottom border", f.heavy_bottom), ("merged title", mw >= 2),
                                     ("title_house", f.bold and mw >= 2 and f.heavy_bottom)) if on]
        if not sig and self.g.blank_row(r - 1, c1, c2):
            sig.append("lone text row")  # Unstructured's leading single-cell row rule
        return bool(sig), sig, mw, str(f.v)

    def header_score(self, r, c1, c2):
        cells = [f for f in self.g.row_cells(r, c1, c2) if not f.note_like]
        if len(cells) < 2:
            return 0
        below = self.g.row_cells(r + 1, c1, c2)
        for f in cells:
            year_like = f.kind == "number" and 1900 <= (f.v if isinstance(f.v, (int, float)) else 0) <= 2100
            if not (f.kind == "text" or (year_like and (f.nf == "@" or f.bold))):
                return 0
        differs = lambda test: all(test(f) for f in cells) and not (below and all(test(f) for f in below))
        text_cols = set(f.c for f in cells)
        score = sum([
            differs(lambda f: f.bold), differs(lambda f: bool(f.b["b"])), differs(lambda f: f.fill),
            differs(lambda f: f.halign == "center"),
            any(f.is_num and f.c in text_cols for f in below),
            any(f.merged and f.merged[3] > f.merged[1] for f in cells),
            0 < self.g.freeze_rows <= 3 and r == self.g.freeze_rows,
            # a bold text row at the top of a region is a header even with nothing under it
            all(f.bold for f in cells) and len(cells) >= 3
            and (self.g.blank_row(r - 1, c1, c2) or self.title_like(r - 1, c1, c2)[0]),
        ])
        return score

    def label_like(self, c, r1, r2):
        cells = self.g.col_cells(c, r1, r2)
        if len(cells) < 2 or sum(1 for f in cells if f.is_text) < 0.8 * len(cells):
            return False
        return any(f.is_num for f in self.g.col_cells(c + 1, r1, r2))

    def total_like(self, r, c1, c2, all_bold=False):
        cells = self.g.row_cells(r, c1, c2)
        if len(cells) < 2:
            return False
        label, values = cells[0], cells[1:]
        if label.kind == "text" and TOTAL_RE.search(str(label.v)):
            return True
        if label.bold and not all_bold and any(f.is_num for f in values):
            return True
        return all(f.is_f and SUM_RE.match(str(f.v)) for f in values)


# --------------------------------------------- 5. classification, confidence


def detect_block(grid, det, box, warnings):
    """Turn a box into a block dict, or return (None, reason) for unclassified."""
    g = grid
    title_row = None
    is_title, tsig, mw, ttext = det.title_like(box.r1, box.c1, box.c2)
    if box.forced and box.forced.get("title_row"):
        title_row, ttext, tsig, mw = box.forced["title_row"], box.forced["title"], ["bold title"], 1
    elif is_title:
        title_row = box.r1
    first = (title_row + 1) if title_row else box.r1
    while first <= box.r2 and title_row:
        lead = g.row_cells(first, box.c1, box.c2)
        if not lead or not all(f.note_like for f in lead) \
                or not any(g.row_cells(r, box.c1, box.c2) for r in range(first + 1, box.r2 + 1)):
            break
        box.subtitle.extend(f for f in lead if not f.copy)  # grey italic line under the title
        first += 1
    if first > box.r2:
        return None, "lone-text"
    content_cols = [c for c in range(box.c1, box.c2 + 1)
                    if any(not f.copy for f in g.col_cells(c, first, box.r2))]
    width = len(content_cols)
    if width == 0:
        return None, "lone-text"
    signals = list(tsig)
    if not title_row:
        signals.append("no title")
    # header rows: up to two for wide boxes; a two-column box only when both top cells are bold text
    header_rows = []
    if box.forced and box.forced.get("header"):
        header_rows = list(range(box.forced["header"], box.forced["header"] + box.forced["header_rows"]))
    elif width >= 3:
        for r in (first, first + 1):
            if r > box.r2 or det.header_score(r, box.c1, box.c2) == 0 or len(header_rows) < r - first:
                break
            header_rows.append(r)
    elif width == 2:
        top = g.row_cells(first, box.c1, box.c2)
        if len(top) == 2 and all(f.bold and f.kind == "text" for f in top) \
                and any(f.is_num for f in g.col_cells(top[1].c, first + 1, box.r2)):
            header_rows = [first]
    body_start = (header_rows[-1] + 1) if header_rows else first
    body_rows = [r for r in range(body_start, box.r2 + 1) if g.row_cells(r, box.c1, box.c2)]
    if header_rows and not body_rows:
        return None, "header row with no body"
    # label column: first text column followed by values; else labels on the right; else a text first column
    body_col = lambda c: g.col_cells(c, body_start, box.r2)
    text_share = lambda fs: sum(1 for f in fs if f.is_text) / float(len(fs) or 1)
    label_col, labels_side = next((c for c in content_cols if det.label_like(c, body_start, box.r2)), None), "left"
    if label_col is None and width >= 2:
        earlier = [f for c in content_cols[:-1] for f in body_col(c)]
        if body_col(content_cols[-1]) and text_share(body_col(content_cols[-1])) == 1.0 and earlier \
                and all(f.is_num for f in earlier):
            label_col, labels_side = content_cols[-1], "right"
        elif body_col(content_cols[0]) and text_share(body_col(content_cols[0])) >= 0.8:
            label_col = content_cols[0]
    value_cols = [c for c in content_cols if c != label_col]
    # totals
    total_row, subtotals, totals = None, 0, []
    if body_rows and (width >= 3 or header_rows):
        all_bold = all(all(f.bold for f in g.row_cells(r, box.c1, box.c2)) for r in body_rows)
        totals = [r for r in body_rows if det.total_like(r, box.c1, box.c2, all_bold)]
        if box.forced and box.forced.get("total"):
            totals = [r for r in totals if r != box.forced["total"]] + [box.forced["total"]]
        if totals and totals[-1] == body_rows[-1] and len(body_rows) > 1:
            total_row = totals[-1]
            subtotals = len(totals) - 1
        elif totals:
            subtotals = len(totals)
            if totals[0] == body_rows[0]:
                signals.append("total row at top")
                warnings.append("%s: total row at the top of %s stays in data" % (g.ws.title, box.range()))
    data_rows = [r for r in body_rows if r != total_row]
    values = [f for c in value_cols for f in g.col_cells(c, body_start, box.r2)
              if f.r in data_rows and not f.copy]
    n = float(len(values)) or 1.0
    fshare = sum(1 for f in values if f.is_f) / n
    blue = sum(1 for f in values if f.blue) / n
    tshare = sum(1 for f in values if f.is_text and not f.numeric_text) / n
    body_cells = [f for c in content_cols for f in g.col_cells(c, body_start, box.r2)]
    # classification
    btype = None
    if box.forced and box.forced.get("header"):
        btype = "calc" if fshare > 0.6 else "table"
    elif width == 1 or not values:
        if any(f.is_num for f in body_cells):
            return None, "single column of values"
        if not body_cells:
            return None, "lone-text"
        btype = "text"
    elif width == 2 and not header_rows:
        btype = "output" if fshare > 0.6 else "inputs"
    else:
        btype = "calc" if fshare > 0.6 else "table"
    # confidence: start at 0.5, add and subtract per the research doc's table, clip to [0, 1]
    conf = [0.5]

    def bump(cond, delta, signal=None):
        if cond:
            conf[0] += delta
            if signal:
                signals.append(signal)
        return cond

    bump(box.forced, 0.30, "listobject")
    bump(title_row, 0.20 if len([x for x in tsig if x != "title_house"]) >= 2 else 0.10)
    if btype == "text":
        merged_wrap = any(f.merged and f.wrap for f in body_cells)
        conf[0] = 0.8 if merged_wrap else (0.6 if any(f.wrap for f in body_cells) else 0.4)
        signals.append("merged wrapped text" if merged_wrap else "prose")
    elif btype in ("table", "calc"):
        if not bump(header_rows, 0.15, "header row"):
            signals.append("no header row")
    elif label_col is not None:
        share = text_share(body_col(label_col))
        bump(share >= 0.9, 0.15)
        signals.append("label share %.2f" % share)
    if btype != "text":
        signals.append("formula share %.2f" % fshare)
        if not bump(fshare <= 0.1 or fshare >= 0.8, 0.10):
            bump(0.4 <= fshare <= 0.6 and not (btype == "inputs" and blue >= 0.5), -0.15, "ambiguous formula share")
        bump(body_cells and all(f.boxed for f in body_cells), 0.10, "bordered body")
        bump(btype == "inputs" and blue >= 0.5, 0.10, "blue literals %.2f" % blue)
        bump(tshare > 0.2, -0.15, "text among values %.2f" % tshare)
        bump(any(f.merged for f in values), -0.10, "merged cells in body")
        bump(len(data_rows) < 2, -0.15, "short body")
    bump("bridged" in box.flags, -0.20, "bridged")
    if not bump("split_weak" in box.flags, -0.20, "weak split") and "split_strong" in box.flags:
        signals.append("strong split")
    bump(len(header_rows) > 1, -0.10, "second header row")
    row = title_row or max(1, first - 1)
    if not title_row and first == 1:
        signals.append("header on row 1, no room for a title")
    hidden_r = sum(1 for r in range(row, box.r2 + 1) if r in g.hidden_rows)
    hidden_c = sum(1 for c in range(box.c1, box.c2 + 1) if c in g.hidden_cols)
    bump(hidden_r or hidden_c, -0.10, "hidden rows %d cols %d" % (hidden_r, hidden_c))
    conf = round(max(0.0, min(0.95 if box.forced else 1.0, conf[0])), 2)
    if conf < 0.35:
        return None, ", ".join(x for x in signals if x.split()[0] in
                               ("no", "ambiguous", "text", "bridged", "weak", "hidden", "short", "prose")) or "low confidence"
    # assemble
    block = {
        "type": btype, "title": ttext if title_row else "", "row": row, "col": box.c1,
        "rows": box.r2 - row + 1, "cols": max(box.width, mw), "columns": [], "rowLabels": [],
        "sample": [], "confidence": conf,
        "detected": {"signals": signals, "reason": "%s: %s" % (btype, "; ".join(signals[:4]))},
        "range": rng(row, box.c1, box.r2, max(box.c2, box.c1 + mw - 1)),
    }
    if labels_side == "right":
        block["labels"] = "right"
    if hidden_r or hidden_c:
        block["hidden_rows"], block["hidden_cols"] = hidden_r, hidden_c
    fill_content(g, box, block, btype, content_cols, label_col, header_rows, data_rows, total_row, subtotals)
    attach_notes(g, box, block, title_row, width)
    block["geometry"] = geometry(box, title_row, header_rows, label_col, content_cols, data_rows, total_row,
                                 [r for r in totals if r != total_row])
    return block, None


def geometry(box, title_row, header_rows, label_col, content_cols, body_rows, total_row, subtotal_rows):
    """Exact cell roles for restyle.py: which rows and columns are title, header, labels, body, totals, notes."""
    return {
        "titleRow": title_row,
        "subtitleCells": [[f.r, f.c] for f in box.subtitle],
        "headerRows": list(header_rows),
        "labelCol": label_col,
        "contentCols": list(content_cols),
        "bodyRows": list(body_rows),
        "totalRow": total_row,
        "subtotalRows": list(subtotal_rows),
        "noteCells": [[f.r, f.c] for f in box.notes if not f.copy],
        "box": [box.r1, box.c1, box.r2, box.c2],
    }


def cell_out(f):
    """Plan-safe cell value: number, string (formulas verbatim), ISO date, or ''."""
    if f is None or f.v is None:
        return ""
    if isinstance(f.v, (datetime, date)):
        return f.v.isoformat()
    return f.v if isinstance(f.v, (int, float, str)) and not isinstance(f.v, bool) else str(f.v)


def fill_content(g, box, block, btype, content_cols, label_col, header_rows, data_rows, total_row, subtotals):
    take = lambda r, c: cell_out(g.cells.get((r, c)))
    if btype == "text":
        paras = [str(f.v) for c in content_cols for f in g.col_cells(c, data_rows[0] if data_rows else box.r1, box.r2)
                 if not f.copy and f.kind == "text"]
        block["text"] = "\n".join(paras)
        block["sample"] = [[p] for p in paras[:SAMPLE_CAP]]
        return
    if btype in ("inputs", "output"):
        lc = label_col if label_col is not None else content_cols[0]
        vc = [c for c in content_cols if c != lc][0]
        for r in data_rows:
            block["rowLabels"].append(str(take(r, lc)))
            block["sample"].append([take(r, lc), take(r, vc)])
        block["columns"] = list(block["rowLabels"])
        block["sample"] = block["sample"][:SAMPLE_CAP]
        return
    if header_rows:
        block["columns"] = [str(take(header_rows[-1], c)) for c in content_cols]
        if len(header_rows) > 1:
            block["header_rows"] = len(header_rows)
    else:
        block["columns"] = [get_column_letter(c) for c in content_cols]
    if box.forced and box.forced.get("columns"):
        block["columns"] = box.forced["columns"]
    lc = label_col if label_col is not None else content_cols[0]
    for r in data_rows[:SAMPLE_CAP]:
        block["sample"].append([take(r, c) for c in content_cols])
        block["rowLabels"].append(str(take(r, lc)))
    block["dataRows"] = len(data_rows)
    if total_row:
        block["totalRow"] = [take(total_row, c) for c in content_cols]
    if subtotals:
        block["subtotals"] = subtotals


def attach_notes(g, box, block, title_row, width):
    row_notes, beside = {}, []
    for f in sorted(box.subtitle + box.notes, key=lambda f: (f.r, f.c)):
        if f.r == title_row or (f.r == box.r1 and not title_row) or f in box.subtitle:
            block["note"] = (block["note"] + "; " + str(f.v)) if block.get("note") else str(f.v)
        else:
            row_notes[f.r] = (row_notes[f.r] + "; " + str(f.v)) if f.r in row_notes else str(f.v)
            if width > 2:
                beside.append(f)
    if row_notes:
        block["rowNotes"] = [row_notes.get(r, "") for r in range(box.r1 + 1, box.r2 + 1)
                             if g.row_cells(r, box.c1, box.c2)]
    if beside:  # the SPEC has no per-row notes on tables, so a rebuild would drop these
        block["extraNotes"] = [{"range": rng(beside[0].r, beside[0].c, beside[-1].r, beside[-1].c),
                                "reason": "notes-beside-table"}]


# ---------------------------------------------- 6. ListObjects and charts


def listobject_boxes(grid, det, warnings):
    boxes, skip = [], set()
    for tbl in grid.ws.tables.values():
        try:
            c1, r1, c2, r2 = range_boundaries(tbl.ref)
        except (TypeError, ValueError):
            warnings.append("%s: table %s has an unreadable ref" % (grid.ws.title, tbl.displayName))
            continue
        hdr = tbl.headerRowCount if tbl.headerRowCount is not None else 1
        tot = tbl.totalsRowCount or 0
        box = Box(r1, c1, r2, c2)
        box.forced = {"header": r1 if hdr else None, "header_rows": max(hdr, 1),
                      "total": r2 if tot else None,
                      "columns": [c.name for c in (tbl.tableColumns or [])] or None}
        is_title, _, _, text = det.title_like(r1 - 1, c1, c2)
        if is_title:
            box.forced["title_row"], box.forced["title"] = r1 - 1, text
            box.r1 = r1 - 1
            skip.update((r1 - 1, c) for c in range(c1, c2 + 1))
        skip.update((r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1))
        boxes.append(box)
    return boxes, skip


def chart_geometry(anchor):
    fr = anchor._from
    row, col = fr.row + 1, fr.col + 1
    to = getattr(anchor, "to", None)
    if to is not None:
        return row, col, max(1, to.row - fr.row + 1), max(1, to.col - fr.col + 1)
    ext = getattr(anchor, "ext", None)
    if ext is not None and ext.width and ext.height:
        return row, col, max(1, int(round(EMU_to_cm(ext.height) / ROW_CM))), max(1, int(round(EMU_to_cm(ext.width) / COL_CM)))
    return row, col, 15, 9


def parse_ref(ref):
    """'Sheet Name'!$B$2:$B$9 -> (sheet, (c1, r1, c2, r2)) or None."""
    if not ref or "!" not in ref:
        return None
    sheet, cells = ref.rsplit("!", 1)
    sheet = sheet.strip()
    if sheet.startswith("'") and sheet.endswith("'"):
        sheet = sheet[1:-1].replace("''", "'")
    try:
        return sheet, range_boundaries(cells.replace("$", ""))
    except (TypeError, ValueError):
        return None


def chart_title(chart):
    try:
        return "".join(r.t for p in chart.title.tx.rich.p for r in (p.r or []))
    except AttributeError:
        return ""


def chart_blocks(wb, tabs, warnings):
    """One chart block per ws._charts entry; series ranges resolve to detected blocks by extent."""
    index = {t["name"]: [b for b in t["blocks"] if b["type"] != "chart"] for t in tabs}
    tabs_by_name = {t["name"]: t for t in tabs}
    for si, ws in enumerate(wb.worksheets):
        tab = tabs_by_name.get(ws.title)
        if tab is None:
            continue
        for ch in ws._charts:
            row, col, rows, cols = chart_geometry(ch.anchor)
            cname = type(ch).__name__
            kind = CHART_KINDS.get(cname, "bar")
            if cname not in CHART_KINDS:
                warnings.append("%s: chart type %s mapped to bar" % (ws.title, cname))
            if len(getattr(ch, "_charts", []) or []) > 1:
                warnings.append("%s: combo chart at %s mapped to %s" % (ws.title, a1(row, col), kind))
            refs, cat_ref = [], None
            for se in ch.series:
                if se.val is not None and se.val.numRef is not None:
                    refs.append(se.val.numRef.f)
                if se.cat is not None and cat_ref is None:
                    cat_ref = (se.cat.numRef.f if se.cat.numRef is not None else
                               se.cat.strRef.f if se.cat.strRef is not None else None)
            src, series_cols, cat_col, src_sheet, resolved = None, [], None, None, 0
            for ref in refs + ([cat_ref] if cat_ref else []):
                parsed = parse_ref(ref)
                if not parsed:
                    continue
                sheet, (c1, r1, c2, r2) = parsed
                inside = [b for b in index.get(sheet, []) if b["row"] <= r1 and r2 <= b["row"] + b["rows"] - 1
                          and b["col"] <= c1 and c2 <= b["col"] + b["cols"] - 1]
                inside.sort(key=lambda b: b["type"] not in ("table", "calc"))
                if not inside or (src is not None and inside[0] is not src):
                    continue
                src, src_sheet, resolved = inside[0], sheet, resolved + 1
                if ref == cat_ref:
                    cat_col = c1 - src["col"]
                else:
                    series_cols.append(c1 - src["col"])
            total = len(refs) + (1 if cat_ref else 0)
            conf = 0.3 if total == 0 else 0.5 if resolved < total else 0.9 if src_sheet == ws.title else 0.7
            block = {"id": "imp_%d_%d" % (si, len(tab["blocks"]) + 1), "type": "chart", "title": chart_title(ch),
                     "row": row, "col": col, "rows": rows, "cols": cols, "columns": [], "rowLabels": [], "sample": [],
                     "confidence": conf, "range": rng(row, col, row + rows - 1, col + cols - 1),
                     "detected": {"signals": ["chart %s" % cname, "series %d" % len(refs), "resolved %d/%d" % (resolved, total)],
                                  "reason": "chart: %d series resolve to %s" % (resolved, src["id"] if src else "no block")},
                     "chart": {"kind": kind, "sourceBlock": src["id"] if src else None, "sourceSheet": src_sheet,
                               "categoriesCol": cat_col if cat_col is not None else 0, "seriesCols": series_cols or [1]}}
            if src:
                block["source"] = src["id"]
            tab["blocks"].append(block)


# --------------------------------------------------------- 7. sheet and file


def scan_sheet(ws, cached_ws, si, warnings):
    grid = Grid(ws, cached_ws)
    det = Detector(grid)
    tab = {"name": ws.title, "purpose": "", "blocks": [], "unclassified": []}
    boxes, skip = listobject_boxes(grid, det, warnings)
    found = components(grid, skip)
    found = bridge(grid, found, det)
    pieces = []
    for box in found:
        peel_notes(grid, box)
        pieces.extend(split_box(grid, det, box))
    for p in pieces:
        peel_notes(grid, p)
    boxes.extend(pieces)
    boxes.sort(key=lambda b: (b.r1, b.c1))
    for box in boxes:
        if box.r1 == box.r2 and box.c1 == box.c2 and not box.forced:
            f = grid.cells[(box.r1, box.c1)]
            if not (f.kind == "text" and f.text_len() >= 40):
                tab["unclassified"].append({"range": box.range(), "reason": "lone-text" if f.kind == "text" else "stray-cell"})
                continue
            sig = "lone wrapped cell" if f.wrap else "prose by length only"
            block, reason = {"type": "text", "title": "", "row": max(1, f.r - 1), "col": f.c, "rows": 2 if f.r > 1 else 1,
                             "cols": 1, "columns": [], "rowLabels": [], "sample": [[str(f.v)]], "text": str(f.v),
                             "confidence": 0.6 if f.wrap else 0.4, "range": box.range(),
                             "detected": {"signals": ["no title", sig], "reason": "text: " + sig},
                             "geometry": geometry(box, None, [], None, [f.c], [f.r], None, [])}, None
        else:
            block, reason = detect_block(grid, det, box, warnings)
        if block is None:
            tab["unclassified"].append({"range": box.range(), "reason": reason})
            continue
        block["id"] = "imp_%d_%d" % (si, len(tab["blocks"]) + 1)
        tab["unclassified"].extend(block.pop("extraNotes", []))
        tab["blocks"].append(block)
    if si == 0 and not any(f.is_num or f.is_f for f in grid.cells.values() if f.has_value):
        warnings.append("readme-like: first sheet '%s' has no numbers or formulas" % ws.title)
    return tab


def scan(path):
    warnings = []
    wb = load_workbook(path)
    try:
        cached = load_workbook(path, data_only=True)
    except Exception:  # cache is optional
        cached = None
    tabs = []
    for si, ws in enumerate(wb.worksheets):
        if ws.sheet_state != "visible":
            warnings.append("hidden sheet skipped: %s" % ws.title)
            continue
        tabs.append(scan_sheet(ws, cached[ws.title] if cached is not None else None, si, warnings))
    chart_blocks(wb, tabs, warnings)
    for tab in tabs:
        tab["blocks"].sort(key=lambda b: (b["row"], b["col"]))
    return {"source": os.path.abspath(path), "tabs": tabs, "warnings": warnings}


def main(argv):
    if len(argv) != 3:
        print(json.dumps({"ok": False, "errors": ["usage: import_scan.py <input.xlsx> <out.json>"]}))
        return 1
    src, out = argv[1], argv[2]
    try:
        plan = scan(src)
        with open(out, "w") as fh:
            json.dump(plan, fh, indent=2, default=str)
    except Exception as exc:
        print(json.dumps({"ok": False, "errors": ["%s: %s" % (type(exc).__name__, exc)]}))
        return 1
    print(json.dumps({
        "ok": True, "sheets": len(plan["tabs"]),
        "blocks": sum(len(t["blocks"]) for t in plan["tabs"]),
        "unclassified": sum(len(t["unclassified"]) for t in plan["tabs"]),
        "warnings": plan["warnings"],
    }))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))

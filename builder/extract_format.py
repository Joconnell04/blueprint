#!/usr/bin/env python3
"""extract_format.py: read .xlsx and .pptx files and derive a formatting skill's tokens.

Usage:
  python3 builder/extract_format.py --out result.json file.xlsx [file.pptx ...]
  python3 builder/extract_format.py --out result.json --merge observations.json file.pptx
  python3 builder/extract_format.py --skill-md "my-format" "When to use it" file.xlsx
  python3 builder/extract_format.py --tokens style.json deck.pptx

Output (result.json, and stdout when --out is absent):

  {"format": {"workbook": {...}, "deck": {...}},
   "confidence": {"workbook.font": 0.97, ...},
   "evidence": {"workbook.font": "Arial on 214 of 214 styled cells", ...},
   "unknowns": ["..."],
   "rules": ["..."],
   "samples": {"palette": ["1E3A5F", ...], "fonts": ["Lato", ...]}}

`format.workbook` uses the keys build_xlsx.py already reads from a skill's
`format:` block, so the block can be pasted into a SKILL.md and built with no
change to the builder. `format.deck` is the deck half of the schema; both are
documented in docs/skill-format-block.md.

Images are not parsed here. Their paths land in `unknowns`; Claude reads the
screenshot in chat and sends its reading back as observations (which the server
hands this script as a merge file). Every `unknowns` entry is written as a
question for Claude to answer from a screenshot or from the user, so it can be
asked as it stands: no entry names a flag of this script.

Python 3.9. openpyxl for .xlsx, python-pptx for .pptx (either may be missing:
the file kinds that need it are then reported as unknowns).
"""

import json
import os
import re
import sys
from collections import Counter

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - reported as an unknown at run time
    load_workbook = None
    get_column_letter = None

try:
    import import_scan                      # cell feature helpers, reused as-is
except ImportError:                         # pragma: no cover
    import_scan = None

XLSX_EXT = (".xlsx", ".xlsm", ".xltx")
PPTX_EXT = (".pptx", ".potx", ".ppsx")
IMAGE_EXT = (".png", ".jpg", ".jpeg", ".gif", ".webp", ".heic", ".tiff", ".bmp", ".pdf")

HEAVY = ("medium", "thick", "double")
BLACKS = ("000000", None)
GREY_RE = re.compile(r"^([0-9A-F])\1([0-9A-F])\2([0-9A-F])\3$")
TOTAL_RE = re.compile(r"^\s*((sub)?total|sum\b|net\s|grand)", re.I)
COVER_RE = re.compile(r"^(read\s*me|readme|cover|about|start\s*here|overview)$", re.I)
INPUT_TAB_RE = re.compile(r"(assumption|input|driver)", re.I)
DASH_RE = re.compile(r"[–—]")
EMU_PT = 12700.0
GRID = 30                                   # the canvas cell, docs/research/deck-planning.md

# The fields the overall workbook / deck score averages. A field with no
# evidence counts as zero, so a file that shows none of them scores low.
WORKBOOK_KEYS = ("font", "size", "input_color", "formula_color", "note_color", "note_italic",
                 "title_rule", "title_rule_color", "label_divider", "body_borders", "gridlines",
                 "fills", "currency", "percent")
DECK_KEYS = ("slideSize.width", "palette.primary", "palette.secondary", "palette.bg",
             "fonts.major", "typeScale.headline.size", "typeScale.kicker.size",
             "typeScale.cardHeader.size", "card.shape", "card.border", "card.borderColor",
             "furniture.headline", "furniture.pageNumber", "backgrounds.content",
             "backgrounds.divider")

TYPE_ROLES = ("headline", "kicker", "body", "caption", "cardHeader", "cardSub", "cardBody",
              "pageNumber", "display")
FURNITURE_ROLES = ("logo", "pageNumber", "kicker", "headline", "body", "source")
# `body` is measured from the content rather than tallied as a repeated box, so it
# is set after this loop; every other role is one shape that repeats.
TALLIED_FURNITURE = tuple(role for role in FURNITURE_ROLES if role != "body")


# --------------------------------------------------------------- small helpers


def hexof(value):
    """Six upper case hex digits from an openpyxl/python-pptx colour, or None."""
    if value is None:
        return None
    text = str(value).strip().lstrip("#").upper()
    if len(text) == 8:                      # openpyxl ARGB
        text = text[2:]
    return text if re.match(r"^[0-9A-F]{6}$", text) else None


def is_grey(colour):
    """A mid grey. Black is body text and white is a background, so neither counts."""
    if not colour:
        return False
    r, g, b = int(colour[0:2], 16), int(colour[2:4], 16), int(colour[4:6], 16)
    return max(r, g, b) - min(r, g, b) <= 16 and 40 <= (r + g + b) / 3.0 <= 216


def score(share, support, ceiling=0.97):
    """Confidence from how dominant a reading is and how many cells or shapes back it.

    One observation never reaches 0.8: a single cell is a coincidence, three
    agreeing cells are a convention.
    """
    if not support:
        return 0.0
    grow = min(1.0, support / 3.0)
    return round(min(ceiling, float(share) * (0.55 + 0.45 * grow)), 2)


def near(value, step=1):
    if value is None:
        return None
    rounded = round(float(value) / step) * step
    return int(rounded) if abs(rounded - int(rounded)) < 1e-6 else round(rounded, 2)


def cell_of(x, y):
    """A pt position as a 1-based cell on the 30 pt canvas grid."""
    return [int(round(float(x) / GRID)) + 1, int(round(float(y) / GRID)) + 1]


def percentile(values, p):
    """The p-th percentile of a list by nearest rank, or None when it is empty.

    Nearest rank rather than interpolated: every answer is an edge some shape
    really sits on, which is what a measure has to be.
    """
    if not values:
        return None
    ordered = sorted(values)
    k = int(round((len(ordered) - 1) * (float(p) / 100.0)))
    return ordered[max(0, min(k, len(ordered) - 1))]


def median(values):
    return percentile(values, 50)


def text_case(text):
    letters = [ch for ch in text if ch.isalpha()]
    if not letters:
        return None
    if text == text.upper():
        return "upper"
    words = [w for w in re.split(r"[\s|]+", text) if w and w[0].isalpha()]
    long_words = [w for w in words if len(w) > 3]
    if len(long_words) >= 2 and all(w[0].isupper() for w in long_words):
        return "title"
    return "sentence"


# Weight and style words a face name carries. "Times New Roman" keeps its Roman:
# only words that name a weight or a slope are stripped.
WEIGHT_WORDS = frozenset((
    "thin", "hairline", "extralight", "ultralight", "light", "book", "regular", "normal",
    "medium", "semibold", "demibold", "demi", "semi", "bold", "extrabold", "ultrabold",
    "extra", "ultra", "black", "heavy", "italic", "oblique"))
FACE_MAJORITY = 0.6                         # one face has to carry this much of its family
TYPEFACE_RE = re.compile(r'typeface="([^"]+)"')


def font_family(name):
    """"Montserrat SemiBold" -> "Montserrat": the family with its weight and slope words off."""
    if not name:
        return name
    parts = [part for part in re.split(r"[\s_-]+", str(name).strip()) if part]
    while len(parts) > 1 and parts[-1].lower() in WEIGHT_WORDS:
        parts.pop()
    return " ".join(parts) or str(name).strip()


# The five weights a token names, and every face-name word that folds onto one.
# A slope word is not a weight, so "Montserrat Medium Italic" is still medium.
WEIGHT_NAMES = {
    "thin": "light", "hairline": "light", "extralight": "light", "ultralight": "light",
    "light": "light", "book": "regular", "regular": "regular", "normal": "regular",
    "medium": "medium", "semibold": "semibold", "demibold": "semibold", "demi": "semibold",
    "semi": "semibold", "bold": "bold", "extrabold": "bold", "ultrabold": "bold",
    "black": "bold", "heavy": "bold"}
SLOPE_WORDS = frozenset(("italic", "oblique"))
BOLD_WEIGHTS = ("semibold", "bold")


def weight_of(name, bold=False):
    """"Montserrat Light" -> "light"; a face that names no weight answers by its bold flag.

    The name wins where it says something, because a deck set in Montserrat Light is
    light whatever the bold flag says; a light face carrying the flag is PowerPoint
    faux-bolding it, so the flag lifts the reading rather than being ignored.
    """
    parts = [part for part in re.split(r"[\s_-]+", str(name or "").strip()) if part]
    tail = []
    while parts and parts[-1].lower() in WEIGHT_WORDS:
        tail.insert(0, parts.pop().lower())
    tail = [word for word in tail if word not in SLOPE_WORDS]
    found = WEIGHT_NAMES.get("".join(tail))
    if found is None and tail:
        found = WEIGHT_NAMES.get(tail[-1])
    if found is None:
        return "bold" if bold else "regular"
    if bold and found not in BOLD_WEIGHTS:
        return "bold"
    return found


def run_weight(run):
    """The weight one run reads as: the face it names, else its bold flag."""
    name = run.faces.most_common(1)[0][0] if run.faces else run.font
    return weight_of(name, bool(run.bold))


def declared_typefaces(path):
    """Every typeface the file names: run and theme declarations, and embedded fonts.

    A weight suffix survives into the tokens only on a name the file declares as a
    face of its own, so "Montserrat SemiBold" can stand and an invented one cannot.
    """
    import zipfile
    out = set()
    try:
        with zipfile.ZipFile(path) as zf:
            for item in zf.namelist():
                if not item.endswith(".xml"):
                    continue
                try:
                    blob = zf.read(item).decode("utf-8", "replace")
                except Exception:           # noqa: BLE001 - one unreadable part proves nothing
                    continue
                for hit in TYPEFACE_RE.findall(blob):
                    face = hit.strip()
                    if face and not face.startswith("+"):
                        out.add(face)
    except Exception:                       # noqa: BLE001 - not a zip, or unreadable
        return out
    return out


class FaceBook(object):
    """The fonts runs actually use, weighted by character count.

    Faces of one family are counted together, so "Montserrat", "Montserrat Light" and
    "Montserrat Medium" elect Montserrat rather than splitting the vote three ways.
    """

    def __init__(self):
        self.chars = Counter()
        self.runs = Counter()

    def add(self, name, chars):
        name = (name or "").strip()
        if not name or chars <= 0:
            return
        self.chars[name] += chars
        self.runs[name] += 1

    def __len__(self):
        return sum(self.runs.values())

    def best(self, declared=()):
        """(name, share, runs, faces) for the family the runs are mostly set in.

        The name keeps its weight suffix only when one face carries the family on
        its own and the file declares that face as a typeface of its own; otherwise
        it is the family name.
        """
        if not self.chars:
            return None
        families = Counter()
        for face, chars in self.chars.items():
            families[font_family(face)] += chars
        family, family_chars = families.most_common(1)[0]
        faces = [(f, n) for f, n in self.chars.most_common() if font_family(f) == family]
        name = family
        top, top_chars = faces[0]
        if top != family and top_chars >= FACE_MAJORITY * family_chars and top in declared:
            name = top
        runs = sum(self.runs[f] for f, _ in faces)
        total = float(sum(self.chars.values())) or 1.0
        return name, family_chars / total, runs, [f for f, _ in faces]


class Tally(object):
    """Counts readings of one field; `best` is the modal value with its share."""

    def __init__(self):
        self.counts = Counter()

    def add(self, value, weight=1):
        if value is not None and value != "":
            self.counts[value] += weight

    def __len__(self):
        return sum(self.counts.values())

    def best(self):
        if not self.counts:
            return None, 0.0, 0
        value, count = self.counts.most_common(1)[0]
        total = sum(self.counts.values())
        return value, count / float(total), total

    def ranked(self, limit=None):
        items = [v for v, _ in self.counts.most_common()]
        return items[:limit] if limit else items


class Draft(object):
    """The result under construction: tokens, per field confidence, evidence, unknowns."""

    def __init__(self):
        self.format = {"workbook": {}, "deck": {}}
        self.confidence = {}
        self.evidence = {}
        self.unknowns = []
        self.rules = []
        self.palette = Counter()
        self.fonts = Counter()
        # Run fonts by role, and every typeface the files declare (deck fonts).
        self.faces = {"major": FaceBook(), "minor": FaceBook()}
        self.typefaces = set()
        # Every body shape on a content slide, and the headline's left edge per
        # slide: the two readings `furniture.body` and the headline column need
        # whole distributions for, rather than a modal value (Q18).
        self.body_boxes = []
        self.headline_lefts = []
        self.content_slides = 0

    def set(self, path, value, conf, why):
        if value is None:
            return
        node = self.format
        parts = path.split(".")
        for part in parts[:-1]:
            node = node.setdefault(part, {})
        node[parts[-1]] = value
        self.confidence[path] = round(float(conf), 2)
        self.evidence[path] = why

    def take(self, path, tally, why, transform=None, ceiling=0.97):
        """Set a field from a tally's modal reading. Returns the value or None."""
        value, share, support = tally.best()
        if value is None:
            return None
        out = transform(value) if transform else value
        if out is None:
            return None
        self.set(path, out, score(share, support, ceiling),
                 why % {"value": value, "share": int(round(share * 100)), "n": support})
        return out

    def unknown(self, text):
        """Record a question. Every entry is asked as it stands, so it reads as one."""
        if text not in self.unknowns:
            self.unknowns.append(text)

    def rule(self, text):
        if text and text not in self.rules:
            self.rules.append(text)

    def overall(self, section, keys):
        block = self.format.get(section) or {}
        if not block:
            return
        total = 0.0
        for key in keys:
            total += self.confidence.get("%s.%s" % (section, key), 0.0)
        self.confidence[section] = round(total / len(keys), 2)
        self.evidence[section] = "mean confidence over %d key fields" % len(keys)

    def result(self):
        out = {"format": {}, "confidence": dict(self.confidence), "evidence": dict(self.evidence),
               "unknowns": list(self.unknowns), "rules": list(self.rules),
               "samples": {"palette": [c for c, _ in self.palette.most_common(16)],
                           "fonts": [f for f, _ in self.fonts.most_common(8)]}}
        for section in ("workbook", "deck"):
            if self.format.get(section):
                out["format"][section] = self.format[section]
        return out


# ------------------------------------------------------------------ xlsx scan


def nf_family(nf):
    """Which of the builder's format slots an Excel number format belongs to."""
    if not nf or nf in ("General", "@"):
        return None
    low = nf.lower()
    if "%" in nf:
        return "percent"
    if "$" in nf or "[$" in nf or "£" in nf or "€" in nf:
        return "currency2" if re.search(r"0\.0+", nf) else "currency"
    if '"x"' in low or low.endswith('x"') or low.endswith("x"):
        return "multiple"
    if re.match(r"^[#0,]+$", nf):
        return "count"
    if re.match(r"^[#0,]+\.0+$", nf):
        return "number"
    return None


class SheetRead(object):
    """One worksheet's cells, already sorted into the roles the house format names."""

    def __init__(self, ws):
        self.ws = ws
        self.feats = {}
        self.names = {}
        self.fills = {}
        self.titles = set()
        self.headers = set()
        self.labels = set()
        self.notes = set()
        self.inputs = set()
        self.formulas = set()
        self.totals = set()
        self.values = set()


def read_sheet(ws):
    """Load one sheet into a SheetRead, classifying every occupied cell."""
    read = SheetRead(ws)
    rows = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            feat = import_scan.feat_from_cell(cell, None)
            key = (cell.row, cell.column)
            read.feats[key] = feat
            # An unstyled cell still renders in the workbook's default font, so read
            # the resolved font rather than skipping the cell.
            read.names[key] = cell.font.name if cell.font is not None else None
            fill = None
            if cell.has_style and cell.fill is not None and cell.fill.fill_type == "solid":
                fill = hexof(getattr(cell.fill.fgColor, "rgb", None))
            read.fills[key] = fill
            rows.setdefault(cell.row, []).append(key)

    for r, keys in rows.items():
        keys.sort()
        first = keys[0]
        line = [read.feats[k] for k in keys]
        bold_text = [k for k in keys if read.feats[k].bold and read.feats[k].is_text]
        numeric = [k for k in keys if read.feats[k].is_num]

        # A block title: bold text under a heavy rule, or bold text alone on its row.
        for key in bold_text:
            feat = read.feats[key]
            if feat.b["b"] in HEAVY or (len(keys) == 1 and not feat.boxed):
                read.titles.add(key)

        # A header row: two or more bold, boxed labels with numbers underneath.
        boxed_bold = [k for k in bold_text if read.feats[k].boxed and k not in read.titles]
        below = rows.get(r + 1) or []
        if len(boxed_bold) >= 2 and any(read.feats[k].is_num for k in below):
            read.headers.update(boxed_bold)

        is_total = read.feats[first].is_text and TOTAL_RE.match(str(read.feats[first].v or ""))
        for key in keys:
            feat = read.feats[key]
            if key in read.titles or key in read.headers:
                continue
            if feat.is_text and not feat.bold and not feat.boxed and (
                    feat.italic or is_grey(feat.color)):
                read.notes.add(key)
                continue
            if feat.is_num:
                read.values.add(key)
                if is_total or (feat.b["t"] in HEAVY and feat.bold):
                    read.totals.add(key)
                if feat.kind == "formula":
                    read.formulas.add(key)
                elif feat.color not in BLACKS:
                    read.inputs.add(key)
                else:
                    read.formulas.add(key)
            elif feat.is_text and key == first and numeric:
                read.labels.add(key)
            elif feat.is_text and str(feat.v).strip().upper() == "TBD" and feat.color not in BLACKS:
                read.inputs.add(key)
    return read


def column_roles(read):
    """Map each column to (label | value | note, how many cells voted for it)."""
    roles = {}
    per_col = {}
    for key in read.labels:
        per_col.setdefault(key[1], Counter())["label"] += 1
    for key in read.values:
        per_col.setdefault(key[1], Counter())["value"] += 1
    for key in read.notes:
        per_col.setdefault(key[1], Counter())["note"] += 1
    for col, counts in per_col.items():
        roles[col] = counts.most_common(1)[0]
    return roles


def scan_xlsx(path, draft, tallies):
    """Fold one workbook into the shared tallies."""
    workbook = load_workbook(path)
    sheets = [ws for ws in workbook.worksheets if ws.sheet_state == "visible"]
    if not sheets:
        draft.unknown("Every sheet in %s is hidden: which tab shows the format, or is there "
                      "another example file?" % os.path.basename(path))
        return
    frozen = []
    input_only = []
    covers = []
    for ws in sheets:
        read = read_sheet(ws)
        grid = ws.sheet_view.showGridLines
        tallies["gridlines"].add(True if grid is None else bool(grid))   # unset means on
        if ws.freeze_panes:
            frozen.append("%s!%s" % (ws.title, ws.freeze_panes))
        if COVER_RE.match(ws.title.strip()):
            covers.append(ws.title)
            sizes = [read.feats[k].size for k in read.titles | set(read.feats)
                     if read.feats[k].bold]
            if sizes:
                tallies["title_size"].add(near(max(sizes)))
        if read.inputs and not read.formulas and len(read.inputs) >= 3:
            input_only.append(ws.title)
        elif INPUT_TAB_RE.search(ws.title) and read.inputs and not read.formulas:
            input_only.append(ws.title)

        for key, feat in read.feats.items():
            name = read.names.get(key)
            if name:
                draft.fonts[name] += 1
                tallies["font"].add(name)
            if key not in read.titles:
                tallies["size"].add(near(feat.size))
            if feat.color:
                draft.palette[feat.color] += 1
            fill = read.fills.get(key)
            if fill:
                draft.palette[fill] += 1
            if DASH_RE.search(str(feat.v or "")):
                tallies["dashes"].add("yes")

        for key in read.titles:
            feat = read.feats[key]
            style = feat.b["b"]
            tallies["title_rule"].add("medium" if style in HEAVY else (style or "none"))
            if style:
                tallies["title_rule_color"].add(hexof(_border_colour(ws, key, "bottom")))
        for key in read.headers:
            tallies["fills"].add(bool(read.fills.get(key)))
            if read.fills.get(key):
                tallies["header_fill"].add(read.fills[key])
        for key in read.labels:
            style = read.feats[key].b["r"]
            tallies["label_divider"].add("dashed" if style == "dashed"
                                         else ("thin" if style else "none"))
        for key in read.inputs:
            tallies["input_color"].add(read.feats[key].color)
        for key in read.formulas:
            tallies["formula_color"].add(read.feats[key].color or "000000")
        for key in read.notes:
            tallies["note_color"].add(read.feats[key].color)
            tallies["note_italic"].add(bool(read.feats[key].italic))
        for key in read.values:
            tallies["body_borders"].add("thin" if read.feats[key].boxed else "none")
            family = nf_family(read.feats[key].nf)
            if family:
                tallies[family].add(read.feats[key].nf)

        roles = column_roles(read)
        for col, (role, votes) in roles.items():
            width = ws.column_dimensions[get_column_letter(col)].width
            if width:
                # A column that holds five notes says more about the note width
                # than a cover subtitle sitting alone in the label column.
                tallies["%s_width" % role].add(near(width), weight=votes)

    if frozen:
        draft.confidence["workbook.freeze_panes"] = score(
            len(frozen) / float(len(sheets)), len(frozen))
        draft.evidence["workbook.freeze_panes"] = "frozen at " + ", ".join(frozen[:3])
        draft.rule("Freeze the header rows: the source file freezes %s." % frozen[0])
    if covers:
        tallies["cover_tab"].add(True, weight=3)
        draft.rule("Open the workbook on a '%s' tab: title, one line of purpose, "
                   "contents, and the rules for reading a cell." % covers[0])
    else:
        tallies["cover_tab"].add(False)
        draft.unknown("Does this format open on a cover tab, and what goes on it? %s has none."
                      % os.path.basename(path))
    if input_only:
        draft.confidence["workbook.inputs_separated"] = 0.85
        draft.evidence["workbook.inputs_separated"] = (
            "%s holds inputs and no formulas" % ", ".join(input_only[:3]))
        draft.rule("Inputs live on their own tab (%s); every other tab reads from it."
                   % input_only[0])


def _border_colour(ws, key, side):
    cell = ws.cell(key[0], key[1])
    edge = getattr(cell.border, side, None)
    if edge is None or edge.color is None:
        return None
    return getattr(edge.color, "rgb", None) if edge.color.type == "rgb" else None


def finish_xlsx(draft, tallies):
    """Turn the workbook tallies into the `format.workbook` block."""
    draft.take("workbook.font", tallies["font"], "%(value)s on %(share)s%% of %(n)d styled cells")
    draft.take("workbook.size", tallies["size"], "%(value)s pt on %(share)s%% of %(n)d body cells")
    draft.take("workbook.title_size", tallies["title_size"], "cover title measured %(value)s pt")
    draft.take("workbook.input_color", tallies["input_color"],
               "%(n)d literal values carry font colour %(value)s")
    draft.take("workbook.formula_color", tallies["formula_color"],
               "%(n)d formula and derived cells are %(value)s")
    draft.take("workbook.note_color", tallies["note_color"],
               "%(n)d unboxed annotation cells are %(value)s")
    draft.take("workbook.note_italic", tallies["note_italic"],
               "%(share)s%% of %(n)d notes are italic")
    draft.take("workbook.title_rule", tallies["title_rule"],
               "%(share)s%% of %(n)d block titles sit on a %(value)s bottom border")
    draft.take("workbook.title_rule_color", tallies["title_rule_color"],
               "the title rule is %(value)s on %(n)d titles", transform=hexof)
    draft.take("workbook.label_divider", tallies["label_divider"],
               "%(share)s%% of %(n)d label cells end in a %(value)s right border")
    draft.take("workbook.body_borders", tallies["body_borders"],
               "%(share)s%% of %(n)d value cells are %(value)s boxed")
    draft.take("workbook.gridlines", tallies["gridlines"],
               "gridlines are %(value)s on %(share)s%% of %(n)d sheets")
    draft.take("workbook.fills", tallies["fills"],
               "%(share)s%% of %(n)d header cells are filled: %(value)s")
    draft.take("workbook.header_fill", tallies["header_fill"], "header fill %(value)s on %(n)d cells")
    draft.take("workbook.cover_tab", tallies["cover_tab"], "cover tab present: %(value)s")
    for key in ("currency", "percent", "multiple", "count"):
        draft.take("workbook.%s" % key, tallies[key],
                   "%(n)d cells use %(value)s")
    draft.take("workbook.label_width", tallies["label_width"], "label columns measure %(value)s")
    draft.take("workbook.value_width", tallies["value_width"], "value columns measure %(value)s")
    draft.take("workbook.note_width", tallies["note_width"], "note columns measure %(value)s")
    if len(tallies["dashes"]):
        draft.set("workbook.dash_replace", False, 0.6,
                  "%d cells contain an em or en dash, so the style keeps them" % len(tallies["dashes"]))
    else:
        draft.unknown("Should em and en dashes be rewritten (dash_replace), and how long may a "
                      "note run (max_note_chars)? Both are authoring rules a file never shows.")
    block = draft.format.get("workbook") or {}
    if "input_color" not in block:
        draft.unknown("Which colour marks an editable input? No coloured input cells were found.")
    if "title_rule" not in block:
        draft.unknown("How is a block title set off: a rule under it, a fill, bold on its own? "
                      "No block titles were found.")
    if "note_color" not in block:
        draft.unknown("How is a note or a source line styled? None were found.")
    draft.overall("workbook", WORKBOOK_KEYS)


# ------------------------------------------------------------------ pptx scan

A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
THEME_SLOTS = ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3",
               "accent4", "accent5", "accent6", "hlink", "folHlink")
SCHEME_MAP = {"DARK_1": "dk1", "TEXT_1": "dk1", "LIGHT_1": "lt1", "BACKGROUND_1": "lt1",
              "DARK_2": "dk2", "TEXT_2": "dk2", "LIGHT_2": "lt2", "BACKGROUND_2": "lt2",
              "ACCENT_1": "accent1", "ACCENT_2": "accent2", "ACCENT_3": "accent3",
              "ACCENT_4": "accent4", "ACCENT_5": "accent5", "ACCENT_6": "accent6",
              "HYPERLINK": "hlink", "FOLLOWED_HYPERLINK": "folHlink"}


def theme_of(prs):
    """The master's colour scheme and font scheme, read straight from the theme part."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    colours, fonts = {}, {}
    try:
        part = prs.slide_masters[0].part.part_related_by(RT.THEME)
        xml = part.blob.decode("utf-8", "replace")
    except Exception:                       # noqa: BLE001 - a theme is optional
        return colours, fonts
    scheme = re.search(r"<a:clrScheme.*?</a:clrScheme>", xml, re.S)
    if scheme:
        for slot in THEME_SLOTS:
            hit = re.search(r"<a:%s>.*?val=\"([0-9A-Fa-f]{6})\".*?</a:%s>" % (slot, slot),
                            scheme.group(0), re.S)
            if hit:
                colours[slot] = hit.group(1).upper()
    for kind in ("major", "minor"):
        hit = re.search(r"<a:%sFont><a:latin typeface=\"([^\"]*)\"" % kind, xml)
        if hit and hit.group(1):
            fonts[kind] = hit.group(1)
    return colours, fonts


def colour_of(fmt, theme):
    """A hex colour from a python-pptx ColorFormat, resolving theme references."""
    if fmt is None:
        return None
    try:
        kind = fmt.type
    except Exception:                       # noqa: BLE001
        return None
    if kind is None:
        return None
    try:
        if str(kind).startswith("RGB") or getattr(kind, "value", None) == 1:
            return hexof(fmt.rgb)
    except Exception:                       # noqa: BLE001
        pass
    try:
        slot = SCHEME_MAP.get(str(fmt.theme_color).split(" ")[0].split(".")[-1])
        if slot:
            return theme.get(slot)
    except Exception:                       # noqa: BLE001
        return None
    return None


def fill_colour(shape, theme):
    try:
        fill = shape.fill
        if fill.type is not None and int(fill.type) == 1:
            return colour_of(fill.fore_color, theme)
    except Exception:                       # noqa: BLE001
        return None
    return None


def preset_of(shape):
    try:
        geom = shape._element.find(".//%sprstGeom" % A_NS)
    except Exception:                       # noqa: BLE001
        return None, None
    if geom is None:
        return None, None
    adj = None
    node = geom.find("%savLst/%sgd" % (A_NS, A_NS))
    if node is not None and node.get("name") == "adj":
        hit = re.search(r"val\s+(-?\d+)", node.get("fmla") or "")
        if hit:
            adj = int(hit.group(1))
    return geom.get("prst"), adj


class Run(object):
    """One text run with the geometry of the shape it sits in."""

    __slots__ = ("text", "size", "bold", "colour", "font", "faces", "x", "y", "w", "h", "para")

    def __init__(self, text, size, bold, colour, font, box, para, faces=None):
        self.text, self.size, self.bold, self.colour, self.font = text, size, bold, colour, font
        self.faces = faces or Counter()     # face name -> characters set in it
        self.x, self.y, self.w, self.h = box
        self.para = para


def shape_runs(shape, theme, box):
    out = []
    if not shape.has_text_frame:
        return out
    for pi, para in enumerate(shape.text_frame.paragraphs):
        text = "".join(r.text for r in para.runs).strip()
        if not text:
            continue
        run = para.runs[0]
        size = run.font.size.pt if run.font.size is not None else (
            para.font.size.pt if para.font.size is not None else None)
        colour = colour_of(run.font.color, theme)
        # Every explicit face in the paragraph with the characters set in it. A run
        # that names none inherits, and the paragraph or the theme answers for it.
        faces = Counter()
        for part in para.runs:
            name = (part.font.name or "").strip()
            chars = len(part.text.strip())
            if name and chars:
                faces[name] += chars
        if not faces and (para.font.name or "").strip():
            faces[para.font.name.strip()] = len(text)
        out.append(Run(text, near(size, 1) if size else None, bool(run.font.bold), colour,
                       run.font.name, box, pi, faces))
    return out


def slide_background(slide, theme, width, height):
    """The slide's own background fill, or a full bleed rectangle standing in for one."""
    try:
        colour = colour_of(slide.background.fill.fore_color, theme)
        if colour:
            return colour
    except Exception:                       # noqa: BLE001
        pass
    for shape in slide.shapes:
        try:
            w, h = shape.width / EMU_PT, shape.height / EMU_PT
        except Exception:                   # noqa: BLE001
            continue
        if w >= width * 0.98 and h >= height * 0.98:
            colour = fill_colour(shape, theme)
            if colour:
                return colour
    return None


def scan_pptx(path, draft, tallies):
    from pptx import Presentation
    prs = Presentation(path)
    width = round(prs.slide_width / EMU_PT, 1)
    height = round(prs.slide_height / EMU_PT, 1)
    tallies["slide_w"].add(near(width))
    tallies["slide_h"].add(near(height))
    theme, theme_fonts = theme_of(prs)
    draft.typefaces.update(declared_typefaces(path))
    for slot, colour in theme.items():
        tallies["theme.%s" % slot].add(colour)
    for kind, name in theme_fonts.items():
        # The theme font is the fallback, not the reading: a run overrides it, and in
        # a real deck every run did (research item 3).
        tallies["font.theme.%s" % kind].add(name)
        draft.fonts[name] += 4

    slides = list(prs.slides)
    for index, slide in enumerate(slides):
        boxes = []
        cards = []
        pictures = []
        for shape in slide.shapes:
            try:
                box = (round(shape.left / EMU_PT, 1), round(shape.top / EMU_PT, 1),
                       round(shape.width / EMU_PT, 1), round(shape.height / EMU_PT, 1))
            except Exception:               # noqa: BLE001 - a shape without geometry
                continue
            if box[2] >= width * 0.98 and box[3] >= height * 0.98:
                continue                    # a background rectangle, not content
            prst, adj = preset_of(shape)
            if shape.shape_type is not None and "PICTURE" in str(shape.shape_type):
                pictures.append(box)
                continue
            runs = shape_runs(shape, theme, box)
            for run in runs:
                if run.colour:
                    draft.palette[run.colour] += 1
                if run.font:
                    draft.fonts[run.font] += 1
            if prst in ("roundRect", "rect", "snip1Rect") and (
                    fill_colour(shape, theme) or _line_pt(shape)):
                cards.append((shape, prst, adj, box, runs))
                continue
            if runs:
                boxes.append((box, runs))

        background = slide_background(slide, theme, width, height)
        if background:
            draft.palette[background] += 2
        kind = slide_kind(index, boxes, cards, pictures, background, height, tallies)
        tallies["bg.%s" % kind].add(background)
        read_slide(kind, boxes, cards, pictures, width, height, draft, tallies, theme)

    if not slides:
        draft.unknown("%s has no slides: is there another example deck, or a screenshot of one?"
                      % os.path.basename(path))


def _line_pt(shape):
    try:
        width = shape.line.width
        return round(width / EMU_PT, 2) if width else None
    except Exception:                       # noqa: BLE001
        return None


def slide_kind(index, boxes, cards, pictures, background, height, tallies):
    sizes = [r.size for _, runs in boxes for r in runs if r.size]
    biggest = max(sizes) if sizes else 0
    if cards or pictures:
        return "content"
    if biggest >= 40 and len(boxes) <= 3:
        return "title" if index == 0 else "divider"
    if index == 0:
        return "title"
    if len(boxes) >= 4:
        return "agenda"
    return "content"


def tally_run(tallies, slot, run, case=True):
    """Record one run under a type slot: size, bold, colour, case, and its weight.

    The weight is tallied on its own rather than inside the reading tuple: a deck
    that sets one role in two sizes still agrees about the weight (Q20).
    """
    tallies["type.%s" % slot].add((run.size, run.bold, run.colour,
                                   text_case(run.text) if case else None))
    tallies["weight.%s" % slot].add(run_weight(run))


def read_slide(kind, boxes, cards, pictures, width, height, draft, tallies, theme):
    """Pull type scale, card style and furniture out of one slide."""

    def faces(slot, runs):
        """Record the fonts these runs are set in, weighted by character count.

        Only the content slides count: a title and a divider are furniture, and the
        page number is its own mono slot.
        """
        if kind in ("title", "divider"):
            return
        for run in runs:
            for name, chars in run.faces.items():
                draft.faces[slot].add(name, chars)

    numbered = [(box, runs) for box, runs in boxes
                if len(runs) == 1 and len(runs[0].text) <= 4
                and re.match(r"^[0-9ivxIVX]+$", runs[0].text)
                and (box[1] < 60 or box[1] > height - 60)]
    for box, runs in numbered:
        # The page number flips to white on a dark slide, so the token comes from
        # the content slides and the flip is reported as a rule.
        slot = "pageNumber" if kind == "content" else "pageNumber.other"
        tally_run(tallies, slot, runs[0], case=False)
        tallies["furn.pageNumber"].add(tuple(near(v) for v in box))
    rest = [(box, runs) for box, runs in boxes if (box, runs) not in numbered]

    if kind in ("title", "divider"):
        slot = "display" if kind == "divider" else "displayTitle"
        for box, runs in rest:
            for run in runs:
                if run.size and run.size >= 28:
                    tally_run(tallies, slot, run)
        return
    draft.content_slides += 1

    body = [(box, runs) for box, runs in rest]
    if body:
        headline_box = max(body, key=lambda item: max(r.size or 0 for r in item[1]))
        above = [item for item in body if item[1][0].y < headline_box[0][1] and item is not headline_box]
        below = [item for item in body if item[1][0].y > headline_box[0][1]]
        run = headline_box[1][0]
        tally_run(tallies, "headline", run)
        tallies["furn.headline"].add(tuple(near(v) for v in headline_box[0]))
        draft.headline_lefts.append(near(headline_box[0][0]))
        faces("major", headline_box[1])
        if above:
            kicker = min(above, key=lambda item: item[0][1])
            tally_run(tallies, "kicker", kicker[1][0])
            tallies["furn.kicker"].add(tuple(near(v) for v in kicker[0]))
            faces("minor", kicker[1])
        for box, runs in below:
            role = "caption" if (box[1] > height * 0.85 or (runs[0].size or 99) <= 11) else "body"
            for r2 in runs:
                tally_run(tallies, role, r2)
            faces("minor", runs)
            if role == "caption" and box[0] < width * 0.5:
                tallies["furn.source"].add(tuple(near(v) for v in box))
            # Q18: the content measure is the union of the body shapes. The source
            # band at the foot is furniture, and so is everything above the headline.
            if box[1] <= height * 0.85:
                draft.body_boxes.append(tuple(near(v) for v in box))
        edges = [b for b, _ in body] + [c[3] for c in cards]
        tallies["margin.left"].add(near(min(b[0] for b in edges)))
        tallies["margin.right"].add(near(width - max(b[0] + b[2] for b in edges)))
        tallies["margin.top"].add(near(min(b[1] for b in edges + pictures)))
        tallies["margin.bottom"].add(near(height - max(b[1] + b[3] for b in edges)))

    for box in pictures:
        if box[1] < height * 0.25 and box[0] < width * 0.25:
            tallies["furn.logo"].add(tuple(near(v) for v in box))

    for shape, prst, adj, box, runs in cards:
        tallies["card.shape"].add(prst)
        if adj is not None:
            tallies["card.radius"].add(adj)
        line = _line_pt(shape)
        if line:
            tallies["card.border"].add(line)
        colour = None
        try:
            colour = colour_of(shape.line.color, theme)
        except Exception:                   # noqa: BLE001
            colour = None
        tallies["card.borderColor"].add(colour)
        tallies["card.fill"].add(fill_colour(shape, theme))
        roles = ("cardHeader", "cardSub", "cardBody")
        faces("minor", runs)
        for i, run in enumerate(runs):
            role = roles[i] if i < len(roles) else "cardBody"
            if len(runs) == 2 and i == 1:
                role = "cardBody"
            tally_run(tallies, role, run)
        # A card is body content. A hairline drawn as a filled rectangle is not, so
        # anything thinner than a line of type stays out of the measure.
        if box[3] >= 4 and box[1] <= height * 0.85:
            draft.body_boxes.append(tuple(near(v) for v in box))


def furniture_of(draft):
    return (draft.format.get("deck") or {}).get("furniture") or {}


def headline_column(draft):
    """`furniture.headline.x` from the median headline left edge, not the modal box.

    Every content slide starts its headline on the same column, but the box around
    it is only modal: one slide whose biggest run is a stat rather than a headline
    moves the mode and cannot move the median.
    """
    head = furniture_of(draft).get("headline")
    if not head or not draft.headline_lefts:
        return
    left = near(median(draft.headline_lefts))
    if left is None or left == head["x"]:
        return
    draft.set("deck.furniture.headline", dict(head, x=left, cell=cell_of(left, head["y"])),
              draft.confidence.get("deck.furniture.headline", 0.7),
              "the headline column is x=%s, the median left edge of %d headlines"
              % (left, len(draft.headline_lefts)))


def body_measure(draft):
    """`furniture.body`: the content measure the builder cuts its column module from.

    The union of the body shapes, trimmed: the 10th percentile left edge and the
    90th percentile right edge, so a legend parked in the margin or one block run
    wide does not widen the measure for the whole deck. The band runs from under
    the headline to the top of the source line, both of which are furniture the
    body has to clear (Q18).
    """
    boxes = draft.body_boxes
    if len(boxes) < 3:
        return
    furniture = furniture_of(draft)
    left = percentile([b[0] for b in boxes], 10)
    right = percentile([b[0] + b[2] for b in boxes], 90)
    head, source = furniture.get("headline"), furniture.get("source")
    top = head["y"] + head["h"] if head else percentile([b[1] for b in boxes], 10)
    bottom = source["y"] if source else percentile([b[1] + b[3] for b in boxes], 90)
    if None in (left, right, top, bottom):
        return
    x, y = near(left), near(top)
    w, h = near(right - left), near(bottom - top)
    if w < GRID or h < GRID:
        return
    inside = [b for b in boxes if b[0] >= left - 1 and b[0] + b[2] <= right + 1]
    draft.set("deck.furniture.body", {"x": x, "y": y, "w": w, "h": h, "cell": cell_of(x, y)},
              score(len(inside) / float(len(boxes)), len(boxes)),
              "content runs x=%s to %s on %d shapes over %d slides, between the headline "
              "and the source line" % (x, x + w, len(boxes), draft.content_slides))


def finish_pptx(draft, tallies):
    """Turn the deck tallies into the `format.deck` block."""
    if not len(tallies["slide_w"]):
        return
    for axis, key in (("width", "slide_w"), ("height", "slide_h")):
        value, share, support = tallies[key].best()
        if value:                            # read straight off the file, not inferred
            draft.set("deck.slideSize.%s" % axis, value, round(0.97 * share, 2),
                      "the deck measures %s pt %s" % (value, "wide" if axis == "width" else "tall"))
    draft.set("deck.slideSize.grid", GRID, 0.9, "the canvas lays slides out on a 30 pt grid")

    usage = draft.palette
    theme_slots = {slot: tallies["theme.%s" % slot].best()[0] for slot in THEME_SLOTS}
    content_bg = tallies["bg.content"].best()[0]
    palette = (("primary", "accent1"), ("secondary", "accent2"), ("accent", "accent3"),
               ("ink", "dk1"), ("muted", "accent4"), ("bg", "lt1"))
    # "muted" is whatever body copy actually is, not whatever the theme reserved
    # for it: a theme slot no shape uses describes nothing.
    body_colour = None
    for role in ("cardBody", "body", "caption"):
        reading = tallies["type.%s" % role].best()[0]
        if reading and reading[2]:
            body_colour = reading[2]
            break
    for name, slot in palette:
        colour = theme_slots.get(slot)
        if name == "bg" and content_bg:
            colour = content_bg
        if name == "muted" and body_colour and not usage.get(colour):
            draft.set("deck.palette.muted", body_colour, 0.85,
                      "body copy and captions are %s (the theme reserves %s and never uses it)"
                      % (body_colour, colour or "nothing"))
            continue
        if not colour:
            continue
        seen = usage.get(colour, 0)
        conf = 0.9 if seen else 0.7
        why = "theme clrScheme %s" % slot
        if seen:
            why += ", used on %d shapes or runs" % seen
        draft.set("deck.palette.%s" % name, colour, conf, why)
    if not theme_slots.get("accent1"):
        ranked = [c for c, _ in usage.most_common() if c not in (content_bg, "FFFFFF")]
        for name, colour in zip(("primary", "secondary", "accent", "ink", "muted"), ranked):
            draft.set("deck.palette.%s" % name, colour, 0.55,
                      "no theme scheme; %s is used on %d shapes or runs" % (colour, usage[colour]))
        draft.unknown("Which colours make up the palette? The deck carries no theme colour "
                      "scheme, so read them off a screenshot or ask.")
    if "muted" not in (draft.format.get("deck", {}).get("palette") or {}) and body_colour:
        draft.set("deck.palette.muted", body_colour, 0.7,
                  "body copy and captions are %s" % body_colour)

    # The fonts are whatever the runs are set in, weighted by characters; the theme
    # answers only when no run names one. A real deck's theme says Calibri and
    # every run overrides it with Montserrat (research item 3).
    fallback = {"major": ("major", "minor"), "minor": ("minor", "major")}
    for kind in ("major", "minor"):
        picked = None
        for slot in fallback[kind]:
            book = draft.faces.get(slot)
            if book and len(book):
                picked = book.best(draft.typefaces)
            if picked:
                break
        if picked:
            name, share, runs, found = picked
            folded = [f for f in found if f != name]
            draft.set("deck.fonts.%s" % kind, name, score(share, runs),
                      "%s on %d%% of the characters in %d runs on the content slides%s"
                      % (name, int(round(share * 100)), runs,
                         ", counting %s" % ", ".join(folded) if folded else ""))
            continue
        value, share, support = tallies["font.theme.%s" % kind].best()
        if value:
            draft.set("deck.fonts.%s" % kind, value, round(0.85 * share, 2),
                      "no run names a font; the theme %s font is %s" % (kind, value))
    mono = [f for f in draft.fonts if "mono" in f.lower()]
    if mono:
        draft.set("deck.fonts.mono", mono[0], 0.7, "%s appears on runs" % mono[0])

    if not len(tallies["type.display"]):
        tallies["type.display"] = tallies["type.displayTitle"]
        tallies["weight.display"] = tallies["weight.displayTitle"]
    other = tallies["type.pageNumber.other"].best()[0]
    page = tallies["type.pageNumber"].best()[0]
    if other and page and other[2] != page[2]:
        draft.rule("The page number is %s on a content slide and %s on a dark one."
                   % (page[2], other[2]))
    elif other and not page:
        tallies["type.pageNumber"] = tallies["type.pageNumber.other"]
        tallies["weight.pageNumber"] = tallies["weight.pageNumber.other"]

    for role in TYPE_ROLES:
        tally = tallies["type.%s" % role]
        value, share, support = tally.best()
        if not value or value[0] is None:
            continue
        size, bold, colour, case = value
        conf = score(share, support)
        draft.set("deck.typeScale.%s.size" % role, size, conf,
                  "%d runs read as %s at %s pt" % (support, role, size))
        # Q20: the weight the face name carries, so a Montserrat Light headline stops
        # being rewritten in bold. The bold flag follows the weight where there is one.
        weight, wshare, wsupport = tallies["weight.%s" % role].best()
        if weight:
            bold = weight in BOLD_WEIGHTS
            draft.set("deck.typeScale.%s.weight" % role, weight, score(wshare, wsupport),
                      "%d%% of %d %s runs are set %s" % (int(round(wshare * 100)), wsupport,
                                                         role, weight))
        draft.set("deck.typeScale.%s.bold" % role, bool(bold), conf,
                  "%s runs are set %s, so bold is %s" % (role, weight, bool(bold)) if weight
                  else "%s runs are bold: %s" % (role, bool(bold)))
        if colour:
            draft.set("deck.typeScale.%s.color" % role, colour, conf,
                      "%s runs are %s" % (role, colour))
        if case:
            draft.set("deck.typeScale.%s.case" % role, case, max(0.5, conf - 0.1),
                      "%s text reads as %s case" % (role, case))

    shape, share, support = tallies["card.shape"].best()
    if shape and support >= 2:
        draft.set("deck.card.shape", shape, score(share, support),
                  "%d repeated %s shapes" % (support, shape))
        draft.take("deck.card.radius", tallies["card.radius"],
                   "corner adj is %(value)s on %(n)d cards")
        draft.take("deck.card.border", tallies["card.border"],
                   "the card border is %(value)s pt on %(n)d cards")
        draft.take("deck.card.borderColor", tallies["card.borderColor"],
                   "the card border is %(value)s on %(n)d cards")
        draft.take("deck.card.fill", tallies["card.fill"], "cards are filled %(value)s on %(n)d cards")
    elif shape:
        draft.unknown("What shape, border and fill does a card carry? Only one card-like shape "
                      "was found, so confirm it from a screenshot.")

    for role in TALLIED_FURNITURE:
        value, share, support = tallies["furn.%s" % role].best()
        if not value:
            continue
        conf = score(share, support)
        x, y, w, h = value
        draft.set("deck.furniture.%s" % role, {"x": x, "y": y, "w": w, "h": h, "cell": cell_of(x, y)},
                  conf, "%s sits at %s, %s pt on %d slides" % (role, x, y, support))
    headline_column(draft)
    body_measure(draft)
    if "logo" not in (draft.format.get("deck", {}).get("furniture") or {}):
        draft.unknown("Where does the logo sit, and which image file is it? None was found on "
                      "the content slides.")

    for kind in ("title", "divider", "content", "agenda"):
        value, share, support = tallies["bg.%s" % kind].best()
        if value:
            draft.set("deck.backgrounds.%s" % kind, value, score(share, support),
                      "%d %s slides have a %s background" % (support, kind, value))
    for side in ("left", "right", "top", "bottom"):
        draft.take("deck.margins.%s" % side, tallies["margin.%s" % side],
                   "content starts %(value)s pt from the " + side)
    draft.unknown("Which slide archetypes does this format use, and when? Slide-by-slide layout "
                  "is not a token; the archetype catalogue covers those choices.")
    draft.overall("deck", DECK_KEYS)


# ----------------------------------------------------------------------- rules


def workbook_rules(block):
    """Plain English restatements of the workbook tokens, for a skill's prose.

    Only the tokens. The findings no token holds live in `result["rules"]` and
    skill_md() writes those under the same heading.
    """
    out = []
    font, size = block.get("font"), block.get("size")
    if font and size:
        title = block.get("title_size")
        out.append("%s %s everywhere%s." % (font, size,
                   ", except the cover title at %s" % title if title else ""))
    colours = []
    if block.get("input_color"):
        colours.append("%s means a cell someone can change" % block["input_color"])
    if block.get("formula_color"):
        colours.append("%s means a formula" % block["formula_color"])
    if colours:
        out.append("Font colour carries one meaning: " + "; ".join(colours) + ".")
    if block.get("note_color"):
        out.append("Notes are %s%s and sit outside the block's border."
                   % (block["note_color"], " italic" if block.get("note_italic") else ""))
    borders = []
    if block.get("body_borders") == "thin":
        borders.append("a thin box around every body cell")
    if block.get("label_divider") in ("dashed", "thin"):
        borders.append("a %s divider after the label column" % block["label_divider"])
    if block.get("title_rule") in ("medium", "thin"):
        borders.append("a %s %s rule under every block title"
                       % (block["title_rule"], block.get("title_rule_color") or ""))
    if borders:
        out.append("Borders do the structural work: " + ", ".join(borders) + ".")
    if block.get("gridlines") is False:
        out.append("Gridlines off on every sheet.")
    if block.get("fills") is False:
        out.append("No fills anywhere, including header rows.")
    elif block.get("header_fill"):
        out.append("Header rows are filled %s; nothing else is." % block["header_fill"])
    formats = [("money", "currency"), ("rates", "percent"), ("multiples", "multiple"),
               ("counts", "count")]
    shown = ["%s as %s" % (label, block[key]) for label, key in formats if block.get(key)]
    if shown:
        out.append("Number formats: " + ", ".join(shown) + ".")
    widths = [("label", "label_width"), ("value", "value_width"), ("note", "note_width")]
    shown = ["%s %s" % (label, block[key]) for label, key in widths if block.get(key)]
    if shown:
        out.append("Column widths: " + ", ".join(shown) + ".")
    return out


def deck_rules(block):
    out = []
    size = block.get("slideSize") or {}
    if size.get("width"):
        out.append("Slides are %s by %s pt, laid out on a %s pt grid."
                   % (size["width"], size.get("height"), size.get("grid", GRID)))
    palette = block.get("palette") or {}
    if palette:
        out.append("Palette: " + ", ".join("%s %s" % (k, v) for k, v in sorted(palette.items())) + ".")
    fonts = block.get("fonts") or {}
    if fonts.get("major"):
        out.append("%s for everything%s." % (fonts["major"],
                   "; %s for numbers and labels" % fonts["mono"] if fonts.get("mono") else ""))
    scale = block.get("typeScale") or {}
    parts = ["%s %s pt" % (role, scale[role]["size"]) for role in TYPE_ROLES
             if scale.get(role, {}).get("size")]
    if parts:
        out.append("Type scale: " + ", ".join(parts) + ".")
    if scale.get("headline", {}).get("case"):
        out.append("Headlines are %s case and state the takeaway, not the topic."
                   % scale["headline"]["case"])
    card = block.get("card") or {}
    if card.get("shape"):
        out.append("Cards are %s, %s pt border in %s, filled %s."
                   % (card["shape"], card.get("border"), card.get("borderColor"), card.get("fill")))
    furniture = block.get("furniture") or {}
    parts = ["%s at %s, %s" % (role, furniture[role]["x"], furniture[role]["y"])
             for role in FURNITURE_ROLES if furniture.get(role)]
    if parts:
        out.append("Furniture never moves: " + ", ".join(parts) + ".")
    backgrounds = block.get("backgrounds") or {}
    if backgrounds:
        out.append("Backgrounds: " + ", ".join("%s %s" % (k, v)
                   for k, v in sorted(backgrounds.items())) + ".")
    return out


# ------------------------------------------------------------- deck builder JSON

# The deck block names its colours and fonts for a reader; build_pptx.py names the
# same things for a writer. These pairs are the whole difference between the two.
TOKEN_PALETTE = (("primary", "accent"), ("secondary", "accent2"), ("ink", "ink"),
                 ("muted", "muted"), ("bg", "bg"))
TOKEN_FONTS = (("major", "display"), ("minor", "body"), ("mono", "mono"))
TOKEN_CARD = (("shape", "shape"), ("radius", "radiusAdj"), ("border", "borderPt"),
              ("borderColor", "borderColor"), ("fill", "fill"))
TOKEN_ROLE = {"display": "dividerTitle"}
# Groups the two vocabularies spell the same way, so they pass straight through.
# A block may carry them even though no file shows them: they are written by hand
# or merged in as observations, and without this they never reach the builder.
TOKEN_SAME = ("table", "chart", "kicker", "divider", "diagram")
TOKEN_TABLE_COLOURS = ("headerFill", "ruleColor")


def luminance(colour):
    if not colour:
        return 255
    r, g, b = int(colour[0:2], 16), int(colour[2:4], 16), int(colour[4:6], 16)
    return 0.299 * r + 0.587 * g + 0.114 * b


def deck_tokens(block):
    """The deck block in the shape builder/build_pptx.py loads as a style JSON.

    Colours come back as palette names rather than literal hex wherever a name
    exists, because a slide kind overrides the name and cannot override a hex.
    """
    palette_in = block.get("palette") or {}
    palette = {}
    for mine, theirs in TOKEN_PALETTE:
        if palette_in.get(mine):
            palette[theirs] = palette_in[mine]
    series = [palette_in.get(k) for k in ("primary", "secondary", "accent", "muted")]
    series = [c for c in series if c]
    if series:
        palette["series"] = series
    card_in = block.get("card") or {}
    if card_in.get("borderColor"):
        palette["cardBorder"] = card_in["borderColor"]
    if card_in.get("fill"):
        palette["cardFill"] = card_in["fill"]

    def name_for(colour):
        """A palette name for a colour, so a dark slide can override it."""
        if not colour:
            return None
        for mine, theirs in TOKEN_PALETTE:
            if palette_in.get(mine) == colour:
                return theirs
        return colour

    tokens = {"name": "Extracted"}
    if palette:
        tokens["palette"] = palette
    fonts = {}
    for mine, theirs in TOKEN_FONTS:
        if (block.get("fonts") or {}).get(mine):
            fonts[theirs] = block["fonts"][mine]
    if fonts:
        tokens["fonts"] = fonts
    card = {}
    for mine, theirs in TOKEN_CARD:
        if card_in.get(mine) is not None:
            card[theirs] = name_for(card_in[mine]) if mine in ("borderColor", "fill") \
                else card_in[mine]
    if card:
        tokens["card"] = card

    scale = {}
    for role, spec in (block.get("typeScale") or {}).items():
        out = {}
        if spec.get("size") is not None:
            out["size"] = spec["size"]
        if spec.get("weight"):
            # Both, because the builder writes a bold flag and the weight says which
            # face the family should be asked for.
            out["weight"] = spec["weight"]
        if spec.get("bold") is not None:
            out["bold"] = bool(spec["bold"])
        if spec.get("case") == "upper":
            out["caps"] = True
        colour = spec.get("color")
        target = TOKEN_ROLE.get(role, role)
        if colour:
            # A divider's title is white because the slide is dark, not because the
            # role is white: keep the role semantic and let the background say so.
            out["color"] = "ink" if target == "dividerTitle" else (name_for(colour) or colour)
        if out:
            scale[target] = out
    if scale:
        tokens["typeScale"] = scale

    furniture = {}
    for role, box in (block.get("furniture") or {}).items():
        furniture[role] = {k: box[k] for k in ("x", "y", "w", "h") if k in box}
    if furniture:
        tokens["furniture"] = furniture

    backgrounds = {}
    display = ((block.get("typeScale") or {}).get("display") or {}).get("color")
    for kind, fill in (block.get("backgrounds") or {}).items():
        entry = {"fill": fill}
        if luminance(fill) < 140:           # a dark slide flips the text colours
            entry["ink"] = display if (kind == "divider" and display) else "FFFFFF"
            entry["muted"] = entry["ink"]
        backgrounds[kind] = entry
    if backgrounds:
        tokens["backgrounds"] = backgrounds
    if block.get("slideSize"):
        tokens["slideSize"] = dict(block["slideSize"])
    for group in TOKEN_SAME:
        value = block.get(group)
        if isinstance(value, dict) and value:
            entry = dict(value)
            if group == "table":
                for key in TOKEN_TABLE_COLOURS:
                    if entry.get(key):
                        entry[key] = name_for(entry[key]) or entry[key]
            tokens[group] = entry
        elif value not in (None, "", {}):
            tokens[group] = value
    if block.get("diagrams"):
        tokens["diagrams"] = block["diagrams"]
    return tokens


# ----------------------------------------------------------------------- merge


def leaves(node, prefix=""):
    """Every scalar or list leaf of a format block, as dotted paths."""
    out = {}
    if isinstance(node, dict):
        for key, value in node.items():
            out.update(leaves(value, "%s.%s" % (prefix, key) if prefix else key))
    else:
        out[prefix] = node
    return out


def conf_for(path, table):
    """A field's confidence, falling back to the nearest ancestor that has one."""
    parts = path.split(".")
    while parts:
        key = ".".join(parts)
        if key in table:
            return table[key]
        parts.pop()
    return 0.0


def put(block, path, value):
    node = block
    parts = path.split(".")
    for part in parts[:-1]:
        node = node.setdefault(part, {})
    node[parts[-1]] = value


def merge(base, obs, threshold=0.8):
    """Merge Claude's screenshot observations over the file-derived block.

    The file wins a field it read with confidence >= threshold. Everything else
    the observations touch is theirs, because a screenshot sees what a file does
    not: what the style looks like rather than what one document happens to hold.
    """
    out = {"format": {}, "confidence": dict(base.get("confidence") or {}),
           "evidence": dict(base.get("evidence") or {}),
           "unknowns": list(base.get("unknowns") or []),
           "rules": list(base.get("rules") or []),
           "samples": {"palette": list((base.get("samples") or {}).get("palette") or []),
                       "fonts": list((base.get("samples") or {}).get("fonts") or [])}}
    base_conf = base.get("confidence") or {}
    obs_conf = obs.get("confidence") or {}
    obs_ev = obs.get("evidence") or {}

    for section in ("workbook", "deck"):
        base_block = (base.get("format") or {}).get(section) or {}
        obs_block = (obs.get("format") or {}).get(section) or {}
        if not base_block and not obs_block:
            continue
        merged = {}
        base_leaves = leaves(base_block)
        obs_leaves = leaves(obs_block)
        for path, value in base_leaves.items():
            put(merged, path, value)
        for path, value in obs_leaves.items():
            full = "%s.%s" % (section, path)
            if path in base_leaves and conf_for(full, base_conf) >= threshold:
                continue
            put(merged, path, value)
            out["confidence"][full] = round(float(conf_for(path, obs_conf)
                                                  or conf_for(full, obs_conf) or 0.75), 2)
            why = obs_ev.get(path) or obs_ev.get(full) or "read from a screenshot"
            out["evidence"][full] = why if why.startswith("observed") else "observed: %s" % why
        if merged:
            out["format"][section] = merged

    for text in obs.get("unknowns") or []:
        if text not in out["unknowns"]:
            out["unknowns"].append(text)
    for text in obs.get("rules") or []:
        if text not in out["rules"]:
            out["rules"].append(text)
    for key in ("palette", "fonts"):
        for item in ((obs.get("samples") or {}).get(key) or []):
            if item not in out["samples"][key]:
                out["samples"][key].append(item)
    # An unknown the observations answered is no longer unknown.
    answered = set()
    for section in ("workbook", "deck"):
        for path in leaves((obs.get("format") or {}).get(section) or {}):
            answered.add(path.split(".")[-1].lower())
    out["unknowns"] = [u for u in out["unknowns"]
                       if not any(word in u.lower() for word in answered if len(word) > 4)]
    for section in ("workbook", "deck"):
        keys = WORKBOOK_KEYS if section == "workbook" else DECK_KEYS
        if out["format"].get(section):
            total = sum(out["confidence"].get("%s.%s" % (section, k), 0.0) for k in keys)
            out["confidence"][section] = round(total / len(keys), 2)
    return out


# -------------------------------------------------------------------- SKILL.md


def yaml_value(value):
    """A scalar as YAML that build_xlsx.front_matter also reads unchanged."""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    if isinstance(value, (int, float)):
        return str(value)
    text = str(value)
    needs = (text == "" or text != text.strip()
             or re.match(r"^[0-9.+\-]", text) is not None
             or any(ch in text for ch in "#:\"'{}[],&*?|<>=!%@`"))
    if not needs:
        return text
    if '"' in text and "'" not in text:
        return "'%s'" % text
    if "'" in text and '"' not in text:
        return '"%s"' % text
    return '"%s"' % text.replace("\\", "\\\\").replace('"', '\\"')


def flat_deck(block, prefix=""):
    """The deck block as dotted `key: value` lines; see docs/skill-format-block.md."""
    out = []
    for key in sorted(block):
        value = block[key]
        path = "%s.%s" % (prefix, key) if prefix else key
        if isinstance(value, dict):
            out.extend(flat_deck(value, path))
        elif isinstance(value, (list, tuple)):
            out.append((path, ",".join(str(v) for v in value)))
        else:
            out.append((path, value))
    return out


def dedupe(items):
    """The list with blanks dropped, each entry stripped, first occurrence kept."""
    out = []
    for item in items:
        text = str(item).strip()
        if text and text not in out:
            out.append(text)
    return out


def skill_md(result, name, description, max_lines=25):
    """A complete SKILL.md: front matter with the tokens, then a "## Rules" section.

    The section carries the tokens restated as prose and then `result["rules"]`:
    what the extractor read that no token holds (a frozen header row, an
    inputs-only tab, a page number that flips colour) together with any rules
    merged in from Claude's observations. The whole prose runs to at most
    `max_lines` lines, and `result["rules"]` keep their places in it, so the
    server never has to append them itself.
    """
    fmt = result.get("format") or {}
    workbook = fmt.get("workbook") or {}
    deck = fmt.get("deck") or {}
    lines = ["---", "name: %s" % name, "description: %s" % yaml_value(description)]
    if workbook:
        lines.append("format:")
        for key in sorted(workbook):
            lines.append("  %s: %s" % (key, yaml_value(workbook[key])))
    if deck:
        lines.append("deck:")
        for path, value in flat_deck(deck):
            lines.append("  %s: %s" % (path, yaml_value(value)))
    lines.append("---")
    title = name.replace("-", " ").replace("_", " ").strip()
    prose = ["", "# %s" % (title[:1].upper() + title[1:]), ""]

    tokens = []
    if workbook:
        tokens.extend(workbook_rules(workbook))
    if deck:
        tokens.extend(deck_rules(deck))
    found = dedupe(result.get("rules") or [])
    unknowns = result.get("unknowns") or []
    asks = ["Still to confirm: %s" % unknowns[0]] if unknowns else []

    # The prose is capped, so the findings and the open question are counted first:
    # a restated token is recoverable from the front matter, a finding is not.
    room = max_lines - len(prose) - 3                 # the heading, its blank line, the closing blank
    keep = max(0, room - len(found) - len(asks))
    bullets = ["- %s" % rule for rule in dedupe(tokens[:keep] + found + asks)[:room]]
    if bullets:
        prose.extend(["## Rules", ""] + bullets)
    prose.append("")
    return "\n".join(lines + prose)


# ------------------------------------------------------------------------ main


def extract(paths):
    draft = Draft()

    class Auto(dict):
        """A dict of Tally that makes one on first touch."""

        def __missing__(self, key):
            self[key] = Tally()
            return self[key]

    tallies = Auto()
    sheets, decks = [], []
    for path in paths:
        low = path.lower()
        if not os.path.exists(path):
            draft.unknown("%s does not exist: where is the file, or should it be left out?" % path)
        elif low.endswith(XLSX_EXT):
            sheets.append(path)
        elif low.endswith(PPTX_EXT):
            decks.append(path)
        elif low.endswith(IMAGE_EXT):
            draft.unknown("%s is a screenshot: what palette, fonts, type sizes and card style "
                          "does it show? Read it and send your reading back as observations."
                          % os.path.basename(path))
        else:
            draft.unknown("%s is not an .xlsx or .pptx: is there an example file in one of those "
                          "formats, or a screenshot of this one?" % os.path.basename(path))

    read_any = False
    for path in sheets:
        if load_workbook is None or import_scan is None:
            draft.unknown("The workbooks could not be read because openpyxl is missing: can you "
                          "install it (pip3 install openpyxl) and try again?")
            break
        try:
            scan_xlsx(path, draft, tallies)
            read_any = True
        except Exception as exc:            # noqa: BLE001 - one bad file never sinks the run
            draft.unknown("%s could not be read (%s: %s): can you re-export it, or send a "
                          "screenshot of it instead?"
                          % (os.path.basename(path), type(exc).__name__, exc))
    if read_any:
        finish_xlsx(draft, tallies)

    read_any = False
    for path in decks:
        try:
            scan_pptx(path, draft, tallies)
            read_any = True
        except ImportError:
            draft.unknown("The decks could not be read because python-pptx is missing: can you "
                          "install it (pip3 install python-pptx) and try again?")
            break
        except Exception as exc:            # noqa: BLE001
            draft.unknown("%s could not be read (%s: %s): can you re-export it, or send a "
                          "screenshot of it instead?"
                          % (os.path.basename(path), type(exc).__name__, exc))
    if read_any:
        finish_pptx(draft, tallies)
    return draft.result()


def parse_args(argv):
    out, merge_path, skill, paths, tokens = None, None, None, [], None
    i = 1
    while i < len(argv):
        arg = argv[i]
        if arg == "--out" and i + 1 < len(argv):
            out = argv[i + 1]
            i += 2
        elif arg == "--tokens" and i + 1 < len(argv):
            tokens = argv[i + 1]
            i += 2
        elif arg == "--merge" and i + 1 < len(argv):
            merge_path = argv[i + 1]
            i += 2
        elif arg == "--skill-md" and i + 2 < len(argv):
            skill = (argv[i + 1], argv[i + 2])
            i += 3
        elif arg.startswith("--"):
            raise ValueError("unknown option %s" % arg)
        else:
            paths.append(arg)
            i += 1
    return out, merge_path, skill, paths, tokens


def main(argv):
    try:
        out, merge_path, skill, paths, tokens = parse_args(argv)
    except ValueError as exc:
        sys.stdout.write(json.dumps({"ok": False, "errors": [str(exc)]}) + "\n")
        return 1
    if not paths and not merge_path:
        sys.stdout.write(json.dumps({"ok": False, "errors": [
            "usage: extract_format.py --out result.json [--merge obs.json] "
            "[--tokens deck.json] [--skill-md name description] path [path ...]"]}) + "\n")
        return 1
    try:
        result = extract(paths)
        if merge_path:
            with open(merge_path, "r") as handle:
                result = merge(result, json.load(handle))
    except Exception as exc:                # noqa: BLE001 - one JSON line, never a traceback
        sys.stdout.write(json.dumps({"ok": False, "errors": ["%s: %s" % (type(exc).__name__, exc)]}) + "\n")
        return 1
    if out:
        with open(out, "w") as handle:
            json.dump(result, handle, indent=2, sort_keys=True)
    if tokens:
        with open(tokens, "w") as handle:
            json.dump(deck_tokens((result.get("format") or {}).get("deck") or {}),
                      handle, indent=2, sort_keys=True)
    if skill:
        sys.stdout.write(skill_md(result, skill[0], skill[1]))
        return 0
    if out:
        sys.stdout.write(json.dumps({
            "ok": True, "out": out,
            "sections": sorted(result["format"].keys()),
            "confidence": {k: v for k, v in result["confidence"].items() if "." not in k},
            "unknowns": result["unknowns"]}) + "\n")
    else:
        sys.stdout.write(json.dumps(result, indent=2, sort_keys=True) + "\n")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))

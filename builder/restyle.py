#!/usr/bin/env python3
"""restyle.py: rewrite the formats of an existing .xlsx, block by block, in a chosen style.

Usage: python3 builder/restyle.py <spec.json>

spec: {
  "inputPath": "/abs/in.xlsx", "outputPath": "/abs/out.xlsx",
  "formatting": "plain" | "house" | "skill" | "template",
  "skill": {"path": "/abs/SKILL.md"},                                      (skill mode)
  "templatePath": "...",                                                   (template mode)
  "scan": <import_scan.py out.json for inputPath>,
  "sheets": ["names to restyle"]                                          (default: all scanned sheets)
}

For every block the scanner detected, fonts, colours, borders, alignment, number formats
(only where the cell has none), column widths, gridlines and the title merge are rewritten
with the same stylers build_xlsx.py uses, following the block geometry the scan recorded
(title row, header rows, label column, body, total row, notes). Values, formulas, charts,
merges outside the title row, data validation and conditional formats are left alone, and
cells outside every detected block keep their formatting. outputPath is never inputPath.

Links follow the style, as in a build (W14): a hyperlink cell and a bare cross-sheet
pass-through take the style's `link_color`, which is the ink colour in house mode because
the house palette is blue, black and grey; the hyperlink keeps its underline and the
pass-through takes one only when the style says `link_underline`.

Block geometry is the scan's, so no fidelity move (W11) belongs here: a restyle rewrites
formats, it never moves a block, widens one or adds a row.

Prints one JSON line: {"ok": true, "path", "sheets": [{"name", "blocks", "cells"}], "warnings"}
or {"ok": false, "errors": [...]}. Exit 0 on success, 1 on failure. Python 3.9, openpyxl only.
"""

import copy
import json
import os
import sys
import zipfile

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import build_xlsx as bx  # noqa: E402
import import_scan  # noqa: E402

GRID_TYPES = ("table", "calc")
PAIR_TYPES = ("inputs", "output")


class Span(object):
    """The row, col, width triple the stylers' block_title reads."""

    def __init__(self, row, col, width):
        self.row, self.col, self.width = row, col, width


def is_formula(cell):
    v = cell.value
    return cell.data_type == "f" or (isinstance(v, str) and v.startswith("="))


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def load_scan(spec, warnings):
    """The scan from the spec, or a fresh one when it is missing or predates block geometry."""
    scan = spec.get("scan")
    ok = isinstance(scan, dict) and isinstance(scan.get("tabs"), list)
    if ok:
        for tab in scan["tabs"]:
            for b in tab.get("blocks") or []:
                if b.get("type") != "chart" and not b.get("geometry"):
                    ok = False
    if not ok:
        warnings.append("scan missing or without block geometry, re-scanned %s" % spec["inputPath"])
        scan = import_scan.scan(spec["inputPath"])
    return scan


def set_width(ws, col, width):
    """Set one column's width, splitting any grouped <col min max> span that covers it first."""
    dims = ws.column_dimensions
    for key, d in list(dims.items()):
        if d.min and d.max and d.min < d.max and d.min <= col <= d.max:
            for c in range(d.min, d.max + 1):
                letter = get_column_letter(c)
                nd = copy.copy(d)
                nd.index, nd.min, nd.max = letter, c, c
                dims[letter] = nd
    dims[get_column_letter(col)].width = width


class SheetRestyler(object):
    def __init__(self, ws, styler, ctx, warnings):
        self.ws = ws
        self.st = styler
        self.ctx = ctx
        self.house = bool(getattr(styler, "colour_inputs", styler.mode == "house"))  # colour inputs, warn on unmerged titles
        self.warnings = warnings
        self.touched = set()
        self.blocks = 0
        self._merges = None  # (row, col) -> MergedCellRange, rebuilt after the title row changes

    # ------------------------------------------------------------ helpers

    def where(self, block):
        return "%s!%s" % (self.ws.title, block.get("range") or block.get("id"))

    def merge_of(self, row, col):
        if self._merges is None:
            self._merges = {}
            for m in self.ws.merged_cells.ranges:
                for r in range(m.min_row, m.max_row + 1):
                    for c in range(m.min_col, m.max_col + 1):
                        self._merges[(r, c)] = m
        return self._merges.get((row, col))

    def reset(self, cell):
        """Back to the style's base: base font, no border, default alignment, no fill."""
        cell.font = self.st.font()
        cell.border = Border()
        cell.alignment = Alignment()
        cell.fill = PatternFill()
        self.touched.add((cell.row, cell.column))
        return cell

    def cell(self, row, col):
        return self.reset(self.ws.cell(row, col))

    def merged_cells(self, row, col):
        """The cell at (row, col) plus every cell of the merge it anchors, reset."""
        m = self.merge_of(row, col)
        if m is None:
            return [self.cell(row, col)]
        return [self.cell(r, c) for r in range(m.min_row, m.max_row + 1)
                for c in range(m.min_col, m.max_col + 1)]

    def valued(self, row, c1, c2):
        """Columns in c1..c2 whose cell on this row holds a value (merged copies hold none)."""
        return [c for c in range(c1, c2 + 1) if self.ws.cell(row, c).value not in (None, "")]

    # ------------------------------------------------------------- title

    def title(self, block, geo):
        row = geo.get("titleRow")
        if not row:
            return
        r1, c1, r2, c2 = geo["box"]
        title_cols = self.valued(row, c1, c2)
        if not title_cols:
            return
        tcol = title_cols[0]
        m = self.merge_of(row, tcol)
        end = max(c2, m.max_col if m else 0)
        # Inputs row notes live inside the block (the title spans them); table row notes sit in
        # the column after the block, outside the title merge, as the builder writes them.
        if block.get("type") not in GRID_TYPES:
            for r, c in geo.get("noteCells") or []:
                if r == row:
                    continue
                nm = self.merge_of(r, c)
                end = max(end, nm.max_col if nm else c)
        for r, c in geo.get("noteCells") or []:  # a note on the title row caps the span
            if r == row and c > tcol:
                end = min(end, c - 1)
        others = [c for c in self.valued(row, c1, end) if c != tcol]
        if others:
            end = min(end, min(others) - 1)
        start = c1
        width = max(1, end - start + 1)
        merges = [x for x in self.ws.merged_cells.ranges
                  if x.min_row <= row <= x.max_row and x.min_col <= end and x.max_col >= start]
        inside = [x for x in merges if x.min_row == x.max_row == row and start <= x.min_col and x.max_col <= end]
        can_merge = tcol == start and width > 1 and len(inside) == len(merges)
        if not can_merge and self.house and width > 1:
            self.warnings.append("%s: title row kept unmerged (title not at the block's first column, or a merge "
                                 "crosses the block edge)" % self.where(block))
        if len(inside) == len(merges):
            for x in inside:
                self.ws.unmerge_cells(str(x))
        for c in range(start, start + width):
            self.cell(row, c)
        self.st.block_title(self.ws, Span(row, start, width), merge=can_merge)
        self.touched.update((row, c) for c in range(start, start + width))
        self._merges = None

    # ----------------------------------------------------------- notes

    def note(self, row, col):
        for cell in self.merged_cells(row, col):
            self.st.note(cell)
        text = self.ws.cell(row, col).value
        if text not in (None, ""):
            self.ctx.fit_note(col, text)

    # ------------------------------------------------------------ grids

    def grid(self, block, geo):
        btype = block["type"]
        cols = geo.get("contentCols") or []
        if not cols:
            return
        c1, c2 = min(cols), max(cols)
        label_col = geo.get("labelCol")
        first_col = label_col if label_col is not None else c1
        self.ctx.want(first_col, self.st.label_width)
        for c in range(c1, c2 + 1):
            if c != first_col:
                self.ctx.want(c, self.st.value_width)
        for r in geo.get("headerRows") or []:
            for c in range(c1, c2 + 1):
                cell = self.cell(r, c)
                self.st.header(cell, first=(c == first_col))
                if cell.value not in (None, ""):
                    self.ctx.fit(c, str(cell.value), bold=True)
        bold_rows = set(geo.get("subtotalRows") or [])
        rows = list(geo.get("bodyRows") or [])
        if geo.get("totalRow"):
            rows.append(geo["totalRow"])
            bold_rows.add(geo["totalRow"])
        for r in rows:
            for c in range(c1, c2 + 1):
                cell = self.cell(r, c)
                value, formula = cell.value, is_formula(cell)
                # The builder's own rule: colour by cell, not by block (W2). A cell stored as
                # text is a year or a code, which the builder reads off the spec's format and
                # here off the number format the file carries.
                role = bx.cell_role(value, formula, link=cell.hyperlink,
                                    label=(c == first_col and btype in GRID_TYPES),
                                    fmt="text" if cell.number_format == "@" else None)
                if isinstance(value, str) and not formula and len(value) > 55:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                if btype in PAIR_TYPES and c == first_col:
                    self.st.pair_label(cell)
                elif btype in PAIR_TYPES:
                    self.st.pair_value(cell, role=role)
                else:
                    self.st.table_cell(cell, first=(c == first_col), bold=(r in bold_rows), role=role)
                if is_number(value) and (cell.number_format or "General") == "General":
                    cell.number_format = bx.number_format(value, None, False)
                if value not in (None, "") and not (btype in PAIR_TYPES and c == first_col):
                    self.ctx.fit(c, value, is_formula=formula)  # pair labels get label_width, as in the builder

    # ------------------------------------------------------------- text

    def text(self, block, geo):
        cols = geo.get("contentCols") or []
        for r in geo.get("bodyRows") or []:
            for c in cols:
                anchor = self.ws.cell(r, c)
                cells = self.merged_cells(r, c)
                for cell in cells:
                    self.st.body_text(cell)
                if len(cells) > 1 and self.ws.row_dimensions[r].height is None and isinstance(anchor.value, str):
                    self.ws.row_dimensions[r].height = bx.wrapped_height(anchor.value, len(cells) * 16)

    # ------------------------------------------------------------ block

    def block(self, block):
        geo = block.get("geometry")
        if block.get("type") == "chart" or not geo:
            return
        before = len(self.touched)
        self.title(block, geo)
        for r, c in geo.get("subtitleCells") or []:
            self.note(r, c)
        if block["type"] == "text":
            self.text(block, geo)
        elif block["type"] in GRID_TYPES + PAIR_TYPES:
            self.grid(block, geo)
        else:
            self.warnings.append("%s: unknown block type %s skipped" % (self.where(block), block["type"]))
            return
        for r, c in geo.get("noteCells") or []:
            self.note(r, c)
        if len(self.touched) > before:
            self.blocks += 1

    def finish(self):
        self.ws.sheet_view.showGridLines = self.st.gridlines
        for col, width in sorted(self.ctx.final_widths().items()):
            set_width(self.ws, col, width)


# ----------------------------------------------------------------- restyle


def media_warning(path, warnings):
    try:
        with zipfile.ZipFile(path) as z:
            media = [n for n in z.namelist() if n.startswith("xl/media/")]
    except Exception:  # noqa: BLE001 - only a warning source
        return
    if not media:
        return
    try:
        import PIL  # noqa: F401
    except ImportError:
        warnings.append("workbook holds %d media file(s); openpyxl drops images without Pillow installed" % len(media))


def restyle(spec):
    errors, warnings = [], []
    src, out = spec.get("inputPath"), spec.get("outputPath")
    if not src or not os.path.isfile(src):
        errors.append("inputPath missing or not a file: %r" % src)
    if not out:
        errors.append("spec has no outputPath")
    elif src and os.path.abspath(out) == os.path.abspath(src):
        errors.append("outputPath must differ from inputPath (the original is never overwritten)")
    if errors:
        return {"ok": False, "errors": errors}

    styler = bx.make_styler(spec, warnings)
    try:
        scan = load_scan(spec, warnings)
    except Exception as exc:  # noqa: BLE001 - scan failure is a spec problem
        return {"ok": False, "errors": ["could not scan %s: %s" % (src, exc)]}
    wanted = spec.get("sheets")
    if wanted is not None and not isinstance(wanted, list):
        return {"ok": False, "errors": ["sheets must be a list of sheet names"]}
    media_warning(src, warnings)

    try:
        wb = load_workbook(src, keep_vba=src.lower().endswith(".xlsm"))
    except Exception as exc:  # noqa: BLE001 - report as an error
        return {"ok": False, "errors": ["could not open %s: %s" % (src, exc)]}
    ctx = bx.Ctx(styler, {}, errors, warnings)
    by_name = {t.get("name"): t for t in scan.get("tabs") or []}
    names = wanted if wanted is not None else [t.get("name") for t in scan.get("tabs") or []]
    report = []
    charts = 0
    has_formulas = False
    for name in names:
        tab = by_name.get(name)
        if tab is None:
            warnings.append("sheet '%s' is not in the scan, skipped" % name)
            continue
        if name not in wb.sheetnames:
            warnings.append("sheet '%s' is not in the workbook, skipped" % name)
            continue
        if wanted is None and bx.norm(name).replace(" ", "") == "readme":
            warnings.append("sheet '%s' left as is: the cover page has its own layout, list it in sheets to force" % name)
            continue
        ws = wb[name]
        charts += len(getattr(ws, "_charts", []) or [])
        ctx.reset_counts()
        sr = SheetRestyler(ws, styler, ctx, warnings)
        for block in tab.get("blocks") or []:
            sr.block(block)
        sr.finish()
        if not has_formulas:
            has_formulas = any(is_formula(c) for row in ws.iter_rows() for c in row if c.value is not None)
        report.append({"name": name, "blocks": sr.blocks, "cells": len(sr.touched)})

    if getattr(styler, "wants_cover", False) and not any(bx.norm(n).replace(" ", "") == "readme" for n in wb.sheetnames):
        warnings.append("%s: no Read Me tab added; a restyle changes formats only, rebuild to get the cover page" % styler.mode)
    if has_formulas:
        warnings.append("formula results are recalculated by Excel on open; openpyxl does not keep cached values")

    out_dir = os.path.dirname(os.path.abspath(out))
    if out_dir and not os.path.isdir(out_dir):
        os.makedirs(out_dir)
    try:
        wb.save(out)
    except Exception as exc:  # noqa: BLE001 - report as an error
        return {"ok": False, "errors": ["could not write %s: %s" % (out, exc)]}
    if charts:
        try:
            kept = sum(len(getattr(ws, "_charts", []) or []) for ws in load_workbook(out).worksheets)
        except Exception:  # noqa: BLE001 - verification only
            kept = None
        if kept is not None and kept < charts:
            warnings.append("%d of %d charts did not survive the openpyxl round trip" % (charts - kept, charts))
    bx.log("restyled %s -> %s (%s, %d sheets)" % (src, out, styler.mode, len(report)))
    return {"ok": True, "path": out, "sheets": report, "warnings": warnings}


def main(argv):
    if len(argv) != 2:
        sys.stdout.write(json.dumps({"ok": False, "errors": ["usage: restyle.py <spec.json>"]}) + "\n")
        return 1
    try:
        with open(argv[1], "r") as handle:
            spec = json.load(handle)
    except Exception as exc:  # noqa: BLE001 - report any read problem as an error
        sys.stdout.write(json.dumps({"ok": False, "errors": ["could not read spec: %s" % exc]}) + "\n")
        return 1
    try:
        result = restyle(spec)
    except Exception as exc:  # noqa: BLE001 - one JSON line, whatever happens
        result = {"ok": False, "errors": ["%s: %s" % (type(exc).__name__, exc)]}
    sys.stdout.write(json.dumps(result, default=str) + "\n")
    return 0 if result.get("ok") else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv))

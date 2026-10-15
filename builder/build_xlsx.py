#!/usr/bin/env python3
"""build_xlsx.py: build an .xlsx workbook from a Workbook Planner spec JSON.

Usage: python3 builder/build_xlsx.py <spec.json>

Writes the workbook, prints exactly one JSON object on stdout, exits 0 on
success or 1 with {"ok": false, "errors": [...]} on failure. All logging goes
to stderr. Python 3.9 compatible, openpyxl only.
"""

import json
import math
import os
import re
import sys
from copy import copy
from datetime import date
from decimal import ROUND_HALF_UP, Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.comments import Comment
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.properties import PageSetupProperties

# --------------------------------------------------------------- constants

NAVY = "1F3864"
GREY = "595959"
BLUE = "0000FF"
BLACK = "000000"
LINK_BLUE = "0563C1"      # plain mode keeps Office's hyperlink blue

# Three-part formats (W10): positive; negative in parentheses; zero as a dash.
FORMATS = {
    "currency": '$#,##0;($#,##0);"-"',
    "currency2": '$#,##0.00;($#,##0.00);"-"',
    "percent": '0.0%;(0.0%);"-"',
    "multiple": '0.0"x";(0.0"x");"-"',
    "count": '#,##0;(#,##0);"-"',
    "number": '#,##0.00;(#,##0.00);"-"',
    "text": "@",
    "year": "@",
}

SNAPSHOT_CELL_CAP = 5000

READ_ME = "Read Me"
BAD_SHEET_CHARS = "[]:*?/\\"
TEMPLATE_RE = re.compile(r"\{\{(.*?)\}\}")
RC_RE = re.compile(r"^(\d+)\s*,\s*(\d+)$")
DASH_RE = re.compile(r"\s*[\u2013\u2014]\s*")

THIN = Side(style="thin", color=BLACK)
DASHED = Side(style="dashed", color=BLACK)
MED_NAVY = Side(style="medium", color=NAVY)
THIN_NAVY = Side(style="thin", color=NAVY)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOX_DASH_R = Border(left=THIN, right=DASHED, top=THIN, bottom=THIN)

DEFAULT_RULES = [
    "Blue cells are inputs. Change these.",
    "Black cells are formulas. Do not overwrite them.",
    "Grey italic text is a note or a source, never a number the model uses.",
    "Every number traces back to an input or to a row in Sources.",
]

# The house format as a configuration. A formatting skill's `format:` front-matter block
# overrides any of these; HouseStyler() with no overrides is the house format itself.
HOUSE_DEFAULTS = {
    "font": "Arial",
    "size": 10,
    "title_size": 16,            # cover page title; block titles use `size`
    "input_color": BLUE,
    "formula_color": BLACK,
    # W14: the house skill's palette is blue, black and grey, so a link takes the ink colour
    # and the underline carries it. "" means formula_color; a style may name its own colour.
    "link_color": "",            # a hyperlink cell and a bare cross-sheet pass-through
    "link_underline": True,      # underline a bare cross-sheet pass-through too
    "note_color": GREY,
    "note_italic": True,
    "title_rule_color": NAVY,    # block title rule, cover title and cover section rules
    "title_rule": "medium",      # medium | thin | none: bottom border of every block title row
    "label_divider": "dashed",   # dashed | thin | none: right border of the label column
    "body_borders": "thin",      # thin | none: box around every body cell
    "gridlines": False,
    "fills": False,              # true fills column header cells with header_fill
    "header_fill": "F2F2F2",
    "cover_tab": True,           # write the Read Me cover into a blockless "Read Me" sheet
    "currency": FORMATS["currency"],
    "percent": FORMATS["percent"],
    "multiple": FORMATS["multiple"],
    "count": FORMATS["count"],
    "label_width": 30,           # minimum; a label column always fits its longest label (W1)
    "label_cap": 60,             # cap on a label or text column
    "value_width": 12,
    "value_cap": 20,             # cap on a numeric column: widths come from the values (W5)
    "note_width": 40,
    "dash_replace": True,        # em and en dashes become a comma
    "max_note_chars": 100,       # warn on cell text longer than this outside Read Me and text blocks
    # Tab colours by sheet role (W16). "" means the style leaves that role alone.
    "tab_cover": "808080",
    "tab_inputs": "0070C0",
    "tab_calc": "",
    "tab_output": "00B050",
}

SKILL_COLOUR_KEYS = ("input_color", "formula_color", "link_color", "note_color", "title_rule_color",
                     "header_fill", "tab_cover", "tab_inputs", "tab_calc", "tab_output")
SKILL_INT_KEYS = ("size", "title_size", "label_width", "label_cap", "value_width", "value_cap",
                  "note_width", "max_note_chars")
SKILL_BOOL_KEYS = ("note_italic", "gridlines", "fills", "cover_tab", "dash_replace",
                   "link_underline")
# Colour keys whose value may be "none": the tab colours leave that role alone, link_color
# falls back to the ink colour (W14).
SKILL_COLOUR_NONE = ("link_color",)
SKILL_CHOICE_KEYS = {
    "title_rule": ("medium", "thin", "none"),
    "label_divider": ("dashed", "thin", "none"),
    "body_borders": ("thin", "none"),
}
SKILL_TEXT_KEYS = ("font", "currency", "percent", "multiple", "count")
HEX_RE = re.compile(r"^#?([0-9A-Fa-f]{6})$")
FRONT_RE = re.compile(r"\A\ufeff?---[ \t]*\r?\n(.*?)\r?\n---[ \t]*(?:\r?\n|\Z)", re.S)

# Rows of a template's "Style" sheet: the label in column A names the role, column B is the example.
STYLE_ROLES = ("title", "header", "label", "input", "formula", "total", "note", "text")


def log(msg):
    sys.stderr.write(str(msg) + "\n")


def a1(row, col):
    return "%s%d" % (get_column_letter(col), row)


def qname(name):
    return "'" + str(name).replace("'", "''") + "'"


def norm(text):
    return str(text).strip().lower()


def as_pos(value, default):
    try:
        n = int(value)
    except (TypeError, ValueError):
        return default
    return n if n > 0 else default


# ---------------------------------------------------------------- geometry

ROLES = ("input", "formula", "link", "note")


def cell_roles(block):
    """A block's `cellRoles` map: {"r,c": "input"} -> {(r, c): "input"}, zero-based in the data area.

    For a table or calc block, r is the data row (0 = the first row under the header) and c
    the column; for inputs and output, r is the pair row and c is 0 for the label, 1 for the
    value. Unparseable keys and unknown roles are dropped here and warned about by the caller.
    """
    raw = block.get("cellRoles")
    out = {}
    if not isinstance(raw, dict):
        return out
    for key, role in raw.items():
        rc = RC_RE.match(str(key).strip())
        if rc and str(role).strip().lower() in ROLES:
            out[(int(rc.group(1)), int(rc.group(2)))] = str(role).strip().lower()
    return out


def bad_cell_roles(block):
    """The `cellRoles` entries cell_roles() had to drop, for a warning naming them."""
    raw = block.get("cellRoles")
    if not isinstance(raw, dict):
        return ["cellRoles must be an object"] if raw not in (None, {}) else []
    bad = []
    for key, role in raw.items():
        if not RC_RE.match(str(key).strip()):
            bad.append("key %r is not \"row,col\"" % key)
        elif str(role).strip().lower() not in ROLES:
            bad.append("role %r is not one of %s" % (role, ", ".join(ROLES)))
    return bad


class Geom:
    """Fixed block geometry, computed for every block before anything is written.

    Row 0 is always the title. inputs/output put label/value pairs from row 1;
    table/calc put the header on row 1 and data from row 2; text puts the
    paragraph on row 1; chart is anchored at (row, col).
    """

    def __init__(self, block, sheet_name):
        self.block = block
        self.sheet = sheet_name
        self.id = str(block.get("id") or "")
        self.type = str(block.get("type") or "")
        self.title = block.get("title") or ""
        self.note = block.get("note")
        self.row = as_pos(block.get("row"), 1)
        self.col = as_pos(block.get("col"), 1)
        self.planned_cols = as_pos(block.get("cols"), 0)
        raw_rows = block.get("rows")
        self.pair_rows = raw_rows if isinstance(raw_rows, list) else []
        self.planned_rows = as_pos(raw_rows, 0) if not isinstance(raw_rows, list) else 0
        self.value_col = self.col + 1
        self.label_rows = {}
        self.header_row = None
        self.first_data_row = None
        self.last_data_row = None
        self.total_row = None
        self.data_cols = 0
        self.row_notes = []          # table/calc: one note or None per data row
        self.total_note = None       # table/calc: note beside the total row
        self.has_row_notes = False
        self.snapped = None          # the move the W13 width pass took on this block
        self.fitted = None           # the move measuring the block took off its planned cols
        self.moves = []              # W11: the fidelity moves taken on this block, in order
        self.total_values = None     # the total row's cells, approved or generated (W11)
        self.wants_total = bool(block.get("total")) if self.type in ("table", "calc") else False
        self.generated_total = False
        self.cell_roles = cell_roles(block)
        self._measure()

    # -------------------------------------------------------- W11 fidelity

    def rect(self):
        """The rectangle the block occupies, its adjacent note column included: the snapshot's."""
        return {"row": self.ext[0], "col": self.ext[1],
                "rows": self.ext[2] - self.ext[0] + 1, "cols": self.ext[3] - self.ext[1] + 1}

    def record(self, reason, before):
        """One fidelity move on this block, for the snapshot and the result JSON.

        A move that carries on where the last one of the same reason stopped extends it
        instead of stacking a second record, and returns None so the caller lists it once.
        """
        last = self.moves[-1] if self.moves else None
        if last is not None and last["reason"] == reason and last["to"] == before:
            last["to"] = self.rect()
            return None
        move = {"blockId": self.id, "from": before, "to": self.rect(), "reason": reason}
        self.moves.append(move)
        return move

    def grow(self, cols):
        """Take `cols` more columns (negative to give them back). Width only, never position."""
        self.width += cols
        self._place_rows()
        self._extents()

    def shift(self, rows):
        """Move the whole block down `rows` rows, header, data and total with it."""
        self.row += rows
        self._place_rows()
        self._extents()

    def set_total(self, values):
        """Add the total row a `total: true` block asks for, growing the block by one row."""
        self.total_values = list(values)
        self.generated_total = True
        self.height += 1
        note = self.block.get("totalNote")
        if note not in (None, ""):
            self.total_note = str(note)
            self.has_row_notes = True
        self._place_rows()
        self._extents()

    def drop_total(self):
        """Give the generated total row back: no room for it after the sheet was spaced."""
        self.total_values = None
        self.generated_total = False
        self.height -= 1
        self.total_note = None
        self.has_row_notes = any(self.row_notes)
        self._place_rows()
        self._extents()

    def _place_rows(self):
        """The rows inside the block, from its current top row: header, data, total, labels."""
        if self.type in ("table", "calc"):
            self.header_row = self.row + 1
            self.first_data_row = self.row + 2
            self.last_data_row = self.row + 1 + self.n_data_rows
            self.total_row = self.last_data_row + 1 if self.total_values is not None else None
        elif self.type in ("inputs", "output"):
            self.value_col = self.col + 1
            self.label_rows = {}
            for i, rowdef in enumerate(self.pair_rows):
                if isinstance(rowdef, dict):
                    self.label_rows[norm(rowdef.get("label", ""))] = self.row + 1 + i
        self.note_col_after = self.col + self.width
        if self.type in ("inputs", "output"):
            self.note_col = self.note_col_after

    # ------------------------------------------------- width and note column

    def snap_width(self, width):
        """W13: widen the block to the sheet's block width. Records the move for the snapshot."""
        if width <= self.width or self.type == "chart":
            return None
        move = {"blockId": self.id, "from": self.width, "to": width, "reason": "snap"}
        self.snapped = move
        self.width = width
        self._place_rows()
        self._extents()
        return move

    def set_note_col(self, col):
        """W4: one note column per row band, immediately right of the band's widest block."""
        self.note_col_after = col
        self.note_col = col
        self._extents()

    def _extents(self):
        self.body_ext = (
            self.row, self.col, self.row + self.height - 1, self.col + self.width - 1,
        )
        has_note = bool(self.note or self.has_row_notes)
        # The note column only grows the recorded extent when it is adjacent to the block;
        # a shared band column further right is recorded as noteCol instead.
        adjacent = has_note and self.note_col_after == self.col + self.width
        self.ext = (
            self.body_ext[0], self.body_ext[1], self.body_ext[2],
            self.body_ext[3] + (1 if adjacent else 0),
        )
        self.note_ext_col = self.note_col_after if has_note and not adjacent else None

    def role_at(self, r, c):
        """A spec `cellRoles` override for the zero-based cell (r, c) of the block's data area."""
        return self.cell_roles.get((r, c))

    def _measure(self):
        if self.type in ("inputs", "output"):
            self._measure_pairs()
        elif self.type in ("table", "calc"):
            self._measure_table()
        elif self.type == "text":
            self.width = max(self.planned_cols, 1)
            self.height = 2
        elif self.type == "chart":
            self.width = max(self.planned_cols, 1)
            self.height = max(self.planned_rows, 10)
        else:
            self.width = max(self.planned_cols, 1)
            self.height = 1
        # Block notes and table row notes share the column right after the block; the band
        # pass moves them to one column per row band (W4).
        self._place_rows()
        self._extents()

    def _measure_pairs(self):
        self.has_row_notes = any(
            isinstance(r, dict) and r.get("note") for r in self.pair_rows
        )
        # A pair block is two columns wide, whatever the plan reserved: the label and its
        # value. Planned columns beyond that are note room, and the note has to sit outside
        # the title rule to read as annotation (W4).
        width = 2
        if self.planned_cols and self.planned_cols != width:
            self.fitted = {"blockId": self.id, "from": self.planned_cols, "to": width,
                           "reason": "fit"}
        self.note_col = self.col + width      # the band pass may move it further right
        self.width = width
        self.height = 1 + len(self.pair_rows)

    def _measure_table(self):
        columns = self.block.get("columns") or []
        data = self.block.get("data") or []
        total = self.block.get("totalRow")
        ncols = len(columns)
        for r in data:
            if isinstance(r, list):
                ncols = max(ncols, len(r))
        if isinstance(total, list):
            ncols = max(ncols, len(total))
        self.data_cols = max(ncols, 1)
        self.n_data_rows = len(data)
        self.total_values = list(total) if isinstance(total, list) else None
        # As wide as the table actually is: a planned column the data does not fill would
        # leave an empty column between the block and its note (W4).
        self.width = self.data_cols
        if self.planned_cols and self.planned_cols != self.width:
            self.fitted = {"blockId": self.id, "from": self.planned_cols, "to": self.width,
                           "reason": "fit"}
        self.height = 2 + len(data) + (1 if self.total_values is not None else 0)
        raw_notes = self.block.get("rowNotes")
        if isinstance(raw_notes, list):
            self.row_notes = [
                (str(n) if n not in (None, "") else None) for n in raw_notes[:len(data)]
            ]
        total_note = self.block.get("totalNote")
        if self.total_values is not None and total_note not in (None, ""):
            self.total_note = str(total_note)
        self.has_row_notes = any(self.row_notes) or self.total_note is not None

    def data_col_index(self, c):
        return self.col + c


# --------------------------------------------------- sheet plan (W13, W4)

SNAP_TOLERANCE = 2      # blocks whose left edges sit within this many columns share a width


def snap_widths(geoms, fidelity=None):
    """W13: one block width per sheet. Blocks sharing a column band grow to the widest.

    A column band is a run of blocks whose left edges sit within SNAP_TOLERANCE columns of
    each other. The widest right edge in the band wins, a block never grows by more than
    SNAP_TOLERANCE columns (a narrow pair block beside a wide table keeps its own width, as
    the house format wants), and a grow that would collide with another block is dropped.
    Returns the moves taken, for the snapshot's `adjusted` list. Under strict the snap is
    off, because matching a neighbour is a preference; measuring a block is not (W11).
    """
    movable = [g for g in geoms if g.type in ("table", "calc", "text")]
    moves = [g.fitted for g in geoms if g.fitted]
    if fidelity == "strict":
        return moves
    band = []
    for g in sorted(movable, key=lambda x: (x.col, x.row)):
        if band and g.col - band[0].col > SNAP_TOLERANCE:
            moves.extend(_snap_band(band, geoms))
            band = []
        band.append(g)
    moves.extend(_snap_band(band, geoms))
    return moves


def _snap_band(band, geoms):
    """Grow the band's blocks to its widest table. A pair block keeps its two columns: it has
    nothing to put in a third, and growing it would leave a gap before its note."""
    grids = [g for g in band if g.type in ("table", "calc")]
    if len(band) < 2 or not grids:
        return []
    target = max(g.col + g.width - 1 for g in grids)
    moves = []
    for g in band:
        want = target - g.col + 1
        if want <= g.width or want - g.width > SNAP_TOLERANCE:
            continue
        before = g.width
        move = g.snap_width(want)
        if any(o is not g and overlaps(g.body_ext, o.body_ext) for o in geoms):
            g.width = before                      # would collide: leave the block alone
            g.snapped = None
            g._extents()
            continue
        if move:
            moves.append(move)
    return moves


# ------------------------------------------------------------- W11 fidelity

FIDELITIES = ("strict", "balanced")
DEFAULT_FIDELITY = "balanced"
GROW_COLS = 2             # columns a balanced widen may take
SHIFT_ROWS = 2            # rows a balanced shift may move a block down
TEXT_COLS = 16            # characters a text block fits in one column of its merge
TEXT_LINES = 3            # lines a paragraph may run to before the block wants more columns
ADDITIVE = ("currency", "currency2", "count", "number")   # formats a total row may sum
# A column whose header reads as a rate, a price or an average is not a column you add up,
# whatever its format: the total row leaves it blank the way a hand-built model does.
NOT_ADDITIVE_RE = re.compile(
    r"\b(per|each|average|avg|median|mean|rate|share|margin|price|ratio|multiple|"
    r"growth|cagr|yield|uplift|discount)\b", re.I)


def read_fidelity(spec, warnings):
    """W11: `fidelity` says how far the builder may move the approved geometry. Default balanced."""
    raw = spec.get("fidelity")
    if raw in (None, ""):
        return DEFAULT_FIDELITY
    value = str(raw).strip().lower()
    if value in FIDELITIES:
        return value
    warnings.append("fidelity '%s' is not one of %s, used %s"
                    % (raw, ", ".join(FIDELITIES), DEFAULT_FIDELITY))
    return DEFAULT_FIDELITY


def apply_fidelity(geoms, fidelity, warnings):
    """W11: the moves `balanced` may take on one sheet's approved geometry, in order.

    Widen a block whose content its columns cannot hold, add the total row a table asks for,
    then open a blank row between blocks the plan left touching. strict takes none of them:
    the approved cells are the cells. Every move comes back as {blockId, from, to, reason}.
    """
    if fidelity != "balanced":
        return []
    # A plan that already overlaps is a plan error, not a spacing problem: those blocks are
    # left exactly where they are so check_overlaps names them.
    clashing = set(id(g) for g in geoms
                   if any(o is not g and overlaps(g.body_ext, o.body_ext) for o in geoms))
    moves = widen_blocks(geoms)
    moves.extend(add_totals(geoms, warnings))
    moves.extend(space_blocks(geoms, clashing))
    moves.extend(drop_crowded_totals(geoms, warnings))
    return moves


def wanted_cols(g, geoms):
    """Columns a block's own content needs beyond the ones it has, capped at GROW_COLS.

    Only `text` and `chart` keep the width the plan gave them (a table is as wide as its
    data), so they are the only two whose columns can be short of their content: a paragraph
    that runs past TEXT_LINES lines, and a chart narrower than the table it plots.
    """
    if g.type == "text":
        chars = len(str(g.block.get("text") or ""))
        if chars <= g.width * TEXT_COLS * TEXT_LINES:
            return 0
        want = int(math.ceil(chars / float(TEXT_COLS * TEXT_LINES)))
        return max(0, min(GROW_COLS, want - g.width))
    if g.type == "chart":
        src = (g.block.get("chart") or {}).get("sourceBlock")
        source = next((o for o in geoms if o.id == src), None)
        if source is None:
            return 0
        return max(0, min(GROW_COLS, source.width - g.width))
    return 0


def widen_blocks(geoms):
    """Balanced (a): take up to GROW_COLS more columns when the content needs them and the
    columns to the right are free. A widen that would touch another block is not taken."""
    moves = []
    for g in sorted(geoms, key=lambda x: (x.row, x.col)):
        want = wanted_cols(g, geoms)
        if not want:
            continue
        before = g.rect()
        g.grow(want)
        if any(o is not g and overlaps(g.body_ext, o.body_ext) for o in geoms):
            g.grow(-want)                       # would collide: leave the block alone
            continue
        move = g.record("widen", before)
        if move:
            moves.append(move)
    return moves


def sums_to_a_number(raw):
    """True when a spec cell holds something a total row could sum: a number or a formula."""
    if isinstance(raw, bool) or raw is None:
        return False
    if isinstance(raw, (int, float)):
        return True
    if isinstance(raw, dict):
        return "f" in raw or isinstance(raw.get("v"), (int, float))
    return False


def total_row_for(g):
    """The total row a `total: true` table gets: "Total", then =SUM() down every additive column.

    A column is additive when its format is a currency, a count or a plain number: a share, a
    multiple, a year and a label are not summed, they are left blank.
    """
    formats = g.block.get("colFormats") or []
    columns = g.block.get("columns") or []
    data = g.block.get("data") or []
    values = [None] * g.data_cols
    values[0] = "Total"
    summed = 0
    for j in range(1, g.data_cols):
        fmt = formats[j] if j < len(formats) else None
        if fmt is not None and fmt not in ADDITIVE:
            continue
        head = str(columns[j]) if j < len(columns) else ""
        if NOT_ADDITIVE_RE.search(head):
            continue
        cells = [r[j] for r in data if isinstance(r, list) and j < len(r)]
        numeric = [c for c in cells if sums_to_a_number(c)]
        if not numeric or (fmt is None and len(numeric) < len(cells)):
            continue        # nothing to add up, or no format and a cell that is not a number
        values[j] = {"f": "=SUM({{%s!col:%d}})" % (g.id, j)}
        summed += 1
    return values if summed else None


def add_totals(geoms, warnings):
    """Balanced (b): the total row a table asks for with `total: true` and its approved rows
    do not carry. Taken on trust here; drop_crowded_totals() gives back any that will not fit."""
    moves = []
    for g in geoms:
        if not g.wants_total or g.total_values is not None or not g.n_data_rows:
            continue
        values = total_row_for(g)
        if values is None:
            warnings.append("%s!%s: total: true but no column of the table adds up, no total "
                            "row added" % (g.sheet, g.id))
            continue
        before = g.rect()
        g.set_total(values)
        move = g.record("total", before)
        if move:
            moves.append(move)
    return moves


def drop_crowded_totals(geoms, warnings):
    """A generated total row that still lands on another block after the sheet was spaced:
    give the row back and say so, rather than failing the build on an overlap nobody planned."""
    moves = []
    for g in geoms:
        if not g.generated_total:
            continue
        if not any(o is not g and overlaps(g.body_ext, o.body_ext) for o in geoms):
            continue
        before = g.rect()
        g.drop_total()
        warnings.append("%s!%s: no free row under the block, the total row it asks for was "
                        "not added" % (g.sheet, g.id))
        move = g.record("no-total", before)
        if move:
            moves.append(move)
    return moves


def stacked(a, b):
    """True when two blocks share a column, so one sits above the other on the page."""
    return a.col <= b.col + b.width - 1 and b.col <= a.col + a.width - 1


def space_blocks(geoms, clashing=()):
    """Balanced (d): shift a block down by up to SHIFT_ROWS rows to open one blank row under
    the block above it.

    Top down, so a shift cascades down the sheet, and a shift onto a block that has not moved
    yet is skipped; the pass runs again while anything moved, which is what lets a whole
    stack open up one block at a time. No block ever moves more than SHIFT_ROWS in total.
    """
    ordered = sorted([g for g in geoms if g.type in WRITERS and id(g) not in clashing],
                     key=lambda x: (x.row, x.col))
    moves = []
    room = dict((id(g), SHIFT_ROWS) for g in ordered)
    for _pass in range(SHIFT_ROWS + 1):
        moved = 0
        ordered.sort(key=lambda x: (x.row, x.col))
        for i, g in enumerate(ordered):
            above = [p for p in ordered[:i] if stacked(p, g)]
            if not above or room[id(g)] <= 0:
                continue
            need = min(max(p.body_ext[2] for p in above) + 2 - g.row, room[id(g)])
            if need <= 0 or blocked(g, need, [o for o in geoms if o is not g]):
                continue
            before = g.rect()
            g.shift(need)
            room[id(g)] -= need
            moved += 1
            move = g.record("shift", before)
            if move:
                moves.append(move)
        if not moved:
            break
    return moves


def blocked(g, rows, others):
    """True when shifting `g` down that many rows would land on one of `others`."""
    moved = (g.body_ext[0] + rows, g.body_ext[1], g.body_ext[2] + rows, g.body_ext[3])
    return any(overlaps(moved, o.body_ext) for o in others)


def note_bands(geoms, fidelity=DEFAULT_FIDELITY):
    """W4: one note column per row band, immediately right of the band's widest block.

    Blocks whose rows overlap (transitively) share one note column, so a sheet never grows
    two note columns and no empty column is left between a block and its note. Under strict
    a note stays beside its own block unless another block is written in that column, which
    is the one case where the band column is not a preference. Returns the moves taken.
    """
    ordered = sorted([g for g in geoms if g.type in WRITERS], key=lambda x: (x.row, x.col))
    moves = []
    band = []
    for g in ordered:
        if band and g.row > max(x.body_ext[2] for x in band):
            moves.extend(_set_band_note(band, geoms, fidelity))
            band = []
        band.append(g)
    moves.extend(_set_band_note(band, geoms, fidelity))
    return moves


def _set_band_note(band, geoms, fidelity):
    if not band:
        return []
    col = max(g.col + g.width for g in band)
    moves = []
    for g in band:
        own = g.col + g.width
        target = col
        if fidelity == "strict" and own != col:
            clear = not any(o is not g and o.col <= own <= o.col + o.width - 1
                            and overlaps((g.row, own, g.row + g.height - 1, own), o.body_ext)
                            for o in geoms)
            target = own if clear else col
        before = g.rect()
        g.set_note_col(target)
        if target != own and (g.note or g.has_row_notes):
            move = g.record("note", before)
            if move:
                moves.append(move)
    return moves


def overlaps(a, b):
    return a[0] <= b[2] and b[0] <= a[2] and a[1] <= b[3] and b[1] <= a[3]


def block_extent(g):
    """One block's snapshot entry: the cells it ended up occupying, notes included."""
    out = {
        "id": g.id,
        "row": g.ext[0],
        "col": g.ext[1],
        "rows": g.ext[2] - g.ext[0] + 1,
        "cols": g.ext[3] - g.ext[1] + 1,
        "type": g.type,
    }
    if g.note_ext_col:                  # a note column the band shares, right of this block
        out["noteCol"] = g.note_ext_col
    if g.moves:                         # W11: what the fidelity pass moved, first from to last to
        out["adjusted"] = {"from": g.moves[0]["from"], "to": g.moves[-1]["to"],
                           "reason": "+".join(m["reason"] for m in g.moves)}
    elif g.snapped or g.fitted:
        out["adjusted"] = (g.snapped or g.fitted)["reason"]
    return out


# ------------------------------------------------------------ template resolver


class Resolver:
    """Resolves {{...}} reference templates inside formula strings."""

    def __init__(self, registry, errors):
        self.registry = registry
        self.errors = errors

    def resolve(self, formula, current_sheet, where):
        def swap(match):
            expr = match.group(1).strip()
            try:
                return self._one(expr, current_sheet)
            except ValueError as exc:
                self.errors.append(
                    "%s: unresolved template {{%s}}: %s" % (where, expr, exc)
                )
                return match.group(0)

        return TEMPLATE_RE.sub(swap, formula)

    def _block(self, block_id):
        g = self.registry.get(block_id.strip())
        if g is None:
            raise ValueError("no block with id '%s'" % block_id.strip())
        return g

    def _one(self, expr, current_sheet):
        if expr.lower().startswith("sheet:"):
            return qname(self._block(expr.split(":", 1)[1]).sheet)
        if "!" not in expr:
            raise ValueError("expected blockId!reference")
        block_id, rest = expr.split("!", 1)
        g = self._block(block_id)
        rest = rest.strip()
        prefix = "" if g.sheet == current_sheet else qname(g.sheet) + "!"

        rc = RC_RE.match(rest)
        if rc:
            return prefix + self._cell_rc(g, int(rc.group(1)), int(rc.group(2)))
        low = rest.lower()
        if low.startswith("col:"):
            return prefix + self._column(g, rest.split(":", 1)[1])
        if low.startswith("total:"):
            return prefix + self._total(g, rest.split(":", 1)[1])
        return prefix + self._label(g, rest)

    @staticmethod
    def _require_table(g):
        if g.type not in ("table", "calc"):
            raise ValueError("block '%s' is a %s block, not table or calc" % (g.id, g.type))

    def _cell_rc(self, g, r, c):
        self._require_table(g)
        if r < 0 or r >= g.n_data_rows:
            raise ValueError("data row %d out of range (%d rows)" % (r, g.n_data_rows))
        if c < 0 or c >= g.data_cols:
            raise ValueError("data col %d out of range (%d cols)" % (c, g.data_cols))
        return a1(g.first_data_row + r, g.data_col_index(c))

    def _column(self, g, raw):
        self._require_table(g)
        try:
            c = int(str(raw).strip())
        except ValueError:
            raise ValueError("column index must be an integer")
        if c < 0 or c >= g.data_cols:
            raise ValueError("data col %d out of range (%d cols)" % (c, g.data_cols))
        if g.n_data_rows < 1:
            raise ValueError("block '%s' has no data rows" % g.id)
        col = g.data_col_index(c)
        return "%s:%s" % (a1(g.first_data_row, col), a1(g.last_data_row, col))

    def _total(self, g, raw):
        self._require_table(g)
        if not g.total_row:
            raise ValueError("block '%s' has no totalRow" % g.id)
        try:
            c = int(str(raw).strip())
        except ValueError:
            raise ValueError("column index must be an integer")
        if c < 0 or c >= g.data_cols:
            raise ValueError("data col %d out of range (%d cols)" % (c, g.data_cols))
        return a1(g.total_row, g.data_col_index(c))

    def _label(self, g, text):
        if g.type not in ("inputs", "output"):
            raise ValueError("block '%s' is a %s block, not inputs or output" % (g.id, g.type))
        row = g.label_rows.get(norm(text))
        if row is None:
            raise ValueError("no row labelled '%s' in block '%s'" % (text.strip(), g.id))
        return a1(row, g.value_col)


# ------------------------------------------------------------------ stylers


class PlainStyler:
    """plain mode: Calibri 11, bold titles, bordered tables, no fills.

    Every styler exposes the same knobs the writers and Ctx read: fonts, the
    four colours, widths, whether literal inputs are coloured, whether dashes
    are replaced, the cell text cap and whether a blockless "Read Me" sheet
    gets the cover page.
    """

    mode = "plain"
    gridlines = True
    label_width = 22
    label_cap = 60             # cap on a label or text column's content fit
    value_width = 12
    value_cap = 20             # cap on a numeric column: its width comes from the values
    note_width = 40            # cap on a note column's content fit, or the width itself when fixed
    note_width_fixed = False
    colour_inputs = False      # literal inputs get input_color
    dash_replace = False       # em and en dashes become a comma
    max_text_chars = None      # cap on cell text outside Read Me and text blocks; None is no cap
    cover_tab = True           # a blockless "Read Me" sheet gets the cover page
    wants_cover = False        # the style expects a cover tab (restyle warns when one is missing)
    title_size = 16
    input_color = BLUE
    formula_color = BLACK
    link_color = LINK_BLUE     # plain mode takes Office's hyperlink blue
    link_underline = True      # underline a bare cross-sheet pass-through, not only a hyperlink
    note_color = GREY
    note_italic = True
    title_rule_color = NAVY
    title_rule_side = MED_NAVY
    formats = FORMATS
    tab_colors = {}            # sheet role -> tab colour; an absent role keeps Excel's default

    def __init__(self, base=("Calibri", 11), title=None):
        self.base_name, self.base_size = base
        self.title_name, self.title_size_block = title if title else base

    @property
    def title_font_size(self):
        """Block title size in plain and template modes (the title font read from a template)."""
        return self.title_size_block

    def font(self, bold=False, italic=False, color=BLACK, size=None, name=None, underline=None):
        return Font(
            name=name or self.base_name,
            size=size or self.base_size,
            bold=bold,
            italic=italic,
            color=color,
            underline=underline,
        )

    # ------------------------------------------------------------ cell roles

    def role_color(self, role):
        """The font colour for a body cell's role (W2): colour by cell, not by block."""
        if role == "link":
            return self.link_color
        if role == "note":
            return self.note_color
        if role == "input" and self.colour_inputs:
            return self.input_color
        return self.formula_color

    def link_underline_for(self, hyperlink):
        """W14: a hyperlink is always underlined; a bare cross-sheet pass-through only when
        the style says `link_underline`, which is the one mark it has when the link colour
        is the ink colour."""
        return "single" if (hyperlink or self.link_underline) else None

    def role_font(self, role, bold=False, hyperlink=False):
        return self.font(bold=bold, color=self.role_color(role),
                         italic=(role == "note" and self.note_italic),
                         underline=self.link_underline_for(hyperlink) if role == "link" else None)

    @staticmethod
    def _role(role, blue):
        """`role` wins; `blue=True` is the old spelling of role "input"."""
        return role or ("input" if blue else None)

    def tab_color(self, role):
        return self.tab_colors.get(role) or None

    def block_title(self, ws, g, merge=True):
        """Style the title row of a block. `g` needs row, col and width; `merge` is used by house."""
        cell = ws.cell(g.row, g.col)
        cell.font = self.font(bold=True, name=self.title_name, size=self.title_size_block)

    def header(self, cell, first=False):
        cell.font = self.font(bold=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    def table_cell(self, cell, first=False, bold=False, blue=False, role=None):
        cell.font = self.role_font(self._role(role, blue), bold=bold,
                                   hyperlink=cell.hyperlink is not None)
        cell.border = BOX

    def pair_label(self, cell):
        cell.font = self.font()

    def pair_value(self, cell, blue=False, role=None):
        cell.font = self.role_font(self._role(role, blue), hyperlink=cell.hyperlink is not None)

    def note(self, cell):
        cell.font = self.font(italic=self.note_italic, color=self.note_color)
        cell.alignment = Alignment(wrap_text=False, vertical="top")

    def body_text(self, cell):
        cell.font = self.font()
        cell.alignment = Alignment(wrap_text=True, vertical="top")


class HouseStyler(PlainStyler):
    """house mode: Arial 10, navy title rules, dashed label divider, no fills.

    Built from HOUSE_DEFAULTS; a config dict (a skill's `format:` block, already
    validated by skill_config) overrides any key.
    """

    mode = "house"
    colour_inputs = True

    def __init__(self, config=None):
        cfg = dict(HOUSE_DEFAULTS)
        cfg.update(config or {})
        self.cfg = cfg
        PlainStyler.__init__(self, base=(cfg["font"], cfg["size"]))
        self.gridlines = bool(cfg["gridlines"])
        self.label_width = cfg["label_width"]
        self.label_cap = cfg["label_cap"]
        self.value_width = cfg["value_width"]
        self.value_cap = max(cfg["value_cap"], cfg["value_width"])
        self.note_width = cfg["note_width"]
        self.title_size = cfg["title_size"]
        self.input_color = cfg["input_color"]
        self.formula_color = cfg["formula_color"]
        # W14: no link colour of its own means the ink colour, which is the house palette.
        self.link_color = cfg["link_color"] or self.formula_color
        self.link_underline = bool(cfg["link_underline"])
        self.tab_colors = {
            "cover": cfg["tab_cover"], "inputs": cfg["tab_inputs"],
            "calc": cfg["tab_calc"], "output": cfg["tab_output"],
        }
        self.note_color = cfg["note_color"]
        self.note_italic = bool(cfg["note_italic"])
        self.title_rule_color = cfg["title_rule_color"]
        self.cover_tab = bool(cfg["cover_tab"])
        self.wants_cover = self.cover_tab
        self.dash_replace = bool(cfg["dash_replace"])
        self.max_text_chars = cfg["max_note_chars"] or None
        self.formats = dict(FORMATS)
        for key in ("currency", "percent", "multiple", "count"):
            self.formats[key] = cfg[key]

        rule = cfg["title_rule"]
        if rule == "none":
            self.title_border = Border()
            self.title_rule_side = Side(style="thin", color=self.title_rule_color)
        else:
            self.title_rule_side = Side(style=rule, color=self.title_rule_color)
            self.title_border = Border(top=THIN, bottom=self.title_rule_side)
        body = THIN if cfg["body_borders"] == "thin" else Side()
        divider = {"dashed": DASHED, "thin": THIN, "none": body}[cfg["label_divider"]]
        self.body_box = Border(left=body, right=body, top=body, bottom=body)
        self.label_box = Border(left=body, right=divider, top=body, bottom=body)
        self.header_fill = (
            PatternFill(fill_type="solid", fgColor=cfg["header_fill"]) if cfg["fills"] else None
        )

    def block_title(self, ws, g, merge=True):
        for c in range(g.col, g.col + g.width):
            cell = ws.cell(g.row, c)
            cell.font = self.font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.title_border
        if merge and g.width > 1:
            ws.merge_cells(
                start_row=g.row,
                start_column=g.col,
                end_row=g.row,
                end_column=g.col + g.width - 1,
            )

    def header(self, cell, first=False):
        cell.font = self.font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = self.label_box if first else self.body_box
        if self.header_fill is not None:
            cell.fill = copy(self.header_fill)

    def table_cell(self, cell, first=False, bold=False, blue=False, role=None):
        cell.font = self.role_font(self._role(role, blue), bold=bold,
                                   hyperlink=cell.hyperlink is not None)
        cell.border = self.label_box if first else self.body_box
        cell.alignment = Alignment(
            horizontal="left" if first else "center", vertical="center"
        )

    def pair_label(self, cell):
        cell.font = self.font()
        cell.border = self.label_box
        cell.alignment = Alignment(horizontal="left", vertical="center")

    def pair_value(self, cell, blue=False, role=None):
        cell.font = self.role_font(self._role(role, blue), hyperlink=cell.hyperlink is not None)
        cell.border = self.body_box
        cell.alignment = Alignment(horizontal="center", vertical="center")


class SkillStyler(HouseStyler):
    """skill mode: the house styler parameterised by a SKILL.md `format:` block."""

    mode = "skill"

    def __init__(self, config, name):
        HouseStyler.__init__(self, config)
        self.skill_name = name


# ------------------------------------------------------- skill front-matter


def split_kv(line):
    i = line.find(":")
    if i <= 0:
        return None, None
    return line[:i].strip(), line[i + 1:].strip()


def unquote(text):
    t = text.strip()
    if len(t) >= 2 and t[0] == t[-1] and t[0] in "\"'":
        return t[1:-1]
    return t


def front_matter(md):
    """The YAML-ish front-matter of a SKILL.md as a dict.

    Top-level `key: value` pairs map to strings. A top-level `key:` with nothing
    after the colon opens a nested block whose two-space-indented `key: value`
    lines map to a dict of strings. Only scalars are read; lists, deeper nesting
    and multi-line values are ignored. Values keep their raw text (quotes
    stripped) so each consumer coerces its own types.
    """
    out = {}
    m = FRONT_RE.match(md)
    if not m:
        return out
    current = None
    for raw in m.group(1).splitlines():
        line = raw.rstrip()
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        indented = line.startswith("  ") or line.startswith("\t")
        if indented:
            if current is not None:
                key, val = split_kv(line.strip())
                if key:
                    out[current][key] = unquote(val)
            continue
        key, val = split_kv(line)
        if key is None:
            current = None
            continue
        if val == "":
            out[key] = {}
            current = key
        else:
            out[key] = unquote(val)
            current = None
    return out


def skill_config(block, name, warnings):
    """Validate a `format:` block against HOUSE_DEFAULTS. Bad values warn and keep the default."""
    cfg = {}
    for raw_key, raw in block.items():
        key = raw_key.replace("colour", "color")
        text = str(raw).strip()
        where = "skill '%s' format.%s" % (name, raw_key)
        if key in SKILL_COLOUR_KEYS:
            hm = HEX_RE.match(text)
            if (key.startswith("tab_") or key in SKILL_COLOUR_NONE) \
                    and text.lower() in ("", "none", "off", "default"):
                cfg[key] = ""      # a tab role Excel colours itself, or the ink for a link
            elif hm:
                cfg[key] = hm.group(1).upper()
            else:
                warnings.append("%s: '%s' is not a 6 digit hex colour, kept the default" % (where, text))
        elif key in SKILL_INT_KEYS:
            try:
                n = float(text)
            except ValueError:
                n = None
            if n is None or n < 0:
                warnings.append("%s: '%s' is not a number, kept the default" % (where, text))
            else:
                cfg[key] = int(n) if n == int(n) else n
        elif key in SKILL_BOOL_KEYS:
            low = text.lower()
            if low in ("true", "yes", "on", "1"):
                cfg[key] = True
            elif low in ("false", "no", "off", "0"):
                cfg[key] = False
            else:
                warnings.append("%s: '%s' is not true or false, kept the default" % (where, text))
        elif key in SKILL_CHOICE_KEYS:
            low = text.lower()
            if low in SKILL_CHOICE_KEYS[key]:
                cfg[key] = low
            else:
                warnings.append("%s: '%s' is not one of %s, kept the default"
                                % (where, text, ", ".join(SKILL_CHOICE_KEYS[key])))
        elif key in SKILL_TEXT_KEYS:
            if text:
                cfg[key] = text
            else:
                warnings.append("%s: empty, kept the default" % where)
        else:
            warnings.append("%s: unknown key ignored" % where)
    return cfg


def skill_styler(path, warnings):
    """skill mode: HouseStyler parameterised by the `format:` block of a SKILL.md."""
    try:
        with open(path, "r", encoding="utf-8") as handle:
            md = handle.read()
    except Exception as exc:  # noqa: BLE001 - any read problem falls back
        warnings.append("skill file could not be read (%s), using house defaults" % exc)
        return HouseStyler()
    fm = front_matter(md)
    name = fm.get("name") if isinstance(fm.get("name"), str) else None
    if not name:
        name = os.path.basename(os.path.dirname(os.path.abspath(path))) or "skill"
    block = fm.get("format")
    if not isinstance(block, dict):
        warnings.append(
            "skill '%s' has no format: block in its front-matter, using house defaults" % name
        )
        return SkillStyler({}, name)
    return SkillStyler(skill_config(block, name, warnings), name)


# ------------------------------------------------------------ template mode


class TemplateStyler(PlainStyler):
    """template mode: plain layout in the template's fonts, plus any roles its Style sheet defines.

    `roles` maps a STYLE_ROLES name to the cloned font, fill, border, alignment
    and number format of the example cell. A role that is missing falls back to
    the plain behaviour in the template's base font.
    """

    mode = "template"

    def __init__(self, base, title=None, roles=None, widths=None):
        PlainStyler.__init__(self, base, title)
        self.roles = roles or {}
        self.colour_inputs = "input" in self.roles
        widths = widths or {}
        self.label_width = widths.get("label", self.label_width)
        self.value_width = widths.get("value", self.value_width)
        self.note_width = widths.get("note", self.note_width)
        self.note_width_fixed = "note" in widths
        if "note" in self.roles:
            fnt = self.roles["note"]["font"]
            self.note_italic = bool(fnt.italic)
            if fnt.color is not None and isinstance(fnt.color.rgb, str):
                self.note_color = fnt.color.rgb[-6:]
        if "input" in self.roles:
            fnt = self.roles["input"]["font"]
            if fnt.color is not None and isinstance(fnt.color.rgb, str):
                self.input_color = fnt.color.rgb[-6:]
        if "formula" in self.roles:
            fnt = self.roles["formula"]["font"]
            if fnt.color is not None and isinstance(fnt.color.rgb, str):
                self.formula_color = fnt.color.rgb[-6:]
        # W14: a template names no link colour, so a link takes the ink and the underline marks it.
        self.link_color = self.formula_color

    def apply(self, cell, *roles):
        """Clone the first role present onto the cell. Returns False when none is."""
        for role in roles:
            ex = self.roles.get(role)
            if ex is None:
                continue
            cell.font = copy(ex["font"])
            cell.fill = copy(ex["fill"])
            cell.border = copy(ex["border"])
            cell.alignment = copy(ex["alignment"])
            nf = ex["number_format"]
            if nf and nf != "General" and (cell.number_format or "General") == "General":
                cell.number_format = nf
            return True
        return False

    def block_title(self, ws, g, merge=True):
        ex = self.roles.get("title")
        if ex is None:
            return PlainStyler.block_title(self, ws, g, merge)
        for c in range(g.col, g.col + g.width):
            self.apply(ws.cell(g.row, c), "title")
        centered = ex["alignment"].horizontal in ("center", "centerContinuous")
        if merge and centered and g.width > 1:
            ws.merge_cells(start_row=g.row, start_column=g.col,
                           end_row=g.row, end_column=g.col + g.width - 1)

    def header(self, cell, first=False):
        if not self.apply(cell, "header"):
            PlainStyler.header(self, cell, first)

    def table_cell(self, cell, first=False, bold=False, blue=False, role=None):
        role = self._role(role, blue)
        if bold:
            roles = ("total",)
        elif role == "input":
            roles = ("input",)
        elif first and isinstance(cell.value, str) and not cell.value.startswith("="):
            roles = ("label", "formula")
        else:
            roles = ("formula",)
        if not self.apply(cell, *roles):
            PlainStyler.table_cell(self, cell, first, bold, role=role)
        elif role == "link":
            self.as_link(cell)

    def pair_label(self, cell):
        if not self.apply(cell, "label"):
            PlainStyler.pair_label(self, cell)

    def pair_value(self, cell, blue=False, role=None):
        role = self._role(role, blue)
        if not self.apply(cell, "input" if role == "input" else "formula"):
            PlainStyler.pair_value(self, cell, role=role)
        elif role == "link":
            self.as_link(cell)

    def as_link(self, cell):
        """Recolour a cloned template role as a hyperlink, keeping its other font settings."""
        fnt = copy(cell.font)
        fnt.color = Font(color=self.link_color).color
        fnt.underline = self.link_underline_for(cell.hyperlink is not None)
        cell.font = fnt

    def note(self, cell):
        if not self.apply(cell, "note"):
            PlainStyler.note(self, cell)

    def body_text(self, cell):
        if not self.apply(cell, "text"):
            return PlainStyler.body_text(self, cell)
        al = copy(cell.alignment)
        al.wrap_text = True
        if al.vertical is None:
            al.vertical = "top"
        cell.alignment = al


def explicit_width(ws, col):
    """The width set on a column, or None when the sheet leaves it at the default."""
    letter = get_column_letter(col)
    for key, d in ws.column_dimensions.items():
        covers = key == letter or (d.min and d.max and d.min <= col <= d.max)
        if covers and d.width:
            return d.width
    return None


def read_style_sheet(ws):
    """(roles, widths) from a template's Style sheet: column A names the role, column B is the example."""
    roles = {}
    for row in ws.iter_rows(min_col=1, max_col=2):
        label = row[0].value
        if not isinstance(label, str):
            continue
        key = label.strip().lower()
        if key in STYLE_ROLES and key not in roles and len(row) > 1:
            ex = row[1]
            roles[key] = {
                "font": copy(ex.font),
                "fill": copy(ex.fill),
                "border": copy(ex.border),
                "alignment": copy(ex.alignment),
                "number_format": ex.number_format,
            }
    widths = {}
    for col, name in ((2, "label"), (3, "value"), (4, "note")):
        w = explicit_width(ws, col)
        if w:
            widths[name] = w
    return roles, widths


def sheet_fonts(ws):
    """(base, title) fonts from a sheet: first non-empty cell, first bold cell. Either may be None."""
    base = None
    title = None
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or str(cell.value).strip() == "":
                continue
            fnt = cell.font
            name = (fnt.name if fnt else None) or "Calibri"
            size = (fnt.sz if fnt else None) or 11
            if base is None:
                base = (name, size)
            if title is None and fnt is not None and fnt.bold:
                title = (name, size)
        if base is not None and title is not None:
            break
    return base, title


def template_styler(path, warnings):
    """template mode: the template's base font and first bold font, plus its Style sheet if it has one."""
    try:
        wb = load_workbook(path)
        style_ws = None
        for ws in wb.worksheets:
            if ws.title.strip().lower() == "style":
                style_ws = ws
                break
        base = None
        title = None
        for ws in wb.worksheets:
            if ws is style_ws:
                continue
            base, title = sheet_fonts(ws)
            if base is not None:
                break
        roles, widths = read_style_sheet(style_ws) if style_ws is not None else ({}, {})
        wb.close()
        if base is None:
            for role in ("text", "label", "formula", "input"):
                if role in roles:
                    fnt = roles[role]["font"]
                    base = (fnt.name or "Calibri", fnt.sz or 11)
                    break
        if base is None:
            raise ValueError("template has no non-empty cells")
        if title is None and "title" in roles:
            fnt = roles["title"]["font"]
            title = (fnt.name or base[0], fnt.sz or base[1])
        return TemplateStyler(base, title or base, roles, widths)
    except Exception as exc:  # noqa: BLE001 - any template problem falls back
        warnings.append("template could not be read (%s), fell back to plain" % exc)
        return PlainStyler()


def make_styler(spec, warnings):
    mode = spec.get("formatting") or "plain"
    if mode == "house":
        return HouseStyler()
    if mode == "skill":
        skill = spec.get("skill")
        path = skill.get("path") if isinstance(skill, dict) else skill
        if not path:
            warnings.append("skill mode with no skill path, using house defaults")
            return HouseStyler()
        return skill_styler(path, warnings)
    if mode == "template":
        path = spec.get("templatePath")
        if not path:
            warnings.append("template mode with no templatePath, fell back to plain")
            return PlainStyler()
        return template_styler(path, warnings)
    return PlainStyler()


# ------------------------------------------------------------------- context


class Ctx:
    def __init__(self, styler, registry, errors, warnings):
        self.styler = styler
        self.registry = registry
        self.errors = errors
        self.warnings = warnings
        self.resolver = Resolver(registry, errors)
        self.names = {}            # W9: defined name -> "'Sheet'!$C$3"
        self.reset_counts()

    def reset_counts(self):
        self.cells = 0
        self.formulas = 0
        self.charts = 0
        self.widths = {}           # floors asked for by want()
        self.fits = {}             # width a column's values need
        self.label_fits = {}       # width a column's labels need (W1)
        self.head_chars = {}       # characters in a column's widest header (W5, wraps)
        self.text_cols = set()     # columns holding a text value, so not capped as numeric
        self.note_widths = {}
        self.caps = {}             # col -> the cap final_widths() applied to it
        self.wrap_cells = {}       # (row, col) -> chars of text that may need wrapping
        self.head_rows = {}        # row -> True for every header row on the sheet
        self.note_cells = []       # (row, col) of every note, for the merge pass
        self.note_spans = {}       # (row, col) -> last column of that note's merge
        self.deferred_charts = []  # (chart, geom) sized once the widths are known
        self.snap_cells = {}
        self.snap_widths = {}
        self.snap_charts = []

    # ------------------------------------------------------------ snapshot

    def rec(self, row, col, value, kind, nf=None, is_formula=False):
        """Record one written cell for the build snapshot."""
        if value is None or value == "":
            return
        entry = {"a": a1(row, col), "v": value, "k": kind}
        if nf:
            entry["fmt"] = nf
            if not is_formula:
                shown = display(value, nf)
                if shown is not None:
                    entry["d"] = shown
        self.snap_cells[(row, col)] = entry

    def mark_comment(self, row, col):
        """N9: the snapshot marks the cell that carries a block's comment thread."""
        entry = self.snap_cells.get((row, col))
        if entry is not None:
            entry["hasComment"] = True

    def rec_chart(self, block_id, anchor, kind, series):
        self.snap_charts.append(
            {"blockId": block_id, "anchor": anchor, "kind": kind, "series": series}
        )

    def set_width(self, ws, col, width):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = width
        self.snap_widths[letter] = width

    def snap_widths_by_col(self):
        return dict((column_index_from_string(k), v) for k, v in self.snap_widths.items())

    def snapshot_sheet(self, ws, geoms, adjusted=None):
        """The sheet's snapshot entry, cells in row-major order, capped at SNAPSHOT_CELL_CAP."""
        ordered = [self.snap_cells[key] for key in sorted(self.snap_cells)]
        truncated = len(ordered) > SNAPSHOT_CELL_CAP
        sheet = {
            "name": ws.title,
            "widths": self.snap_widths,
            "maxRow": ws.max_row,
            "maxCol": ws.max_column,
            "cells": ordered[:SNAPSHOT_CELL_CAP],
            "merges": sorted(str(r) for r in ws.merged_cells.ranges),
            "blocks": [block_extent(g) for g in geoms],
            "charts": self.snap_charts,
        }
        if adjusted:
            sheet["adjusted"] = adjusted
        if truncated:
            sheet["truncated"] = True
        return sheet

    def fit(self, col, value, fmt=None, is_formula=False, bold=False):
        """Ask for a column wide enough to show this value without clipping (W5).

        A number is measured at the width its own number format shows, not at its raw
        digits, and a header (bold=True) does not widen the column at all: it wraps.
        """
        if value is None or value == "":
            return
        if bold:
            n = len(str(value))
            if n > self.head_chars.get(col, 0):
                self.head_chars[col] = n
            return
        if is_formula:
            n = 12
        elif isinstance(value, (int, float)) and not isinstance(value, bool):
            nf = number_format(value, fmt, False, self.styler.formats)
            shown = display(value, nf)
            n = len(shown) if shown else len("{:,.2f}".format(float(value)))
            n += 2
        else:
            n = int(len(str(value)) * 1.05) + 2
            self.text_cols.add(col)
        if n > self.fits.get(col, 0):
            self.fits[col] = n

    def fit_label(self, col, text):
        """W1: a label column is never narrower than its longest label, capped at label_cap."""
        if text is None or text == "":
            return
        n = int(len(str(text)) * 1.05) + 2
        if n > self.label_fits.get(col, 0):
            self.label_fits[col] = n

    def fit_note(self, col, text):
        st = self.styler
        n = st.note_width if st.note_width_fixed else min(len(str(text)) + 2, max(st.note_width, 8))
        if n > self.note_widths.get(col, 0):
            self.note_widths[col] = n

    def want_wrap(self, row, col, text):
        """Remember a text cell that wraps if its column ends up narrower than the text."""
        n = len(str(text))
        if n > self.wrap_cells.get((row, col), 0):
            self.wrap_cells[(row, col)] = n

    def final_widths(self):
        """Column widths for the sheet (W1, W5): labels fit their labels, numbers their values.

        A label column floors at the style's label_width and caps at label_cap; a column of
        numbers caps at value_cap so header text cannot stretch it; a column holding text
        caps at label_cap; a note-only column keeps its note fit. A header too long for two
        lines at that width widens the column, still inside the cap.
        """
        st = self.styler
        cols = (set(self.widths) | set(self.fits) | set(self.label_fits)
                | set(self.head_chars) | set(self.note_widths))
        out = {}
        for col in cols:
            floor = self.widths.get(col, 0)
            label = self.label_fits.get(col, 0)
            vals = self.fits.get(col, 0)
            head = self.head_chars.get(col, 0)
            if label:
                w, cap = max(floor, label, vals), st.label_cap
            elif vals:
                cap = st.label_cap if col in self.text_cols else st.value_cap
                w = max(floor, vals)
            elif col in self.note_widths:
                w, cap = max(floor, self.note_widths[col]), st.label_cap
            else:
                w, cap = max(floor, head), st.label_cap
            w = min(w, cap)
            if head and head > 2 * max(w - 1, 1):       # wrap a header over two lines, not three
                w = min(cap, max(w, int(math.ceil(head / 2.0)) + 1))
            out[col] = max(8, w)
            self.caps[col] = cap
        return out

    def clean(self, text, where):
        if not self.styler.dash_replace or not isinstance(text, str):
            return text
        fixed = DASH_RE.sub(", ", text)
        if fixed != text:
            self.warnings.append("%s: dash replaced with a comma" % where)
        return fixed

    def check_length(self, text, where, exempt=False):
        """House rule: no cell text over the style's cap (100 chars) outside Read Me.

        Reported as a warning, not an error, so a long label never kills a
        build. Read Me and the body of a text block are exempt: both are
        merged wrapped paragraphs by design.
        """
        limit = self.styler.max_text_chars
        if not limit or exempt or not isinstance(text, str):
            return
        if len(text) > limit:
            self.warnings.append(
                "%s: cell text is %d chars, over the %d char %s limit"
                % (where, len(text), limit, self.styler.mode)
            )

    def want(self, col, width):
        if width > self.widths.get(col, 0):
            self.widths[col] = width

    # -------------------------------------------------------- named ranges

    def name_input(self, ws, row, col, role, label):
        """W9: a workbook-scoped defined name for one input cell, from its label.

        The label wins when it is free; a clash takes the sheet name as a prefix and then a
        counter, so every input cell ends up with one unique, readable name.
        """
        if role != "input":
            return None
        base = defined_name(label)
        if not base:
            return None
        target = "%s!$%s$%d" % (qname(ws.title), get_column_letter(col), row)
        for candidate in name_candidates(base, defined_name(ws.title)):
            if self.names.get(candidate) == target:
                return candidate
            if candidate not in self.names:
                self.names[candidate] = target
                return candidate
        return None


URL_RE = re.compile(r"^(?:https?://|mailto:)\S+$", re.I)
INTERNAL_LINK_RE = re.compile(r"^#?(?:'[^']+'|[^'!]+)!\$?[A-Za-z]{1,3}\$?[0-9]+$")
LINK_FORMULA_RE = re.compile(r"^=\s*(?:'[^']+'|[A-Za-z0-9_. ]+)!\$?[A-Za-z]{1,3}\$?[0-9]+\s*$")


class Spec(object):
    """One spec cell value, unpacked: the value, its format, and what kind of cell it is."""

    __slots__ = ("value", "fmt", "is_formula", "link", "role")

    def __init__(self, value, fmt, is_formula=False, link=None, role=None):
        self.value, self.fmt, self.is_formula = value, fmt, is_formula
        self.link, self.role = link, role


def unpack(raw, default_fmt):
    """Unpack a spec cell value: a scalar, {"f": "=.."}, {"v": ..} or {"link": ..} (W14)."""
    if isinstance(raw, dict):
        fmt = raw.get("format", default_fmt)
        role = str(raw["role"]).strip().lower() if str(raw.get("role") or "").strip() else None
        if "f" in raw:
            return Spec(raw["f"], fmt, True, None, role)
        link = raw.get("link")
        if link not in (None, ""):
            value = raw["v"] if raw.get("v") not in (None, "") else str(link).lstrip("#")
            return Spec(value, fmt, False, str(link), role or "link")
        if "v" in raw:
            return Spec(raw["v"], fmt, False, None, role)
        return Spec(None, fmt, False, None, role)
    if isinstance(raw, str) and URL_RE.match(raw.strip()):
        return Spec(raw.strip(), default_fmt, False, raw.strip(), "link")
    return Spec(raw, default_fmt, False, None, None)


NUMERIC_TEXT_RE = re.compile(r"^[-(]?\$?\d[\d,]*(?:\.\d+)?%?\)?$")


def numeric_text(value):
    """True for a literal string that reads as a number, such as a year written as text."""
    return isinstance(value, str) and bool(NUMERIC_TEXT_RE.match(value.strip()))


def cell_role(value, is_formula, link=None, picklist=False, label=False, fmt=None):
    """The role of one body cell (W2): colour by cell, never by block type.

    A literal number anywhere is an input; so is `TBD` and a picklist value a dropdown
    covers. A bare cross-sheet reference is a link, any other formula is a formula, and a
    label, a year and any other text are neither.
    """
    if link:
        return "link"
    if is_formula:
        return "link" if LINK_FORMULA_RE.match(str(value).strip()) else "formula"
    if isinstance(value, str) and value.strip() == "TBD":
        return "input"
    if label:
        return "label"
    if value is None or value == "" or isinstance(value, bool):
        return "text"
    if picklist:
        return "input"       # a dropdown value is adjustable whatever it looks like
    if fmt in ("year", "text"):
        return "text"        # stored as text: a year or a code, not a number the model reads
    if isinstance(value, (int, float)) or numeric_text(value):
        return "input"
    return "text"


NAME_BAD_RE = re.compile(r"[^0-9A-Za-z_.]+")
NAME_CELL_RE = re.compile(r"^(?:[A-Za-z]{1,3}[0-9]{1,7}|[RrCc][0-9]+[Cc][0-9]+|[RrCc])$")
NAME_RESERVED = ("TRUE", "FALSE", "R", "C")
NAME_CAP = 60


def defined_name(label):
    """A label as a legal Excel defined name: word characters, no leading digit, not a cell ref."""
    text = NAME_BAD_RE.sub("_", str(label or "").strip()).strip("_.")
    while "__" in text:
        text = text.replace("__", "_")
    if not text:
        return ""
    if text[0].isdigit():
        text = "_" + text
    if NAME_CELL_RE.match(text) or text.upper() in NAME_RESERVED:
        text += "_"
    return text[:NAME_CAP]


def name_candidates(base, sheet):
    """The names to try for one input cell: the label, then the sheet plus the label, then _2.."""
    yield base
    if sheet:
        yield ("%s_%s" % (sheet, base))[:NAME_CAP]
    for n in range(2, 60):
        suffix = "_%d" % n
        yield base[:NAME_CAP - len(suffix)] + suffix


def cell_format(col_formats, row_formats, i, j):
    """W3: which format one table cell takes.

    A `rowFormats` entry wins over the column format on every cell of its row, except where
    the column format is `text`: the label column stays text. A null row entry, or no row
    entry at all, falls back to the column format. A per-cell `{"v":.., "format":..}` wins
    over both, in unpack().
    """
    col = col_formats[j] if j < len(col_formats) else None
    if not row_formats or col == "text":
        return col
    row = row_formats[i] if i < len(row_formats) else None
    return row or col


def is_list_rule(rule):
    """True when a validation rule is a dropdown, so its literal string values are inputs (W2)."""
    return isinstance(rule, dict) and rule.get("list") not in (None, "", [])


def table_name(rowdata, columns, j):
    """The name for a table input cell: its row label and its column header."""
    label = rowdata[0] if rowdata and not isinstance(rowdata[0], dict) else None
    head = columns[j] if j < len(columns) else ""
    return " ".join(str(x) for x in (label, head) if x not in (None, "", True, False))


def write_link(ctx, cell, target, where):
    """Attach a hyperlink to a cell: a URL, or an internal 'Sheet'!A1 reference (W14)."""
    text = str(target).strip()
    if URL_RE.match(text):
        cell.hyperlink = text
        return True
    if INTERNAL_LINK_RE.match(text):
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=text.lstrip("#"))
        return True
    ctx.warnings.append(
        "%s: link %r is not a URL or a Sheet!A1 reference, written as plain text" % (where, target)
    )
    return False


def number_format(value, fmt, is_formula, formats=None):
    formats = formats or FORMATS
    if fmt:
        return formats.get(fmt, formats["number"])
    if is_formula or isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return formats["number"] if float(value) != int(value) else formats["count"]


def _fixed(value, decimals):
    """abs(value) rounded half up to `decimals` places with thousands separators, plus its sign."""
    quantum = Decimal(1).scaleb(-decimals)
    d = Decimal(str(value)).quantize(quantum, rounding=ROUND_HALF_UP)
    negative = d < 0
    return ("{:,.%df}" % decimals).format(abs(d)), negative


# The builder's own formats, as (decimals, prefix, suffix) for the display string.
_SHOWN = {
    "currency": (0, "$", ""),
    "currency2": (2, "$", ""),
    "percent": (1, "", "%"),
    "multiple": (1, "", "x"),
    "count": (0, "", ""),
    "number": (2, "", ""),
}
ZERO_DASH = "-"


def display(value, nf):
    """The string Excel shows for a literal number under one of the builder's formats.

    The formats are three part (W10), so zero shows as a dash and a negative shows in
    parentheses. Returns None for anything that is not a plain number or for a format the
    builder does not emit, so the snapshot simply omits "d".
    """
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    name = None
    for key, spec in _SHOWN.items():
        if nf == FORMATS[key]:
            name = key
            break
    if name is None:
        return None
    decimals, prefix, suffix = _SHOWN[name]
    scale = 100 if name == "percent" else 1
    text, negative = _fixed(value * scale, decimals)
    if not negative and float(Decimal(text.replace(",", ""))) == 0:
        return ZERO_DASH
    body = prefix + text + suffix
    return "(%s)" % body if negative else body


def put(ctx, ws, row, col, raw, fmt, where, exempt_length=False, kind=None, block_type=None,
        role=None, label=False, picklist=False):
    """Write one spec value into a cell. Returns (cell, value, role).

    `kind` fixes the snapshot kind; when None it is the cell's role (W2). `role` overrides
    the derived role, which is how a spec's `cellRoles` and a row's `role` reach the cell.
    """
    spec = unpack(raw, fmt)
    value, fmt, is_formula = spec.value, spec.fmt, spec.is_formula
    if fmt and fmt not in ctx.styler.formats:
        ctx.warnings.append("%s: format '%s' is not one of %s, used number"
                            % (where, fmt, ", ".join(sorted(ctx.styler.formats))))
    cell = ws.cell(row, col)
    if is_formula:
        text = value if isinstance(value, str) else str(value)
        if not text.strip().startswith("="):
            ctx.errors.append("%s: formula does not start with '=': %r" % (where, text))
        else:
            text = ctx.resolver.resolve(text.strip(), ws.title, where)
        cell.value = text
        value = text
        ctx.formulas += 1
        ctx.cells += 1
    else:
        if fmt == "year" and value is not None:
            value = str(value)
        if isinstance(value, str):
            value = ctx.clean(value, where)
            ctx.check_length(value, where, exempt_length)
        cell.value = value
        if value is not None and value != "":
            ctx.cells += 1
    nf = number_format(value, fmt, is_formula, ctx.styler.formats)
    if nf:
        cell.number_format = nf
    if label:
        ctx.fit_label(col, value)
    else:
        ctx.fit(col, value, fmt, is_formula)
    if isinstance(value, str) and not is_formula:
        ctx.want_wrap(row, col, value)
    role = role or spec.role or cell_role(value, is_formula, spec.link, picklist, label, fmt)
    if spec.link and not write_link(ctx, cell, spec.link, where) and role == "link":
        role = "text"
    ctx.rec(row, col, cell.value, kind or role, nf, is_formula)
    return cell, value, role


def is_input(value, is_formula, block_type=None, picklist=False, label=False, fmt=None):
    """True when the cell's role is `input`. The older spelling; `cell_role` says more."""
    return cell_role(value, is_formula, None, picklist, label, fmt) == "input"


# --------------------------------------------------------- data validation


def _num_text(value):
    """A spec number as the text Excel wants in a validation formula, or None when it is not a number."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return str(int(value)) if float(value) == int(value) else repr(float(value))


def make_validation(ctx, ws, rule, where):
    """An openpyxl DataValidation for a spec `validation` rule, or None (with a warning) when malformed.

    Shapes: {"list": ["Low", "Base", "High"]} or {"list": "=<range>"} (templates resolved),
    {"min": 0, "max": 1} (decimal), {"type": "whole"|"decimal", "min": .., "max": ..}.
    Either bound may be left out. Every rule allows blanks and shows a short error message.
    """
    if not isinstance(rule, dict):
        ctx.warnings.append("%s: validation must be an object, ignored" % where)
        return None
    if "list" in rule:
        items = rule["list"]
        if isinstance(items, str) and items.strip().startswith("="):
            formula = ctx.resolver.resolve(items.strip(), ws.title, where)[1:]
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        elif isinstance(items, list) and items:
            texts = [ctx.clean(str(x), where) for x in items]
            bad = [t for t in texts if "," in t or '"' in t]
            if bad:
                ctx.warnings.append("%s: list item %r holds a comma or quote, validation skipped" % (where, bad[0]))
                return None
            joined = '"%s"' % ",".join(texts)
            if len(joined) > 255:
                ctx.warnings.append("%s: list is over 255 characters, validation skipped" % where)
                return None
            dv = DataValidation(type="list", formula1=joined, allow_blank=True)
        else:
            ctx.warnings.append("%s: list must be a non-empty array or a =range, ignored" % where)
            return None
        dv.error = "Pick a value from the list."
    else:
        kind = str(rule.get("type") or "decimal").lower()
        if kind not in ("whole", "decimal"):
            ctx.warnings.append("%s: type must be whole or decimal, ignored" % where)
            return None
        lo, hi = _num_text(rule.get("min")), _num_text(rule.get("max"))
        if lo is None and hi is None:
            ctx.warnings.append("%s: needs a list, a min or a max, ignored" % where)
            return None
        noun = "a whole number" if kind == "whole" else "a number"
        if lo is not None and hi is not None:
            dv = DataValidation(type=kind, operator="between", formula1=lo, formula2=hi, allow_blank=True)
            dv.error = "Enter %s between %s and %s." % (noun, lo, hi)
        elif lo is not None:
            dv = DataValidation(type=kind, operator="greaterThanOrEqual", formula1=lo, allow_blank=True)
            dv.error = "Enter %s of at least %s." % (noun, lo)
        else:
            dv = DataValidation(type=kind, operator="lessThanOrEqual", formula1=hi, allow_blank=True)
            dv.error = "Enter %s of at most %s." % (noun, hi)
    dv.showErrorMessage = True
    dv.errorTitle = "Invalid entry"
    return dv


def apply_validation(ctx, ws, sqref, rule, where, cells):
    """Attach a validation to `sqref` and flag the recorded snapshot cells. No-op when rule is None."""
    if rule is None:
        return
    dv = make_validation(ctx, ws, rule, where)
    if dv is None:
        return
    dv.add(sqref)
    ws.add_data_validation(dv)
    for key in cells:
        entry = ctx.snap_cells.get(key)
        if entry is not None:
            entry["validation"] = True


# ------------------------------------------------------------ block writers


def write_title(ctx, ws, g):
    if g.title:
        where = "%s!%s title" % (g.sheet, g.id)
        cell = ws.cell(g.row, g.col)
        cell.value = ctx.clean(str(g.title), where)
        ctx.check_length(cell.value, where)
        ctx.cells += 1
        ctx.rec(g.row, g.col, cell.value, "title")
    ctx.styler.block_title(ws, g)
    if g.note:
        write_note(ctx, ws, g.row, g.note_col_after, g.note, "%s!%s block note" % (g.sheet, g.id))


def write_note(ctx, ws, row, col, text, where):
    """Write one grey italic note (block note, inputs row note, table row note) and record it.

    The note lands in its row band's one note column (W4); merge_notes() spans it across the
    empty columns to the sheet's right edge once every block on the sheet is written.
    """
    cell = ws.cell(row, col)
    text = ctx.clean(str(text), where)
    if cell.value not in (None, "") and cell.value != text:
        ctx.warnings.append("%s: a note is already in %s, the earlier one was replaced"
                            % (where, a1(row, col)))
    cell.value = text
    ctx.check_length(cell.value, where)
    ctx.styler.note(cell)
    ctx.cells += 1
    ctx.fit_note(col, cell.value)
    ctx.note_cells.append((row, col))
    ctx.rec(row, col, cell.value, "note")
    return cell


def write_pairs(ctx, ws, g):
    st = ctx.styler
    write_title(ctx, ws, g)
    ctx.want(g.col, st.label_width)
    for c in range(g.col + 1, g.col + g.width):
        ctx.want(c, st.value_width)
    for i, rowdef in enumerate(g.pair_rows):
        if not isinstance(rowdef, dict):
            continue
        row = g.row + 1 + i
        where = "%s!%s row %d" % (g.sheet, g.id, i)
        label = ctx.clean(str(rowdef.get("label", "")), where)
        ctx.check_length(label, where)
        lcell = ws.cell(row, g.col)
        lcell.value = label
        if label:
            ctx.cells += 1
            ctx.fit_label(g.col, label)      # W1: the label column fits its longest label
            ctx.want_wrap(row, g.col, label)
            ctx.rec(row, g.col, label, g.role_at(i, 0) or "label")
        st.pair_label(lcell)
        validation = rowdef.get("validation")
        row_role = str(rowdef.get("role") or "").strip().lower() or None
        cell, value, role = put(
            ctx, ws, row, g.value_col, rowdef.get("value"), rowdef.get("format"), where,
            block_type=g.type, role=g.role_at(i, 1) or (row_role if row_role in ROLES else None),
            picklist=is_list_rule(validation),
        )
        st.pair_value(cell, role=role)
        ctx.name_input(ws, row, g.value_col, role, label)
        apply_validation(ctx, ws, cell.coordinate, validation, where + " validation",
                         [(row, g.value_col)])
        note = rowdef.get("note")
        if note:
            write_note(ctx, ws, row, g.note_col, note, where + " note")


def write_table(ctx, ws, g):
    st = ctx.styler
    write_title(ctx, ws, g)
    columns = g.block.get("columns") or []
    col_formats = g.block.get("colFormats") or []
    data = g.block.get("data") or []
    total = g.total_values          # the approved row, or the one balanced generated (W11)

    ctx.want(g.col, st.label_width)
    for c in range(g.col + 1, g.col + g.width):
        ctx.want(c, st.value_width)

    col_validation = g.block.get("colValidation")
    picklists = [
        is_list_rule(col_validation[j]) if isinstance(col_validation, list) and j < len(col_validation)
        else False
        for j in range(g.data_cols)
    ]
    row_formats = g.block.get("rowFormats")
    if row_formats is not None and not isinstance(row_formats, list):
        ctx.warnings.append("%s!%s: rowFormats must be a list, ignored" % (g.sheet, g.id))
        row_formats = None

    for j in range(g.data_cols):
        where = "%s!%s header %d" % (g.sheet, g.id, j)
        cell = ws.cell(g.header_row, g.col + j)
        ctx.fit(g.col + j, str(columns[j]) if j < len(columns) else "", bold=True)
        text = ctx.clean(str(columns[j]) if j < len(columns) else "", where)
        ctx.check_length(text, where)
        cell.value = text
        if text:
            ctx.cells += 1
            ctx.want_wrap(g.header_row, g.col + j, text)   # W5: headers wrap, never widen
            ctx.rec(g.header_row, g.col + j, text, "header")
        st.header(cell, first=(j == 0))
    ctx.head_rows[g.header_row] = True

    for i, rowdata in enumerate(data):
        if not isinstance(rowdata, list):
            continue
        for j in range(g.data_cols):
            raw = rowdata[j] if j < len(rowdata) else None
            fmt = cell_format(col_formats, row_formats, i, j)
            where = "%s!%s r%d c%d" % (g.sheet, g.id, i, j)
            cell, value, role = put(ctx, ws, g.first_data_row + i, g.col + j, raw, fmt, where,
                                    block_type=g.type, role=g.role_at(i, j),
                                    label=(j == 0), picklist=picklists[j])
            st.table_cell(cell, first=(j == 0), role=role)
            if j > 0 and role == "input":
                ctx.name_input(ws, g.first_data_row + i, g.col + j, role,
                               table_name(rowdata, columns, j))
        note = g.row_notes[i] if i < len(g.row_notes) else None
        if note:
            write_note(ctx, ws, g.first_data_row + i, g.note_col_after, note,
                       "%s!%s r%d note" % (g.sheet, g.id, i))

    if g.total_row:
        for j in range(g.data_cols):
            raw = total[j] if j < len(total) else None
            fmt = col_formats[j] if j < len(col_formats) else None
            where = "%s!%s total c%d" % (g.sheet, g.id, j)
            cell, value, role = put(ctx, ws, g.total_row, g.col + j, raw, fmt, where,
                                    kind="total", block_type=g.type, label=(j == 0))
            st.table_cell(cell, first=(j == 0), bold=True, role=role)
        if g.total_note:
            write_note(ctx, ws, g.total_row, g.note_col_after, g.total_note,
                       "%s!%s total note" % (g.sheet, g.id))

    if isinstance(col_validation, list) and g.n_data_rows > 0:
        for j, rule in enumerate(col_validation[:g.data_cols]):
            if not rule:
                continue
            col = g.col + j
            target = "%s:%s" % (a1(g.first_data_row, col), a1(g.last_data_row, col))
            where = "%s!%s colValidation %d" % (g.sheet, g.id, j)
            apply_validation(ctx, ws, target, rule, where,
                             [(r, col) for r in range(g.first_data_row, g.last_data_row + 1)])
    elif col_validation not in (None, []):
        ctx.warnings.append("%s!%s: colValidation must be a list, ignored" % (g.sheet, g.id))


def write_text(ctx, ws, g):
    st = ctx.styler
    write_title(ctx, ws, g)
    where = "%s!%s text" % (g.sheet, g.id)
    body = ctx.clean(str(g.block.get("text") or ""), where)
    ctx.check_length(body, where, exempt=True)  # a paragraph is allowed to be long
    row = g.row + 1
    for c in range(g.col, g.col + g.width):
        st.body_text(ws.cell(row, c))
    cell = ws.cell(row, g.col)
    cell.value = body
    if body:
        ctx.cells += 1
        ctx.rec(row, g.col, body, "text")
    if g.width > 1:
        ws.merge_cells(start_row=row, start_column=g.col,
                       end_row=row, end_column=g.col + g.width - 1)
    ws.row_dimensions[row].height = wrapped_height(body, g.width * 16)


def wrapped_height(text, chars_per_line):
    """Room for `text` wrapped at that many width units, one character per unit.

    Deliberately pessimistic: a line too many is a blank line, a line too few clips.
    """
    lines = max(1, int(math.ceil(len(text or "") / float(max(chars_per_line, 10)))))
    return max(15, lines * 13 + 4)


def write_chart(ctx, ws, g):
    spec = g.block.get("chart") or {}
    where = "%s!%s chart" % (g.sheet, g.id)
    src = ctx.registry.get(str(spec.get("sourceBlock") or ""))
    if src is None:
        ctx.errors.append("%s: sourceBlock '%s' not found" % (where, spec.get("sourceBlock")))
        return
    if src.type not in ("table", "calc"):
        ctx.errors.append("%s: sourceBlock '%s' is not a table or calc block" % (where, src.id))
        return
    if src.sheet != g.sheet:
        ctx.errors.append("%s: sourceBlock '%s' is on another sheet" % (where, src.id))
        return
    if src.n_data_rows < 1:
        ctx.errors.append("%s: sourceBlock '%s' has no data rows" % (where, src.id))
        return

    kind = str(spec.get("kind") or "bar").lower()
    if kind == "line":
        chart = LineChart()
    elif kind == "pie":
        chart = PieChart()
    else:
        chart = BarChart()
        chart.type = "col"

    series_cols = [int(c) for c in (spec.get("seriesCols") or [1])]
    cat_col = int(spec.get("categoriesCol") or 0)
    bad = [c for c in series_cols + [cat_col] if c < 0 or c >= src.data_cols]
    if bad:
        ctx.errors.append("%s: column index %s out of range for '%s'" % (where, bad[0], src.id))
        return
    if kind == "pie":
        series_cols = series_cols[:1]

    for c in series_cols:
        ref = Reference(ws, min_col=src.data_col_index(c), min_row=src.header_row,
                        max_col=src.data_col_index(c), max_row=src.last_data_row)
        chart.add_data(ref, titles_from_data=True)
    cats = Reference(ws, min_col=src.data_col_index(cat_col), min_row=src.first_data_row,
                     max_col=src.data_col_index(cat_col), max_row=src.last_data_row)
    chart.set_categories(cats)

    if g.title:
        chart.title = ctx.clean(str(g.title), where)
    if spec.get("yTitle") and kind != "pie":
        chart.y_axis.title = ctx.clean(str(spec["yTitle"]), where)
    if len(series_cols) < 2:
        chart.legend = None                      # one series needs no legend
    if kind != "pie":
        chart.y_axis.majorGridlines = None       # the axis carries the scale, not the grid
        chart.x_axis.majorGridlines = None
    # W12: the plot area is inset so the drawing does not run into its own frame.
    chart.layout = Layout(manualLayout=ManualLayout(
        xMode="edge", yMode="edge", x=0.06, y=0.16, w=0.90, h=0.70))
    chart.width = g.width * 1.7                  # provisional: finish_charts sizes it
    chart.height = g.height * 0.5
    # W12: anchored from the block's top-left to its bottom-right, so the drawing covers
    # exactly the cells the plan reserved and never spills over a written one.
    ws.add_chart(chart, TwoCellAnchor(
        editAs="oneCell",
        _from=AnchorMarker(col=g.col - 1, row=g.row - 1),
        to=AnchorMarker(col=g.col - 1 + g.width, row=g.row - 1 + g.height),
    ))
    ctx.charts += 1
    ctx.deferred_charts.append((chart, g))
    ctx.rec_chart(g.id, a1(g.row, g.col), kind, len(series_cols))


# ----------------------------------------------------------- comments (N9)

COMMENT_AUTHOR = "Blueprint"
COMMENT_WIDTH = 260               # points: wide enough for a thread, narrow enough to read
COMMENT_LINE = 15                 # points a line of comment text takes


def block_threads(spec, warnings):
    """N9: {blockId: [thread, ...]} from `spec.threads`, or {} when `comments` is off.

    Off by default: a workbook only carries its threads when the spec asks for them. The
    server resolves them first, so what arrives here is the closed history of each block.
    """
    if not spec.get("comments"):
        return {}
    raw = spec.get("threads")
    if raw in (None, []):
        return {}
    if not isinstance(raw, list):
        warnings.append("threads must be a list of {blockId, entries}, ignored")
        return {}
    out = {}
    for thread in raw:
        if not isinstance(thread, dict):
            warnings.append("a thread is not an object, ignored")
            continue
        block_id = str(thread.get("blockId") or "").strip()
        if not block_id:
            warnings.append("a thread names no blockId, ignored")
            continue
        out.setdefault(block_id, []).append(thread)
    return out


def thread_lines(thread):
    """One line per entry of one thread, in the order given: oldest first, newest last."""
    lines = []
    for entry in thread.get("entries") or []:
        if not isinstance(entry, dict):
            continue
        text = str(entry.get("text") or "").strip()
        if not text:
            continue
        by = str(entry.get("by") or "").strip()
        status = str(entry.get("status") or "").strip().lower()
        line = "%s: %s" % (by, text) if by else text
        if status and status != "resolved":
            line += " [%s]" % status
        lines.append(line)
    return lines


def comment_text(threads):
    """The body of one block's comment: its threads, a blank line between them, newest last."""
    parts = ["\n".join(lines) for lines in (thread_lines(t) for t in threads) if lines]
    return "\n\n".join(parts)


def write_comments(ctx, ws, geoms, threads):
    """N9: one Excel comment per block, on its title cell, holding that block's threads."""
    if not threads:
        return
    for g in geoms:
        body = comment_text(threads.get(g.id) or [])
        if not body:
            continue
        cell = ws.cell(g.row, g.col)
        if cell.value in (None, ""):
            ctx.warnings.append("%s!%s: no title cell to hang its %d thread(s) on, no comment "
                                "written" % (g.sheet, g.id, len(threads[g.id])))
            continue
        note = Comment(body, COMMENT_AUTHOR)
        note.width = COMMENT_WIDTH
        note.height = max(60, COMMENT_LINE * (body.count("\n") + 2))
        cell.comment = note
        ctx.mark_comment(g.row, g.col)


# ----------------------------------------------- finishing passes per sheet

CHAR_CM = 0.185            # one column width unit in centimetres at 96 dpi
POINT_CM = 0.0353          # one row height point in centimetres
DEFAULT_WIDTH = 8.43       # Excel's width for a column nobody set
DEFAULT_HEIGHT = 15.0      # and its row height in points
PORTRAIT_WIDTH = 90        # width units that still fit a portrait page
PAGE_ROWS = 45             # rows that still fit one page, so title rows start repeating


def last_used_col(ws):
    cols = [c for (_, c), cell in ws._cells.items() if cell.value is not None]
    return max(cols) if cols else 0


def merge_notes(ctx, ws):
    """W4: span every note from its band's note column to the sheet's last used column plus two.

    The note keeps wrap off and one row, so the merge is what gives it room to read; it stops
    before any cell that holds a value or is already merged, so nothing is overwritten.
    """
    if not ctx.note_cells:
        return
    target = last_used_col(ws) + 2
    taken = set()
    for rng in ws.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                taken.add((r, c))
    for row, col in sorted(set(ctx.note_cells)):
        end = col
        while end + 1 <= target:
            nxt = ws.cell(row, end + 1)
            if nxt.value not in (None, "") or (row, end + 1) in taken:
                break
            end += 1
        if end == col:
            continue
        for c in range(col + 1, end + 1):
            ctx.styler.note(ws.cell(row, c))
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=end)
        for c in range(col, end + 1):
            taken.add((row, c))
        ctx.note_spans[(row, col)] = end


def finish_widths(ctx, ws):
    """Set every column width, giving a merged note's empty columns room for the note's text."""
    widths = ctx.final_widths()
    for (row, col), end in sorted(ctx.note_spans.items()):
        need = len(str(ws.cell(row, col).value or "")) + 2
        have = widths.get(col, DEFAULT_WIDTH)
        for c in range(col + 1, end + 1):
            if have >= need:
                break
            if c in widths:                      # another block owns that column
                have += widths[c]
                continue
            extra = max(DEFAULT_WIDTH, min(25, need - have))
            widths[c] = max(widths.get(c, 0), extra)
            have += extra
    for col, width in sorted(widths.items()):
        ctx.set_width(ws, col, width)
    return widths


def with_wrap(cell, vertical=None):
    al = cell.alignment
    cell.alignment = Alignment(
        horizontal=al.horizontal, vertical=vertical or al.vertical or "center",
        wrap_text=True, indent=al.indent or 0,
    )


def wrapped_lines(chars, width):
    return max(1, int(math.ceil(chars / float(max(width - 1, 1)))))


def apply_wraps(ctx, ws, widths):
    """W5: wrap what does not fit, then give the row the height its wrapped lines need.

    Header cells always wrap, which is what keeps a long header from widening its column;
    a data text cell wraps once its column has hit the style's cap and the text is still
    longer. Notes never wrap: they have their merge instead.
    """
    notes = set(ctx.note_cells)
    heights = {}
    for (row, col), chars in sorted(ctx.wrap_cells.items()):
        if (row, col) in notes:
            continue
        width = widths.get(col, DEFAULT_WIDTH)
        header = row in ctx.head_rows
        lines = wrapped_lines(chars, width)
        if not header and (lines < 2 or width < ctx.caps.get(col, ctx.styler.label_cap)):
            continue
        cell = ws.cell(row, col)
        if cell.value in (None, ""):
            continue
        with_wrap(cell)
        if lines > 1:
            heights[row] = max(heights.get(row, 0), lines)
    for row, lines in heights.items():
        height = max(DEFAULT_HEIGHT, lines * 13 + 4)
        current = ws.row_dimensions[row].height
        if current is None or current < height:
            ws.row_dimensions[row].height = height


def finish_charts(ctx, ws, widths):
    """W12: size every chart to the cells its block reserves, once those widths are known."""
    for chart, g in ctx.deferred_charts:
        cols = sum(widths.get(c, DEFAULT_WIDTH) for c in range(g.col, g.col + g.width))
        rows = sum(ws.row_dimensions[r].height or DEFAULT_HEIGHT
                   for r in range(g.row, g.row + g.height))
        chart.width = max(6.0, round(cols * CHAR_CM, 2))
        chart.height = max(4.0, round(rows * POINT_CM, 2))


def sheet_role(geoms, name):
    """The sheet's role, for its tab colour (W16): cover, inputs, output or calc.

    A tab is an inputs tab or an outputs tab only when that is the whole job: text and
    chart blocks do not count, and a mix is a calc tab. Every sheet in the house format
    ends with the number it produces, so colouring any tab with an `output` block would
    colour almost all of them.
    """
    if not geoms:
        return "cover" if norm(name or "").replace(" ", "") == "readme" else None
    types = set(g.type for g in geoms) - set(("text", "chart"))
    if types == set(("inputs",)):
        return "inputs"
    if types == set(("output",)):
        return "output"
    return "calc"


def colour_tab(ws, geoms, styler):
    colour = styler.tab_color(sheet_role(geoms, ws.title))
    if colour:
        ws.sheet_properties.tabColor = colour


def computed_setup(ws, geoms, widths):
    """W6: the sheet setup the builder works out for itself. Spec values override each field.

    freeze holds the label column and the header rows of the first table or calc block, the
    print area is the used range, and the page lands one page wide, landscape once the used
    range is wider than a portrait page.
    """
    out = {}
    blocks = [g for g in geoms if g.type in WRITERS]
    tables = sorted([g for g in blocks if g.type in ("table", "calc")], key=lambda g: (g.row, g.col))
    pairs = sorted([g for g in blocks if g.type in ("inputs", "output")], key=lambda g: (g.row, g.col))
    if tables:
        out["freeze"] = a1(tables[0].first_data_row, tables[0].col + 1)
    elif pairs:
        out["freeze"] = a1(pairs[0].row + 1, pairs[0].col + 1)
    used = used_range(ws)
    if not used:
        return out
    first, last = used.split(":")
    r1, c1 = ws[first].row, ws[first].column
    r2, c2 = ws[last].row, ws[last].column
    for g in blocks:
        # A chart is a drawing, so it holds no values: without this the print area would
        # stop above it and the chart would never print.
        r2, c2 = max(r2, g.body_ext[2]), max(c2, g.body_ext[3])
    out["printArea"] = "%s:%s" % (a1(r1, c1), a1(r2, c2))
    total = sum(widths.get(c, DEFAULT_WIDTH) for c in range(c1, c2 + 1))
    out["landscape"] = total > PORTRAIT_WIDTH
    out["fitToWidth"] = True
    if tables and r2 - r1 + 1 > PAGE_ROWS:
        out["titleRows"] = "%d:%d" % (tables[0].row, tables[0].header_row)
    return out


WRITERS = {
    "inputs": write_pairs,
    "output": write_pairs,
    "table": write_table,
    "calc": write_table,
    "text": write_text,
    "chart": write_chart,
}


# ------------------------------------------------------------------ read me


COVER_WIDTHS = (30, 45, 14)      # 89 units: the purpose paragraph merges to the portrait width


def cover_contents(spec, warnings=None):
    """W11: one Contents row per tab, in sheet order, whatever the spec's own list holds.

    `holds` comes from the matching `readme.contents` entry, else from the sheet's own
    `purpose`. A contents entry naming no tab is reported and dropped, so the cover can
    never disagree with the workbook.
    """
    readme = spec.get("readme") or {}
    given = {}
    for item in readme.get("contents") or []:
        if isinstance(item, dict) and item.get("tab"):
            given[norm(item["tab"])] = str(item.get("holds") or "")
    rows = []
    for sheet in spec.get("sheets") or []:
        name = str(sheet.get("name") or "")
        if not name or norm(name).replace(" ", "") == "readme":
            continue
        holds = given.pop(norm(name), None)
        if holds is None:
            holds = str(sheet.get("purpose") or "")
        rows.append({"tab": name, "holds": holds})
    if given and warnings is not None:
        warnings.append("Read Me contents names %s, which is not a tab in this workbook"
                        % ", ".join("'%s'" % t for t in sorted(given)))
    return rows


def cover_results(spec):
    """W11: the cover's Answer section, from the build payload.

    `readme.answer` (or `project.takeaway`) is the one line the model exists to produce.
    `readme.results` adds a row per result with its live value: `true` takes every row of
    every `output` block, and a list of `{label, tab, ref, format}` says exactly which.
    """
    readme = spec.get("readme") or {}
    project = spec.get("project") or {}
    answer = readme.get("answer") or project.get("takeaway") or project.get("answer") or ""
    want = readme.get("results")
    rows = []
    if isinstance(want, list):
        for item in want:
            if not isinstance(item, dict) or not item.get("label"):
                continue
            rows.append({"label": str(item["label"]), "tab": str(item.get("tab") or ""),
                         "ref": str(item.get("ref") or ""), "format": item.get("format")})
        return str(answer), [r for r in rows if r["ref"].startswith("=")]
    if not want:
        return str(answer), []
    for sheet in spec.get("sheets") or []:
        for block in sheet.get("blocks") or []:
            if block.get("type") != "output":
                continue
            g = Geom(block, sheet.get("name"))
            for i, rowdef in enumerate(g.pair_rows):
                if not isinstance(rowdef, dict) or not rowdef.get("label"):
                    continue
                rows.append({
                    "label": str(rowdef["label"]),
                    "tab": str(sheet.get("name") or ""),
                    "ref": "=%s!%s" % (qname(sheet.get("name")), a1(g.row + 1 + i, g.value_col)),
                    "format": rowdef.get("format"),
                })
    return str(answer), rows


def readme_data(spec):
    readme = spec.get("readme") or {}
    project = spec.get("project") or {}
    purpose = readme.get("purpose") or project.get("brief") or "This workbook has no stated purpose yet."
    rules = readme.get("rules") or DEFAULT_RULES
    name = project.get("name") or "Workbook"
    return name, purpose, rules


def write_readme(ctx, ws, spec):
    st = ctx.styler
    name, purpose, rules = readme_data(spec)
    contents = cover_contents(spec, ctx.warnings)
    answer, results = cover_results(spec)
    for i, width in enumerate(COVER_WIDTHS):
        ctx.set_width(ws, i + 1, width)

    title = ws.cell(1, 1)
    title.value = ctx.clean(str(name), "Read Me title")
    title.font = Font(name=st.base_name, size=st.title_size, bold=True, color=st.title_rule_color)
    sub = ws.cell(2, 1)
    sub.value = "v1, built %s" % date.today().isoformat()
    sub.font = Font(name=st.base_name, size=st.base_size, italic=True, color=st.note_color)
    ctx.cells += 2
    ctx.rec(1, 1, title.value, "readme")
    ctx.rec(2, 1, sub.value, "readme")
    for c in range(1, 4):
        ws.cell(3, c).border = Border(bottom=st.title_rule_side)
    ws.row_dimensions[3].height = 6

    row = 5
    row = readme_section(ctx, ws, row, "Purpose")
    row = readme_paragraph(ctx, ws, row, purpose, "Read Me purpose")

    if answer or results:
        row = readme_section(ctx, ws, row + 1, "Answer")
        if answer:
            row = readme_paragraph(ctx, ws, row, answer, "Read Me answer")
        for item in results:
            readme_result(ctx, ws, row, item)
            row += 1

    row = readme_section(ctx, ws, row + 1, "Contents")
    for item in contents:
        if not isinstance(item, dict):
            continue
        tab = ctx.clean(str(item.get("tab", "")), "Read Me contents")
        holds = ctx.clean(str(item.get("holds", "")), "Read Me contents")
        readme_pair(ctx, ws, row, tab, holds)
        row += 1

    row = readme_section(ctx, ws, row + 1, "Reading a Cell")
    for label, holds, color, italic in (
        ("Blue", "an input, change it", st.input_color, False),
        ("Black", "a formula, do not overwrite it", st.formula_color, False),
        ("Grey italic", "a note or a source, not a number the model uses",
         st.note_color, st.note_italic),
    ):
        readme_pair(ctx, ws, row, label, holds, color=color, italic=italic)
        row += 1

    row = readme_section(ctx, ws, row + 1, "Rules")
    for rule in rules:
        row = readme_paragraph(ctx, ws, row, str(rule), "Read Me rules")


def readme_section(ctx, ws, row, label):
    st = ctx.styler
    for c in range(1, 4):
        cell = ws.cell(row, c)
        cell.font = st.font(bold=True)
        cell.border = Border(bottom=Side(style="thin", color=st.title_rule_color))
    ws.cell(row, 1).value = label
    ctx.cells += 1
    ctx.rec(row, 1, label, "readme")
    return row + 1


def readme_paragraph(ctx, ws, row, text, where):
    st = ctx.styler
    body = ctx.clean(str(text), where)
    for c in range(1, 4):
        st.body_text(ws.cell(row, c))
    ws.cell(row, 1).value = body
    ctx.cells += 1
    ctx.rec(row, 1, body, "readme")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    ws.row_dimensions[row].height = wrapped_height(body, sum(COVER_WIDTHS))
    return row + 1


def readme_pair(ctx, ws, row, left, right, color=BLACK, italic=False):
    st = ctx.styler
    lcell = ws.cell(row, 1)
    lcell.value = left
    lcell.font = st.font(color=color, italic=italic)
    lcell.alignment = Alignment(horizontal="left", vertical="top")
    rcell = ws.cell(row, 2)
    rcell.value = right
    rcell.font = st.font()
    rcell.alignment = Alignment(wrap_text=True, vertical="top")
    ctx.cells += 2
    ctx.rec(row, 1, left, "readme")
    ctx.rec(row, 2, right, "readme")
    ws.row_dimensions[row].height = wrapped_height(right, COVER_WIDTHS[1])


def readme_result(ctx, ws, row, item):
    """One Result row on the cover: the output's label, its live value, the tab it lives on."""
    st = ctx.styler
    lcell = ws.cell(row, 1)
    lcell.value = str(item["label"])
    lcell.font = st.font()
    lcell.alignment = Alignment(horizontal="left", vertical="top")
    vcell = ws.cell(row, 2)
    vcell.value = item["ref"]
    vcell.font = st.font()
    vcell.alignment = Alignment(horizontal="right", vertical="top")
    nf = number_format(None, item.get("format"), True, st.formats)
    if nf:
        vcell.number_format = nf
    tcell = ws.cell(row, 3)
    tcell.value = str(item["tab"])
    tcell.font = st.font(italic=st.note_italic, color=st.note_color)
    tcell.alignment = Alignment(horizontal="left", vertical="top")
    ctx.cells += 3
    ctx.formulas += 1
    ctx.rec(row, 1, lcell.value, "readme")
    ctx.rec(row, 2, vcell.value, "readme", nf, True)
    ctx.rec(row, 3, tcell.value, "readme")


# ------------------------------------------------------------- sheet setup

CELL_REF_RE = re.compile(r"^\$?[A-Za-z]{1,3}\$?\d+$")
RANGE_REF_RE = re.compile(r"^\$?[A-Za-z]{1,3}\$?\d+:\$?[A-Za-z]{1,3}\$?\d+$")


def used_range(ws):
    """The range from the first to the last cell holding a value, or None on an empty sheet.

    Reads values rather than ws.dimensions because openpyxl creates empty cells whenever a
    range is iterated, which would drag the range up to A1. A merged span counts as used:
    a paragraph and a note read across their whole merge, so printing has to cover it.
    """
    rows = [r for (r, c), cell in ws._cells.items() if cell.value is not None]
    cols = [c for (r, c), cell in ws._cells.items() if cell.value is not None]
    if not rows:
        return None
    r1, r2, c1, c2 = min(rows), max(rows), min(cols), max(cols)
    for rng in ws.merged_cells.ranges:
        if r1 <= rng.min_row <= r2 and c1 <= rng.min_col <= c2:
            r2 = max(r2, rng.max_row)
            c2 = max(c2, rng.max_col)
    return "%s:%s" % (a1(r1, c1), a1(r2, c2))


def sheet_setup(ws, sheet, warnings, computed=None):
    """Freeze panes, print area, orientation, fit to width and zoom for one sheet.

    Runs after the sheet is written so `"printArea": "auto"` can read the used range. Every
    field the spec sets wins; the rest come from `computed`, the defaults W6 works out from
    what was written. A bad spec value warns and falls back to the computed default.
    """
    where = "sheet '%s'" % ws.title
    computed = computed or {}
    freeze = sheet.get("freeze")
    if freeze not in (None, "", False):
        text = str(freeze).strip().upper()
        if CELL_REF_RE.match(text):
            ws.freeze_panes = text.replace("$", "")
        else:
            warnings.append("%s: freeze %r is not a cell reference, ignored" % (where, freeze))
            freeze = None
    if freeze in (None, "") and computed.get("freeze"):
        ws.freeze_panes = computed["freeze"]

    area = sheet.get("printArea")
    if area not in (None, "", False):
        text = str(area).strip()
        if text.lower() == "auto":
            used = used_range(ws)
            if used:
                ws.print_area = used
            else:
                warnings.append("%s: printArea auto skipped, the sheet is empty" % where)
        elif RANGE_REF_RE.match(text):
            ws.print_area = text.upper()
        else:
            warnings.append("%s: printArea %r is not auto or a range, ignored" % (where, area))
            area = None
    if area in (None, "") and computed.get("printArea"):
        ws.print_area = computed["printArea"]

    landscape = sheet.get("landscape")
    if landscape is not None and not isinstance(landscape, bool):
        warnings.append("%s: landscape must be true or false, ignored" % where)
        landscape = None
    if landscape is None:
        landscape = computed.get("landscape")
    if landscape is not None:
        ws.page_setup.orientation = "landscape" if landscape else "portrait"

    fit = sheet.get("fitToWidth")
    if fit is not None and not isinstance(fit, bool):
        warnings.append("%s: fitToWidth must be true or false, ignored" % where)
        fit = None
    if fit is None:
        fit = computed.get("fitToWidth")
    if fit:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

    rows = sheet.get("titleRows", computed.get("titleRows"))
    if rows:
        text = str(rows).strip()
        if re.match(r"^\d+:\d+$", text):
            ws.print_title_rows = text
        else:
            warnings.append("%s: titleRows %r is not a row range such as 2:3, ignored" % (where, rows))

    zoom = sheet.get("zoom")
    if zoom is not None:
        ok = isinstance(zoom, (int, float)) and not isinstance(zoom, bool) and 10 <= zoom <= 400
        if ok:
            ws.sheet_view.zoomScale = int(zoom)
        else:
            warnings.append("%s: zoom must be a number from 10 to 400, ignored" % where)


# ----------------------------------------------------------------------- qa


def check_sheet_names(names, errors):
    seen = set()
    for name in names:
        if not name or not str(name).strip():
            errors.append("sheet name is empty")
            continue
        if len(name) > 31:
            errors.append("sheet name '%s' is over 31 chars" % name)
        bad = [ch for ch in BAD_SHEET_CHARS if ch in name]
        if bad:
            errors.append("sheet name '%s' contains %s" % (name, " ".join(bad)))
        if norm(name) in seen:
            errors.append("duplicate sheet name '%s'" % name)
        seen.add(norm(name))


def check_overlaps(geoms, errors):
    """No two blocks may occupy the same cell.

    Compared on the blocks' own rectangles: a note column is shared by its whole row band
    and always sits right of every block in it, so it can never collide with one.
    """
    for i in range(len(geoms)):
        for j in range(i + 1, len(geoms)):
            a, b = geoms[i], geoms[j]
            if overlaps(a.body_ext, b.body_ext):
                errors.append(
                    "blocks '%s' (%s:%s) and '%s' (%s:%s) overlap on sheet '%s'"
                    % (a.id, a1(a.body_ext[0], a.body_ext[1]), a1(a.body_ext[2], a.body_ext[3]),
                       b.id, a1(b.body_ext[0], b.body_ext[1]), a1(b.body_ext[2], b.body_ext[3]),
                       a.sheet)
                )


def check_unresolved(ws, errors):
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "{{" in cell.value:
                errors.append("%s!%s still holds an unresolved template" % (ws.title, cell.coordinate))


# ---------------------------------------------------------------------- build


def build(spec):
    errors = []
    warnings = []
    out_path = spec.get("outputPath")
    if not out_path:
        return {"ok": False, "path": None, "sheets": [], "warnings": warnings,
                "errors": ["spec has no outputPath"]}

    styler = make_styler(spec, warnings)
    fidelity = read_fidelity(spec, warnings)
    threads = block_threads(spec, warnings)
    sheets = spec.get("sheets") or []
    if not isinstance(sheets, list) or any(not isinstance(s, dict) for s in sheets):
        return {"ok": False, "path": None, "sheets": [], "warnings": warnings,
                "errors": ["spec sheets must be a list of sheet objects"]}
    if not sheets:
        errors.append("spec has no sheets")

    # The cover page is written into a sheet the plan calls "Read Me" (any case) that has no
    # blocks of its own. Never add a tab the user did not approve.
    check_sheet_names([s.get("name") for s in sheets], errors)

    registry = {}
    per_sheet = []
    for sheet in sheets:
        geoms = []
        for block in sheet.get("blocks") or []:
            g = Geom(block, sheet.get("name"))
            if not g.id:
                errors.append("a block on '%s' has no id" % sheet.get("name"))
            elif g.id in registry:
                errors.append("duplicate block id '%s'" % g.id)
            else:
                registry[g.id] = g
            if g.type not in WRITERS:
                errors.append("block '%s' has unknown type '%s'" % (g.id, g.type))
            for problem in bad_cell_roles(block):
                warnings.append("block '%s' cellRoles: %s, ignored" % (g.id, problem))
            geoms.append(g)
        # W13, then W11's fidelity moves, then W4: one block width per sheet, what balanced
        # is allowed to move, then one note column per row band.
        adjusted = snap_widths(geoms, fidelity)
        adjusted.extend(apply_fidelity(geoms, fidelity, warnings))
        adjusted.extend(note_bands(geoms, fidelity))
        check_overlaps(geoms, errors)
        per_sheet.append((sheet, geoms, adjusted))

    if errors:
        return {"ok": False, "path": None, "sheets": [], "warnings": warnings, "errors": errors}

    ctx = Ctx(styler, registry, errors, warnings)
    wb = Workbook()
    wb.remove(wb.active)
    try:
        wb._named_styles["Normal"].font = Font(name=styler.base_name, size=styler.base_size)
    except Exception as exc:  # noqa: BLE001 - cosmetic only
        log("could not set the default font: %s" % exc)

    report = []
    moved = []                     # W11: every fidelity and width move, sheet by sheet
    snapshot = {"sheets": []}
    for sheet, geoms, adjusted in per_sheet:
        moved.extend(dict(m, sheet=str(sheet.get("name"))) for m in adjusted)
        ws = wb.create_sheet(str(sheet.get("name")))
        ws.sheet_view.showGridLines = styler.gridlines
        ctx.reset_counts()
        if not geoms and norm(sheet.get("name", "")).replace(" ", "") == "readme":
            if styler.cover_tab:
                write_readme(ctx, ws, spec)
            else:
                warnings.append("cover_tab is off for this style: '%s' left empty" % ws.title)
            colour_tab(ws, geoms, styler)
            sheet_setup(ws, sheet, warnings, computed_setup(ws, geoms, ctx.snap_widths_by_col()))
            report.append({"name": ws.title, "cells": ctx.cells,
                           "formulas": ctx.formulas, "charts": 0})
            snapshot["sheets"].append(ctx.snapshot_sheet(ws, geoms, adjusted))
            continue
        for g in geoms:
            WRITERS[g.type](ctx, ws, g)
        write_comments(ctx, ws, geoms, threads)
        merge_notes(ctx, ws)
        widths = finish_widths(ctx, ws)
        apply_wraps(ctx, ws, widths)
        finish_charts(ctx, ws, widths)
        colour_tab(ws, geoms, styler)
        check_unresolved(ws, errors)
        sheet_setup(ws, sheet, warnings, computed_setup(ws, geoms, widths))
        report.append({"name": ws.title, "cells": ctx.cells,
                       "formulas": ctx.formulas, "charts": ctx.charts})
        snapshot["sheets"].append(ctx.snapshot_sheet(ws, geoms, adjusted))

    for block_id in sorted(set(threads) - set(registry)):
        warnings.append("threads: no block '%s' in this workbook, its comment was not written"
                        % block_id)

    for name, target in sorted(ctx.names.items()):
        wb.defined_names.add(DefinedName(name, attr_text=target))

    if errors:
        return {"ok": False, "path": None, "sheets": report, "warnings": warnings, "errors": errors}

    wb.save(out_path)
    snap_path = snapshot_path(out_path)
    with open(snap_path, "w") as handle:
        json.dump(snapshot, handle, separators=(",", ":"))
    log("wrote %s (%d sheets) and %s" % (out_path, len(report), snap_path))
    result = {"ok": True, "path": out_path, "snapshot": snap_path, "sheets": report,
              "fidelity": fidelity, "warnings": warnings, "errors": []}
    if moved:
        result["adjusted"] = moved
    return result


def snapshot_path(out_path):
    """<outputPath minus .xlsx>.snapshot.json, next to the workbook."""
    base = str(out_path)
    if base.lower().endswith(".xlsx"):
        base = base[:-5]
    return base + ".snapshot.json"


def main(argv):
    if len(argv) != 2:
        sys.stdout.write(json.dumps(
            {"ok": False, "errors": ["usage: build_xlsx.py <spec.json>"]}))
        return 1
    try:
        with open(argv[1], "r") as handle:
            spec = json.load(handle)
    except Exception as exc:  # noqa: BLE001 - report any read problem as an error
        sys.stdout.write(json.dumps(
            {"ok": False, "errors": ["could not read spec: %s" % exc]}))
        return 1
    result = build(spec)
    sys.stdout.write(json.dumps(result))
    return 0 if result.get("ok") else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv))

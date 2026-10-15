#!/usr/bin/env python3
"""Tests for build_pptx.py and check_deck.py. Run: python3 builder/test_build_pptx.py"""

import copy
import json
import os
import shutil
import struct
import subprocess
import sys
import tempfile
import unittest
import zlib

from pptx import Presentation
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import build_pptx  # noqa: E402
import check_deck  # noqa: E402
import diagram_layout  # noqa: E402
import fit  # noqa: E402

from pptx.enum.dml import MSO_FILL  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

PT = 12700.0
HARBOR_TOKENS = os.path.join(HERE, "styles", "harbor_deck.json")


def make_png(path, width, height, color=(59, 91, 219)):
    """A valid RGB PNG of the exact size, without PIL."""
    raw = b""
    for _ in range(height):
        raw += b"\x00" + bytes(bytearray(color)) * width

    def chunk(tag, data):
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    blob = b"\x89PNG\r\n\x1a\n"
    blob += chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
    blob += chunk(b"IDAT", zlib.compress(raw))
    blob += chunk(b"IEND", b"")
    with open(path, "wb") as handle:
        handle.write(blob)


def make_template(path, width=960, height=540):
    """A .pptx template: the stock master and layouts, no slides."""
    prs = Presentation()
    prs.slide_width = build_pptx.emu(width)
    prs.slide_height = build_pptx.emu(height)
    prs.save(path)


def sample_spec(out_path, mode, template_path=None, image_path=None):
    """Five slides: title, divider, cards and stats, text and visuals, everything else."""
    style = {"mode": mode}
    if mode == "skill":
        style["tokensPath"] = HARBOR_TOKENS
    if mode == "template":
        style["templatePath"] = template_path
    return {
        "name": "Builder Test Deck",
        "outputPath": out_path,
        "exportPdf": False,
        "thumbnails": False,
        # This fixture pins the approved geometry and brings its own chrome:
        # the fidelity pass and the generated cover and agenda have their own
        # suites (FidelityTests, DeckTypeTests, ChromeTests).
        "fidelity": "strict",
        "cover": False,
        "agenda": False,
        "style": style,
        "slideSize": {"w": 960, "h": 540},
        "grid": {"cols": 32, "rows": 18, "cell": 30},
        "brief": {"audience": "The test suite", "goal": "Exercise every block",
                  "takeaway": "Every block type builds", "emphasis": ["cards"]},
        "slides": [
            {
                "id": "t1", "kind": "title", "name": "Cover",
                "headline": "A deck the tests can read",
                "subtitle": "Builder fixture", "date": "Draft",
                "blocks": [],
            },
            {
                "id": "t2", "name": "Section one", "section": "Section one",
                "subtitle": "Where the evidence starts.",
                "blocks": [],
            },
            {
                "id": "t3", "name": "Cards and stats", "kicker": "Evidence | Cards",
                "headline": "Three cards and three stats carry the argument",
                "notes": "Slide level notes.",
                "source": "Test fixture",
                "blocks": [
                    {"id": "b_cards", "type": "cards", "role": "card",
                     "col": 3, "row": 7, "cols": 28, "rows": 5, "emphasis": True,
                     "content": {"items": [
                         {"header": "First", "sub": "One", "body": "The first card body."},
                         {"header": "Second", "sub": "Two", "body": "The second card body."},
                         {"header": "Third", "sub": "Three", "body": "The third card body."},
                     ]}},
                    {"id": "b_stats", "type": "stat", "role": "stat",
                     "col": 3, "row": 13, "cols": 28, "rows": 4,
                     "content": {"items": [
                         {"value": "4.1x", "label": "Better"},
                         {"value": "38%", "label": "Of the total"},
                         {"value": "11", "label": "Days saved"},
                     ]}},
                ],
            },
            {
                "id": "t4", "name": "Text and visuals", "kicker": "Evidence | Visuals",
                "headline": "Text, bullets, a chart and a table all land on one slide",
                "blocks": [
                    {"id": "b_text", "type": "text", "role": "text",
                     "col": 3, "row": 7, "cols": 12, "rows": 5,
                     "content": {"paragraphs": ["A short paragraph.", "A second one."]}},
                    {"id": "b_bullets", "type": "bullets", "role": "list",
                     "col": 3, "row": 13, "cols": 12, "rows": 4,
                     "content": {"items": [
                         {"text": "First point", "children": [{"text": "A nested point"}]},
                         "Second point",
                     ]}},
                    {"id": "b_chart", "type": "chart", "role": "visual",
                     "col": 17, "row": 7, "cols": 14, "rows": 5,
                     "content": {"kind": "column", "categories": ["A", "B", "C"],
                                 "series": [{"name": "Now", "values": [3, 5, 2]},
                                            {"name": "Next", "values": [4, 6, 3]}]}},
                    {"id": "b_table", "type": "table", "role": "table",
                     "col": 17, "row": 13, "cols": 14, "rows": 4,
                     "content": {"headers": ["Move", "Owner"],
                                 "rows": [["Script the ask", "Partner lead"],
                                          ["Name the owner", "Account lead"]]}},
                ],
            },
            {
                "id": "t5", "name": "Everything else", "section": "The rest",
                "blocks": [
                    {"id": "b_kicker", "type": "kicker",
                     "col": 3, "row": 2, "cols": 20, "rows": 1,
                     "content": {"text": "Close | The Rest"}},
                    {"id": "b_headline", "type": "headline",
                     "col": 3, "row": 3, "cols": 28, "rows": 3,
                     "content": {"text": "The rest of the palette renders too"}},
                    {"id": "b_quote", "type": "quote", "role": "quote",
                     "col": 3, "row": 7, "cols": 14, "rows": 4,
                     "content": {"text": "Nobody ever asked us at an easy moment.",
                                 "attribution": "Partner interview"}},
                    {"id": "b_image", "type": "image", "role": "visual",
                     "col": 19, "row": 7, "cols": 12, "rows": 4,
                     "content": {"path": image_path, "placeholder": "Product shot"}},
                    {"id": "b_ph", "type": "placeholder", "role": "visual",
                     "col": 3, "row": 12, "cols": 14, "rows": 3,
                     "content": {"label": "Designed process strip"}},
                    {"id": "b_take", "type": "takeaway", "role": "statement",
                     "col": 3, "row": 16, "cols": 28, "rows": 2,
                     "content": {"text": "Every block type has a writer"}},
                    {"id": "b_notes", "type": "notes",
                     "col": 19, "row": 12, "cols": 12, "rows": 1,
                     "content": {"text": "Block level notes."}},
                ],
            },
        ],
    }


def block_map(spec):
    out = {}
    for slide in spec["slides"]:
        for block in slide.get("blocks") or []:
            out[block["id"]] = block
    return out


def shapes_for(snapshot, block_id):
    found = []
    for slide in snapshot["slides"]:
        for shape in slide["shapes"]:
            if shape.get("blockId") == block_id:
                found.append(shape)
    return found


def ovals(slide):
    """The oval shapes on a slide (table rating marks)."""
    out = []
    for shape in slide.shapes:
        try:
            kind = shape.auto_shape_type
        except (ValueError, AttributeError):
            continue
        if kind is not None and "OVAL" in str(kind):
            out.append(shape)
    return out


def rounded(slide):
    """The rounded rectangles on a slide (cards, takeaway bars, placeholders)."""
    out = []
    for shape in slide.shapes:
        try:
            kind = shape.auto_shape_type      # raises on a text box
        except (ValueError, AttributeError):
            continue
        if kind is not None and "ROUNDED" in str(kind):
            out.append(shape)
    return out


def read_snapshot(path):
    with open(build_pptx.snapshot_path(path)) as handle:
        return json.load(handle)


class DeckBuildTests(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="blueprint_deck_test_")
        cls.image = os.path.join(cls.tmp, "shot.png")
        make_png(cls.image, 200, 100)
        cls.template = os.path.join(cls.tmp, "template.pptx")
        make_template(cls.template)
        cls.results = {}
        cls.decks = {}
        cls.snapshots = {}
        cls.specs = {}
        for mode in ("plain", "house", "skill", "template"):
            path = os.path.join(cls.tmp, mode + ".pptx")
            spec = sample_spec(path, mode, cls.template, cls.image)
            cls.specs[mode] = spec
            result = build_pptx.build(copy.deepcopy(spec))
            cls.results[mode] = result
            if result["ok"]:
                cls.decks[mode] = Presentation(path)
                cls.snapshots[mode] = read_snapshot(path)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    # ------------------------------------------------------------- all modes

    def test_all_four_modes_build(self):
        for mode in ("plain", "house", "skill", "template"):
            result = self.results[mode]
            self.assertTrue(result["ok"], "%s failed: %s" % (mode, result.get("errors")))
            self.assertEqual(result["errors"], [])
            self.assertTrue(os.path.exists(result["path"]), mode)
            self.assertTrue(os.path.exists(result["snapshotPath"]), mode)

    def test_slide_count_and_result_shape(self):
        for mode, result in self.results.items():
            self.assertEqual(len(list(self.decks[mode].slides)), 5, mode)
            self.assertEqual([s["index"] for s in result["slides"]], [1, 2, 3, 4, 5], mode)
            for entry in result["slides"]:
                self.assertTrue(entry["title"], mode)
                self.assertGreater(entry["shapes"], 0, mode)

    def test_every_block_type_reaches_the_snapshot(self):
        wanted = set(block_map(self.specs["house"]))
        for mode, snapshot in self.snapshots.items():
            for block_id in wanted:
                if mode == "plain" and block_id == "b_kicker":
                    continue   # plain has no kicker furniture; it warns instead
                self.assertTrue(shapes_for(snapshot, block_id),
                                "%s lost block %s" % (mode, block_id))
        self.assertTrue(any("b_kicker" in w for w in self.results["plain"]["warnings"]))

    def test_snapshot_has_one_shape_per_simple_block(self):
        snapshot = self.snapshots["house"]
        for block_id in ("b_text", "b_bullets", "b_chart", "b_table", "b_quote",
                         "b_image", "b_ph", "b_take", "b_notes", "b_headline"):
            self.assertEqual(len(shapes_for(snapshot, block_id)), 1, block_id)
        self.assertEqual(len(shapes_for(snapshot, "b_cards")), 3)
        self.assertEqual(len(shapes_for(snapshot, "b_stats")), 3)

    def test_geometry_lands_within_a_point(self):
        blocks = block_map(self.specs["house"])
        snapshot = self.snapshots["house"]
        for block_id in ("b_text", "b_bullets", "b_chart", "b_table", "b_quote",
                         "b_ph", "b_take", "b_cards", "b_stats"):
            block = blocks[block_id]
            shape = shapes_for(snapshot, block_id)[0]
            self.assertAlmostEqual(shape["x"], (block["col"] - 1) * 30, delta=1.0)
            self.assertAlmostEqual(shape["y"], (block["row"] - 1) * 30, delta=1.0)
        for block_id in ("b_text", "b_chart", "b_quote", "b_ph", "b_take"):
            block = blocks[block_id]
            shape = shapes_for(snapshot, block_id)[0]
            self.assertAlmostEqual(shape["w"], block["cols"] * 30, delta=1.0)
            self.assertAlmostEqual(shape["h"], block["rows"] * 30, delta=1.0)
        # A table is as tall as its rows need (Q14), never stretched to the block.
        table = shapes_for(snapshot, "b_table")[0]
        self.assertAlmostEqual(table["w"], blocks["b_table"]["cols"] * 30, delta=1.0)
        self.assertLess(table["h"], blocks["b_table"]["rows"] * 30)

    def test_pptx_geometry_matches_the_snapshot(self):
        deck = self.decks["house"]
        snapshot = self.snapshots["house"]
        drawn = []
        for slide in deck.slides:
            for shape in slide.shapes:
                drawn.append((round(shape.left / PT, 1), round(shape.top / PT, 1)))
        for slide in snapshot["slides"]:
            for entry in slide["shapes"]:
                if entry["kind"] == "notes":
                    continue
                self.assertTrue(
                    any(abs(x - entry["x"]) <= 1 and abs(y - entry["y"]) <= 1 for x, y in drawn),
                    "%s at %s,%s is not in the deck" % (entry["kind"], entry["x"], entry["y"]))

    def test_snapshot_header(self):
        snapshot = self.snapshots["house"]
        self.assertEqual(snapshot["slideSize"], {"w": 960.0, "h": 540.0})
        self.assertEqual(snapshot["grid"], {"cols": 32, "rows": 18, "cell": 30.0})
        self.assertEqual(snapshot["style"]["mode"], "house")
        self.assertEqual([s["index"] for s in snapshot["slides"]], [1, 2, 3, 4, 5])
        self.assertEqual(snapshot["slides"][2]["headline"],
                         "Three cards and three stats carry the argument")

    def test_nothing_runs_off_the_slide(self):
        for mode, deck in self.decks.items():
            width, height = deck.slide_width / PT, deck.slide_height / PT
            for i, slide in enumerate(deck.slides):
                for shape in slide.shapes:
                    self.assertGreaterEqual(shape.left / PT, -0.5, "%s slide %d" % (mode, i + 1))
                    self.assertGreaterEqual(shape.top / PT, -0.5, "%s slide %d" % (mode, i + 1))
                    self.assertLessEqual((shape.left + shape.width) / PT, width + 0.5,
                                         "%s slide %d" % (mode, i + 1))
                    self.assertLessEqual((shape.top + shape.height) / PT, height + 0.5,
                                         "%s slide %d" % (mode, i + 1))

    # --------------------------------------------------------- slide kinds

    def test_divider_slide_renders_divider_furniture(self):
        snapshot = self.snapshots["house"]
        divider = snapshot["slides"][1]
        self.assertEqual(divider["kind"], "divider")
        roles = [s["role"] for s in divider["shapes"]]
        self.assertIn("dividerTitle", roles)
        self.assertIn("dividerSubtitle", roles)
        self.assertNotIn("kicker", roles)
        title = [s for s in divider["shapes"] if s["role"] == "dividerTitle"][0]
        self.assertEqual(title["text"], "01 Section one")
        self.assertGreaterEqual(title["fontPt"], 54)

    def test_section_row_without_blocks_is_a_divider(self):
        spec = sample_spec(os.path.join(self.tmp, "sectiononly.pptx"), "house")
        spec["slides"] = [dict(spec["slides"][1])]
        spec["slides"][0].pop("kind", None)
        result = build_pptx.build(spec)
        self.assertTrue(result["ok"], result.get("errors"))
        self.assertEqual(read_snapshot(result["path"])["slides"][0]["kind"], "divider")

    def test_title_slide_has_no_page_number(self):
        snapshot = self.snapshots["house"]
        roles = [s["role"] for s in snapshot["slides"][0]["shapes"]]
        self.assertNotIn("pageNumber", roles)
        self.assertIn("coverTitle", roles)
        self.assertIn("coverSubtitle", roles)

    def test_page_numbers_count_the_slides(self):
        snapshot = self.snapshots["house"]
        numbers = []
        for slide in snapshot["slides"]:
            for shape in slide["shapes"]:
                if shape["role"] == "pageNumber":
                    numbers.append(shape["text"])
        self.assertEqual(numbers, ["2", "3", "4", "5"])

    def test_headline_block_folds_into_the_slide_headline(self):
        snapshot = self.snapshots["house"]
        last = snapshot["slides"][4]
        headlines = [s for s in last["shapes"] if s["role"] == "headline"]
        self.assertEqual(len(headlines), 1)
        self.assertEqual(headlines[0]["blockId"], "b_headline")
        self.assertEqual(last["headline"], "The rest of the palette renders too")
        # furniture geometry, not the block's own row 3
        self.assertAlmostEqual(headlines[0]["y"], 66.0, delta=1.0)

    def test_a_headline_block_the_slide_already_has_is_dropped(self):
        spec = sample_spec(os.path.join(self.tmp, "twohead.pptx"), "house")
        spec["slides"] = [spec["slides"][4]]
        spec["slides"][0]["headline"] = "The slide already says it"
        result = build_pptx.build(spec)
        self.assertTrue(any("b_headline" in w and "repeats" in w for w in result["warnings"]),
                        result["warnings"])
        snapshot = read_snapshot(result["path"])
        headlines = [s for s in snapshot["slides"][0]["shapes"] if s["role"] == "headline"]
        self.assertEqual(len(headlines), 1)
        self.assertEqual(headlines[0]["text"], "The slide already says it")

    def test_speaker_notes_from_slide_and_block(self):
        deck = self.decks["house"]
        third = list(deck.slides)[2]
        self.assertIn("Slide level notes.", third.notes_slide.notes_text_frame.text)
        fifth = list(deck.slides)[4]
        self.assertIn("Block level notes.", fifth.notes_slide.notes_text_frame.text)

    def test_notes_block_draws_nothing(self):
        snapshot = self.snapshots["house"]
        entry = shapes_for(snapshot, "b_notes")[0]
        self.assertEqual(entry["kind"], "notes")
        deck = self.decks["house"]
        fifth = list(deck.slides)[4]
        for shape in fifth.shapes:
            if shape.has_text_frame:
                self.assertNotIn("Block level notes.", shape.text_frame.text)

    # -------------------------------------------------------------- styles

    def test_plain_drops_the_kicker_with_a_warning(self):
        result = self.results["plain"]
        self.assertTrue(any("no kicker" in w for w in result["warnings"]), result["warnings"])
        roles = []
        for slide in self.snapshots["plain"]["slides"]:
            roles.extend(s["role"] for s in slide["shapes"])
        self.assertNotIn("kicker", roles)
        self.assertIn("headline", roles)
        self.assertIn("pageNumber", roles)

    def test_house_cards_are_rounded_with_a_one_point_border(self):
        deck = self.decks["house"]
        third = list(deck.slides)[2]
        cards = rounded(third)
        self.assertEqual(len(cards), 3)
        for card in cards:
            # b_cards carries `emphasis: true`, so the border is the accent at
            # borderPt + 0.75 (Q13); the fill is still the card fill.
            self.assertAlmostEqual(card.line.width / PT, 1.75, delta=0.05)
            self.assertEqual(str(card.line.color.rgb), "3B5BDB")
            self.assertEqual(str(card.fill.fore_color.rgb), "FFFFFF")

    def test_skill_tokens_drive_the_look(self):
        deck = self.decks["skill"]
        third = list(deck.slides)[2]
        cards = rounded(third)
        self.assertEqual(len(cards), 3)
        for card in cards:
            self.assertAlmostEqual(card.line.width / PT, 3.0, delta=0.05)
            self.assertEqual(str(card.line.color.rgb), "1E3A5F")
        snapshot = self.snapshots["skill"]
        kicker = [s for s in snapshot["slides"][2]["shapes"] if s["role"] == "kicker"][0]
        self.assertEqual(kicker["fontPt"], 16.0)
        self.assertAlmostEqual(kicker["x"], 32.0, delta=0.5)
        page = [s for s in snapshot["slides"][2]["shapes"] if s["role"] == "pageNumber"][0]
        self.assertLess(page["y"], 60, "the Harbor page number sits top right")

    def test_skill_backgrounds_per_slide_kind(self):
        deck = self.decks["skill"]
        slides = list(deck.slides)
        self.assertEqual(str(slides[0].background.fill.fore_color.rgb), "1E3A5F")
        self.assertEqual(str(slides[1].background.fill.fore_color.rgb), "2A9D8F")
        self.assertEqual(str(slides[2].background.fill.fore_color.rgb), "FFFFFF")

    def test_missing_font_falls_back(self):
        warnings = self.results["skill"]["warnings"]
        installed = build_pptx.FONTS.has("Lato")
        if installed:
            self.assertFalse(any("Lato" in w for w in warnings))
        else:
            self.assertTrue(any("Lato" in w and "Arial" in w for w in warnings), warnings)
            deck = self.decks["skill"]
            third = list(deck.slides)[2]
            names = set()
            for shape in third.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            if run.font.name:
                                names.add(run.font.name)
            self.assertNotIn("Lato", names)

    def test_template_mode_uses_the_template_layouts(self):
        deck = self.decks["template"]
        template = Presentation(self.template)
        self.assertEqual([lay.name for lay in deck.slide_layouts],
                         [lay.name for lay in template.slide_layouts])
        self.assertEqual(len(list(deck.slides)), 5)

    def test_template_headline_uses_the_title_placeholder(self):
        """The measure comes from the layout a content slide takes, not the cover (Q8)."""
        spec = sample_spec(os.path.join(self.tmp, "tpl2.pptx"), "template", self.template)
        warnings = []
        style, path = build_pptx.make_style(spec, warnings)
        prs, spec = build_pptx.open_presentation(spec, style, path, warnings)
        profiles = build_pptx.layout_profiles(prs)
        body = next(p for p in profiles if p["kind"] == "titleBody")
        titles = [ph for ph in body["layout"].placeholders
                  if "TITLE" in str(ph.placeholder_format.type)]
        self.assertTrue(titles)
        slot = style.slot("headline", "content")
        self.assertAlmostEqual(slot["x"], titles[0].left / PT, delta=1.0)
        self.assertAlmostEqual(slot["y"], titles[0].top / PT, delta=1.0)

    def test_template_theme_fonts_are_adopted(self):
        spec = sample_spec(os.path.join(self.tmp, "tpl3.pptx"), "template", self.template)
        warnings = []
        style, path = build_pptx.make_style(spec, warnings)
        build_pptx.open_presentation(spec, style, path, warnings)
        self.assertTrue(style.tokens["fonts"]["body"].startswith("Calibri"))

    def test_template_with_a_different_slide_size_spreads_the_grid(self):
        narrow = os.path.join(self.tmp, "narrow.pptx")
        make_template(narrow, 720, 540)
        spec = sample_spec(os.path.join(self.tmp, "narrow_out.pptx"), "template", narrow)
        result = build_pptx.build(spec)
        self.assertTrue(result["ok"], result.get("errors"))
        self.assertTrue(any("720 x 540" in w for w in result["warnings"]), result["warnings"])
        snapshot = read_snapshot(result["path"])
        self.assertEqual(snapshot["slideSize"]["w"], 720.0)
        table = shapes_for(snapshot, "b_table")[0]
        self.assertAlmostEqual(table["x"], 16 * (720.0 / 32), delta=1.0)

    def test_unknown_template_falls_back_to_house(self):
        spec = sample_spec(os.path.join(self.tmp, "notpl.pptx"), "template", "/nope/none.pptx")
        result = build_pptx.build(spec)
        self.assertTrue(result["ok"])
        self.assertTrue(any("template not found" in w for w in result["warnings"]))
        self.assertEqual(read_snapshot(result["path"])["style"]["mode"], "house")

    def test_unknown_style_mode_falls_back_to_house(self):
        spec = sample_spec(os.path.join(self.tmp, "weird.pptx"), "house")
        spec["style"] = {"mode": "neon"}
        result = build_pptx.build(spec)
        self.assertTrue(result["ok"])
        self.assertTrue(any("unknown style mode" in w for w in result["warnings"]))

    # ------------------------------------------------------- fit and limits

    def test_font_steps_down_within_the_range(self):
        spec = sample_spec(os.path.join(self.tmp, "longhead.pptx"), "house")
        spec["slides"] = [spec["slides"][3]]
        spec["slides"][0]["headline"] = (
            "This headline is far too long for two lines at the default size and has to "
            "step down through the range before the builder gives up on it entirely")
        result = build_pptx.build(spec)
        snapshot = read_snapshot(result["path"])
        headline = [s for s in snapshot["slides"][0]["shapes"] if s["role"] == "headline"][0]
        floor = build_pptx.DEFAULT_TYPE_SCALE["headline"]["min"]
        self.assertLess(headline["fontPt"], 28.0)
        self.assertGreaterEqual(headline["fontPt"], floor)

    def test_overflow_warning_fires_at_the_floor(self):
        spec = sample_spec(os.path.join(self.tmp, "overflow.pptx"), "house")
        spec["slides"] = [spec["slides"][3]]
        spec["slides"][0]["blocks"] = [{
            "id": "b_over", "type": "text", "col": 3, "row": 7, "cols": 6, "rows": 2,
            "content": {"paragraphs": [" ".join(["overflowing copy"] * 40)]},
        }]
        result = build_pptx.build(spec)
        self.assertTrue(result["ok"])
        self.assertTrue(any("b_over" in w and "overflows" in w for w in result["warnings"]),
                        result["warnings"])
        snapshot = read_snapshot(result["path"])
        entry = shapes_for(snapshot, "b_over")[0]
        self.assertEqual(entry["fontPt"], float(build_pptx.DEFAULT_TYPE_SCALE["body"]["min"]))

    def test_off_slide_block_is_clamped_and_warned(self):
        spec = sample_spec(os.path.join(self.tmp, "clamp.pptx"), "house")
        spec["slides"] = [spec["slides"][3]]
        spec["slides"][0]["blocks"] = [
            {"id": "b_wide", "type": "text", "col": 28, "row": 7, "cols": 12, "rows": 4,
             "content": {"paragraphs": ["Runs off the right edge."]}},
            {"id": "b_tall", "type": "text", "col": 3, "row": 16, "cols": 8, "rows": 6,
             "content": {"paragraphs": ["Runs off the bottom."]}},
        ]
        result = build_pptx.build(spec)
        self.assertTrue(any("b_wide" in w and "runs off the slide" in w
                            for w in result["warnings"]), result["warnings"])
        snapshot = read_snapshot(result["path"])
        wide = shapes_for(snapshot, "b_wide")[0]
        tall = shapes_for(snapshot, "b_tall")[0]
        self.assertAlmostEqual(wide["x"] + wide["w"], 960.0, delta=0.5)
        self.assertAlmostEqual(tall["y"] + tall["h"], 540.0, delta=0.5)

    def test_empty_text_becomes_muted_placeholder_copy(self):
        spec = sample_spec(os.path.join(self.tmp, "empty.pptx"), "house")
        spec["slides"] = [spec["slides"][3]]
        spec["slides"][0]["blocks"] = [{
            "id": "b_empty", "type": "text", "title": "Evidence",
            "col": 3, "row": 7, "cols": 12, "rows": 4, "content": {"paragraphs": []},
        }]
        result = build_pptx.build(spec)
        self.assertTrue(any("b_empty" in w and "no text" in w for w in result["warnings"]),
                        result["warnings"])
        deck = Presentation(result["path"])
        found = None
        for shape in list(deck.slides)[0].shapes:
            if shape.has_text_frame and "[Evidence]" in shape.text_frame.text:
                found = shape
        self.assertIsNotNone(found, "no placeholder copy")
        run = found.text_frame.paragraphs[-1].runs[0]
        self.assertTrue(run.font.italic)
        self.assertEqual(str(run.font.color.rgb), "6B7280")

    def test_fit_never_goes_below_the_floor(self):
        spec = build_pptx.DEFAULT_TYPE_SCALE["cardBody"]
        para = build_pptx.Para("word " * 400, "cardBody", dict(build_pptx.ROLE_DEFAULTS, **spec))
        overflow, _ = build_pptx.fit_paragraphs([para], 120, 40)
        self.assertTrue(overflow)
        self.assertEqual(para.size, float(spec["min"]))

    def test_widow_is_pulled_up_when_a_step_helps(self):
        spec = dict(build_pptx.ROLE_DEFAULTS, **build_pptx.DEFAULT_TYPE_SCALE["headline"])
        para = build_pptx.Para("Referrals stall at three moments and each has its own fix",
                               "headline", spec)
        overflow, widow = build_pptx.fit_paragraphs([para], 620, 90)
        self.assertFalse(overflow)
        self.assertFalse(widow)
        self.assertGreaterEqual(para.size, spec["min"])

    def test_image_keeps_its_aspect_ratio(self):
        entry = shapes_for(self.snapshots["house"], "b_image")[0]
        self.assertEqual(entry["kind"], "image")
        self.assertAlmostEqual(entry["w"] / entry["h"], 2.0, delta=0.05)

    def test_missing_image_becomes_a_placeholder(self):
        spec = sample_spec(os.path.join(self.tmp, "noimg.pptx"), "house", None, "/nope/x.png")
        spec["slides"] = [spec["slides"][4]]
        result = build_pptx.build(spec)
        self.assertTrue(any("image not found" in w for w in result["warnings"]),
                        result["warnings"])
        entry = shapes_for(read_snapshot(result["path"]), "b_image")[0]
        self.assertEqual(entry["kind"], "placeholder")

    def test_charts_and_tables_are_native(self):
        deck = self.decks["house"]
        fourth = list(deck.slides)[3]
        charts = [s for s in fourth.shapes if getattr(s, "has_chart", False)]
        tables = [s for s in fourth.shapes if getattr(s, "has_table", False)]
        self.assertEqual(len(charts), 1)
        self.assertEqual(len(tables), 1)
        self.assertEqual(len(charts[0].chart.plots[0].series), 2)
        self.assertEqual(list(charts[0].chart.plots[0].categories), ["A", "B", "C"])
        table = tables[0].table
        self.assertEqual(len(table.rows), 3)
        self.assertEqual(table.cell(0, 0).text, "MOVE")      # caps heads (Q14)
        self.assertEqual(table.cell(2, 1).text, "Account lead")

    def test_bullets_nest(self):
        deck = self.decks["house"]
        fourth = list(deck.slides)[3]
        target = None
        for shape in fourth.shapes:
            if shape.has_text_frame and "First point" in shape.text_frame.text:
                target = shape
        self.assertIsNotNone(target)
        levels = [p.level for p in target.text_frame.paragraphs]
        self.assertIn(1, levels, "the child bullet keeps its level")

    # --------------------------------------------------------------- errors

    def test_missing_output_path_is_an_error(self):
        result = build_pptx.build({"slides": []})
        self.assertFalse(result["ok"])
        self.assertIn("outputPath", result["errors"][0])

    def test_duplicate_block_ids_are_an_error(self):
        spec = sample_spec(os.path.join(self.tmp, "dupe.pptx"), "house")
        spec["slides"][3]["blocks"][1]["id"] = "b_text"
        result = build_pptx.build(spec)
        self.assertFalse(result["ok"])
        self.assertTrue(any("duplicate block id" in e for e in result["errors"]))

    def test_cli_prints_one_json_object(self):
        spec_path = os.path.join(self.tmp, "cli_spec.json")
        out = os.path.join(self.tmp, "cli.pptx")
        with open(spec_path, "w") as handle:
            json.dump(sample_spec(out, "house", self.template, self.image), handle)
        proc = subprocess.run([sys.executable, os.path.join(HERE, "build_pptx.py"), spec_path],
                              stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=120)
        self.assertEqual(proc.returncode, 0, proc.stderr.decode())
        payload = json.loads(proc.stdout.decode())
        self.assertTrue(payload["ok"])
        self.assertEqual(payload["path"], out)
        self.assertTrue(os.path.exists(out))


class RenderTests(unittest.TestCase):
    """PDF and PNG steps, with and without LibreOffice."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="blueprint_render_test_")
        self.spec = sample_spec(os.path.join(self.tmp, "render.pptx"), "house")
        self.spec["slides"] = self.spec["slides"][:2]
        self.spec["exportPdf"] = True
        self.spec["thumbnails"] = True

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_pdf_and_png_are_skipped_without_soffice(self):
        saved_paths = build_pptx.SOFFICE_PATHS
        saved_env = os.environ.get("PATH")
        build_pptx.SOFFICE_PATHS = []
        os.environ["PATH"] = self.tmp
        try:
            result = build_pptx.build(copy.deepcopy(self.spec))
        finally:
            build_pptx.SOFFICE_PATHS = saved_paths
            if saved_env is not None:
                os.environ["PATH"] = saved_env
        self.assertTrue(result["ok"], result.get("errors"))
        self.assertNotIn("pdfPath", result)
        self.assertNotIn("thumbnails", result)
        self.assertTrue(any("LibreOffice not found" in w for w in result["warnings"]),
                        result["warnings"])
        self.assertTrue(os.path.exists(result["path"]))

    @unittest.skipIf(build_pptx.find_soffice() is None, "LibreOffice is not installed")
    def test_pdf_and_png_are_written_when_soffice_is_present(self):
        result = build_pptx.build(copy.deepcopy(self.spec))
        self.assertTrue(result["ok"], result.get("errors"))
        self.assertIn("pdfPath", result)
        self.assertTrue(os.path.exists(result["pdfPath"]))
        self.assertIn("thumbnails", result)
        self.assertEqual(len(result["thumbnails"]), 2)
        for path in result["thumbnails"]:
            self.assertTrue(os.path.exists(path))
            self.assertGreater(os.path.getsize(path), 1000)
        self.assertTrue(os.path.isdir(build_pptx.slides_dir_for(result["path"])))


class CheckDeckTests(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="blueprint_check_test_")
        cls.image = os.path.join(cls.tmp, "shot.png")
        make_png(cls.image, 200, 100)
        cls.path = os.path.join(cls.tmp, "clean.pptx")
        cls.spec = sample_spec(cls.path, "house", None, cls.image)
        cls.result = build_pptx.build(copy.deepcopy(cls.spec))
        cls.spec_path = os.path.join(cls.tmp, "clean_spec.json")
        with open(cls.spec_path, "w") as handle:
            json.dump(cls.spec, handle)
        cls.report = check_deck.check_deck(cls.path, cls.spec_path)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_a_clean_deck_has_no_errors(self):
        errors = [f for f in self.report["findings"] if f["severity"] == "error"]
        self.assertEqual(errors, [])
        self.assertEqual(self.report["ok"], True)
        self.assertEqual(self.report["slides"], 5)
        self.assertGreater(self.report["shapes"], 20)
        self.assertEqual(self.report["summary"]["errors"], 0)

    def test_findings_carry_slide_kind_and_severity(self):
        for finding in self.report["findings"]:
            self.assertIn(finding["kind"], check_deck.SEVERITY)
            self.assertEqual(finding["severity"], check_deck.SEVERITY[finding["kind"]])
            self.assertIn("slide", finding)
            self.assertIn("message", finding)
        self.assertEqual(self.report["summary"]["errors"]
                         + self.report["summary"]["warnings"],
                         len(self.report["findings"]))

    def test_planted_overlap_is_found(self):
        spec = sample_spec(os.path.join(self.tmp, "overlap.pptx"), "house", None, self.image)
        spec["slides"] = [spec["slides"][3]]
        spec["slides"][0]["blocks"] = [
            {"id": "b_a", "type": "text", "col": 3, "row": 7, "cols": 12, "rows": 6,
             "content": {"paragraphs": ["Left block."]}},
            {"id": "b_b", "type": "text", "col": 5, "row": 8, "cols": 12, "rows": 6,
             "content": {"paragraphs": ["Overlapping block."]}},
        ]
        result = build_pptx.build(spec)
        report = check_deck.check_deck(result["path"])
        overlaps = [f for f in report["findings"] if f["kind"] == "overlap"]
        self.assertTrue(overlaps, report["findings"])
        self.assertEqual(overlaps[0]["severity"], "error")
        self.assertIn("b_", overlaps[0]["message"])

    def test_off_slide_shape_is_found(self):
        path = os.path.join(self.tmp, "offslide.pptx")
        shutil.copyfile(self.path, path)
        deck = Presentation(path)
        moved = list(list(deck.slides)[3].shapes)[-1]
        moved.left = build_pptx.emu(900)
        moved.top = build_pptx.emu(520)
        deck.save(path)
        report = check_deck.check_deck(path)
        off = [f for f in report["findings"] if f["kind"] == "off-slide"]
        self.assertTrue(off, report["findings"])
        self.assertEqual(off[0]["severity"], "error")
        self.assertEqual(off[0]["slide"], 4)

    def test_missing_headline_is_found(self):
        spec = sample_spec(os.path.join(self.tmp, "nohead.pptx"), "house", None, self.image)
        spec["slides"] = [spec["slides"][3]]
        spec["slides"][0].pop("headline")
        result = build_pptx.build(spec)
        self.assertTrue(any("no headline" in w for w in result["warnings"]))
        report = check_deck.check_deck(result["path"])
        kinds = [f["kind"] for f in report["findings"]]
        self.assertIn("missing-headline", kinds)

    def test_overflow_and_font_floor_are_found(self):
        spec = sample_spec(os.path.join(self.tmp, "over.pptx"), "house", None, self.image)
        spec["slides"] = [spec["slides"][3]]
        spec["slides"][0]["blocks"] = [{
            "id": "b_over", "type": "text", "col": 3, "row": 7, "cols": 5, "rows": 2,
            "content": {"paragraphs": [" ".join(["far too much copy"] * 30)]},
        }]
        result = build_pptx.build(spec)
        report = check_deck.check_deck(result["path"], None)
        kinds = set(f["kind"] for f in report["findings"])
        self.assertIn("overflow", kinds)
        self.assertIn("font-floor", kinds)
        over = [f for f in report["findings"] if f["kind"] == "overflow"][0]
        self.assertEqual(over["blockId"], "b_over")
        self.assertEqual(over["severity"], "warning")

    def test_empty_box_is_found(self):
        spec = sample_spec(os.path.join(self.tmp, "blank.pptx"), "house", None, self.image)
        spec["slides"] = [spec["slides"][3]]
        spec["slides"][0]["blocks"] = [{
            "id": "b_blank", "type": "text", "title": "Evidence",
            "col": 3, "row": 7, "cols": 12, "rows": 4, "content": {},
        }]
        result = build_pptx.build(spec)
        report = check_deck.check_deck(result["path"])
        empty = [f for f in report["findings"] if f["kind"] == "empty"]
        self.assertTrue(empty, report["findings"])
        self.assertEqual(empty[0]["blockId"], "b_blank")

    def test_widow_is_found(self):
        spec = sample_spec(os.path.join(self.tmp, "widow.pptx"), "house", None, self.image)
        spec["slides"] = [spec["slides"][3]]
        spec["slides"][0]["blocks"] = [{
            "id": "b_widow", "type": "text", "col": 3, "row": 7, "cols": 10, "rows": 4,
            "content": {"paragraphs":
                        ["Referrals stall at the ask and nowhere else in the relationship"]},
        }]
        result = build_pptx.build(spec)
        report = check_deck.check_deck(result["path"])
        widows = [f for f in report["findings"] if f["kind"] == "widow"]
        self.assertTrue(widows, report["findings"])
        self.assertEqual(widows[0]["severity"], "warning")

    def test_cli_writes_json_and_exits_zero(self):
        out = os.path.join(self.tmp, "findings.json")
        proc = subprocess.run([sys.executable, os.path.join(HERE, "check_deck.py"),
                               self.path, out, "--spec", self.spec_path],
                              stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=120)
        self.assertEqual(proc.returncode, 0, proc.stderr.decode())
        payload = json.loads(proc.stdout.decode())
        self.assertTrue(payload["ok"])
        self.assertTrue(os.path.exists(out))
        with open(out) as handle:
            self.assertEqual(json.load(handle)["slides"], 5)


class SampleDeckTests(unittest.TestCase):
    """The shipped builder/sample_deck_spec.json, in house and in the Harbor skill style."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="blueprint_sample_test_")
        cls.spec_path = os.path.join(HERE, "sample_deck_spec.json")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def load(self, mode):
        with open(self.spec_path) as handle:
            spec = json.load(handle)
        spec["outputPath"] = os.path.join(self.tmp, "sample_%s.pptx" % mode)
        spec["thumbnails"] = False
        spec["exportPdf"] = False
        if mode == "harbor":
            spec["style"] = {"mode": "skill", "tokensPath": HARBOR_TOKENS}
        return spec

    def test_sample_builds_in_house_with_no_warnings(self):
        spec = self.load("house")
        result = build_pptx.build(spec)
        self.assertTrue(result["ok"], result.get("errors"))
        self.assertEqual(result["warnings"], [])
        self.assertEqual(len(result["slides"]), 9)
        deck = Presentation(result["path"])
        self.assertEqual(len(list(deck.slides)), 9)

    def test_sample_builds_in_the_harbor_style(self):
        spec = self.load("harbor")
        result = build_pptx.build(spec)
        self.assertTrue(result["ok"], result.get("errors"))
        for warning in result["warnings"]:
            self.assertIn("is not installed", warning)
        snapshot = read_snapshot(result["path"])
        kinds = [s["kind"] for s in snapshot["slides"]]
        self.assertEqual(kinds[0], "title")
        self.assertEqual(kinds[1], "agenda")
        self.assertEqual(kinds[5], "divider")

    def test_sample_passes_the_deck_check(self):
        spec = self.load("house")
        result = build_pptx.build(spec)
        spec_path = os.path.join(self.tmp, "sample_spec.json")
        with open(spec_path, "w") as handle:
            json.dump(spec, handle)
        report = check_deck.check_deck(result["path"], spec_path)
        errors = [f for f in report["findings"] if f["severity"] == "error"]
        self.assertEqual(errors, [], "sample deck has errors: %s" % errors)
        self.assertEqual(report["slides"], 9)


DARK_TOKENS = {
    "name": "Dark",
    "palette": {"bg": "000000", "ink": "FFFFFF", "muted": "FFFFFF", "accent": "D9A066",
                "cardFill": "2D2D2D", "cardBorder": "262626"},
    "fonts": {"display": "Montserrat Light", "body": "Calibri"},
    "furniture": {"headline": {"x": 336, "y": 63, "w": 593, "h": 158},
                  "kicker": {"x": 336, "y": 29, "w": 281, "h": 16},
                  "source": {"x": 336, "y": 499, "w": 593, "h": 22},
                  "pageNumber": {"x": 785, "y": 29, "w": 144, "h": 16}},
}


def diagram_spec(out_path, kind, nodes, edges=None, mode="house", tokens=None,
                 col=3, row=7, cols=28, rows=7, direction=None):
    """One content slide carrying one diagram block."""
    content = {"kind": kind, "nodes": nodes}
    if edges is not None:
        content["edges"] = edges
    if direction:
        content["direction"] = direction
    style = {"mode": mode}
    if tokens:
        style["tokens"] = tokens
    return {
        "name": "Diagram", "outputPath": out_path, "thumbnails": False, "exportPdf": False,
        "fidelity": "strict", "cover": False, "agenda": False,
        "style": style,
        "slideSize": {"w": 960, "h": 540}, "grid": {"cols": 32, "rows": 18, "cell": 30},
        "slides": [{
            "id": "d1", "name": "The flow", "section": "Diagrams",
            "headline": "the diagram carries the argument",
            "blocks": [{"id": "b_diag", "type": "diagram", "role": "visual",
                        "col": col, "row": row, "cols": cols, "rows": rows,
                        "content": content}],
        }],
    }


def nodes_of(*labels):
    return [{"id": "n%d" % (i + 1), "label": label} for i, label in enumerate(labels)]


def tail_ends(slide):
    """Every shape on the slide that carries an arrowhead."""
    found = []
    for shape in slide.shapes:
        try:
            ln = shape.line._get_or_add_ln()
        except (AttributeError, ValueError):
            continue
        if ln.findall(qn("a:tailEnd")):
            found.append(shape)
    return found


class DiagramBlockTests(unittest.TestCase):
    """F18: the diagram block draws nodes, arrowed connectors and labels."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="blueprint_diagram_test_")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def build(self, name, **kwargs):
        spec = diagram_spec(os.path.join(self.tmp, name + ".pptx"), **kwargs)
        result = build_pptx.build(copy.deepcopy(spec))
        self.assertTrue(result["ok"], result.get("errors"))
        return spec, result, read_snapshot(result["path"])

    def test_every_kind_builds_without_errors(self):
        for kind in diagram_layout.KINDS:
            labels = ["Spot the moment", "Script the ask", "Hand off", "Follow up",
                      "Review", "Renew"]
            ids = ["n%d" % (i + 1) for i in range(len(labels))]
            edges = [{"from": a, "to": b} for a, b in zip(ids, ids[1:])]
            _, result, snapshot = self.build("kind_" + kind, kind=kind,
                                             nodes=nodes_of(*labels), edges=edges)
            shapes = shapes_for(snapshot, "b_diag")
            self.assertTrue(shapes, kind)
            for entry in shapes:
                self.assertEqual(entry["blockId"], "b_diag")
                self.assertLessEqual(entry["x"] + entry["w"], 960.5, kind)
                self.assertLessEqual(entry["y"] + entry["h"], 540.5, kind)

    def test_one_shape_per_node_and_one_per_edge(self):
        edges = [{"from": "n1", "to": "n2", "label": "prompts"},
                 {"from": "n2", "to": "n3", "label": "warms"}]
        _, result, snapshot = self.build("counts", kind="flow",
                                         nodes=nodes_of("Spot", "Ask", "Hand off"),
                                         edges=edges)
        shapes = shapes_for(snapshot, "b_diag")
        self.assertEqual(len([s for s in shapes if s["kind"] == "node"]), 3)
        self.assertEqual(len([s for s in shapes if s["kind"] == "edge"]), 2)
        self.assertEqual(result["warnings"], [])
        texts = [s["text"] for s in shapes if s["kind"] == "node"]
        self.assertEqual(texts, ["Spot", "Ask", "Hand off"])

    def test_connectors_carry_a_triangle_arrowhead(self):
        _, result, _ = self.build("arrows", kind="flow",
                                  nodes=nodes_of("Spot", "Ask", "Hand off"),
                                  edges=[{"from": "n1", "to": "n2"},
                                         {"from": "n1", "to": "n3"}])
        deck = Presentation(result["path"])
        arrows = tail_ends(list(deck.slides)[0])
        self.assertEqual(len(arrows), 2)
        for shape in arrows:
            ln = shape.line._get_or_add_ln()
            self.assertEqual(ln.findall(qn("a:tailEnd"))[0].get("type"), "triangle")
            self.assertFalse(shape._element.findall(qn("p:style")))

    def test_nodes_are_card_shapes_from_the_tokens(self):
        _, result, _ = self.build("cards", kind="flow", nodes=nodes_of("Spot", "Ask"),
                                  edges=[{"from": "n1", "to": "n2"}])
        deck = Presentation(result["path"])
        boxes = rounded(list(deck.slides)[0])
        self.assertEqual(len(boxes), 2)
        for box in boxes:
            self.assertAlmostEqual(box.line.width / PT, 1.0, delta=0.05)
            self.assertEqual(str(box.line.color.rgb), "D1D5DB")
            self.assertEqual(str(box.fill.fore_color.rgb), "FFFFFF")

    def test_edge_labels_are_caption_sized_and_clear_of_the_line(self):
        _, result, snapshot = self.build("edgelabels", kind="flow",
                                         nodes=nodes_of("Spot", "Ask"),
                                         edges=[{"from": "n1", "to": "n2",
                                                 "label": "prompts"}])
        shapes = shapes_for(snapshot, "b_diag")
        label = [s for s in shapes if s["kind"] == "text"][0]
        line = [s for s in shapes if s["kind"] == "edge"][0]
        self.assertEqual(label["text"], "prompts")
        self.assertEqual(label["fontPt"],
                         float(build_pptx.DEFAULT_TYPE_SCALE["caption"]["size"]))
        self.assertLess(label["y"] + label["h"], line["y"] + 1.0)

    def test_matrix_axes_are_text_and_the_cells_are_cards(self):
        _, result, snapshot = self.build("matrix", kind="matrix",
                                         nodes=nodes_of("Effort", "Value", "Quick wins",
                                                        "Big bets", "Fill ins",
                                                        "Money pits"),
                                         edges=[])
        shapes = shapes_for(snapshot, "b_diag")
        nodes = [s for s in shapes if s["kind"] == "node"]
        self.assertEqual(len(nodes), 6)
        labels = [s for s in nodes if s["role"] == "diagramLabel"]
        cells = [s for s in nodes if s["role"] == "diagramNode"]
        self.assertEqual(len(labels), 2)
        self.assertEqual(len(cells), 4)
        deck = Presentation(result["path"])
        self.assertEqual(len(rounded(list(deck.slides)[0])), 4)
        self.assertAlmostEqual(cells[0]["w"], cells[0]["h"], delta=1.0)

    def test_timeline_draws_a_spine_and_alternates_the_boxes(self):
        _, result, snapshot = self.build("timeline", kind="timeline",
                                         nodes=nodes_of("Q1 pilot", "Q2 rollout",
                                                        "Q3 review"),
                                         edges=[])
        shapes = shapes_for(snapshot, "b_diag")
        spine = [s for s in shapes if s["role"] == "diagramAxis"]
        self.assertEqual(len(spine), 1)
        self.assertEqual(spine[0]["kind"], "rule")
        boxes = [s for s in shapes if s["role"] == "diagramNode"]
        self.assertEqual(len(boxes), 3)
        self.assertLess(boxes[0]["y"], boxes[1]["y"])
        self.assertLess(boxes[2]["y"], boxes[1]["y"])

    def test_placeholder_mode_labels_the_gap_instead_of_drawing_it(self):
        spec = diagram_spec(os.path.join(self.tmp, "ph.pptx"), "flow",
                            nodes_of("Prepare", "Ask", "Hand off"),
                            edges=[{"from": "n1", "to": "n2"}],
                            mode="skill", tokens={"diagrams": "placeholder"})
        result = build_pptx.build(copy.deepcopy(spec))
        self.assertTrue(result["ok"], result.get("errors"))
        snapshot = read_snapshot(result["path"])
        shapes = shapes_for(snapshot, "b_diag")
        self.assertEqual(len(shapes), 1)
        self.assertEqual(shapes[0]["kind"], "placeholder")
        self.assertIn("flow", shapes[0]["text"])
        self.assertIn("Prepare", shapes[0]["text"])
        self.assertIn("Hand off", shapes[0]["text"])
        deck = Presentation(result["path"])
        boxes = rounded(list(deck.slides)[0])
        self.assertEqual(len(boxes), 1)
        self.assertEqual(str(boxes[0].line.dash_style), "DASH (4)")

    def test_the_harbor_style_never_draws_shape_diagrams(self):
        spec = diagram_spec(os.path.join(self.tmp, "harbor.pptx"), "cycle",
                            nodes_of("Plan", "Build", "Learn"), edges=[], mode="skill")
        spec["style"] = {"mode": "skill", "tokensPath": HARBOR_TOKENS}
        result = build_pptx.build(copy.deepcopy(spec))
        snapshot = read_snapshot(result["path"])
        kinds = [s["kind"] for s in shapes_for(snapshot, "b_diag")]
        self.assertEqual(kinds, ["placeholder"])

    def test_over_the_cap_is_truncated_with_a_warning(self):
        labels = ["Step %d" % (i + 1) for i in range(14)]
        _, result, snapshot = self.build("cap", kind="flow", nodes=nodes_of(*labels),
                                         edges=[])
        nodes = [s for s in shapes_for(snapshot, "b_diag") if s["kind"] == "node"]
        self.assertEqual(len(nodes), diagram_layout.NODE_CAP)
        self.assertTrue(any("cap" in w and "b_diag" in w for w in result["warnings"]),
                        result["warnings"])

    def test_an_edge_to_a_missing_node_warns_and_is_dropped(self):
        _, result, snapshot = self.build("badedge", kind="flow",
                                         nodes=nodes_of("Spot", "Ask"),
                                         edges=[{"from": "n1", "to": "nope"}])
        self.assertTrue(any("does not exist" in w for w in result["warnings"]),
                        result["warnings"])
        edges = [s for s in shapes_for(snapshot, "b_diag") if s["kind"] == "edge"]
        self.assertEqual(edges, [])

    def test_an_empty_diagram_becomes_a_placeholder(self):
        _, result, snapshot = self.build("empty", kind="flow", nodes=[], edges=[])
        shapes = shapes_for(snapshot, "b_diag")
        self.assertEqual([s["kind"] for s in shapes], ["placeholder"])
        self.assertTrue(any("no nodes" in w for w in result["warnings"]),
                        result["warnings"])

    def test_labels_step_down_to_the_floor_and_warn(self):
        long_labels = ["Partners describe the offer in their own words"] * 4
        _, result, snapshot = self.build("floor", kind="flow", nodes=nodes_of(*long_labels),
                                         edges=[], col=3, row=7, cols=8, rows=3)
        nodes = [s for s in shapes_for(snapshot, "b_diag") if s["kind"] == "node"]
        floor = float(build_pptx.DEFAULT_TYPE_SCALE["diagramNode"]["min"])
        self.assertTrue(nodes)
        for entry in nodes:
            self.assertGreaterEqual(entry["fontPt"], floor)
        self.assertEqual(min(n["fontPt"] for n in nodes), floor)
        self.assertTrue(any("floor" in w for w in result["warnings"]), result["warnings"])

    def test_check_deck_reports_diagram_findings(self):
        spec = diagram_spec(os.path.join(self.tmp, "findings.pptx"), "flow",
                            nodes_of("Partners describe the offer in their own words",
                                     "Ask", "Hand off"),
                            edges=[{"from": "n1", "to": "ghost"}],
                            col=3, row=7, cols=8, rows=3)
        result = build_pptx.build(copy.deepcopy(spec))
        spec_path = os.path.join(self.tmp, "findings_spec.json")
        with open(spec_path, "w") as handle:
            json.dump(spec, handle)
        report = check_deck.check_deck(result["path"], spec_path)
        kinds = [f["kind"] for f in report["findings"]]
        self.assertIn("diagram-edge", kinds)
        self.assertIn("diagram-font", kinds)
        edge = [f for f in report["findings"] if f["kind"] == "diagram-edge"][0]
        self.assertEqual(edge["severity"], "error")
        self.assertEqual(edge["blockId"], "b_diag")
        self.assertIn("ghost", edge["message"])
        font = [f for f in report["findings"] if f["kind"] == "diagram-font"][0]
        self.assertEqual(font["severity"], "warning")

    def test_a_clean_diagram_has_no_findings(self):
        spec = diagram_spec(os.path.join(self.tmp, "clean.pptx"), "flow",
                            nodes_of("Spot the moment", "Script the ask", "Hand off"),
                            edges=[{"from": "n1", "to": "n2", "label": "prompts"},
                                   {"from": "n2", "to": "n3", "label": "owns"}])
        result = build_pptx.build(copy.deepcopy(spec))
        spec_path = os.path.join(self.tmp, "clean_spec.json")
        with open(spec_path, "w") as handle:
            json.dump(spec, handle)
        report = check_deck.check_deck(result["path"], spec_path)
        self.assertEqual([f for f in report["findings"]
                          if f["kind"].startswith("diagram")], [])
        self.assertEqual(report["summary"]["errors"], 0, report["findings"])

    def test_planted_node_overlap_is_found(self):
        spec = diagram_spec(os.path.join(self.tmp, "overlap.pptx"), "flow",
                            nodes_of("Spot", "Ask"), edges=[])
        result = build_pptx.build(copy.deepcopy(spec))
        path = result["path"]
        deck = Presentation(path)
        boxes = rounded(list(deck.slides)[0])
        boxes[1].left = boxes[0].left
        boxes[1].top = boxes[0].top
        deck.save(path)
        report = check_deck.check_deck(path)
        kinds = [f["kind"] for f in report["findings"]]
        self.assertIn("diagram-overlap", kinds)
        found = [f for f in report["findings"] if f["kind"] == "diagram-overlap"][0]
        self.assertEqual(found["severity"], "error")
        self.assertEqual(found["blockId"], "b_diag")


class FurnitureTests(unittest.TestCase):
    """F5 kicker pattern, F11 headline casing, F16 numbered dividers, the guard."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="blueprint_furniture_test_")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def build(self, name, slides, style=None):
        spec = {"name": "Furniture", "outputPath": os.path.join(self.tmp, name + ".pptx"),
                "thumbnails": False, "exportPdf": False, "fidelity": "strict",
                "cover": False, "agenda": False, "style": style or {"mode": "house"},
                "slideSize": {"w": 960, "h": 540},
                "grid": {"cols": 32, "rows": 18, "cell": 30}, "slides": slides}
        result = build_pptx.build(copy.deepcopy(spec))
        self.assertTrue(result["ok"], result.get("errors"))
        return result, read_snapshot(result["path"])

    @staticmethod
    def role_text(snapshot, index, role):
        return [s["text"] for s in snapshot["slides"][index]["shapes"] if s["role"] == role]

    def content(self, **over):
        slide_id = over.pop("id", "c1")
        slide = {"id": slide_id, "name": "The prize", "section": "Why referrals",
                 "headline": "a warm introduction converts four times better",
                 "blocks": [{"id": "b_%s" % slide_id, "type": "text", "col": 3, "row": 7,
                             "cols": 12, "rows": 4,
                             "content": {"paragraphs": ["Body copy."]}}]}
        slide.update(over)
        return slide

    def test_kicker_follows_the_section_and_the_slide_name(self):
        _, snapshot = self.build("kicker", [self.content()])
        self.assertEqual(self.role_text(snapshot, 0, "kicker"),
                         ["Why referrals | The prize"])

    def test_an_explicit_kicker_wins(self):
        _, snapshot = self.build("kicker2", [self.content(kicker="Objective | The Prize")])
        self.assertEqual(self.role_text(snapshot, 0, "kicker"), ["Objective | The Prize"])

    def test_a_style_can_ask_for_the_topic_alone(self):
        _, snapshot = self.build("kicker3", [self.content()],
                                 style={"mode": "skill",
                                        "tokens": {"kicker": {"pattern": "{topic}"}}})
        self.assertEqual(self.role_text(snapshot, 0, "kicker"), ["The prize"])

    def test_a_missing_section_drops_the_separator(self):
        _, snapshot = self.build("kicker4", [self.content(section="")])
        self.assertEqual(self.role_text(snapshot, 0, "kicker"), ["The prize"])

    def test_headlines_are_capitalised_and_otherwise_untouched(self):
        _, snapshot = self.build("caps", [
            self.content(headline="referrals stall at the ask"),
            self.content(id="c2", headline="iPhone referrals stall at the ask"),
        ])
        self.assertEqual(self.role_text(snapshot, 0, "headline"),
                         ["Referrals stall at the ask"])
        self.assertEqual(self.role_text(snapshot, 1, "headline"),
                         ["iPhone referrals stall at the ask"])
        self.assertEqual(snapshot["slides"][0]["headline"], "Referrals stall at the ask")

    def test_dividers_are_numbered_in_deck_order(self):
        slides = [
            {"id": "d1", "kind": "divider", "name": "One", "section": "Why referrals",
             "subtitle": "Where the evidence starts."},
            self.content(),
            {"id": "d2", "kind": "divider", "name": "Two", "section": "What we do next"},
            {"id": "d3", "kind": "thanks", "name": "Thank you",
             "headline": "thank you"},
        ]
        _, snapshot = self.build("dividers", slides)
        self.assertEqual(self.role_text(snapshot, 0, "dividerTitle"), ["01 Why referrals"])
        self.assertEqual(self.role_text(snapshot, 0, "dividerSubtitle"),
                         ["Where the evidence starts."])
        self.assertEqual(self.role_text(snapshot, 2, "dividerTitle"), ["02 What we do next"])
        self.assertEqual(self.role_text(snapshot, 3, "dividerTitle"), ["Thank you"])
        self.assertEqual(snapshot["slides"][0]["ordinal"], 1)

    def test_a_style_can_turn_the_numbers_off(self):
        slides = [{"id": "d1", "kind": "divider", "name": "One", "section": "Why referrals"}]
        _, snapshot = self.build("plainnums", slides,
                                 style={"mode": "skill", "tokensPath": HARBOR_TOKENS})
        self.assertEqual(self.role_text(snapshot, 0, "dividerTitle"), ["Why referrals"])

    def test_a_furniture_box_is_only_as_tall_as_its_text(self):
        _, snapshot = self.build("shrink", [self.content(headline="One short line")])
        headline = [s for s in snapshot["slides"][0]["shapes"]
                    if s["role"] == "headline"][0]
        slot = build_pptx.make_style({"style": {"mode": "house"}}, []
                                     )[0].slot("headline", "content")
        self.assertLess(headline["h"], slot["h"])
        self.assertGreater(headline["h"], 20)

    def test_a_block_on_the_headline_band_is_moved_down(self):
        slide = self.content()
        slide["blocks"] = [{"id": "b_high", "type": "text", "col": 3, "row": 2,
                            "cols": 10, "rows": 4,
                            "content": {"paragraphs": ["Sits in the furniture."]}}]
        result, snapshot = self.build("guardtop", [slide])
        entry = shapes_for(snapshot, "b_high")[0]
        self.assertGreaterEqual(entry["y"], 120.0)
        self.assertTrue(any("b_high" in w and "headline" in w and "moved to row" in w
                            for w in result["warnings"]), result["warnings"])

    def test_a_block_over_the_source_line_moves_up_when_it_can(self):
        slide = self.content(source="Partner CRM")
        slide["blocks"] = [{"id": "b_low", "type": "text", "col": 3, "row": 8,
                            "cols": 20, "rows": 11,
                            "content": {"paragraphs": ["Runs into the footer."]}}]
        result, snapshot = self.build("guardbottom", [slide])
        entry = shapes_for(snapshot, "b_low")[0]
        self.assertLessEqual(entry["y"] + entry["h"], 510.5)
        self.assertAlmostEqual(entry["h"], 330.0, delta=0.5)   # kept its height
        self.assertTrue(any("b_low" in w and "source" in w and "moved to row" in w
                            for w in result["warnings"]), result["warnings"])

    def test_a_block_that_cannot_move_up_is_trimmed(self):
        slide = self.content(source="Partner CRM")
        slide["blocks"] = [{"id": "b_tall", "type": "text", "col": 3, "row": 5,
                            "cols": 20, "rows": 14,
                            "content": {"paragraphs": ["Covers the whole slide."]}}]
        result, snapshot = self.build("guardtrim", [slide])
        entry = shapes_for(snapshot, "b_tall")[0]
        self.assertLessEqual(entry["y"] + entry["h"], 510.5)
        self.assertGreaterEqual(entry["y"], 120.0)
        self.assertTrue(any("b_tall" in w and "trimmed to row" in w
                            for w in result["warnings"]), result["warnings"])

    def test_the_source_band_is_only_reserved_when_a_source_exists(self):
        slide = self.content()
        slide["blocks"] = [{"id": "b_low", "type": "text", "col": 3, "row": 8,
                            "cols": 20, "rows": 11,
                            "content": {"paragraphs": ["No source on this slide."]}}]
        result, snapshot = self.build("nosource", [slide])
        entry = shapes_for(snapshot, "b_low")[0]
        self.assertAlmostEqual(entry["y"] + entry["h"], 540.0, delta=0.5)
        self.assertFalse(any("b_low" in w and "trimmed" in w for w in result["warnings"]),
                         result["warnings"])

    def test_the_guarded_slide_passes_the_overlap_check(self):
        slide = self.content(source="Partner CRM")
        slide["blocks"] = [
            {"id": "b_high", "type": "text", "col": 3, "row": 2, "cols": 10, "rows": 4,
             "content": {"paragraphs": ["Sits in the furniture."]}},
            {"id": "b_low", "type": "text", "col": 3, "row": 8, "cols": 20, "rows": 11,
             "content": {"paragraphs": ["Runs into the footer."]}},
        ]
        result, _ = self.build("guardcheck", [slide])
        report = check_deck.check_deck(result["path"])
        overlaps = [f for f in report["findings"] if f["kind"] == "overlap"]
        self.assertEqual(overlaps, [], report["findings"])


class DarkStyleTests(unittest.TestCase):
    """A partial token file with a dark palette derives its own defaults."""

    def style(self, tokens=None):
        warnings = []
        spec = {"style": {"mode": "skill", "tokens": tokens or copy.deepcopy(DARK_TOKENS)}}
        style, _ = build_pptx.make_style(spec, warnings)
        return style, warnings

    def test_a_dark_style_that_names_no_table_gets_no_header_fill(self):
        """Q14: a header band is non-data ink, so a style has to ask for it."""
        style, _ = self.style()
        self.assertIsNone(style.table_header_fill())
        self.assertNotIn("headerFill", style.table)
        self.assertEqual(style.hex_of(style.table["ruleColor"]), "262626")
        self.assertEqual(style.hex_of(style.role("tableHeader")["color"]), "FFFFFF")

    def test_a_style_that_names_a_header_fill_keeps_it(self):
        style, _ = self.style({"palette": {"bg": "000000", "ink": "FFFFFF"},
                               "table": {"headerFill": "111111"}})
        self.assertEqual(style.hex_of(style.table_header_fill()), "111111")

    def test_the_takeaway_bar_is_not_a_light_grey_box(self):
        style, _ = self.style()
        self.assertEqual(style.hex_of(style.palette["takeawayFill"]), "2D2D2D")

    def test_muted_gets_its_own_value_when_it_equals_the_ink(self):
        style, _ = self.style()
        self.assertNotEqual(style.hex_of("muted"), style.hex_of("ink"))
        self.assertEqual(style.hex_of("muted"), "A6A6A6")

    def test_a_style_with_its_own_furniture_drops_the_house_hairline(self):
        style, _ = self.style()
        self.assertIsNone(style.slot("hairline", "content"))
        house, _ = build_pptx.make_style({"style": {"mode": "house"}}, [])
        self.assertIsNotNone(house.slot("hairline", "content"))

    def test_a_light_palette_keeps_the_house_defaults(self):
        style, _ = self.style({"palette": {"bg": "FFFFFF", "ink": "1F2937"}})
        self.assertEqual(style.hex_of(style.table["headerFill"]), "F3F4F6")
        self.assertIsNone(style.table_header_fill(), "house named it, not the style")

    def test_a_weight_suffix_resolves_to_the_family(self):
        if not build_pptx.FONTS.has("Montserrat"):
            self.skipTest("Montserrat is not installed on this machine")
        style, warnings = self.style()
        self.assertEqual(style.font("display"), "Montserrat")
        self.assertTrue(any("Montserrat Light" in w for w in warnings), warnings)

    def test_a_missing_family_falls_back_to_arial_never_helvetica(self):
        style, warnings = self.style({"palette": {"bg": "FFFFFF"},
                                      "fonts": {"display": "Nonesuch Grotesk",
                                                "body": "Nonesuch Grotesk"}})
        self.assertEqual(style.font("body"), "Arial")
        self.assertTrue(any("is not installed" in w for w in warnings), warnings)
        self.assertFalse(any("Helvetica" in w for w in warnings), warnings)

    def test_a_placeholder_on_a_dark_palette_is_a_dashed_outline(self):
        tmp = tempfile.mkdtemp(prefix="blueprint_darkph_test_")
        try:
            spec = {"name": "Dark", "outputPath": os.path.join(tmp, "ph.pptx"),
                    "thumbnails": False, "exportPdf": False, "fidelity": "strict",
                    "cover": False, "agenda": False,
                    "style": {"mode": "skill", "tokens": copy.deepcopy(DARK_TOKENS)},
                    "slideSize": {"w": 960, "h": 540},
                    "grid": {"cols": 32, "rows": 18, "cell": 30},
                    "slides": [{
                        "id": "s1", "name": "Visual", "section": "Offering",
                        "headline": "the placeholder is honest about the gap",
                        "blocks": [{"id": "b_ph", "type": "placeholder", "col": 3,
                                    "row": 8, "cols": 26, "rows": 6,
                                    "content": {"label": "Designed process visual"}}],
                    }]}
            result = build_pptx.build(copy.deepcopy(spec))
            self.assertTrue(result["ok"], result.get("errors"))
            deck = Presentation(result["path"])
            box = rounded(list(deck.slides)[0])[0]
            self.assertEqual(str(box.line.dash_style), "DASH (4)")
            self.assertEqual(str(box.line.color.rgb), "262626")
            self.assertEqual(box.fill.type, MSO_FILL.BACKGROUND)
            run = box.text_frame.paragraphs[0].runs[0]
            self.assertTrue(run.font.italic)
            self.assertEqual(str(run.font.color.rgb), "A6A6A6")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_the_dark_deck_builds_with_readable_tables(self):
        tmp = tempfile.mkdtemp(prefix="blueprint_dark_test_")
        try:
            spec = {"name": "Dark", "outputPath": os.path.join(tmp, "dark.pptx"),
                    "thumbnails": False, "exportPdf": False, "fidelity": "strict",
                    "cover": False, "agenda": False,
                    "style": {"mode": "skill", "tokens": copy.deepcopy(DARK_TOKENS)},
                    "slideSize": {"w": 960, "h": 540},
                    "grid": {"cols": 32, "rows": 18, "cell": 30},
                    "slides": [{
                        "id": "s1", "name": "Table", "section": "Appendix",
                        "headline": "the table reads on a black slide",
                        "blocks": [{"id": "b_t", "type": "table", "col": 3, "row": 8,
                                    "cols": 26, "rows": 8,
                                    "content": {"headers": ["Firm", "Offer"],
                                                "rows": [["Interbrand", "Strategy"]]}}],
                    }]}
            result = build_pptx.build(copy.deepcopy(spec))
            self.assertTrue(result["ok"], result.get("errors"))
            deck = Presentation(result["path"])
            table = None
            for shape in list(deck.slides)[0].shapes:
                if getattr(shape, "has_table", False):
                    table = shape.table
            self.assertIsNotNone(table)
            header = table.cell(0, 0)
            self.assertIn("BACKGROUND", str(header.fill.type), "no header band (Q14)")
            run = header.text_frame.paragraphs[0].runs[0]
            self.assertEqual(run.text, "FIRM")
            self.assertEqual(str(run.font.color.rgb), "D9A066")
            report = check_deck.check_deck(result["path"])
            self.assertEqual([f for f in report["findings"] if f["severity"] == "error"],
                             [], report["findings"])
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------------- Q3 fidelity


# House with no `furniture.body`, so there is no column module to snap to: the
# suites that test grow, rebalance and align on their own use this.
NO_MODULE = {"mode": "house", "tokens": {"furniture": {"body": None}}}


def mini_spec(out_path, slides, **over):
    """A bare deck: house style, no generated chrome, whatever slides you pass."""
    spec = {"name": "Fidelity Deck", "outputPath": out_path, "thumbnails": False,
            "exportPdf": False, "style": {"mode": "house"},
            "slideSize": {"w": 960, "h": 540},
            "grid": {"cols": 32, "rows": 18, "cell": 30},
            "cover": False, "agenda": False, "slides": slides}
    spec.update(over)
    return spec


def content_slide(slide_id, blocks, headline="The geometry has to earn its keep", **over):
    slide = {"id": slide_id, "name": slide_id, "kind": "content",
             "headline": headline, "blocks": blocks}
    slide.update(over)
    return slide


def text_block(block_id, col, row, cols, rows, text, **over):
    block = {"id": block_id, "type": "text", "role": "text", "col": col, "row": row,
             "cols": cols, "rows": rows, "content": {"paragraphs": [text]}}
    block.update(over)
    return block


def cells_of(snapshot, block_id, cell=30.0):
    """The drawn rectangle of a block, back in grid cells."""
    entry = shapes_for(snapshot, block_id)[0]
    return {"col": int(round(entry["x"] / cell)) + 1,
            "row": int(round(entry["y"] / cell)) + 1,
            "cols": int(round(entry["w"] / cell)),
            "rows": int(round(entry["h"] / cell))}


LONG = "overflowing copy that will not fit in two rows however hard the builder tries " * 3
SHORT = "Short."


class FidelityTests(unittest.TestCase):
    """Q3: strict draws the plan, balanced settles it, loose re-picks the slot."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="blueprint_fidelity_test_")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def build(self, name, slides, **over):
        spec = mini_spec(os.path.join(self.tmp, name + ".pptx"), slides, **over)
        result = build_pptx.build(copy.deepcopy(spec))
        self.assertTrue(result["ok"], result.get("errors"))
        return result, read_snapshot(result["path"])

    def crowded(self):
        """A block that overflows at the floor over a block with rows to spare."""
        return [content_slide("s1", [
            text_block("b_tight", 3, 7, 6, 2, LONG),
            text_block("b_slack", 3, 9, 6, 6, SHORT),
            text_block("b_foot", 3, 16, 6, 2, SHORT),
        ])]

    # -- strict

    def test_strict_leaves_the_geometry_alone(self):
        result, snapshot = self.build("strict", self.crowded(), fidelity="strict")
        self.assertEqual(cells_of(snapshot, "b_tight"),
                         {"col": 3, "row": 7, "cols": 6, "rows": 2})
        self.assertEqual(cells_of(snapshot, "b_slack"),
                         {"col": 3, "row": 9, "cols": 6, "rows": 6})
        self.assertNotIn("adjusted", result)
        self.assertEqual(result["fidelity"], "strict")

    def test_strict_still_reports_the_overflow_it_refuses_to_fix(self):
        result, _ = self.build("strict2", self.crowded(), fidelity="strict")
        self.assertTrue(any("b_tight" in w and "overflows" in w
                            for w in result["warnings"]), result["warnings"])

    def test_an_unknown_fidelity_falls_back_to_balanced(self):
        result, _ = self.build("weird", self.crowded(), fidelity="sideways")
        self.assertEqual(result["fidelity"], "balanced")
        self.assertTrue(any("unknown fidelity" in w for w in result["warnings"]),
                        result["warnings"])

    # -- balanced: grow

    def test_balanced_grows_the_block_that_does_not_fit(self):
        _, snapshot = self.build("grow", self.crowded())
        grown = cells_of(snapshot, "b_tight")
        self.assertEqual(grown["row"], 7)
        self.assertEqual(grown["rows"], 4, "the block should take its two rows")

    def test_balanced_takes_the_rows_from_the_neighbour_with_slack(self):
        _, snapshot = self.build("shrink", self.crowded())
        shrunk = cells_of(snapshot, "b_slack")
        self.assertEqual(shrunk["row"], 11)
        self.assertEqual(shrunk["rows"], 4)

    def test_a_neighbour_with_no_slack_keeps_its_cells(self):
        slides = [content_slide("s1", [
            text_block("b_tight", 3, 7, 6, 2, LONG),
            text_block("b_full", 3, 9, 6, 2, LONG),
            text_block("b_foot", 3, 12, 6, 6, SHORT),
        ])]
        _, snapshot = self.build("noslack", slides)
        self.assertEqual(cells_of(snapshot, "b_full")["row"], 9)
        self.assertEqual(cells_of(snapshot, "b_tight"),
                         {"col": 3, "row": 7, "cols": 6, "rows": 2})

    def test_a_grown_block_never_lands_on_the_furniture(self):
        slides = [content_slide("s1", [text_block("b_up", 3, 7, 6, 2, LONG)],
                                source="Test fixture")]
        _, snapshot = self.build("guarded", slides)
        grown = cells_of(snapshot, "b_up")
        self.assertGreaterEqual(grown["row"], 7, "the hairline band is row 6")
        self.assertLessEqual(grown["row"] + grown["rows"] - 1, 17)

    # -- balanced: rebalance

    def test_balanced_centres_one_row_of_blocks_in_an_empty_band(self):
        slides = [content_slide("s1", [text_block("b_one", 3, 7, 20, 3, SHORT)])]
        result, snapshot = self.build("centre", slides)
        moved = cells_of(snapshot, "b_one")
        self.assertEqual(moved["row"], 11)
        self.assertEqual(moved["rows"], 3)
        reasons = [a["reason"] for a in result["adjusted"]]
        self.assertTrue(any("centred in the body band" in r for r in reasons), reasons)

    def test_balanced_spreads_two_rows_of_blocks_over_the_band(self):
        slides = [content_slide("s1", [
            text_block("b_top", 3, 7, 12, 3, SHORT),
            text_block("b_bottom", 3, 11, 12, 3, SHORT),
        ])]
        result, snapshot = self.build("spread", slides)
        top = cells_of(snapshot, "b_top")
        bottom = cells_of(snapshot, "b_bottom")
        self.assertEqual(top["row"], 7)
        self.assertGreater(top["rows"], 3)
        self.assertGreater(bottom["row"], 11)
        self.assertEqual(bottom["row"] + bottom["rows"] - 1, 17)
        reasons = [a["reason"] for a in result["adjusted"]]
        self.assertTrue(any("spread to fill the body band" in r for r in reasons),
                        reasons)

    def test_a_full_body_is_left_alone(self):
        """A block already on the module and filling the band does not move."""
        slides = [content_slide("s1", [text_block("b_full", 3, 7, 28, 11, SHORT)])]
        result, snapshot = self.build("nogap", slides)
        self.assertEqual(cells_of(snapshot, "b_full"),
                         {"col": 3, "row": 7, "cols": 28, "rows": 11})
        self.assertNotIn("adjusted", result)

    # -- balanced: align

    def test_balanced_snaps_edges_within_one_cell(self):
        slides = [content_slide("s1", [
            text_block("b_wide", 3, 7, 28, 5, SHORT),
            text_block("b_narrow", 3, 13, 27, 5, SHORT),
        ])]
        result, snapshot = self.build("align", slides, style=NO_MODULE)
        wide = cells_of(snapshot, "b_wide")
        narrow = cells_of(snapshot, "b_narrow")
        self.assertEqual(wide["col"] + wide["cols"], narrow["col"] + narrow["cols"])
        self.assertTrue(any("edge aligned" in a["reason"] for a in result["adjusted"]),
                        result["adjusted"])

    def test_edges_further_than_one_cell_apart_are_left_alone(self):
        slides = [content_slide("s1", [
            text_block("b_wide", 3, 7, 28, 5, SHORT),
            text_block("b_narrow", 3, 13, 24, 5, SHORT),
        ])]
        _, snapshot = self.build("noalign", slides, style=NO_MODULE)
        self.assertEqual(cells_of(snapshot, "b_narrow")["cols"], 24)

    # -- the record

    def test_every_deviation_is_recorded_in_the_result_and_the_snapshot(self):
        result, snapshot = self.build("record", self.crowded())
        self.assertIn("adjusted", result)
        moves = dict((m["blockId"], m) for m in result["adjusted"])
        self.assertIn("b_tight", moves)
        move = moves["b_tight"]
        for key in ("from", "to", "reason", "slide"):
            self.assertIn(key, move)
        self.assertEqual(move["from"], {"col": 3, "row": 7, "cols": 6, "rows": 2})
        self.assertEqual(move["to"]["rows"], 4)
        slide = snapshot["slides"][0]
        self.assertIn("adjusted", slide)
        self.assertIn("b_tight", [m["blockId"] for m in slide["adjusted"]])
        stamped = shapes_for(snapshot, "b_tight")[0]
        self.assertEqual(stamped["adjusted"]["reason"], move["reason"])

    def test_the_snapshot_header_carries_the_fidelity(self):
        _, snapshot = self.build("header", self.crowded())
        self.assertEqual(snapshot["fidelity"], "balanced")

    def test_a_notes_block_is_not_moved_and_never_blocks(self):
        slides = [content_slide("s1", [
            text_block("b_tight", 3, 7, 6, 2, LONG),
            {"id": "b_notes", "type": "notes", "col": 3, "row": 8, "cols": 6, "rows": 2,
             "content": {"text": "Block level notes."}},
        ])]
        result, snapshot = self.build("notesblock", slides)
        self.assertEqual(cells_of(snapshot, "b_tight")["rows"], 4)
        self.assertNotIn("b_notes", [m["blockId"] for m in result.get("adjusted") or []])

    # -- loose

    def test_loose_repicks_a_seven_column_table_into_the_full_width_slot(self):
        table = {"id": "b_table", "type": "table", "role": "table",
                 "col": 3, "row": 7, "cols": 12, "rows": 8,
                 "content": {"headers": ["Role", "Rate", "Build", "Rollout", "Office",
                                         "Hours", "Cost"],
                             "rows": [["Design Director", "$500", "24", "6", "2", "32",
                                       "$18,500"]]}}
        result, snapshot = self.build("loosetable", [content_slide("s1", [table])],
                                      fidelity="loose")
        moved = cells_of(snapshot, "b_table")
        # The catalogue slot is col 2; the module then pulls it onto the body's
        # own left edge and full measure (Q9).
        self.assertEqual(moved["col"], 3)
        self.assertEqual(moved["cols"], 28)
        self.assertIn("full-width table slot",
                      " ".join(a["reason"] for a in result["adjusted"]))

    def test_balanced_leaves_that_table_where_the_plan_put_it(self):
        table = {"id": "b_table", "type": "table", "role": "table",
                 "col": 3, "row": 7, "cols": 12, "rows": 8,
                 "content": {"headers": ["Role", "Rate", "Build", "Rollout", "Office",
                                         "Hours", "Cost"],
                             "rows": [["Design Director", "$500", "24", "6", "2", "32",
                                       "$18,500"]]}}
        _, snapshot = self.build("baltable", [content_slide("s1", [table])])
        self.assertEqual(cells_of(snapshot, "b_table")["col"], 3)

    def test_loose_repicks_five_cards_out_of_a_three_card_slot(self):
        cards = {"id": "b_cards", "type": "cards", "role": "card",
                 "col": 3, "row": 7, "cols": 12, "rows": 8,
                 "content": {"items": [{"header": "One", "body": "First card body."},
                                       {"header": "Two", "body": "Second card body."},
                                       {"header": "Three", "body": "Third card body."},
                                       {"header": "Four", "body": "Fourth card body."},
                                       {"header": "Five", "body": "Fifth card body."}]}}
        result, snapshot = self.build("loosecards", [content_slide("s1", [cards])],
                                      fidelity="loose")
        first = shapes_for(snapshot, "b_cards")[0]
        self.assertLess(first["x"], 3 * 30.0)
        self.assertIn("wide card box", " ".join(a["reason"] for a in result["adjusted"]))

    def test_loose_moves_a_stat_that_is_really_a_sentence(self):
        stat = {"id": "b_stat", "type": "stat", "role": "stat",
                "col": 3, "row": 7, "cols": 8, "rows": 5,
                "content": {"items": [{"value": "Eight words is not a number anyone reads",
                                       "label": "Not a stat"}]}}
        result, snapshot = self.build("loosestat", [content_slide("s1", [stat])],
                                      fidelity="loose")
        moved = cells_of(snapshot, "b_stat")
        self.assertEqual(moved["cols"], 23)      # 22 cells, snapped to the module
        self.assertIn("reads as a statement",
                      " ".join(a["reason"] for a in result["adjusted"]))

    def test_a_repick_that_would_collide_is_refused(self):
        table = {"id": "b_table", "type": "table", "role": "table",
                 "col": 3, "row": 7, "cols": 12, "rows": 8,
                 "content": {"headers": ["Role", "Rate", "Build", "Rollout", "Office",
                                         "Hours", "Cost"],
                             "rows": [["Design Director", "$500", "24", "6", "2", "32",
                                       "$18,500"]]}}
        slides = [content_slide("s1", [table, text_block("b_rail", 17, 7, 14, 8, SHORT)])]
        _, snapshot = self.build("norepick", slides, fidelity="loose")
        self.assertEqual(cells_of(snapshot, "b_table")["col"], 3)
        self.assertEqual(cells_of(snapshot, "b_rail")["col"], 17)

    def test_a_block_the_guard_trimmed_keeps_its_trimmed_height(self):
        """The guard trims off the row grid; the pass must not hand the row back."""
        slides = [content_slide("s1", [
            text_block("b_table", 5, 8, 26, 9, SHORT),
            text_block("b_legend", 5, 17, 26, 1, SHORT),
        ], source="Partner CRM")]
        result, snapshot = self.build("trimmed", slides)
        legend = shapes_for(snapshot, "b_legend")[0]
        self.assertLess(legend["y"] + legend["h"], 514.0,
                        "the source line starts at 514 in house")
        report = check_deck.check_deck(result["path"])
        self.assertEqual([f for f in report["findings"] if f["severity"] == "error"],
                         [], report["findings"])

    def test_the_adjusted_slide_still_passes_the_overlap_check(self):
        result, _ = self.build("checked", self.crowded())
        report = check_deck.check_deck(result["path"])
        errors = [f for f in report["findings"] if f["severity"] == "error"]
        self.assertEqual(errors, [], report["findings"])


class FitModuleTests(unittest.TestCase):
    """builder/fit.py on its own: cells, bands, ladders."""

    def test_the_ladder_starts_at_the_style_size_and_ends_at_the_floor(self):
        ladder = fit.size_ladder({"size": 28, "min": 24, "max": 32})
        self.assertEqual(ladder[0], 28.0)
        self.assertEqual(ladder[-1], 24.0)
        self.assertEqual(ladder, [28.0, 26.0, 24.0])

    def test_the_ladder_steps_one_point_below_fourteen(self):
        ladder = fit.size_ladder({"size": 16, "min": 11, "max": 16})
        self.assertEqual(ladder, [16.0, 14.0, 13.0, 12.0, 11.0])

    def test_a_ladder_at_its_floor_has_one_rung(self):
        self.assertEqual(fit.size_ladder({"size": 11, "min": 11, "max": 14}), [11.0])
        self.assertEqual(fit.size_ladder({"size": 9, "min": 12}), [12.0])

    def test_a_left_edge_ties_to_the_outer_column_and_a_right_edge_to_the_outer(self):
        values = {30: ["a"], 31: ["b"]}
        self.assertEqual(fit.Plan.target_edge("col", [30, 31], values), 30)
        self.assertEqual(fit.Plan.target_edge("right", [30, 31], values), 31)
        crowd = {30: ["a", "b"], 31: ["c"]}
        self.assertEqual(fit.Plan.target_edge("right", [30, 31], crowd), 30)

    def test_a_rectangle_becomes_the_cells_it_fills_not_the_cells_it_touches(self):
        cells = fit.rect_to_cells((60.0, 210.0, 180.0, 289.0), 30.0, 30.0)
        self.assertEqual(cells, {"col": 3, "row": 8, "cols": 6, "rows": 9})

    def test_a_furniture_band_takes_every_cell_it_reaches(self):
        band = fit.rect_to_band((336.0, 499.0, 593.0, 12.0), 30.0, 30.0)
        self.assertEqual(band["col"], 12)
        self.assertEqual(band["row"], 17)
        self.assertEqual(fit.bottom(band) - 1, 18)

    def test_edges_cluster_within_one_cell(self):
        self.assertEqual(fit.Plan.clusters([3, 4, 9, 10, 11, 20]),
                         [[3, 4], [9, 10, 11], [20]])

    def test_abutting_blocks_do_not_count_as_overlapping(self):
        left = fit.cells(3, 7, 6, 4)
        right = fit.cells(9, 7, 6, 4)
        self.assertFalse(fit.cells_overlap(left, right))
        self.assertTrue(fit.cells_overlap(left, fit.cells(8, 7, 6, 4)))


# ------------------------------------------------------------ Q4 deck-wide type


class DeckTypeTests(unittest.TestCase):
    """Q4: one headline size for the deck, one size per body role."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="blueprint_decktype_test_")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def build(self, name, slides, **over):
        spec = mini_spec(os.path.join(self.tmp, name + ".pptx"), slides, **over)
        result = build_pptx.build(copy.deepcopy(spec))
        self.assertTrue(result["ok"], result.get("errors"))
        return result, read_snapshot(result["path"])

    HEADLINES = [
        "Referrals stall at three moments",
        "This headline is a good deal longer than the one above it and has to wrap "
        "onto a second line before it fits the slot at all",
        "The prize is a point of share",
    ]

    def mixed(self):
        return [content_slide("s%d" % (i + 1),
                              [text_block("b%d" % (i + 1), 3, 7, 12, 11, SHORT)],
                              headline=text)
                for i, text in enumerate(self.HEADLINES)]

    def sizes_for(self, snapshot, role):
        return [s["fontPt"] for slide in snapshot["slides"] for s in slide["shapes"]
                if s["role"] == role and s.get("fontPt")]

    def test_one_headline_size_on_every_slide(self):
        result, snapshot = self.build("headsize", self.mixed())
        sizes = self.sizes_for(snapshot, "headline")
        self.assertEqual(len(sizes), 3)
        self.assertEqual(len(set(sizes)), 1, sizes)
        self.assertEqual(sizes[0], result["typeScale"]["headline"])

    def test_the_deck_headline_size_fits_the_longest_headline(self):
        result, _ = self.build("headfit", self.mixed())
        size = result["typeScale"]["headline"]
        spec = build_pptx.DEFAULT_TYPE_SCALE["headline"]
        for text in self.HEADLINES:
            lines = build_pptx.wrap_lines(
                text, build_pptx.cpl(840.0, size, spec["k"]))
            self.assertLessEqual(len(lines), spec["maxLines"], text)

    def test_the_deck_headline_size_is_never_below_the_floor(self):
        long_one = "word " * 120
        slides = self.mixed() + [content_slide("s4", [], headline=long_one)]
        result, _ = self.build("headfloor", slides)
        self.assertGreaterEqual(result["typeScale"]["headline"],
                                build_pptx.DEFAULT_TYPE_SCALE["headline"]["min"])

    def test_a_headline_that_cannot_fit_does_not_shrink_the_others(self):
        slides = self.mixed() + [content_slide("s4", [], headline="word " * 120)]
        result, _ = self.build("headhog", slides)
        clean, _ = self.build("headclean", self.mixed())
        self.assertEqual(result["typeScale"]["headline"], clean["typeScale"]["headline"])

    def test_one_bullet_size_across_the_deck(self):
        def bullets(block_id, items):
            return {"id": block_id, "type": "bullets", "role": "list", "col": 3,
                    "row": 7, "cols": 12, "rows": 11, "content": {"items": items}}
        # Four bullets that need a second line at 14 pt and one line below it.
        tight = ["A bullet long enough to need a second line at 14pt"] * 4
        slides = [
            content_slide("s1", [bullets("b1", ["One", "Two"])]),
            content_slide("s2", [dict(bullets("b2", tight), rows=4)]),
        ]
        result, snapshot = self.build("bullets", slides, fidelity="strict")
        sizes = self.sizes_for(snapshot, "bullets")
        self.assertEqual(len(set(sizes)), 1, sizes)
        self.assertEqual(sizes[0], result["typeScale"]["bullets"])
        self.assertLess(sizes[0], build_pptx.DEFAULT_TYPE_SCALE["bullets"]["size"])

    def test_one_stat_value_size_across_the_deck(self):
        def stat(block_id, items, cols=28):
            return {"id": block_id, "type": "stat", "role": "stat", "col": 3, "row": 7,
                    "cols": cols, "rows": 11, "content": {"items": items}}
        slides = [
            content_slide("s1", [stat("b1", [{"value": "4.1x", "label": "Better"}])]),
            content_slide("s2", [stat("b2", [
                {"value": "7.7%", "label": "of company revenue goes to marketing, "
                                           "flat from 2024"},
                {"value": "59%", "label": "of CMOs say they lack the budget"},
                {"value": "49%", "label": "of CMOs report time savings"}], cols=14)]),
        ]
        result, snapshot = self.build("stats", slides, fidelity="strict")
        sizes = self.sizes_for(snapshot, "stat")
        self.assertEqual(len(set(sizes)), 1, sizes)
        self.assertEqual(sizes[0], result["typeScale"]["statValue"])

    def test_the_card_roles_step_together(self):
        cards = {"id": "b_cards", "type": "cards", "role": "card", "col": 3, "row": 7,
                 "cols": 28, "rows": 5,
                 "content": {"items": [{"header": "A header that is quite long indeed",
                                        "sub": "Subhead",
                                        "body": "A card body long enough to need "
                                                "several lines of copy before it is "
                                                "done saying what it came to say."}] * 3}}
        result, _ = self.build("cards", [content_slide("s1", [cards])], fidelity="strict")
        scale = result["typeScale"]
        self.assertLessEqual(scale["cardHeader"],
                             build_pptx.DEFAULT_TYPE_SCALE["cardHeader"]["size"])
        self.assertLessEqual(scale["cardBody"],
                             build_pptx.DEFAULT_TYPE_SCALE["cardBody"]["size"])

    def test_the_chosen_sizes_reach_the_snapshot(self):
        _, snapshot = self.build("scalesnap", self.mixed())
        self.assertIn("typeScale", snapshot)
        self.assertIn("headline", snapshot["typeScale"])

    def test_a_role_at_its_floor_is_named_once_in_the_result(self):
        def bullets(block_id):
            return {"id": block_id, "type": "bullets", "role": "list", "col": 3,
                    "row": 7, "cols": 8, "rows": 3,
                    "content": {"items": ["A bullet with a good many words in it "
                                          "so that it cannot fit"] * 4}}
        result, _ = self.build("floor", [content_slide("s1", [bullets("b1")]),
                                         content_slide("s2", [bullets("b2")])],
                               fidelity="strict")
        self.assertEqual(result["typeScale"]["bullets"],
                         build_pptx.DEFAULT_TYPE_SCALE["bullets"]["min"])
        self.assertEqual(result["typeScale"]["atFloor"], ["bullets"])

    def test_check_deck_reports_a_deck_wide_floor_once(self):
        def bullets(block_id):
            return {"id": block_id, "type": "bullets", "role": "list", "col": 3,
                    "row": 7, "cols": 8, "rows": 3,
                    "content": {"items": ["A bullet with a good many words in it "
                                          "so that it cannot fit"] * 4}}
        result, _ = self.build("floorcheck", [content_slide("s1", [bullets("b1")]),
                                              content_slide("s2", [bullets("b2")])],
                               fidelity="strict")
        report = check_deck.check_deck(result["path"])
        floors = [f for f in report["findings"] if f["kind"] == "font-floor"]
        self.assertEqual(len(floors), 1, floors)
        self.assertEqual(floors[0]["slide"], 0)
        self.assertEqual(floors[0]["scope"], "deck")
        self.assertEqual(floors[0]["severity"], "warning")
        self.assertIn("deck-wide bullets", floors[0]["message"])

    def test_only_the_deck_wide_finding_carries_a_scope(self):
        def bullets(block_id):
            return {"id": block_id, "type": "bullets", "role": "list", "col": 3,
                    "row": 7, "cols": 8, "rows": 3,
                    "content": {"items": ["A bullet with a good many words in it "
                                          "so that it cannot fit"] * 4}}
        result, _ = self.build("scoped", [content_slide("s1", [bullets("b1")]),
                                          content_slide("s2", [bullets("b2")])],
                               fidelity="strict")
        report = check_deck.check_deck(result["path"])
        for finding in report["findings"]:
            if finding.get("scope") == "deck":
                self.assertEqual(finding["slide"], 0)
            else:
                self.assertGreater(finding["slide"], 0)
                self.assertNotIn("scope", finding)

    def test_the_deck_size_takes_one_step_to_clear_a_widow(self):
        text = "Referrals stall at three moments, and each one has a different fix"
        result, _ = self.build("widowstep", [content_slide("s1", [], headline=text)])
        size = result["typeScale"]["headline"]
        spec = build_pptx.DEFAULT_TYPE_SCALE["headline"]
        lines = build_pptx.wrap_lines(text, build_pptx.cpl(840.0, size, spec["k"]))
        self.assertLess(size, spec["size"])
        self.assertNotEqual(len(lines[-1].split()), 1, lines)


# --------------------------------------------------------- Q5 notes and chrome


class ChromeTests(unittest.TestCase):
    """Q5: the story becomes the notes, and the deck gets a cover and an agenda."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="blueprint_chrome_test_")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def build(self, name, slides, **over):
        spec = mini_spec(os.path.join(self.tmp, name + ".pptx"), slides, **over)
        spec.pop("cover", None)
        spec.pop("agenda", None)
        spec["brief"] = {"audience": "The test suite",
                         "goal": "Land the referral recommendation: who we ask, what we "
                                 "say, who owns the reply and when it lands",
                         "takeaway": "Referrals are the cheapest growth left"}
        spec.update(over)
        result = build_pptx.build(copy.deepcopy(spec))
        self.assertTrue(result["ok"], result.get("errors"))
        return result, read_snapshot(result["path"])

    def sectioned(self):
        return [
            {"id": "d1", "kind": "divider", "name": "Why referrals",
             "section": "Why referrals"},
            content_slide("s1", [text_block("b1", 3, 7, 12, 11, SHORT)],
                          section="Why referrals",
                          story="Partners already refer. The gap is the moment."),
            {"id": "d2", "kind": "divider", "name": "The ask", "section": "The ask"},
            content_slide("s2", [text_block("b2", 3, 7, 12, 11, SHORT)],
                          section="The ask",
                          story="Script the moment and name the owner."),
        ]

    def notes_of(self, path, index):
        deck = Presentation(path)
        slide = list(deck.slides)[index]
        if not slide.has_notes_slide:
            return ""
        return slide.notes_slide.notes_text_frame.text

    # -- notes from the story

    def test_the_story_becomes_the_speaker_notes(self):
        result, _ = self.build("story", self.sectioned())
        self.assertIn("The gap is the moment", self.notes_of(result["path"], 3))

    def test_explicit_notes_win_over_the_story(self):
        slides = self.sectioned()
        slides[1]["notes"] = "Say this instead."
        result, _ = self.build("storynotes", slides)
        notes = self.notes_of(result["path"], 3)
        self.assertIn("Say this instead.", notes)
        self.assertNotIn("The gap is the moment", notes)

    def test_notes_from_story_can_be_turned_off(self):
        result, _ = self.build("nostory", self.sectioned(), notesFromStory=False)
        self.assertEqual(self.notes_of(result["path"], 3), "")

    # -- the cover

    def test_a_deck_with_no_title_slide_gets_a_cover(self):
        result, snapshot = self.build("cover", self.sectioned())
        self.assertEqual(snapshot["slides"][0]["kind"], "title")
        self.assertEqual(snapshot["slides"][0]["name"], "Fidelity Deck")
        titles = [s["text"] for s in snapshot["slides"][0]["shapes"]
                  if s["role"] == "coverTitle"]
        self.assertEqual(titles, ["Fidelity Deck"])

    def test_the_cover_subtitle_comes_from_the_brief(self):
        _, snapshot = self.build("coversub", self.sectioned())
        subs = [s["text"] for s in snapshot["slides"][0]["shapes"]
                if s["role"] == "coverSubtitle"]
        self.assertEqual(subs, ["Land the referral recommendation"])

    def test_the_cover_can_carry_a_date(self):
        _, snapshot = self.build("coverdate", self.sectioned(),
                                 date="Draft for discussion")
        dates = [s["text"] for s in snapshot["slides"][0]["shapes"]
                 if s["role"] == "coverDate"]
        self.assertEqual(dates, ["Draft for discussion"])

    def test_a_deck_that_brings_its_own_title_slide_keeps_it(self):
        slides = [{"id": "t1", "kind": "title", "name": "My cover",
                   "headline": "My own cover"}] + self.sectioned()
        result, snapshot = self.build("owncover", slides)
        kinds = [s["kind"] for s in snapshot["slides"]]
        self.assertEqual(kinds[0], "title")
        self.assertEqual(snapshot["slides"][0]["id"], "t1")
        self.assertEqual(kinds.count("title"), 1)
        self.assertEqual(len(result["slides"]), len(slides) + 1)

    def test_the_cover_can_be_turned_off(self):
        result, snapshot = self.build("nocover", self.sectioned(), cover=False)
        self.assertEqual(snapshot["slides"][0]["kind"], "agenda")
        self.assertEqual(len(result["slides"]), 5)

    # -- the agenda

    def test_the_agenda_lists_the_dividers_with_a_line_each(self):
        _, snapshot = self.build("agenda", self.sectioned())
        agenda = snapshot["slides"][1]
        self.assertEqual(agenda["kind"], "agenda")
        entries = [s["text"] for s in agenda["shapes"] if s["role"] == "body"]
        self.assertEqual(len(entries), 2)
        self.assertIn("01 Why referrals", entries[0])
        self.assertIn("Partners already refer.", entries[0])
        self.assertIn("02 The ask", entries[1])

    def test_the_agenda_statement_comes_from_the_brief(self):
        _, snapshot = self.build("agendastate", self.sectioned())
        statements = [s["text"] for s in snapshot["slides"][1]["shapes"]
                      if s["role"] == "agendaStatement"]
        self.assertEqual(statements, ["Land the referral recommendation"])

    def test_the_agenda_statement_can_be_given(self):
        _, snapshot = self.build("agendagiven", self.sectioned(),
                                 agendaStatement="What we will cover")
        statements = [s["text"] for s in snapshot["slides"][1]["shapes"]
                      if s["role"] == "agendaStatement"]
        self.assertEqual(statements, ["What we will cover"])

    def test_a_deck_that_brings_its_own_agenda_keeps_it(self):
        slides = [{"id": "a1", "kind": "agenda", "name": "Agenda",
                   "headline": "My own agenda"}] + self.sectioned()
        _, snapshot = self.build("ownagenda", slides)
        kinds = [s["kind"] for s in snapshot["slides"]]
        self.assertEqual(kinds.count("agenda"), 1)
        self.assertEqual(snapshot["slides"][1]["id"], "a1")

    def test_a_deck_with_no_dividers_gets_no_agenda(self):
        slides = [content_slide("s1", [text_block("b1", 3, 7, 12, 11, SHORT)])]
        result, snapshot = self.build("noagenda", slides)
        kinds = [s["kind"] for s in snapshot["slides"]]
        self.assertEqual(kinds, ["title", "content"])
        self.assertTrue(any("no section dividers" in w for w in result["warnings"]),
                        result["warnings"])

    def test_too_many_sections_warns(self):
        slides = []
        for i in range(5):
            name = "Section %d" % (i + 1)
            slides.append({"id": "d%d" % i, "kind": "divider", "name": name,
                           "section": name})
            slides.append(content_slide("s%d" % i,
                                        [text_block("b%d" % i, 3, 7, 12, 11, SHORT)],
                                        section=name, story="A story for %s." % name))
        result, _ = self.build("manysections", slides)
        self.assertTrue(any("this geometry reads four" in w for w in result["warnings"]),
                        result["warnings"])

    # -- numbering

    def test_page_numbers_skip_the_generated_cover(self):
        result, snapshot = self.build("pages", self.sectioned())
        numbers = {}
        for slide in snapshot["slides"]:
            for shape in slide["shapes"]:
                if shape["role"] == "pageNumber":
                    numbers[slide["index"]] = shape["text"]
        self.assertNotIn(1, numbers, "the cover carries no page number")
        self.assertEqual(numbers[2], "1")
        self.assertEqual(numbers[3], "2")
        self.assertEqual(len(result["slides"]), 6)

    def test_the_dividers_are_still_numbered_in_deck_order(self):
        _, snapshot = self.build("ordinals", self.sectioned())
        ordinals = [(s["index"], s.get("ordinal")) for s in snapshot["slides"]
                    if s["kind"] == "divider"]
        self.assertEqual(ordinals, [(3, 1), (5, 2)])

    def test_the_generated_chrome_passes_the_deck_check(self):
        result, _ = self.build("chromecheck", self.sectioned())
        report = check_deck.check_deck(result["path"])
        errors = [f for f in report["findings"] if f["severity"] == "error"]
        self.assertEqual(errors, [], report["findings"])


# ------------------------------------------------- Q10 the style's type scale


# The shape of a real extracted token file: sizes and no floors, a dark palette,
# backgrounds that scope a fill and an ink and nothing else.
EXTRACTED_TOKENS = {
    "name": "Extracted",
    "palette": {"bg": "000000", "ink": "FFFFFF", "muted": "FFFFFF", "accent": "D9A066",
                "accent2": "BE854A", "cardFill": "2D2D2D", "cardBorder": "262626"},
    "fonts": {"display": "Montserrat Light", "body": "Calibri"},
    "typeScale": {"kicker": {"size": 9.5}, "headline": {"size": 27},
                  "body": {"size": 10.5}, "caption": {"size": 8.5},
                  "cardHeader": {"size": 14, "caps": True, "color": "accent"},
                  "pageNumber": {"size": 10}},
    "furniture": {"headline": {"x": 336, "y": 63, "w": 593, "h": 158},
                  "kicker": {"x": 336, "y": 29, "w": 281, "h": 16},
                  "body": {"x": 336, "y": 240, "w": 593, "h": 250},
                  "source": {"x": 336, "y": 499, "w": 593, "h": 22},
                  "pageNumber": {"x": 785, "y": 29, "w": 144, "h": 16}},
    "backgrounds": {"content": {"fill": "000000", "ink": "FFFFFF", "muted": "FFFFFF"},
                    "divider": {"fill": "000000", "ink": "FFFFFF", "muted": "FFFFFF"}},
}


def styled(tokens, mode="skill"):
    """(style, warnings) for a partial token file, as make_style builds it."""
    warnings = []
    spec = {"style": {"mode": mode, "tokens": copy.deepcopy(tokens)}}
    style, _ = build_pptx.make_style(spec, warnings)
    return style, warnings


class TypeScaleTests(unittest.TestCase):
    """Q10: a style's own type scale wins, and the house floors stay at home."""

    def test_a_token_size_under_the_house_floor_is_kept(self):
        style, _ = styled(EXTRACTED_TOKENS)
        self.assertEqual(style.role("kicker")["size"], 9.5)
        house, _ = build_pptx.make_style({"style": {"mode": "house"}}, [])
        self.assertEqual(house.role("kicker")["min"], 11.0, "the house floor still exists")

    def test_a_role_with_no_floor_falls_to_the_body_readability_floor(self):
        style, _ = styled(EXTRACTED_TOKENS)
        self.assertEqual(style.role("body")["min"], build_pptx.BODY_FLOOR)
        self.assertEqual(style.role("kicker")["min"], build_pptx.BODY_FLOOR)

    def test_a_caption_rank_role_falls_to_the_caption_floor(self):
        style, _ = styled(EXTRACTED_TOKENS)
        self.assertEqual(style.role("caption")["min"], build_pptx.CAPTION_FLOOR)
        self.assertEqual(style.role("pageNumber")["min"], build_pptx.CAPTION_FLOOR)

    def test_the_styles_own_floor_wins_when_it_names_one(self):
        tokens = copy.deepcopy(EXTRACTED_TOKENS)
        tokens["typeScale"]["body"] = {"size": 10.5, "min": 10}
        style, _ = styled(tokens)
        self.assertEqual(style.role("body")["min"], 10.0)

    def test_a_floor_never_lifts_the_size_the_style_asked_for(self):
        tokens = copy.deepcopy(EXTRACTED_TOKENS)
        tokens["typeScale"]["caption"] = {"size": 7}
        style, _ = styled(tokens)
        self.assertEqual(style.role("caption")["size"], 7.0)
        self.assertEqual(style.role("caption")["min"], 7.0)

    def test_a_ceiling_never_cuts_the_size_the_style_asked_for(self):
        tokens = copy.deepcopy(EXTRACTED_TOKENS)
        tokens["typeScale"]["headline"] = {"size": 40}      # house caps at 32
        style, _ = styled(tokens)
        self.assertEqual(style.role("headline")["size"], 40.0)
        self.assertGreaterEqual(style.role("headline")["max"], 40.0)

    def test_the_house_style_keeps_its_own_ranges(self):
        house, _ = build_pptx.make_style({"style": {"mode": "house"}}, [])
        self.assertEqual(house.role("body")["min"], 11.0)
        self.assertEqual(house.role("headline")["min"], 24.0)
        self.assertEqual(house.role("cardHeader")["min"], 15.0)

    def test_the_deck_wide_pass_walks_the_styles_own_range(self):
        """A body that will not fit steps below the house floor, down its own range."""
        tmp = tempfile.mkdtemp(prefix="blueprint_scale_test_")
        try:
            text = ("The fee is flat because a lean team can build it in four weeks, "
                    "the rollout runs on the client's own IT people rather than on "
                    "ours, and the office hours are capped at six.")
            slides = [content_slide("s1", [text_block("b_long", 3, 7, 8, 2, text)])]
            spec = mini_spec(os.path.join(tmp, "scale.pptx"), slides,
                             style={"mode": "skill",
                                    "tokens": copy.deepcopy(EXTRACTED_TOKENS)})
            result = build_pptx.build(copy.deepcopy(spec))
            self.assertTrue(result["ok"], result.get("errors"))
            self.assertEqual(result["typeScale"]["body"], 9.5)
            self.assertLess(result["typeScale"]["body"], 11.0, "the house floor")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------- Q11 the font chain


class FontChainTests(unittest.TestCase):
    """Q11: every role resolves from the style's own fonts."""

    def setUp(self):
        if not build_pptx.FONTS.has("Montserrat"):
            self.skipTest("Montserrat is not installed on this machine")

    def test_an_uninstalled_body_font_falls_back_to_the_styles_display_family(self):
        style, warnings = styled(EXTRACTED_TOKENS)
        self.assertEqual(style.font("body"), "Montserrat")
        self.assertTrue(any("Calibri" in w for w in warnings), warnings)

    def test_mono_uses_the_body_font_when_the_style_names_none(self):
        style, _ = styled(EXTRACTED_TOKENS)
        self.assertEqual(style.font("mono"), "Montserrat")
        self.assertNotEqual(style.font("mono"), "Menlo")

    def test_a_style_that_names_a_mono_face_gets_it(self):
        tokens = copy.deepcopy(EXTRACTED_TOKENS)
        tokens["fonts"]["mono"] = "Menlo"
        style, _ = styled(tokens)
        self.assertEqual(style.font("mono"), "Menlo")

    def test_the_house_style_keeps_its_own_mono(self):
        house, _ = build_pptx.make_style({"style": {"mode": "house"}}, [])
        self.assertEqual(house.font("mono"), "Menlo")

    def test_the_chain_puts_the_styles_own_fonts_first(self):
        style, _ = styled(EXTRACTED_TOKENS)
        chain = style.font_chain("body")
        self.assertEqual(chain[0], "Calibri")
        self.assertIn("Montserrat Light", chain)
        self.assertLess(chain.index("Montserrat Light"), chain.index("Arial"))
        self.assertNotIn("Menlo", style.font_chain("mono"))

    def test_no_run_in_a_built_deck_falls_through_to_a_house_font(self):
        tmp = tempfile.mkdtemp(prefix="blueprint_fontrun_test_")
        try:
            table = {"id": "b_t", "type": "table", "col": 3, "row": 7, "cols": 14,
                     "rows": 4, "content": {"headers": ["Phase", "Cost"],
                                            "rows": [["Build", "$18,500"]]}}
            stat = {"id": "b_s", "type": "stat", "col": 18, "row": 7, "cols": 12,
                    "rows": 4, "content": {"items": [{"value": "38%", "label": "Of it"}]}}
            slides = [content_slide("s1", [table, stat], source="Northwind rate card")]
            spec = mini_spec(os.path.join(tmp, "fonts.pptx"), slides,
                             style={"mode": "skill",
                                    "tokens": copy.deepcopy(EXTRACTED_TOKENS)})
            result = build_pptx.build(copy.deepcopy(spec))
            self.assertTrue(result["ok"], result.get("errors"))
            seen = set()
            for slide in Presentation(result["path"]).slides:
                for shape in slide.shapes:
                    frames = []
                    if getattr(shape, "has_text_frame", False):
                        frames.append(shape.text_frame)
                    if getattr(shape, "has_table", False):
                        frames.extend(cell.text_frame for row in shape.table.rows
                                      for cell in row.cells)
                    for frame in frames:
                        for para in frame.paragraphs:
                            for run in para.runs:
                                if run.text.strip() and run.font.name:
                                    seen.add(run.font.name)
            self.assertEqual(seen, {"Montserrat"}, seen)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


# ------------------------------------------------------------- Q13 emphasis


def emph_spec(out_path, blocks, **over):
    """One content slide of emphasised blocks in the house style."""
    return mini_spec(out_path, [content_slide("s1", blocks, **over)])


class EmphasisTests(unittest.TestCase):
    """Q13: `emphasis: true` gets the style's accent treatment, per role."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="blueprint_emph_test_")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def build(self, name, blocks, **over):
        spec = emph_spec(os.path.join(self.tmp, name + ".pptx"), blocks, **over)
        spec.update(over)
        result = build_pptx.build(copy.deepcopy(spec))
        self.assertTrue(result["ok"], result.get("errors"))
        return (result, read_snapshot(result["path"]),
                list(Presentation(result["path"]).slides)[0])

    def cards(self, **over):
        block = {"id": "b_cards", "type": "cards", "col": 3, "row": 7, "cols": 28,
                 "rows": 6, "content": {"items": [
                     {"header": "One", "body": "First card body."},
                     {"header": "Two", "body": "Second card body."}]}}
        block.update(over)
        return block

    def test_an_emphasised_cards_block_takes_the_accent_border(self):
        _, _, slide = self.build("cards", [self.cards(emphasis=True)])
        cards = rounded(slide)
        self.assertEqual(len(cards), 2)
        for card in cards:
            self.assertAlmostEqual(card.line.width / PT, 1.0 + build_pptx.EMPH_BORDER_PT,
                                   delta=0.05)
            self.assertEqual(str(card.line.color.rgb), "3B5BDB")

    def test_an_emphasised_card_header_is_the_accent(self):
        _, _, slide = self.build("cardhead", [self.cards(emphasis=True)])
        head = None
        for card in rounded(slide):
            head = card.text_frame.paragraphs[0].runs[0]
            break
        self.assertEqual(str(head.font.color.rgb), "3B5BDB")

    def test_a_plain_cards_block_is_left_alone(self):
        _, _, slide = self.build("cardsplain", [self.cards()])
        for card in rounded(slide):
            self.assertAlmostEqual(card.line.width / PT, 1.0, delta=0.05)
            self.assertEqual(str(card.line.color.rgb), "D1D5DB")

    def test_an_emphasised_stat_gets_a_hairline_under_the_value(self):
        stat = {"id": "b_stat", "type": "stat", "col": 3, "row": 7, "cols": 12,
                "rows": 4, "emphasis": True,
                "content": {"items": [{"value": "$60K", "label": "flat fee"}]}}
        _, snapshot, _ = self.build("stat", [stat])
        rules = [s for s in shapes_for(snapshot, "b_stat") if s["role"] == "hairline"]
        self.assertEqual(len(rules), 1)
        value = [s for s in shapes_for(snapshot, "b_stat") if s["role"] == "stat"][0]
        self.assertGreater(rules[0]["y"], value["y"])
        self.assertLess(rules[0]["y"], value["y"] + 80)
        self.assertAlmostEqual(rules[0]["w"], value["w"], delta=1.0)

    def test_an_emphasised_text_block_gets_an_accent_bar_and_a_bold_lead(self):
        block = text_block("b_text", 3, 7, 14, 4,
                           "The fee is flat. Everything else follows from that.",
                           emphasis=True)
        _, snapshot, slide = self.build("text", [block])
        bar = [s for s in shapes_for(snapshot, "b_text") if s["role"] == "hairline"][0]
        body = [s for s in shapes_for(snapshot, "b_text") if s["role"] == "body"][0]
        self.assertAlmostEqual(bar["w"], build_pptx.EMPH_BAR_PT, delta=0.01)
        self.assertAlmostEqual(bar["x"], 60.0, delta=0.5)
        self.assertGreater(body["x"], bar["x"] + build_pptx.EMPH_BAR_PT)
        runs = None
        for shape in slide.shapes:
            if shape.has_text_frame and "flat" in shape.text_frame.text:
                runs = shape.text_frame.paragraphs[0].runs
        self.assertEqual(len(runs), 2)
        self.assertTrue(runs[0].font.bold)
        self.assertEqual(runs[0].text, "The fee is flat.")
        self.assertFalse(bool(runs[1].font.bold))

    def test_an_emphasised_bullets_block_gets_the_bar_too(self):
        block = {"id": "b_bul", "type": "bullets", "col": 3, "row": 7, "cols": 14,
                 "rows": 4, "emphasis": True,
                 "content": {"items": ["Below the usual rate. A lean team can build it.",
                                       "No subscription."]}}
        _, snapshot, _ = self.build("bullets", [block])
        bars = [s for s in shapes_for(snapshot, "b_bul") if s["role"] == "hairline"]
        self.assertEqual(len(bars), 1)
        body = [s for s in shapes_for(snapshot, "b_bul") if s["role"] == "bullets"][0]
        self.assertGreater(body["x"], bars[0]["x"])

    def test_an_emphasised_takeaway_is_inverse(self):
        block = {"id": "b_take", "type": "takeaway", "col": 3, "row": 15, "cols": 28,
                 "rows": 2, "emphasis": True,
                 "content": {"text": "One flat fee, four weeks, 90 days of office hours"}}
        _, _, slide = self.build("take", [block])
        bar = rounded(slide)[0]
        self.assertEqual(str(bar.fill.fore_color.rgb), "3B5BDB")
        run = bar.text_frame.paragraphs[0].runs[0]
        self.assertEqual(str(run.font.color.rgb), "FFFFFF")

    def test_an_emphasised_table_head_rule_is_heavier(self):
        block = {"id": "b_t", "type": "table", "col": 3, "row": 7, "cols": 14,
                 "rows": 5, "emphasis": True,
                 "content": {"headers": ["Phase", "Weeks"],
                             "rows": [["Build", "2"], ["Rollout", "2"]]}}
        _, _, slide = self.build("table", [block])
        table = [s for s in slide.shapes if getattr(s, "has_table", False)][0].table
        width = rule_width(table.cell(0, 0))
        self.assertAlmostEqual(width, build_pptx.EMPH_RULE_PT, delta=0.01)

    def test_an_emphasised_chart_puts_one_series_in_the_accent(self):
        block = {"id": "b_c", "type": "chart", "col": 3, "row": 7, "cols": 14,
                 "rows": 6, "emphasis": True,
                 "content": {"kind": "column", "categories": ["A", "B"],
                             "series": [{"name": "Northwind", "values": [3, 5]},
                                        {"name": "Others", "values": [2, 2]}]}}
        _, _, slide = self.build("chart", [block])
        chart = [s for s in slide.shapes if getattr(s, "has_chart", False)][0].chart
        series = list(chart.plots[0].series)
        self.assertEqual(str(series[0].format.fill.fore_color.rgb), "3B5BDB")
        self.assertEqual(str(series[1].format.fill.fore_color.rgb), "6B7280")

    def test_the_named_series_can_be_the_emphasised_one(self):
        block = {"id": "b_c", "type": "chart", "col": 3, "row": 7, "cols": 14,
                 "rows": 6, "emphasis": True,
                 "content": {"kind": "column", "categories": ["A", "B"],
                             "emphasisSeries": "Others",
                             "series": [{"name": "Northwind", "values": [3, 5]},
                                        {"name": "Others", "values": [2, 2]}]}}
        _, _, slide = self.build("chart2", [block])
        chart = [s for s in slide.shapes if getattr(s, "has_chart", False)][0].chart
        series = list(chart.plots[0].series)
        self.assertEqual(str(series[1].format.fill.fore_color.rgb), "3B5BDB")

    def test_the_snapshot_marks_the_emphasised_shapes(self):
        blocks = [self.cards(emphasis=True),
                  text_block("b_plain", 3, 14, 14, 3, SHORT)]
        _, snapshot, _ = self.build("marks", blocks)
        for shape in shapes_for(snapshot, "b_cards"):
            self.assertTrue(shape.get("emphasis"), shape)
        for shape in shapes_for(snapshot, "b_plain"):
            self.assertNotIn("emphasis", shape)

    def test_an_emphasised_slide_still_passes_the_deck_check(self):
        blocks = [self.cards(emphasis=True),
                  text_block("b_text", 3, 14, 14, 3, SHORT, emphasis=True)]
        result, _, _ = self.build("emphcheck", blocks)
        report = check_deck.check_deck(result["path"])
        errors = [f for f in report["findings"] if f["severity"] == "error"]
        self.assertEqual(errors, [], report["findings"])


def rule_width(cell):
    """The width in points of the bottom rule on a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    line = tcPr.findall(qn("a:lnB"))
    if not line:
        return 0.0
    return float(line[0].get("w") or 0) / PT


def rule_color(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    line = tcPr.findall(qn("a:lnB"))
    if not line:
        return None
    found = line[0].findall(qn("a:solidFill"))
    if not found:
        return None
    return found[0].findall(qn("a:srgbClr"))[0].get("val")


# --------------------------------------------------------------- Q14 tables


class TableStyleTests(unittest.TestCase):
    """Q14: Tufte defaults, so a table is rules and figures, not a grey box."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="blueprint_table_test_")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    FEE = {"headers": ["Phase", "Duration", "Cost"],
           "rows": [["Build", "2 weeks", "$18,500"],
                    ["Rollout and training", "2 weeks", "$9,500"],
                    ["Office hours", "90 days", "$4,000"]]}

    def build(self, name, content=None, style=None, **over):
        block = {"id": "b_t", "type": "table", "col": 3, "row": 7, "cols": 26,
                 "rows": 8, "content": content or copy.deepcopy(self.FEE)}
        block.update(over)
        spec = mini_spec(os.path.join(self.tmp, name + ".pptx"),
                         [content_slide("s1", [block])])
        if style:
            spec["style"] = style
        result = build_pptx.build(copy.deepcopy(spec))
        self.assertTrue(result["ok"], result.get("errors"))
        slide = list(Presentation(result["path"]).slides)[0]
        table = [s for s in slide.shapes if getattr(s, "has_table", False)][0].table
        return result, read_snapshot(result["path"]), slide, table

    def test_no_header_fill_when_the_style_names_none(self):
        style = {"mode": "skill", "tokens": copy.deepcopy(EXTRACTED_TOKENS)}
        _, _, _, table = self.build("nofill", style=style)
        self.assertIn("BACKGROUND", str(table.cell(0, 0).fill.type))
        self.assertFalse(table.first_row)

    def test_a_named_header_fill_is_drawn(self):
        _, _, _, table = self.build("fill")        # house names one
        self.assertEqual(str(table.cell(0, 0).fill.fore_color.rgb), "F3F4F6")
        self.assertTrue(table.first_row)

    def test_the_column_head_is_caps_in_the_accent_at_the_card_sub_size(self):
        result, _, _, table = self.build("head")
        run = table.cell(0, 0).text_frame.paragraphs[0].runs[0]
        self.assertEqual(run.text, "PHASE")
        self.assertEqual(str(run.font.color.rgb), "3B5BDB")
        cell_size = result["typeScale"]["tableCell"]
        self.assertLessEqual(run.font.size.pt, cell_size)

    def test_a_low_contrast_accent_gives_way_to_ink(self):
        tokens = copy.deepcopy(EXTRACTED_TOKENS)
        tokens["palette"]["bg"] = "FFFFFF"
        tokens["palette"]["ink"] = "1F2937"
        tokens["backgrounds"] = {"content": {"fill": "FFFFFF"}}
        style, _ = styled(tokens)
        self.assertEqual(style.accent_key("content"), "ink",
                         "copper on white cannot carry a column head")
        _, _, _, table = self.build("contrast",
                                    style={"mode": "skill", "tokens": tokens})
        run = table.cell(0, 0).text_frame.paragraphs[0].runs[0]
        self.assertEqual(str(run.font.color.rgb), "1F2937")

    def test_one_accent_rule_under_the_header(self):
        _, _, _, table = self.build("headrule")
        self.assertAlmostEqual(rule_width(table.cell(0, 0)),
                               build_pptx.TABLE_HEAD_RULE_PT, delta=0.01)
        self.assertEqual(rule_color(table.cell(0, 0)), "3B5BDB")

    def test_hairlines_run_between_the_rows_and_not_under_the_last(self):
        _, _, _, table = self.build("rowrules")
        self.assertAlmostEqual(rule_width(table.cell(1, 0)),
                               build_pptx.TABLE_ROW_RULE_PT, delta=0.01)
        self.assertEqual(rule_color(table.cell(1, 0)), "D1D5DB")
        self.assertEqual(rule_width(table.cell(3, 0)), 0.0, "the last row closes with air")

    def test_numeric_columns_go_right_and_text_goes_left(self):
        _, _, _, table = self.build("align")
        money = table.cell(1, 2).text_frame.paragraphs[0]
        words = table.cell(1, 1).text_frame.paragraphs[0]
        self.assertEqual(str(money.alignment), "RIGHT (3)")
        self.assertEqual(str(words.alignment), "LEFT (1)")

    def test_one_figure_in_a_column_of_phrases_stays_left(self):
        content = {"headers": ["Structure", "Year-one revenue"],
                   "rows": [["Fixed fee", "$60K"],
                            ["Share of parent", "$50K to $75K"],
                            ["Three tiers", "Depends on tier"]]}
        _, _, _, table = self.build("ragged", content=content)
        for row in (1, 2, 3):
            para = table.cell(row, 1).text_frame.paragraphs[0]
            self.assertEqual(str(para.alignment), "LEFT (1)", row)
        self.assertEqual(build_pptx.numeric_columns(content["headers"],
                                                    content["rows"]), set())

    def test_what_counts_as_a_number(self):
        for value in ("12", "7.7%", "$60K", "\u00a31,200", "4.1x", "(3)", "-12", 42, 3.5):
            self.assertTrue(build_pptx.is_numeric_cell(value), value)
        for value in ("2 weeks", "90 days", "Build", "", None, True,
                      "6 hours over 90 days", "Design lead, strategist"):
            self.assertFalse(build_pptx.is_numeric_cell(value), value)

    def test_the_first_column_is_medium_when_it_reads_as_a_label_column(self):
        _, _, _, table = self.build("labels")
        self.assertTrue(table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.bold)
        self.assertFalse(bool(table.cell(1, 1).text_frame.paragraphs[0].runs[0].font.bold))

    def test_a_column_of_prose_is_not_a_label_column(self):
        content = {"headers": ["Firm", "What they sell"],
                   "rows": [["A translation of an existing identity without foundational "
                             "work here", "Strategy"],
                            ["A point solution product or a subscription service too",
                             "Software"]]}
        _, _, _, table = self.build("noplabels", content=content)
        self.assertFalse(bool(table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.bold))

    def test_rating_cells_become_circles_in_the_accent(self):
        content = {"headers": ["Firm", "Depth"],
                   "rows": [["Interbrand", "\u25cf\u25cf\u25d0\u25cb"],
                            ["MetaDesign", "\u25cf\u25cb\u25cb\u25cb"]]}
        _, snapshot, slide, table = self.build("ratings", content=content)
        marks = [s for s in shapes_for(snapshot, "b_t") if s["role"] == "tableRating"]
        self.assertEqual(len(marks), 8)
        self.assertEqual(table.cell(1, 1).text, build_pptx.EMPTY_CELL,
                         "the glyphs give way to the drawn marks")
        drawn = ovals(slide)
        self.assertEqual(len(drawn), 8)
        self.assertEqual(str(drawn[0].line.color.rgb), "3B5BDB")
        self.assertEqual(str(drawn[0].fill.fore_color.rgb), "3B5BDB")
        self.assertIn("BACKGROUND", str(drawn[3].fill.type), "the empty mark is hollow")

    def test_a_fraction_and_a_lone_digit_read_as_ratings(self):
        self.assertEqual(build_pptx.rating_of("3/5"), (3.0, 5))
        self.assertEqual(build_pptx.rating_of("4"), (4.0, 5))
        self.assertEqual(build_pptx.rating_of("\u25cf\u25cb"), (1.0, 2))
        self.assertIsNone(build_pptx.rating_of("12"))
        self.assertIsNone(build_pptx.rating_of("Build"))
        self.assertEqual(build_pptx.rating_marks(2.5, 4), [1.0, 1.0, 0.5, 0.0])

    def test_a_lone_digit_in_a_column_of_numbers_stays_a_number(self):
        content = {"headers": ["Role", "Hours", "Cost"],
                   "rows": [["Design lead", "2", "$1,000"],
                            ["Analyst", "24", "$9,500"]]}
        _, snapshot, _, table = self.build("hours", content=content)
        self.assertEqual([s for s in shapes_for(snapshot, "b_t")
                          if s["role"] == "tableRating"], [])
        self.assertEqual(table.cell(1, 1).text, "2")

    def test_the_row_height_comes_from_the_deck_wide_cell_size(self):
        result, snapshot, _, table = self.build("rowheight")
        size = result["typeScale"]["tableCell"]
        lead = 1.25
        want = size * lead + 2 * build_pptx.TABLE_PAD_Y
        self.assertAlmostEqual(table.rows[1].height / PT, want, delta=1.0)
        entry = shapes_for(snapshot, "b_t")[0]
        self.assertLess(entry["h"], 8 * 30, "the table is its own height, not the block's")

    def test_a_wrapping_cell_gets_a_taller_row(self):
        content = {"headers": ["Phase", "Team"],
                   "rows": [["Build", "Design lead, strategist and analyst with "
                                      "the client IT team alongside them"],
                            ["Office hours", "Whoever the question needs"]]}
        _, _, _, table = self.build("wrap", content=content)
        self.assertGreater(table.rows[1].height, table.rows[2].height)

    def test_an_empty_cell_keeps_its_row_the_height_it_was_measured_at(self):
        """A bare empty run grows the row in the renderer, which moves the marks."""
        content = {"headers": ["", "One"], "rows": [["Interbrand", ""], ["Landor", "x"]]}
        _, _, _, table = self.build("emptycell", content=content)
        self.assertEqual(table.cell(1, 1).text, build_pptx.EMPTY_CELL)
        self.assertEqual(table.cell(0, 0).text, build_pptx.EMPTY_CELL)
        heights = [round(r.height / PT, 2) for r in table.rows]
        self.assertEqual(len(set(heights)), 1, heights)

    def test_a_rating_mark_sits_inside_its_own_row(self):
        content = {"headers": ["Firm", "Depth"],
                   "rows": [["Interbrand", "\u25cf\u25cb"], ["Landor", "\u25cb\u25cb"],
                            ["Column Five", "\u25cf\u25cf"]]}
        _, snapshot, _, table = self.build("markrows", content=content)
        entry = [s for s in shapes_for(snapshot, "b_t") if s["kind"] == "table"][0]
        heights = [r.height / PT for r in table.rows]
        marks = [s for s in shapes_for(snapshot, "b_t") if s["role"] == "tableRating"]
        self.assertEqual(len(marks), 6)
        for i in range(3):
            top = entry["y"] + sum(heights[:i + 1])
            for mark in marks[i * 2:i * 2 + 2]:
                self.assertGreaterEqual(mark["y"], top - 0.5)
                self.assertLessEqual(mark["y"] + mark["h"], top + heights[i + 1] + 0.5)

    def test_the_table_still_passes_the_deck_check(self):
        result, _, _, _ = self.build("tablecheck")
        report = check_deck.check_deck(result["path"])
        errors = [f for f in report["findings"] if f["severity"] == "error"]
        self.assertEqual(errors, [], report["findings"])


# --------------------------------------------------------- Q9 the column module


class ModuleTests(unittest.TestCase):
    """Q9: every left edge and width lands on the style's 12-column module."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="blueprint_module_test_")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def build(self, name, slides, **over):
        spec = mini_spec(os.path.join(self.tmp, name + ".pptx"), slides, **over)
        result = build_pptx.build(copy.deepcopy(spec))
        self.assertTrue(result["ok"], result.get("errors"))
        return result, read_snapshot(result["path"])

    def test_the_module_comes_from_the_body_slot(self):
        module = fit.module_of({"x": 60, "y": 180, "w": 840, "h": 300}, 30.0)
        self.assertEqual(module["cols"], 12)
        self.assertAlmostEqual(module["module"], (840 - 11 * 30) / 12.0, delta=0.01)
        self.assertAlmostEqual(module["lefts"][0], 60.0, delta=0.01)
        self.assertAlmostEqual(module["lefts"][4], 60 + 4 * (42.5 + 30), delta=0.01)
        self.assertAlmostEqual(module["rights"][-1], 900.0, delta=0.01)

    def test_a_style_with_no_body_slot_has_no_module(self):
        self.assertIsNone(fit.module_of(None, 30.0))
        self.assertIsNone(fit.module_of({"x": 60, "w": 0}, 30.0))

    def test_three_columns_land_on_the_module(self):
        slides = [content_slide("s1", [
            text_block("b_a", 2, 7, 10, 5, SHORT),
            text_block("b_b", 12, 7, 10, 5, SHORT),
            text_block("b_c", 22, 7, 10, 5, SHORT),
        ])]
        _, snapshot = self.build("three", slides)
        lefts = [shapes_for(snapshot, b)[0]["x"] for b in ("b_a", "b_b", "b_c")]
        self.assertEqual([round(x, 1) for x in lefts], [60.0, 350.0, 640.0])
        widths = [shapes_for(snapshot, b)[0]["w"] for b in ("b_a", "b_b", "b_c")]
        self.assertEqual([round(w, 1) for w in widths], [260.0, 260.0, 260.0])

    def test_the_move_is_recorded_as_module_with_the_points(self):
        slides = [content_slide("s1", [text_block("b_a", 2, 7, 10, 11, SHORT)])]
        result, snapshot = self.build("record", slides)
        move = [m for m in result["adjusted"] if m["blockId"] == "b_a"][0]
        self.assertEqual(move["reason"], "module")
        self.assertEqual(move["from"], {"col": 2, "row": 7, "cols": 10, "rows": 11})
        self.assertEqual(move["snappedPt"], {"x": 60.0, "w": 260.0})
        shape = shapes_for(snapshot, "b_a")[0]
        self.assertEqual(shape["adjusted"]["snappedPt"], {"x": 60.0, "w": 260.0})

    def test_an_edge_already_on_the_module_is_not_recorded(self):
        slides = [content_slide("s1", [text_block("b_a", 3, 7, 28, 11, SHORT)])]
        result, _ = self.build("ongrid", slides)
        self.assertNotIn("adjusted", result)

    def test_strict_leaves_the_module_alone(self):
        slides = [content_slide("s1", [text_block("b_a", 2, 7, 10, 5, SHORT)])]
        result, snapshot = self.build("strictmod", slides, fidelity="strict")
        self.assertEqual(shapes_for(snapshot, "b_a")[0]["x"], 30.0)
        self.assertNotIn("adjusted", result)

    def test_an_edge_more_than_two_cells_off_is_left_where_it_is(self):
        # The 12-column module's pitch is 72.5 pt, so every edge is within two
        # cells of one of its edges; the tolerance bites on a coarser module.
        wide = fit.module_of({"x": 60, "y": 180, "w": 840, "h": 300}, 30.0, columns=3)
        self.assertEqual([round(x, 1) for x in wide["lefts"]], [60.0, 350.0, 640.0])
        self.assertIsNone(fit.nearest_edge(wide["lefts"], 190.0, 2 * 30.0))
        self.assertAlmostEqual(fit.nearest_edge(wide["lefts"], 320.0, 2 * 30.0),
                               350.0, delta=0.01)
        module = fit.module_of({"x": 60, "y": 180, "w": 840, "h": 300}, 30.0)
        self.assertAlmostEqual(fit.nearest_edge(module["lefts"], 320.0, 2 * 30.0),
                               350.0, delta=0.01)

    def test_snapped_blocks_do_not_land_on_each_other(self):
        slides = [content_slide("s1", [
            text_block("b_a", 3, 7, 9, 5, SHORT),
            text_block("b_b", 13, 7, 9, 5, SHORT),
        ])]
        result, snapshot = self.build("nocollide", slides)
        first = shapes_for(snapshot, "b_a")[0]
        second = shapes_for(snapshot, "b_b")[0]
        self.assertLessEqual(first["x"] + first["w"], second["x"] + 0.01)
        report = check_deck.check_deck(result["path"])
        self.assertEqual([f for f in report["findings"] if f["severity"] == "error"],
                         [], report["findings"])

    def test_the_module_runs_before_the_rebalance(self):
        slides = [content_slide("s1", [text_block("b_a", 2, 7, 10, 3, SHORT)])]
        result, snapshot = self.build("order", slides)
        move = [m for m in result["adjusted"] if m["blockId"] == "b_a"][0]
        self.assertTrue(move["reason"].startswith("module"), move["reason"])
        self.assertGreater(shapes_for(snapshot, "b_a")[0]["y"], 6 * 30)


# ------------------------------------------------------- Q16 background merge


class BackgroundMergeTests(unittest.TestCase):
    """Q16: a scoped background takes its missing keys from the style's palette."""

    def test_a_scoped_background_without_an_accent_uses_the_palette(self):
        style, _ = styled(EXTRACTED_TOKENS)
        self.assertEqual(style.hex_of("accent", "divider"), "D9A066")
        self.assertEqual(style.hex_of("accent2", "divider"), "BE854A")
        house, _ = build_pptx.make_style({"style": {"mode": "house"}}, [])
        self.assertEqual(house.hex_of("accent", "divider"), "FFFFFF",
                         "house still dresses its own dividers")

    def test_a_scoped_muted_is_derived_from_the_kinds_ink_and_fill(self):
        tokens = copy.deepcopy(EXTRACTED_TOKENS)
        tokens["backgrounds"]["divider"] = {"fill": "000000", "ink": "FFFFFF"}
        style, _ = styled(tokens)
        self.assertEqual(style.hex_of("muted", "divider"),
                         build_pptx.blend("FFFFFF", "000000", 0.65))

    def test_a_key_the_style_names_is_kept(self):
        tokens = copy.deepcopy(EXTRACTED_TOKENS)
        tokens["backgrounds"]["divider"] = {"fill": "000000", "ink": "FFFFFF",
                                            "accent": "CDC5B1", "muted": "FFFFFF"}
        style, _ = styled(tokens)
        self.assertEqual(style.hex_of("accent", "divider"), "CDC5B1")
        self.assertEqual(style.hex_of("muted", "divider"), "FFFFFF")

    def test_a_scoped_line_falls_through_to_the_palette(self):
        style, _ = styled(EXTRACTED_TOKENS)
        self.assertEqual(style.hex_of("line", "divider"), style.hex_of("line", "content"))

    def test_the_palette_bg_becomes_the_content_fill(self):
        tokens = copy.deepcopy(EXTRACTED_TOKENS)
        tokens.pop("backgrounds")
        style, _ = styled(tokens)
        self.assertEqual(style.background("content"), "000000")

    def test_a_style_that_scopes_its_own_fill_keeps_it(self):
        tokens = copy.deepcopy(EXTRACTED_TOKENS)
        tokens["backgrounds"]["content"] = {"fill": "111111"}
        style, _ = styled(tokens)
        self.assertEqual(style.background("content"), "111111")

# ------------------------------------------------------- Q7 consistency findings


class ConsistencyTests(unittest.TestCase):
    """Q7: the findings a reader would call "this deck was not proofread"."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="blueprint_consist_test_")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def check(self, name, slides, spec_out=False, **over):
        spec = mini_spec(os.path.join(self.tmp, name + ".pptx"), slides, **over)
        result = build_pptx.build(copy.deepcopy(spec))
        self.assertTrue(result["ok"], result.get("errors"))
        spec_path = None
        if spec_out:
            spec_path = os.path.join(self.tmp, name + ".spec.json")
            with open(spec_path, "w") as handle:
                json.dump(spec, handle)
        report = check_deck.check_deck(result["path"], spec_path)
        return result, report

    def kinds(self, report, kind):
        return [f for f in report["findings"] if f["kind"] == kind]

    def content(self, slide_id, blocks, **over):
        one = content_slide(slide_id, blocks, **over)
        one.setdefault("section", "Competitive analysis")
        return one

    def test_a_clean_deck_fires_no_consistency_findings(self):
        cards = {"id": "b_c", "type": "cards", "col": 3, "row": 7, "cols": 28, "rows": 6,
                 "content": {"items": [{"header": "One", "body": "The first card body."},
                                       {"header": "Two", "body": "The second card body."}]}}
        slides = [self.content("s1", [cards], headline="The offer holds at one flat fee",
                               source="Northwind rate card, 2021")]
        _, report = self.check("clean", slides)
        for kind in ("kicker-drift", "headline-case", "not-parallel", "empty-band",
                     "source-missing", "bullet-count", "duplicate-headline"):
            self.assertEqual(self.kinds(report, kind), [], kind)

    def test_a_kicker_off_the_pattern_is_found(self):
        slides = [self.content("s1", [text_block("b_t", 3, 7, 28, 10, SHORT)],
                               kicker="Competitive analysis / Strategy consultancies")]
        _, report = self.check("kicker", slides)
        found = self.kinds(report, "kicker-drift")
        self.assertEqual(len(found), 1, report["findings"])
        self.assertEqual(found[0]["slide"], 1)
        self.assertIn("{section} | {topic}", found[0]["message"])

    def test_a_kicker_of_one_field_is_not_drift(self):
        slides = [self.content("s1", [text_block("b_t", 3, 7, 28, 10, SHORT)],
                               kicker="Competitive analysis")]
        _, report = self.check("onefield", slides)
        self.assertEqual(self.kinds(report, "kicker-drift"), [])

    def test_a_title_case_headline_is_found(self):
        slides = [self.content("s1", [text_block("b_t", 3, 7, 28, 10, SHORT)],
                               headline="Three Ways To Grow The Referral Pipeline")]
        _, report = self.check("titlecase", slides)
        found = self.kinds(report, "headline-case")
        self.assertEqual(len(found), 1, report["findings"])
        self.assertIn("sentence case", found[0]["message"])

    def test_a_headline_that_ends_in_a_full_stop_is_found(self):
        slides = [self.content("s1", [text_block("b_t", 3, 7, 28, 10, SHORT)],
                               headline="The fee is flat and the team is lean.")]
        _, report = self.check("stop", slides)
        found = self.kinds(report, "headline-case")
        self.assertEqual(len(found), 1, report["findings"])
        self.assertIn("full stop", found[0]["message"])

    def test_a_headline_with_proper_nouns_is_not_title_case(self):
        slides = [self.content("s1", [text_block("b_t", 3, 7, 28, 10, SHORT)],
                               headline="Interbrand and MetaDesign sell AI brand strategy")]
        _, report = self.check("propernouns", slides)
        self.assertEqual(self.kinds(report, "headline-case"), [])

    def test_cards_that_do_not_close_alike_are_found(self):
        cards = {"id": "b_c", "type": "cards", "col": 3, "row": 7, "cols": 28, "rows": 6,
                 "content": {"items": [{"header": "One", "body": "A card body."},
                                       {"header": "Two", "body": "A card body"}]}}
        _, report = self.check("stops", [self.content("s1", [cards])])
        found = self.kinds(report, "not-parallel")
        self.assertTrue(found, report["findings"])
        self.assertEqual(found[0]["blockId"], "b_c")
        self.assertIn("full stop", found[0]["message"])

    def test_cards_whose_line_counts_diverge_are_found(self):
        long_body = ("A card body that runs on for long enough to take four lines of "
                     "copy in its own column, which its neighbour does not, and then "
                     "keeps going for another sentence so the two cards cannot be "
                     "read as one set of anything.")
        cards = {"id": "b_c", "type": "cards", "col": 3, "row": 7, "cols": 28, "rows": 6,
                 "content": {"items": [{"header": "One", "body": long_body},
                                       {"header": "Two", "body": "Short."}]}}
        _, report = self.check("lines", [self.content("s1", [cards])])
        found = [f for f in self.kinds(report, "not-parallel") if "lines" in f["message"]]
        self.assertTrue(found, report["findings"])

    def test_more_than_seven_bullets_are_found(self):
        block = {"id": "b_b", "type": "bullets", "col": 3, "row": 7, "cols": 28,
                 "rows": 10, "content": {"items": ["Point %d" % i for i in range(1, 10)]}}
        _, report = self.check("bullets", [self.content("s1", [block])])
        found = self.kinds(report, "bullet-count")
        self.assertEqual(len(found), 1, report["findings"])
        self.assertIn("9 bullets", found[0]["message"])
        self.assertEqual(found[0]["blockId"], "b_b")

    def test_seven_bullets_are_allowed(self):
        block = {"id": "b_b", "type": "bullets", "col": 3, "row": 7, "cols": 28,
                 "rows": 10, "content": {"items": ["Point %d" % i for i in range(1, 8)]}}
        _, report = self.check("sevenbullets", [self.content("s1", [block])])
        self.assertEqual(self.kinds(report, "bullet-count"), [])

    def test_an_empty_band_is_found_under_strict(self):
        slides = [self.content("s1", [text_block("b_t", 3, 7, 28, 2, SHORT)])]
        _, report = self.check("band", slides, fidelity="strict")
        found = self.kinds(report, "empty-band")
        self.assertEqual(len(found), 1, report["findings"])
        self.assertEqual(found[0]["blockId"], "b_t")
        self.assertIn("balanced", found[0]["message"])

    def test_balanced_fills_the_band_so_nothing_is_found(self):
        slides = [self.content("s1", [text_block("b_t", 3, 7, 28, 2, SHORT)])]
        _, report = self.check("bandbalanced", slides)
        self.assertEqual(self.kinds(report, "empty-band"), [])

    def test_a_deck_of_figures_with_no_sources_is_found_once(self):
        slides = []
        for i in range(3):
            stat = {"id": "b_s%d" % i, "type": "stat", "col": 3, "row": 7, "cols": 28,
                    "rows": 5, "content": {"items": [{"value": "38%",
                                                      "label": "of buyers say so"}]}}
            slides.append(self.content("s%d" % i, [stat],
                                       headline="The market moved %d ways" % (i + 2)))
        _, report = self.check("nosource", slides)
        found = self.kinds(report, "source-missing")
        self.assertEqual(len(found), 1, report["findings"])
        self.assertEqual(found[0]["slide"], 0)
        self.assertEqual(found[0]["scope"], "deck")
        self.assertIn("1, 2, 3", found[0]["message"])

    def test_sources_on_most_of_them_is_enough(self):
        slides = []
        for i in range(3):
            stat = {"id": "b_s%d" % i, "type": "stat", "col": 3, "row": 7, "cols": 28,
                    "rows": 5, "content": {"items": [{"value": "38%",
                                                      "label": "of buyers say so"}]}}
            slides.append(self.content("s%d" % i, [stat],
                                       headline="The market moved %d ways" % (i + 2),
                                       source=None if i == 2 else "Gartner, 2026"))
        _, report = self.check("somesource", slides)
        self.assertEqual(self.kinds(report, "source-missing"), [])

    def test_two_consecutive_slides_with_one_headline_are_found(self):
        head = "The fee is flat and the team is lean"
        slides = [self.content("s1", [text_block("b_a", 3, 7, 28, 10, SHORT)],
                               headline=head),
                  self.content("s2", [text_block("b_b", 3, 7, 28, 10, SHORT)],
                               headline=head)]
        _, report = self.check("repeat", slides)
        found = self.kinds(report, "duplicate-headline")
        self.assertEqual(len(found), 1, report["findings"])
        self.assertEqual(found[0]["slide"], 2)
        self.assertIn("slide 1", found[0]["message"])

    def test_a_table_of_figures_counts_as_a_slide_with_numbers(self):
        rows = [["Build", "$18,500"], ["Rollout", "$9,500"]]
        slides = []
        for i in range(2):
            table = {"id": "b_t%d" % i, "type": "table", "col": 3, "row": 7,
                     "cols": 28, "rows": 6,
                     "content": {"headers": ["Phase", "Cost"], "rows": rows}}
            slides.append(self.content("s%d" % i, [table],
                                       headline="The cost lands at %d thousand" % (i + 2)))
        _, report = self.check("tablefigures", slides)
        self.assertEqual(len(self.kinds(report, "source-missing")), 1, report["findings"])

    def test_every_consistency_finding_is_a_warning(self):
        slides = [self.content("s1", [text_block("b_t", 3, 7, 28, 2, SHORT)],
                               headline="Three Ways To Grow.",
                               kicker="Competitive analysis / Strategy")]
        _, report = self.check("severity", slides, fidelity="strict")
        kinds = set(f["kind"] for f in report["findings"] if f["severity"] == "warning")
        self.assertTrue({"headline-case", "kicker-drift", "empty-band"} <= kinds, kinds)
        self.assertEqual(report["summary"]["errors"], 0, report["findings"])


class SpecSlideMappingTests(unittest.TestCase):
    """The spec slide is found by id, not by the built deck's position."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="blueprint_specmap_test_")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def spec(self, name):
        """A deck that generates both a cover and an agenda, with a broken diagram."""
        diagram = {"id": "b_d", "type": "diagram", "col": 3, "row": 7, "cols": 28,
                   "rows": 8,
                   "content": {"kind": "flow", "direction": "LR",
                               "nodes": [{"id": "a", "label": "Editors"},
                                         {"id": "b", "label": "Client team"}],
                               "edges": [{"from": "a", "to": "nowhere"}]}}
        slides = [
            {"id": "d1", "name": "Section one", "section": "Section one", "blocks": []},
            content_slide("s1", [text_block("b_t", 3, 7, 28, 10, SHORT)],
                          headline="The first slide carries the plan"),
            content_slide("s2", [diagram], headline="The flow runs left to right"),
        ]
        spec = {"name": "Mapping", "outputPath": os.path.join(self.tmp, name + ".pptx"),
                "thumbnails": False, "exportPdf": False, "style": {"mode": "house"},
                "slideSize": {"w": 960, "h": 540},
                "grid": {"cols": 32, "rows": 18, "cell": 30},
                "brief": {"goal": "Land the offer in one page"},
                "slides": slides}
        return spec

    def build(self, name):
        spec = self.spec(name)
        result = build_pptx.build(copy.deepcopy(spec))
        self.assertTrue(result["ok"], result.get("errors"))
        path = os.path.join(self.tmp, name + ".spec.json")
        with open(path, "w") as handle:
            json.dump(spec, handle)
        return result, path

    def test_the_generated_cover_and_agenda_shift_the_slides(self):
        result, _ = self.build("shift")
        snapshot = read_snapshot(result["path"])
        kinds = [s["kind"] for s in snapshot["slides"]]
        self.assertEqual(kinds[:2], ["title", "agenda"])
        self.assertEqual(len(kinds), 5)

    def test_a_broken_diagram_edge_is_found_on_the_right_slide(self):
        result, spec_path = self.build("mapped")
        report = check_deck.check_deck(result["path"], spec_path)
        found = [f for f in report["findings"] if f["kind"] == "diagram-edge"]
        self.assertEqual(len(found), 1, report["findings"])
        self.assertEqual(found[0]["slide"], 5, "the diagram is the fifth built slide")
        self.assertEqual(found[0]["blockId"], "b_d")

    def test_a_generated_slide_has_no_spec_slide(self):
        result, spec_path = self.build("generated")
        with open(spec_path) as handle:
            spec = json.load(handle)
        checker = check_deck.Checker(result["path"], spec)
        self.assertEqual(checker.spec_slide(0), {}, "the generated cover")
        self.assertEqual(checker.spec_slide(1), {}, "the generated agenda")
        self.assertEqual(checker.spec_slide(2).get("id"), "d1")
        self.assertEqual(checker.spec_slide(4).get("id"), "s2")

    def test_a_deck_with_no_generated_chrome_still_maps(self):
        spec = self.spec("nochrome")
        spec["cover"] = False
        spec["agenda"] = False
        result = build_pptx.build(copy.deepcopy(spec))
        self.assertTrue(result["ok"], result.get("errors"))
        path = os.path.join(self.tmp, "nochrome.spec.json")
        with open(path, "w") as handle:
            json.dump(spec, handle)
        report = check_deck.check_deck(result["path"], path)
        found = [f for f in report["findings"] if f["kind"] == "diagram-edge"]
        self.assertEqual(len(found), 1, report["findings"])
        self.assertEqual(found[0]["slide"], 3)


# ----------------------------------------------- Q19-Q23, F-1, F-3 (round two)


class FakeFonts(build_pptx.FontBook):
    """A font book with exactly the faces a test names, so a weight is testable."""

    def __init__(self, names):
        self.stems = set(self.key(name) for name in names)


class FontPatch(object):
    """`with FontPatch(["Montserrat", "Montserrat SemiBold"]):` for one test."""

    def __init__(self, names):
        self.names = names
        self.saved = None

    def __enter__(self):
        self.saved = build_pptx.FONTS
        build_pptx.FONTS = FakeFonts(self.names)
        return build_pptx.FONTS

    def __exit__(self, *exc):
        build_pptx.FONTS = self.saved
        return False


def stat_block(block_id, col, row, cols, rows, items, **over):
    values = [i if isinstance(i, dict) else {"value": i, "label": ""} for i in items]
    block = {"id": block_id, "type": "stat", "role": "stat", "col": col, "row": row,
             "cols": cols, "rows": rows, "content": {"items": values}}
    block.update(over)
    return block


def bullets_block(block_id, col, row, cols, rows, items, **over):
    block = {"id": block_id, "type": "bullets", "role": "list", "col": col, "row": row,
             "cols": cols, "rows": rows, "content": {"items": list(items)}}
    block.update(over)
    return block


def table_block(block_id, col, row, cols, rows, headers, table_rows, **over):
    block = {"id": block_id, "type": "table", "role": "table", "col": col, "row": row,
             "cols": cols, "rows": rows,
             "content": {"headers": list(headers), "rows": [list(r) for r in table_rows]}}
    block.update(over)
    return block


def runs_of(path):
    """[(text, font name, bold, size)] for every non-empty run in a built deck."""
    out = []
    for slide in Presentation(path).slides:
        for shape in slide.shapes:
            frames = []
            if getattr(shape, "has_text_frame", False):
                frames.append(shape.text_frame)
            if getattr(shape, "has_table", False):
                frames.extend(cell.text_frame for row in shape.table.rows
                              for cell in row.cells)
            for frame in frames:
                for para in frame.paragraphs:
                    for run in para.runs:
                        if run.text.strip():
                            out.append((run.text, run.font.name, bool(run.font.bold),
                                        run.font.size.pt if run.font.size else None))
    return out


class BuildsOnce(unittest.TestCase):
    """One temp directory per suite, and a build helper that asserts success."""

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="blueprint_round2_")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def build(self, name, slides, **over):
        spec = mini_spec(os.path.join(self.tmp, name + ".pptx"), slides, **over)
        result = build_pptx.build(copy.deepcopy(spec))
        self.assertTrue(result["ok"], result.get("errors"))
        return result


# ---------------------------------------------------- Q19 the display floor


DISPLAY_TOKENS = {
    "name": "Display",
    "palette": {"bg": "000000", "ink": "FFFFFF", "accent": "D9A066"},
    "fonts": {"display": "Montserrat", "body": "Montserrat"},
    "typeScale": {"headline": {"size": 27}, "dividerTitle": {"size": 66},
                  "coverTitle": {"size": 36}, "statValue": {"size": 44},
                  "agendaStatement": {"size": 40}, "body": {"size": 10.5},
                  "kicker": {"size": 9.5}, "caption": {"size": 8.5}},
}


class DisplayFloorTests(BuildsOnce):
    """Q19: a display role floors at 0.85 of its own size, not at 9 pt."""

    def test_a_headline_floors_at_85_percent_of_the_styles_size(self):
        style, _ = styled(DISPLAY_TOKENS)
        self.assertEqual(style.role("headline")["min"], round(27 * 0.85, 2))

    def test_every_display_role_takes_the_display_floor(self):
        style, _ = styled(DISPLAY_TOKENS)
        for role in build_pptx.DISPLAY_ROLES:
            size = style.role(role)["size"]
            self.assertAlmostEqual(style.role(role)["min"], round(size * 0.85, 2),
                                   places=2, msg=role)

    def test_a_text_role_keeps_the_readability_floor(self):
        style, _ = styled(DISPLAY_TOKENS)
        self.assertEqual(style.role("body")["min"], build_pptx.BODY_FLOOR)
        self.assertEqual(style.role("kicker")["min"], build_pptx.BODY_FLOOR)
        self.assertEqual(style.role("caption")["min"], build_pptx.CAPTION_FLOOR)

    def test_the_styles_own_floor_still_wins_over_the_display_floor(self):
        tokens = copy.deepcopy(DISPLAY_TOKENS)
        tokens["typeScale"]["headline"] = {"size": 27, "min": 18}
        style, _ = styled(tokens)
        self.assertEqual(style.role("headline")["min"], 18.0)

    def test_the_display_floor_never_lifts_a_small_size(self):
        tokens = copy.deepcopy(DISPLAY_TOKENS)
        tokens["typeScale"]["headline"] = {"size": 10}
        style, _ = styled(tokens)
        self.assertEqual(style.role("headline")["size"], 10.0)
        self.assertLessEqual(style.role("headline")["min"], 10.0)

    def test_the_deck_wide_headline_never_steps_under_the_display_floor(self):
        long_headline = ("Interbrand and MetaDesign sell an AI brand strategy that "
                         "stops well short of a working system the client can run")
        slides = [content_slide("s1", [text_block("b_t", 3, 7, 26, 6, SHORT)],
                                headline=long_headline)]
        result = self.build("displayfloor", slides,
                            style={"mode": "skill",
                                   "tokens": copy.deepcopy(DISPLAY_TOKENS)})
        self.assertGreaterEqual(result["typeScale"]["headline"], round(27 * 0.85, 2))

    def test_the_house_style_keeps_its_own_headline_floor(self):
        house, _ = build_pptx.make_style({"style": {"mode": "house"}}, [])
        self.assertEqual(house.role("headline")["min"], 24.0)


# ------------------------------------------------ Q20 weights and bold faces


class WeightTests(BuildsOnce):
    """Q20: `typeScale.<role>.bold` and `weight` name the face, not a guess."""

    FACES = ("Montserrat", "Montserrat SemiBold", "Montserrat Medium", "Montserrat Bold")

    def faced(self, scale, role, names=None, bold=None):
        """(typeface, bold) for one role on a machine with exactly `names` installed."""
        tokens = {"name": "Weighted",
                  "fonts": {"display": "Montserrat", "body": "Montserrat"},
                  "typeScale": scale}
        with FontPatch(list(names if names is not None else self.FACES)):
            style, warnings = styled(tokens)
            return (style.face(style.role(role), bold=bold), warnings)

    def test_an_installed_weight_is_written_as_the_face(self):
        face, _ = self.faced({"cardHeader": {"size": 14, "weight": "semibold"}},
                             "cardHeader")
        self.assertEqual(face, ("Montserrat SemiBold", False))

    def test_a_weight_with_no_installed_face_falls_back_to_the_family(self):
        face, _ = self.faced({"headline": {"size": 27, "weight": "light"}}, "headline")
        self.assertEqual(face, ("Montserrat", False))

    def test_the_bold_weight_is_the_family_plus_the_attribute(self):
        """"Montserrat Bold" is a face of Montserrat, not a family: never write it."""
        for names in (("Montserrat",), ("Montserrat", "Montserrat Bold")):
            face, _ = self.faced({"takeaway": {"size": 13, "weight": "bold"}},
                                 "takeaway", names=names)
            self.assertEqual(face, ("Montserrat", True), names)

    def test_the_regular_weight_is_the_bare_family(self):
        """A renderer has no family called "Montserrat Regular" to resolve."""
        for names in (("Montserrat",), ("Montserrat", "Montserrat Regular")):
            face, _ = self.faced({"kicker": {"size": 9.5, "weight": "regular"}},
                                 "kicker", names=names)
            self.assertEqual(face, ("Montserrat", False), names)

    def test_no_built_run_is_set_in_a_ribbi_face_name(self):
        if not build_pptx.FONTS.has("Montserrat"):
            self.skipTest("Montserrat is not installed on this machine")
        tokens = {"fonts": {"display": "Montserrat", "body": "Montserrat"},
                  "typeScale": {"kicker": {"size": 9.5, "weight": "regular"},
                                "headline": {"size": 27, "weight": "light"},
                                "takeaway": {"size": 13, "weight": "bold"}}}
        block = {"id": "b_take", "type": "takeaway", "col": 2, "row": 15, "cols": 30,
                 "rows": 2, "content": {"text": "The fee is flat."}}
        slides = [content_slide("s1", [block], section="Pricing", name="The fee")]
        result = self.build("ribbi", slides, style={"mode": "skill", "tokens": tokens})
        names = {name for _, name, _, _ in runs_of(result["path"])}
        self.assertNotIn("Montserrat Regular", names)
        self.assertNotIn("Montserrat Bold", names)
        self.assertEqual(result["typeScale"]["faces"]["kicker"], "Montserrat")

    def test_a_missing_medium_face_is_not_bolded(self):
        face, _ = self.faced({"body": {"size": 11, "weight": "medium"}}, "body",
                             names=("Montserrat",))
        self.assertEqual(face, ("Montserrat", False))

    def test_an_unknown_weight_warns_and_keeps_the_family(self):
        face, warnings = self.faced({"body": {"size": 11, "weight": "ultra"}}, "body")
        self.assertEqual(face, ("Montserrat", False))
        self.assertTrue(any("ultra" in w for w in warnings), warnings)

    def test_house_says_bold_and_a_silent_style_does_not(self):
        house, _ = build_pptx.make_style({"style": {"mode": "house"}}, [])
        self.assertTrue(house.role("headline")["bold"])
        style, _ = styled({"typeScale": {"headline": {"size": 27}}})
        self.assertFalse(style.role("headline")["bold"],
                         "a style that names no weight gets the family as it comes")

    def test_a_style_that_says_bold_false_is_not_bold(self):
        style, _ = styled({"typeScale": {"headline": {"size": 27, "bold": False}}})
        self.assertFalse(style.face(style.role("headline"))[1])

    def test_a_run_asking_for_bold_drops_to_the_base_family(self):
        scale = {"body": {"size": 11, "weight": "semibold"}}
        self.assertEqual(self.faced(scale, "body")[0], ("Montserrat SemiBold", False))
        self.assertEqual(self.faced(scale, "body", bold=True)[0], ("Montserrat", True))
        self.assertEqual(self.faced(scale, "body", bold=False)[0], ("Montserrat", False))

    def test_the_built_deck_writes_the_weighted_face(self):
        if not build_pptx.FONTS.has("Montserrat SemiBold"):
            self.skipTest("Montserrat SemiBold is not installed on this machine")
        tokens = {"fonts": {"display": "Montserrat", "body": "Montserrat"},
                  "typeScale": {"headline": {"size": 27, "weight": "semibold"},
                                "tableCell": {"size": 10, "weight": "medium"}}}
        table = table_block("b_tbl", 3, 7, 26, 5, ["Phase", "Cost"],
                            [["Build", "$18,500"], ["Rollout", "$9,500"]])
        slides = [content_slide("s1", [table], headline="The fee is flat")]
        result = self.build("weightface", slides,
                            style={"mode": "skill", "tokens": tokens})
        names = dict((text, name) for text, name, _, _ in runs_of(result["path"]))
        self.assertEqual(names.get("The fee is flat"), "Montserrat SemiBold")
        self.assertEqual(names.get("Build"), "Montserrat Medium")

    def test_the_result_records_the_resolved_face_per_role(self):
        slides = [content_slide("s1", [text_block("b_t", 3, 7, 26, 6, SHORT)])]
        result = self.build("faces", slides)
        faces = result["typeScale"]["faces"]
        self.assertEqual(faces["headline"], "Arial")
        self.assertEqual(faces["pageNumber"], "Menlo", "house names its own mono")
        snapshot = read_snapshot(result["path"])
        self.assertEqual(snapshot["typeFaces"], faces)


# ------------------------------------------ Q21 a stat band boxed in by a rail


RAIL = ["39% plan to cut labor spend", "22% rely less on outside agencies",
        "76% fix AI output every week", "More than half weigh AI readiness"]
FIGURES = [{"value": "7.7%", "label": "of revenue goes to marketing"},
           {"value": "59%", "label": "of CMOs lack the budget"},
           {"value": "49%", "label": "report time savings"}]


class StatBandTests(BuildsOnce):
    """Q21: loose frees a half-width stat band that a text rail boxed in."""

    def pair(self, name, fidelity="loose", rail=None):
        rail = rail or bullets_block("b_rail", 16, 7, 16, 5, RAIL)
        stats = stat_block("b_stat", 2, 7, 13, 5, FIGURES)
        slides = [content_slide("s1", [stats, rail],
                                headline="Budgets are flat and teams lean on AI")]
        result = self.build(name, slides, fidelity=fidelity)
        moves = dict((m["blockId"], m) for m in result.get("adjusted") or [])
        return (result, moves)

    def test_loose_gives_the_stats_the_full_width_band(self):
        _, moves = self.pair("statband")
        move = moves.get("b_stat")
        self.assertIsNotNone(move, moves)
        self.assertIn(fit.STAT_BAND_REASON, move["reason"])
        self.assertGreaterEqual(move["to"]["cols"], 26)
        self.assertGreater(move["to"]["cols"], move["from"]["cols"])

    def test_the_rail_moves_below_the_stats(self):
        _, moves = self.pair("statrail")
        stats, rail = moves["b_stat"], moves["b_rail"]
        self.assertGreaterEqual(rail["to"]["row"],
                                stats["to"]["row"] + stats["to"]["rows"])
        self.assertIn(fit.STAT_BAND_REASON, rail["reason"])

    def test_balanced_leaves_the_pair_side_by_side(self):
        _, moves = self.pair("statbalanced", fidelity="balanced")
        for move in moves.values():
            self.assertNotIn(fit.STAT_BAND_REASON, move["reason"])

    def test_a_stat_beside_a_table_is_not_a_boxed_in_band(self):
        table = table_block("b_rail", 16, 7, 16, 5, ["Phase", "Cost"],
                            [["Build", "$18,500"], ["Rollout", "$9,500"]])
        _, moves = self.pair("stattable", rail=table)
        self.assertNotIn(fit.STAT_BAND_REASON,
                         (moves.get("b_stat") or {}).get("reason", ""))

    def test_a_band_with_no_room_under_it_is_left_alone(self):
        stat = fit.Item({"id": "s", "type": "stat"}, fit.cells(2, 16, 13, 2))
        rail = fit.Item({"id": "r", "type": "bullets"}, fit.cells(16, 16, 16, 2))
        moves = fit.plan_slide([stat, rail], [], 32, 18, "loose")
        for move in moves:
            self.assertNotIn(fit.STAT_BAND_REASON, move["reason"])

    def test_a_stat_that_already_spans_the_band_is_left_alone(self):
        stats = stat_block("b_stat", 2, 7, 30, 4, FIGURES)
        rail = bullets_block("b_rail", 2, 12, 30, 5, RAIL)
        slides = [content_slide("s1", [stats, rail])]
        result = self.build("statwide", slides, fidelity="loose")
        for move in result.get("adjusted") or []:
            self.assertNotIn(fit.STAT_BAND_REASON, move["reason"])


# ------------------------------------------------------- Q22 the figure grid


GRID_FIGURES = [{"value": "39%", "label": "of CMOs plan to cut labor spend"},
                {"value": "22%", "label": "say GenAI lets them rely less on agencies"},
                {"value": "76%", "label": "spend three hours a week fixing AI output"},
                {"value": "65%", "label": "would pause a rollout to add guardrails"}]


class StatGridTests(BuildsOnce):
    """Q22: `layout: "grid"` sets number and label pairs with a hairline each."""

    def grid(self, name, block, **over):
        slides = [content_slide("s1", [block], headline="The supporting figures")]
        result = self.build(name, slides, fidelity="strict", **over)
        return read_snapshot(result["path"])

    def shapes(self, snapshot, block_id="b_fig"):
        return shapes_for(snapshot, block_id)

    def test_four_figures_in_a_tall_block_pick_the_grid(self):
        snap = self.grid("gridauto", stat_block("b_fig", 2, 7, 30, 8, GRID_FIGURES))
        self.assertEqual({s["layout"] for s in self.shapes(snap)}, {"grid"})

    def test_three_figures_stay_a_band(self):
        snap = self.grid("gridthree",
                         stat_block("b_fig", 2, 7, 30, 8, GRID_FIGURES[:3]))
        self.assertEqual({s["layout"] for s in self.shapes(snap)}, {"band"})

    def test_four_figures_in_a_shallow_block_stay_a_band(self):
        snap = self.grid("gridshallow",
                         stat_block("b_fig", 2, 7, 30, 4, GRID_FIGURES))
        self.assertEqual({s["layout"] for s in self.shapes(snap)}, {"band"})

    def test_the_block_can_ask_for_the_grid(self):
        snap = self.grid("gridnamed",
                         stat_block("b_fig", 2, 7, 30, 4, GRID_FIGURES[:2],
                                    layout="grid"))
        self.assertEqual({s["layout"] for s in self.shapes(snap)}, {"grid"})

    def test_the_block_can_ask_for_the_band(self):
        snap = self.grid("gridband",
                         stat_block("b_fig", 2, 7, 30, 8, GRID_FIGURES,
                                    layout="band"))
        self.assertEqual({s["layout"] for s in self.shapes(snap)}, {"band"})

    def test_the_number_is_the_accent_at_06_of_the_stat_value(self):
        snap = self.grid("gridsize", stat_block("b_fig", 2, 7, 30, 8, GRID_FIGURES))
        values = [s for s in self.shapes(snap) if s["role"] == "statValue"]
        self.assertEqual(len(values), 4)
        for shape in values:
            self.assertAlmostEqual(shape["fontPt"], round(44 * 0.6, 1), places=1)

    def test_the_label_is_set_in_the_body_role(self):
        snap = self.grid("gridlabel", stat_block("b_fig", 2, 7, 30, 8, GRID_FIGURES))
        labels = [s for s in self.shapes(snap) if s["role"] == "statLabel"]
        self.assertEqual(len(labels), 4)
        self.assertAlmostEqual(labels[0]["fontPt"], 12.0, places=1)

    def test_each_pair_gets_a_hairline_under_it(self):
        snap = self.grid("gridrule", stat_block("b_fig", 2, 7, 30, 8, GRID_FIGURES))
        rules = [s for s in self.shapes(snap) if s["role"] == "statRule"]
        self.assertEqual(len(rules), 4)
        for rule in rules:
            self.assertLessEqual(rule["h"], 1.0)

    def test_the_pairs_sit_in_two_columns_in_reading_order(self):
        snap = self.grid("gridcols", stat_block("b_fig", 2, 7, 30, 8, GRID_FIGURES))
        values = [s for s in self.shapes(snap) if s["role"] == "statValue"]
        self.assertEqual(len({round(s["x"], 1) for s in values}), 2, "two columns")
        self.assertEqual(len({round(s["y"], 1) for s in values}), 2, "two rows")
        self.assertEqual(values[0]["text"], "39%")
        self.assertEqual(values[1]["text"], "22%", "reading order, not column order")
        self.assertEqual(values[0]["y"], values[1]["y"])

    def test_an_unknown_layout_warns_and_draws_the_band(self):
        block = stat_block("b_fig", 2, 7, 30, 8, GRID_FIGURES, layout="ladder")
        slides = [content_slide("s1", [block])]
        result = self.build("gridunknown", slides, fidelity="strict")
        self.assertTrue(any("ladder" in w for w in result["warnings"]),
                        result["warnings"])
        snap = read_snapshot(result["path"])
        self.assertEqual({s["layout"] for s in self.shapes(snap)}, {"band"})


# ------------------------------------------------------- F-5 the stat rule


def ruled(rule=None):
    """A skill style that is house plus whatever `stat.rule` says."""
    tokens = {"name": "Ruled"}
    if rule is not None:
        tokens["stat"] = {"rule": rule}
    return {"mode": "skill", "tokens": tokens}


class StatRuleTests(BuildsOnce):
    """F-5: `stat.rule` says which side of a figure its rule sits on."""

    FIGURES = [{"value": "7.7%", "label": "of revenue goes to marketing"},
               {"value": "59%", "label": "of CMOs lack the budget to deliver"}]

    def band(self, name, **over):
        block = stat_block("b_stat", 2, 7, 20, 5, self.FIGURES, emphasis=True)
        slides = [content_slide("s1", [block], headline="The figures that carry it")]
        result = self.build(name, slides, fidelity="strict", **over)
        return (result, read_snapshot(result["path"]))

    def grid(self, name, **over):
        block = stat_block("b_fig", 2, 7, 30, 8, GRID_FIGURES)
        slides = [content_slide("s1", [block], headline="The supporting figures")]
        result = self.build(name, slides, fidelity="strict", **over)
        return (result, read_snapshot(result["path"]))

    def parts(self, snapshot, block_id, role):
        return [s for s in shapes_for(snapshot, block_id) if s["role"] == role]

    def test_the_default_is_below_and_house_says_so(self):
        house, _ = build_pptx.make_style({"style": {"mode": "house"}}, [])
        bare, _ = build_pptx.make_style({"style": {"mode": "plain"}}, [])
        self.assertEqual(house.stat_rule(), "below")
        self.assertEqual(bare.stat_rule(), "below")
        with open(os.path.join(HERE, "styles", "house_deck.json")) as handle:
            self.assertEqual(json.load(handle)["stat"]["rule"], "below")

    def test_a_style_can_ask_for_the_rule_above(self):
        style, warnings = styled({"name": "Ruled", "stat": {"rule": "above"}})
        self.assertEqual(style.stat_rule(), "above")
        self.assertEqual(warnings, [])

    def test_an_unknown_rule_warns_and_draws_it_under_the_figure(self):
        style, warnings = styled({"name": "Ruled", "stat": {"rule": "beside"}})
        self.assertEqual(style.stat_rule(), "below")
        self.assertTrue(any("beside" in w for w in warnings), warnings)

    def test_an_emphasised_band_rules_under_the_value_by_default(self):
        _, snap = self.band("rule_band_below")
        rules = self.parts(snap, "b_stat", "hairline")
        values = self.parts(snap, "b_stat", "stat")
        self.assertEqual(len(rules), 2)
        for rule, value in zip(rules, values):
            self.assertGreater(rule["y"], value["y"])
            self.assertEqual(rule["rulePosition"], "below")

    def test_above_heads_each_figure_with_its_rule(self):
        _, snap = self.band("rule_band_above", style=ruled("above"))
        rules = self.parts(snap, "b_stat", "hairline")
        values = self.parts(snap, "b_stat", "stat")
        self.assertEqual(len(rules), 2)
        for rule, value in zip(rules, values):
            self.assertLess(rule["y"], value["y"])
            self.assertAlmostEqual(rule["x"], value["x"], delta=0.5)
            self.assertEqual(rule["rulePosition"], "above")

    def test_none_draws_no_rule_on_an_emphasised_band(self):
        _, snap = self.band("rule_band_none", style=ruled("none"))
        self.assertEqual(self.parts(snap, "b_stat", "hairline"), [])
        self.assertEqual(len(self.parts(snap, "b_stat", "stat")), 2)

    def test_the_grid_hairlines_close_each_pair_by_default(self):
        _, snap = self.grid("rule_grid_below")
        rules = self.parts(snap, "b_fig", "statRule")
        values = self.parts(snap, "b_fig", "statValue")
        self.assertEqual(len(rules), 4)
        for rule, value in zip(rules, values):
            self.assertGreater(rule["y"], value["y"])
            self.assertEqual(rule["rulePosition"], "below")

    def test_the_grid_hairlines_head_each_pair_on_above(self):
        _, snap = self.grid("rule_grid_above", style=ruled("above"))
        rules = self.parts(snap, "b_fig", "statRule")
        values = self.parts(snap, "b_fig", "statValue")
        self.assertEqual(len(rules), 4)
        for rule, value in zip(rules, values):
            self.assertLess(rule["y"], value["y"])
            self.assertEqual(rule["rulePosition"], "above")

    def test_the_grid_takes_the_same_band_either_way(self):
        """The rule moves inside the pair; the block's footprint does not move."""
        def extent(snap):
            shapes = shapes_for(snap, "b_fig")
            return (min(s["y"] for s in shapes),
                    max(s["y"] + s["h"] for s in shapes))
        _, below = self.grid("rule_grid_in_below")
        _, above = self.grid("rule_grid_in_above", style=ruled("above"))
        self.assertEqual(len(shapes_for(below, "b_fig")),
                         len(shapes_for(above, "b_fig")))
        for mine, theirs in zip(extent(above), extent(below)):
            self.assertAlmostEqual(mine, theirs, delta=0.5)

    def test_none_leaves_the_grid_pairs_unruled(self):
        _, snap = self.grid("rule_grid_none", style=ruled("none"))
        self.assertEqual(self.parts(snap, "b_fig", "statRule"), [])
        self.assertEqual(len(self.parts(snap, "b_fig", "statValue")), 4)

    def test_a_plain_band_is_unruled_whatever_the_token_says(self):
        block = stat_block("b_stat", 2, 7, 20, 5, self.FIGURES)
        result = self.build("rule_plain", [content_slide("s1", [block])],
                            fidelity="strict", style=ruled("above"))
        snap = read_snapshot(result["path"])
        self.assertEqual(self.parts(snap, "b_stat", "hairline"), [])


# ----------------------------------------------------- Q23 the label column


SPEC_ROWS = [["What they sell", "AI brand strategy", "Branded AI"],
             ["How they sell it", "A line item", "Software at flat pricing"],
             ["Where they stop", "A strategy document", "Personality only"]]


class LabelColumnTests(BuildsOnce):
    """Q23: `table.labelColumn` says how the row labels read."""

    def table(self, name, tokens=None):
        block = table_block("b_tbl", 2, 7, 30, 6, ["", "Interbrand", "MetaDesign"],
                            SPEC_ROWS)
        slides = [content_slide("s1", [block], headline="The two firms compared")]
        style = ({"mode": "skill", "tokens": tokens} if tokens else {"mode": "house"})
        result = self.build(name, slides, fidelity="strict", style=style)
        return result

    def cells(self, path, text):
        return [r for r in runs_of(path) if r[0] == text]

    def test_house_sets_the_label_column_in_ink_at_the_cell_size(self):
        house, _ = build_pptx.make_style({"style": {"mode": "house"}}, [])
        self.assertEqual(house.label_column(),
                         {"case": "none", "color": "ink", "bold": True})

    def test_the_harbor_style_ships_the_same_default(self):
        harbor, _ = build_pptx.make_style(
            {"style": {"mode": "skill", "tokensPath": HARBOR_TOKENS}}, [])
        self.assertEqual(harbor.label_column()["case"], "none")
        self.assertEqual(harbor.label_column()["color"], "ink")
        self.assertTrue(harbor.label_column()["bold"])

    def test_the_default_writes_the_label_as_given(self):
        result = self.table("labeldefault")
        found = self.cells(result["path"], "What they sell")
        self.assertEqual(len(found), 1, "the label is not recased")
        self.assertTrue(found[0][2], "and carries the medium weight")

    def test_a_style_can_set_caps_in_the_accent(self):
        tokens = {"palette": {"accent": "D9A066", "ink": "1F2937"},
                  "table": {"labelColumn": {"case": "caps", "color": "accent",
                                            "size": 8.5, "bold": False}}}
        result = self.table("labelcaps", tokens=tokens)
        found = self.cells(result["path"], "WHAT THEY SELL")
        self.assertEqual(len(found), 1, [r[0] for r in runs_of(result["path"])])
        self.assertFalse(found[0][2])
        self.assertAlmostEqual(found[0][3], 8.5, places=1)
        self.assertEqual(self.cells(result["path"], "What they sell"), [])

    def test_the_label_colour_follows_the_token(self):
        tokens = {"palette": {"bg": "000000", "ink": "FFFFFF", "accent": "D9A066"},
                  "backgrounds": {"content": {"fill": "000000", "ink": "FFFFFF"}},
                  "table": {"labelColumn": {"case": "none", "color": "accent"}}}
        style, _ = styled(tokens)
        spec = style.label_column_spec({"size": 10, "color": "ink", "font": "body"})[0]
        self.assertEqual(spec["color"], "accent")
        self.assertEqual(style.hex_of(spec["color"]), "D9A066")

    def test_a_label_accent_that_fails_contrast_gives_way_to_ink(self):
        tokens = {"palette": {"bg": "FFFFFF", "ink": "1F2937", "accent": "D9A066"},
                  "table": {"labelColumn": {"case": "none", "color": "accent"}}}
        style, _ = styled(tokens)
        self.assertEqual(style.label_column()["color"], "accent")
        spec = style.label_column_spec({"size": 10, "color": "ink", "font": "body"})[0]
        self.assertEqual(spec["color"], "ink", "copper on white does not read")

    def test_a_table_whose_first_column_is_data_is_untouched(self):
        block = table_block("b_tbl", 2, 7, 30, 5, ["Week", "Owner", "Cost"],
                            [["1", "Northwind", "$18,500"], ["2", "Client", "$9,500"]])
        slides = [content_slide("s1", [block])]
        result = self.build("labelnone", slides, fidelity="strict")
        found = self.cells(result["path"], "Northwind")
        self.assertEqual(len(found), 1)
        self.assertFalse(found[0][2], "a figure column is not a label column")

    def test_an_unknown_case_warns_and_leaves_the_label_alone(self):
        tokens = {"table": {"labelColumn": {"case": "smallcaps"}}}
        warnings = []
        style, _ = build_pptx.make_style(
            {"style": {"mode": "skill", "tokens": tokens}}, warnings)
        self.assertEqual(style.label_column()["case"], "none")
        self.assertTrue(any("smallcaps" in w for w in warnings), warnings)


# --------------------------------------------------------- F-1 kicker casing


class KickerCaseTests(BuildsOnce):
    """F-1: `kicker.case` recases a derived kicker, never an explicit one."""

    def style_with(self, case):
        tokens = {"kicker": {"pattern": "{section} | {topic}", "case": case}}
        return styled(tokens)[0]

    def test_house_leaves_a_derived_kicker_alone(self):
        house, _ = build_pptx.make_style({"style": {"mode": "house"}}, [])
        self.assertEqual(house.kicker_case(), "none")
        self.assertEqual(
            build_pptx.kicker_for(house, {"section": "why referrals",
                                          "name": "the prize"}),
            "why referrals | the prize")

    def test_the_harbor_style_title_cases_its_kickers(self):
        harbor, _ = build_pptx.make_style(
            {"style": {"mode": "skill", "tokensPath": HARBOR_TOKENS}}, [])
        self.assertEqual(harbor.kicker_case(), "title")
        self.assertEqual(
            build_pptx.kicker_for(harbor, {"section": "why referrals",
                                          "name": "the prize"}),
            "Why Referrals | The Prize")

    def test_title_case_keeps_the_small_words_low(self):
        style = self.style_with("title")
        self.assertEqual(
            build_pptx.kicker_for(style, {"section": "the cost of delay",
                                          "name": "a case for referrals"}),
            "The Cost of Delay | A Case for Referrals")

    def test_title_case_keeps_an_acronym(self):
        style = self.style_with("title")
        self.assertEqual(
            build_pptx.kicker_for(style, {"section": "GTM and pricing",
                                          "name": "the AI market"}),
            "GTM and Pricing | The AI Market")

    def test_sentence_case_lowers_all_but_the_first_word(self):
        style = self.style_with("sentence")
        self.assertEqual(
            build_pptx.kicker_for(style, {"section": "GTM And Pricing",
                                          "name": "The First Buyers"}),
            "GTM and pricing | The first buyers")

    def test_caps_upper_cases_the_whole_line(self):
        style = self.style_with("caps")
        self.assertEqual(
            build_pptx.kicker_for(style, {"section": "gtm", "name": "market"}),
            "GTM | MARKET")

    def test_an_explicit_kicker_is_written_as_given(self):
        style = self.style_with("title")
        self.assertEqual(
            build_pptx.kicker_for(style, {"kicker": "why referrals | the prize",
                                          "section": "why referrals",
                                          "name": "the prize"}),
            "why referrals | the prize")

    def test_an_unknown_case_warns_and_changes_nothing(self):
        warnings = []
        style, _ = build_pptx.make_style(
            {"style": {"mode": "skill", "tokens": {"kicker": {"case": "smallcaps"}}}},
            warnings)
        self.assertEqual(style.kicker_case(), "none")
        self.assertTrue(any("smallcaps" in w for w in warnings), warnings)

    def test_the_built_deck_carries_the_recased_kicker(self):
        slides = [content_slide("s1", [text_block("b_t", 3, 7, 26, 6, SHORT)],
                                section="why referrals", name="the prize")]
        result = self.build("kickercase", slides,
                            style={"mode": "skill", "tokensPath": HARBOR_TOKENS})
        snap = read_snapshot(result["path"])
        kickers = [s["text"] for s in snap["slides"][0]["shapes"]
                   if s["role"] == "kicker"]
        self.assertEqual(kickers, ["Why Referrals | The Prize"])


# -------------------------------------------------- F-3 row 17 and the source


class SourceRowTests(BuildsOnce):
    """F-3: a slide with a source line keeps row 17 clear."""

    def last_band(self, name, row=13, rows=5, source="Northwind rate card, 2021",
                  fidelity="balanced"):
        block = text_block("b_low", 2, row, 30, rows, SHORT)
        slides = [content_slide("s1", [block], source=source)]
        result = self.build(name, slides, fidelity=fidelity)
        moves = dict((m["blockId"], m) for m in result.get("adjusted") or [])
        return (result, moves, cells_of(read_snapshot(result["path"]), "b_low"))

    def test_a_band_that_reaches_row_17_moves_up(self):
        _, moves, cells = self.last_band("sourceup")
        self.assertEqual(cells["row"] + cells["rows"] - 1, 16)
        self.assertIn("source line", moves["b_low"]["reason"])

    def test_a_slide_with_no_source_keeps_row_17(self):
        _, moves, cells = self.last_band("nosource", source=None)
        self.assertEqual(cells["row"] + cells["rows"] - 1, 17)
        self.assertNotIn("source line", (moves.get("b_low") or {}).get("reason", ""))

    def test_a_block_that_cannot_move_up_is_named_in_a_warning(self):
        result, _, _ = self.last_band("sourcestuck", row=5, rows=13)
        self.assertTrue(any("b_low" in w and "source line" in w
                            for w in result["warnings"]), result["warnings"])

    def test_the_whole_row_group_moves_together(self):
        left = text_block("b_left", 2, 13, 14, 5, SHORT)
        right = text_block("b_right", 18, 13, 14, 5, SHORT)
        slides = [content_slide("s1", [left, right], source="Northwind rate card")]
        result = self.build("sourcegroup", slides, fidelity="balanced")
        snap = read_snapshot(result["path"])
        for block_id in ("b_left", "b_right"):
            cells = cells_of(snap, block_id)
            self.assertEqual(cells["row"] + cells["rows"] - 1, 16, block_id)

    def test_strict_draws_the_plan_and_the_guard_still_runs(self):
        _, moves, _ = self.last_band("sourcestrict", fidelity="strict")
        self.assertEqual(moves, {})

    def test_the_body_band_stops_at_row_16_with_a_source(self):
        plan = fit.Plan([], [], 32, 18, "balanced", source=True)
        self.assertEqual(plan.body_bottom(), fit.SOURCE_ROW - 1)
        self.assertEqual(fit.Plan([], [], 32, 18, "balanced").body_bottom(),
                         fit.BODY_BOTTOM)

    def test_growing_down_never_takes_the_row_the_source_holds(self):
        item = fit.Item({"id": "b", "type": "text"}, fit.cells(2, 13, 30, 4),
                        need_rows=2)
        moves = fit.plan_slide([item], [], 32, 18, "balanced", source=True)
        self.assertLessEqual(item.cells["row"] + item.cells["rows"] - 1,
                             fit.SOURCE_ROW - 1, moves)

# ================================================== Q15 chrome, furniture, copy


def notes_of(path):
    """The speaker notes of every slide in a built deck, in order."""
    out = []
    for slide in Presentation(path).slides:
        out.append(slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else "")
    return out


def chart_block(block_id, col, row, cols, rows, kind, categories, series, **over):
    block = {"id": block_id, "type": "chart", "role": "visual", "col": col, "row": row,
             "cols": cols, "rows": rows,
             "content": {"kind": kind, "categories": list(categories),
                         "series": [dict(s) for s in series]}}
    block.update(over)
    return block


def charts_of(path):
    """Every native chart in a built deck, in slide order."""
    out = []
    for slide in Presentation(path).slides:
        for shape in slide.shapes:
            if getattr(shape, "has_chart", False):
                out.append(shape.chart)
    return out


def reasons(result, reason):
    return [m for m in result.get("adjusted") or [] if m.get("reason") == reason]


# A format with its own marks and its own measure: a dark-deck shape, with the
# three dots top left and the hairline under the kicker row.
MARK_TOKENS = {
    "name": "Marks",
    "palette": {"bg": "000000", "ink": "FFFFFF", "accent": "D9A066", "muted": "919191",
                "line": "262626"},
    "fonts": {"display": "Arial", "body": "Arial"},
    "chrome": [
        {"kind": "dots", "x": 22, "y": 22, "w": 34, "h": 8, "count": 3, "color": "muted"},
        {"kind": "line", "x": 32, "y": 45, "w": 896, "h": 0.75, "color": "line"},
    ],
    "furniture": {
        "kicker": {"x": 336, "y": 29, "w": 281, "h": 16},
        "headline": {"x": 336, "y": 63, "w": 593, "h": 115},
        "pageNumber": {"x": 785, "y": 29, "w": 144, "h": 16},
        "source": {"x": 336, "y": 499, "w": 593, "h": 22},
        "body": {"x": 32, "y": 178, "w": 890, "h": 321},
    },
    "typeScale": {"headline": {"size": 27}, "dividerTitle": {"size": 44},
                  "body": {"size": 10.5}, "caption": {"size": 8.5}},
}


def marked(**over):
    tokens = copy.deepcopy(MARK_TOKENS)
    tokens.update(over)
    return {"mode": "skill", "tokens": tokens}


class ChromeTokenTests(BuildsOnce):
    """Q15: the marks a format repeats live in `chrome`, not in the builder."""

    def test_dots_draw_one_circle_per_count(self):
        result = self.build("chrome_dots", [content_slide("s1", [])], style=marked())
        slide = list(Presentation(result["path"]).slides)[0]
        circles = sorted(ovals(slide), key=lambda s: s.left)
        self.assertEqual(len(circles), 3)
        self.assertAlmostEqual(circles[0].left / PT, 22.0, delta=0.5)
        self.assertAlmostEqual(circles[0].width / PT, 8.0, delta=0.5)
        self.assertAlmostEqual(circles[2].left / PT, 48.0, delta=0.5)

    def test_the_marks_land_in_the_snapshot_at_the_chrome_role(self):
        result = self.build("chrome_snap", [content_slide("s1", [])], style=marked())
        shapes = read_snapshot(result["path"])["slides"][0]["shapes"]
        marks = [s for s in shapes if s["role"] == "chrome"]
        self.assertEqual(len(marks), 4)          # three dots and the hairline
        line = [s for s in marks if s["w"] > 100][0]
        self.assertAlmostEqual(line["y"], 45.0, delta=0.5)
        self.assertAlmostEqual(line["w"], 896.0, delta=0.5)

    def test_house_and_harbor_define_no_chrome(self):
        for tokens in (build_pptx.house_tokens([]),
                       json.load(open(HARBOR_TOKENS))):
            self.assertFalse(tokens.get("chrome"))

    def test_chrome_draws_on_content_and_not_on_a_divider(self):
        slides = [{"id": "d", "kind": "divider", "section": "One", "headline": "One"},
                  content_slide("s1", [])]
        result = self.build("chrome_kind", slides, style=marked())
        snapshot = read_snapshot(result["path"])
        first = [s for s in snapshot["slides"][0]["shapes"] if s["role"] == "chrome"]
        second = [s for s in snapshot["slides"][1]["shapes"] if s["role"] == "chrome"]
        self.assertEqual(first, [])
        self.assertTrue(second)

    def test_an_entry_may_name_the_kinds_it_draws_on(self):
        marks = [dict(MARK_TOKENS["chrome"][0], kinds=["content", "divider"])]
        slides = [{"id": "d", "kind": "divider", "section": "One", "headline": "One"}]
        result = self.build("chrome_on", slides, style=marked(chrome=marks))
        shapes = read_snapshot(result["path"])["slides"][0]["shapes"]
        self.assertEqual(len([s for s in shapes if s["role"] == "chrome"]), 3)

    def test_a_wordmark_mark_draws_its_own_text(self):
        marks = [{"kind": "wordmark", "x": 32, "y": 22, "w": 200, "h": 16,
                  "text": "HARBOR", "size": 9}]
        result = self.build("chrome_word", [content_slide("s1", [])],
                            style=marked(chrome=marks))
        texts = [t for t, _, _, _ in runs_of(result["path"])]
        self.assertIn("HARBOR", texts)

    def test_a_missing_logo_warns_and_draws_nothing(self):
        marks = [{"kind": "logo", "x": 32, "y": 22, "w": 60, "h": 20,
                  "path": os.path.join(self.tmp, "nope.png")}]
        result = self.build("chrome_logo", [content_slide("s1", [])],
                            style=marked(chrome=marks))
        self.assertTrue(any("chrome logo not found" in w for w in result["warnings"]))
        shapes = read_snapshot(result["path"])["slides"][0]["shapes"]
        self.assertEqual([s for s in shapes if s["role"] == "chrome"], [])

    def test_a_logo_file_is_placed_inside_its_box(self):
        png = os.path.join(self.tmp, "logo.png")
        make_png(png, 80, 20)
        marks = [{"kind": "logo", "x": 32, "y": 22, "w": 60, "h": 40, "path": png}]
        result = self.build("chrome_logo2", [content_slide("s1", [])],
                            style=marked(chrome=marks))
        shapes = read_snapshot(result["path"])["slides"][0]["shapes"]
        mark = [s for s in shapes if s["role"] == "chrome"][0]
        self.assertEqual(mark["kind"], "image")
        self.assertLessEqual(mark["w"], 60.5)
        self.assertLessEqual(mark["h"], 40.5)

    def test_an_unknown_chrome_kind_warns(self):
        marks = [{"kind": "sparkle", "x": 10, "y": 10, "w": 10, "h": 10}]
        result = self.build("chrome_unknown", [content_slide("s1", [])],
                            style=marked(chrome=marks))
        self.assertTrue(any("unknown chrome kind" in w for w in result["warnings"]))

    def test_a_block_on_the_chrome_band_is_guarded_off_it(self):
        block = text_block("b1", 2, 1, 10, 4, "Copy that starts on the top row")
        result = self.build("chrome_guard", [content_slide("s1", [block])],
                            style=marked(), fidelity="strict")
        entry = shapes_for(read_snapshot(result["path"]), "b1")[0]
        self.assertGreaterEqual(entry["y"], 45.0)


class DerivedFurnitureTests(BuildsOnce):
    """Q15: divider, cover and agenda slots from the style's own measure."""

    def setUp(self):
        self.style, _ = styled(copy.deepcopy(MARK_TOKENS))
        self.house, _ = build_pptx.make_style({"style": {"mode": "house"}}, [])

    def test_the_divider_title_takes_the_bodys_left_edge(self):
        slot = self.style.slot("dividerTitle", "divider")
        self.assertAlmostEqual(slot["x"], 32.0, delta=0.5)
        self.assertAlmostEqual(slot["w"], 890.0, delta=0.5)

    def test_the_divider_title_is_not_houses_geometry(self):
        mine = self.style.slot("dividerTitle", "divider")
        theirs = self.house.slot("dividerTitle", "divider")
        self.assertNotAlmostEqual(mine["x"], theirs["x"], delta=0.5)
        self.assertNotAlmostEqual(mine["y"], theirs["y"], delta=0.5)

    def test_the_divider_group_sits_at_the_top_of_the_body_band(self):
        title = self.style.slot("dividerTitle", "divider")
        sub = self.style.slot("dividerSubtitle", "divider")
        self.assertAlmostEqual(title["y"], 178.0, delta=0.5)
        self.assertLess(title["y"] + title["h"], sub["y"])
        self.assertLess(sub["y"] + sub["h"], 178.0 + 321.0)

    def test_the_divider_title_heads_the_band_it_opens(self):
        """F-5: a section title is the head of its band, not a caption on its floor."""
        title = self.style.slot("dividerTitle", "divider")
        body = self.style.slot("body", "content")
        sub = self.style.slot("dividerSubtitle", "divider")
        self.assertAlmostEqual(title["y"], body["y"], delta=0.5)
        self.assertAlmostEqual(sub["y"], title["y"] + title["h"] + build_pptx.SLOT_GAP,
                               delta=0.5)
        # Not bottomed out: the band's floor is a long way under the subtitle.
        self.assertGreater(body["y"] + body["h"] - (sub["y"] + sub["h"]), 60.0)

    def test_a_style_with_no_body_heads_its_headline_band(self):
        tokens = copy.deepcopy(MARK_TOKENS)
        del tokens["furniture"]["body"]
        style, _ = styled(tokens)
        title = style.slot("dividerTitle", "divider")
        head = tokens["furniture"]["headline"]
        self.assertAlmostEqual(title["y"], head["y"] + head["h"] + build_pptx.SLOT_GAP,
                               delta=0.5)

    def test_a_named_subtitle_slot_keeps_its_place_under_a_derived_title(self):
        tokens = copy.deepcopy(MARK_TOKENS)
        tokens["furniture"]["dividerSubtitle"] = {"x": 32, "y": 430, "w": 500, "h": 40}
        style, _ = styled(tokens)
        self.assertAlmostEqual(style.slot("dividerSubtitle", "divider")["y"], 430.0,
                               delta=0.5)
        self.assertAlmostEqual(style.slot("dividerTitle", "divider")["y"], 178.0,
                               delta=0.5)

    def test_the_cover_group_bottoms_out_together(self):
        title = self.style.slot("coverTitle", "title")
        sub = self.style.slot("coverSubtitle", "title")
        stamp = self.style.slot("coverDate", "title")
        self.assertAlmostEqual(stamp["y"] + stamp["h"], 499.0, delta=1.0)
        self.assertLess(sub["y"] + sub["h"], stamp["y"] + 0.1)
        self.assertLess(title["y"] + title["h"], sub["y"] + 0.1)
        self.assertAlmostEqual(title["x"], 32.0, delta=0.5)

    def test_a_slot_the_style_names_is_left_alone(self):
        tokens = copy.deepcopy(MARK_TOKENS)
        tokens["furniture"]["dividerTitle"] = {"x": 336, "y": 200, "w": 500, "h": 120}
        style, _ = styled(tokens)
        slot = style.slot("dividerTitle", "divider")
        self.assertAlmostEqual(slot["x"], 336.0, delta=0.5)
        self.assertAlmostEqual(slot["y"], 200.0, delta=0.5)

    def test_a_derived_subtitle_follows_the_title_the_style_named(self):
        tokens = copy.deepcopy(MARK_TOKENS)
        tokens["furniture"]["dividerTitle"] = {"x": 336, "y": 200, "w": 500, "h": 120}
        style, _ = styled(tokens)
        sub = style.slot("dividerSubtitle", "divider")
        self.assertGreaterEqual(sub["y"], 320.0)

    def test_house_keeps_its_own_coordinates(self):
        slot = self.house.slot("dividerTitle", "divider")
        self.assertAlmostEqual(slot["x"], 60.0, delta=0.5)
        self.assertAlmostEqual(slot["y"], 180.0, delta=0.5)

    def test_a_style_with_no_body_derives_from_its_headline_column(self):
        tokens = copy.deepcopy(MARK_TOKENS)
        del tokens["furniture"]["body"]
        style, _ = styled(tokens)
        slot = style.slot("dividerTitle", "divider")
        self.assertAlmostEqual(slot["x"], 336.0, delta=0.5)
        self.assertAlmostEqual(slot["w"], 593.0, delta=0.5)

    def test_the_divider_title_size_is_the_styles_own(self):
        self.assertEqual(self.style.role("dividerTitle")["size"], 44)

    def test_the_agenda_statement_takes_the_headline_slot(self):
        slot = self.style.slot("agendaStatement", "agenda")
        head = self.style.slot("headline", "content")
        self.assertAlmostEqual(slot["x"], head["x"], delta=0.5)
        self.assertAlmostEqual(slot["y"], head["y"], delta=0.5)

    def test_the_agenda_list_places_the_generated_entries(self):
        slot = self.style.slot("agendaList", "agenda")
        places = build_pptx.agenda_rows(self.style, 3, 30.0, 30.0)
        self.assertEqual(len(places), 3)
        tops = [(row - 1) * 30.0 for _, row, _, _ in places]
        self.assertGreaterEqual(tops[0], slot["y"] - 30.0)
        self.assertLess(tops[-1], slot["y"] + slot["h"])
        self.assertLess(tops[0], tops[-1])

    def test_a_built_divider_draws_its_title_at_the_derived_slot(self):
        slides = [{"id": "d", "kind": "divider", "section": "Why referrals",
                   "headline": "Why referrals", "subtitle": "The three moments"}]
        result = self.build("derived_divider", slides, style=marked())
        shapes = read_snapshot(result["path"])["slides"][0]["shapes"]
        title = [s for s in shapes if s["role"] == "dividerTitle"][0]
        sub = [s for s in shapes if s["role"] == "dividerSubtitle"][0]
        self.assertAlmostEqual(title["x"], 32.0, delta=0.5)
        self.assertAlmostEqual(title["y"], 178.0, delta=0.5)
        self.assertGreater(sub["y"], title["y"])
        self.assertLess(sub["y"] + sub["h"], 499.0)


class SourceLineTests(BuildsOnce):
    """Q15: a source line is one line, and the full wording is in the notes."""

    BRACKET = ("Gartner 2026 CMO Spend Survey (a global study of 2,003 marketing "
               "leaders fielded in June 2026); Forrester 2026 B2B Brand and "
               "Communications Survey; Northwind rate card, 2021")
    CLAUSES = ("Gartner 2026 CMO Spend Survey of 2,003 marketing leaders; Forrester "
               "2026 B2B Brand and Communications Survey of 1,200 buyers; Northwind rate "
               "card 2021 and the seven pricing structures in the appendix")

    def drawn_source(self, path):
        shapes = read_snapshot(path)["slides"][0]["shapes"]
        return [s for s in shapes if s["role"] == "source"][0]["text"]

    def test_a_source_that_fits_is_left_alone(self):
        result = self.build("src_short", [content_slide("s1", [], source="Partner CRM")])
        self.assertEqual(reasons(result, "source"), [])
        self.assertIn("Partner CRM", self.drawn_source(result["path"]))

    def test_the_detail_in_brackets_goes_first(self):
        result = self.build("src_bracket",
                            [content_slide("s1", [], source=self.BRACKET)])
        drawn = self.drawn_source(result["path"])
        self.assertNotIn("(", drawn)
        self.assertIn("Forrester", drawn)

    def test_the_trailing_clauses_go_next(self):
        result = self.build("src_clause",
                            [content_slide("s1", [], source=self.CLAUSES)])
        drawn = self.drawn_source(result["path"])
        self.assertNotIn("appendix", drawn)
        self.assertIn("Gartner", drawn)
        self.assertEqual(drawn.count(";"), 1)

    def test_the_full_text_is_kept_in_the_notes(self):
        result = self.build("src_notes",
                            [content_slide("s1", [], source=self.BRACKET)])
        self.assertIn("2,003 marketing leaders", notes_of(result["path"])[0])

    def test_the_shortening_is_reported_as_an_adjustment(self):
        result = self.build("src_report",
                            [content_slide("s1", [], source=self.BRACKET)])
        moves = reasons(result, "source")
        self.assertEqual(len(moves), 1)
        self.assertIn("(", moves[0]["from"]["text"])
        self.assertNotIn("(", moves[0]["to"]["text"])
        self.assertIsNone(moves[0]["blockId"])

    def test_one_line_at_the_caption_floor_is_the_test(self):
        style, _ = build_pptx.make_style({"style": {"mode": "house"}}, [])
        ctx = build_pptx.Ctx({}, style, [])
        short = build_pptx.shorten_source(ctx, "Source: " + self.CLAUSES, 600.0)
        spec = style.role("caption")
        chars = build_pptx.cpl(600.0, spec["min"], spec["k"])
        self.assertEqual(len(build_pptx.wrap_lines(short, chars)), 1)


class DuplicateContentTests(BuildsOnce):
    """Q15: a paragraph that repeats the headline is redundancy, not content."""

    HEAD = "Referrals stall at three moments, and each one has a different fix"

    def test_a_first_paragraph_repeating_the_headline_is_dropped(self):
        block = text_block("b1", 2, 7, 12, 4, self.HEAD)
        block["content"]["paragraphs"].append("The first is the moment of the ask.")
        result = self.build("dup_drop", [content_slide("s1", [block], headline=self.HEAD)])
        entry = shapes_for(read_snapshot(result["path"]), "b1")[0]
        self.assertNotIn("three moments", entry["text"])
        self.assertIn("moment of the ask", entry["text"])

    def test_the_drop_is_recorded(self):
        block = text_block("b1", 2, 7, 12, 4, self.HEAD)
        block["content"]["paragraphs"].append("The first is the moment of the ask.")
        result = self.build("dup_report", [content_slide("s1", [block], headline=self.HEAD)])
        moves = reasons(result, "duplicate")
        self.assertEqual(len(moves), 1)
        self.assertEqual(moves[0]["blockId"], "b1")

    def test_a_block_that_only_repeats_the_headline_goes_entirely(self):
        block = {"id": "b1", "type": "takeaway", "role": "takeaway", "col": 2, "row": 7,
                 "cols": 20, "rows": 4, "content": {"text": self.HEAD}}
        result = self.build("dup_whole", [content_slide("s1", [block], headline=self.HEAD)])
        self.assertEqual(shapes_for(read_snapshot(result["path"]), "b1"), [])
        self.assertTrue(any("says nothing else" in w for w in result["warnings"]))

    def test_a_paragraph_that_merely_shares_words_is_kept(self):
        block = text_block("b1", 2, 7, 12, 4, "Referrals stall.")
        result = self.build("dup_keep", [content_slide("s1", [block], headline=self.HEAD)])
        self.assertEqual(reasons(result, "duplicate"), [])
        self.assertIn("Referrals stall", shapes_for(read_snapshot(result["path"]), "b1")[0]["text"])

    def test_two_blocks_with_the_same_title_warn(self):
        one = text_block("b1", 2, 7, 12, 4, "Alpha", title="What it is")
        two = text_block("b2", 16, 7, 12, 4, "Beta", title="What it is")
        result = self.build("dup_title", [content_slide("s1", [one, two])])
        self.assertTrue(any("carry the same title" in w for w in result["warnings"]))

    def test_strict_keeps_the_duplicate(self):
        block = text_block("b1", 2, 7, 12, 4, self.HEAD)
        block["content"]["paragraphs"].append("The first is the moment of the ask.")
        result = self.build("dup_strict", [content_slide("s1", [block], headline=self.HEAD)],
                            fidelity="strict")
        self.assertEqual(reasons(result, "duplicate"), [])
        self.assertIn("three moments", shapes_for(read_snapshot(result["path"]), "b1")[0]["text"])


class EmptySlideTests(BuildsOnce):
    """F-5: a content slide with nothing left on it is a hole in the deck."""

    HEAD = "Referrals stall at three moments, and each one has a different fix"

    def only_repeat(self):
        return {"id": "b1", "type": "takeaway", "role": "takeaway", "col": 2, "row": 7,
                "cols": 20, "rows": 4, "content": {"text": self.HEAD}}

    def checked(self, name, slides, **over):
        result = self.build(name, slides, **over)
        return (result, check_deck.check_deck(result["path"]))

    def empties(self, report):
        return [f for f in report["findings"] if f["kind"] == "empty-slide"]

    def test_a_slide_the_duplicate_pass_emptied_is_an_error(self):
        slide = content_slide("s1", [self.only_repeat()], headline=self.HEAD,
                              name="Why referrals stall")
        _, report = self.checked("empty_dup", [slide])
        found = self.empties(report)
        self.assertEqual(len(found), 1, report["findings"])
        self.assertEqual(found[0]["severity"], "error")
        self.assertEqual(found[0]["slide"], 1)
        self.assertIn("Why referrals stall", found[0]["message"])
        self.assertIn("add a block or drop the slide", found[0]["message"])
        self.assertGreaterEqual(report["summary"]["errors"], 1)

    def test_the_builders_warning_still_fires(self):
        slide = content_slide("s1", [self.only_repeat()], headline=self.HEAD)
        result, _ = self.checked("empty_warn", [slide])
        self.assertTrue(any("says nothing else" in w for w in result["warnings"]),
                        result["warnings"])

    def test_a_slide_that_kept_a_block_is_not_empty(self):
        block = text_block("b1", 2, 7, 12, 4, self.HEAD)
        block["content"]["paragraphs"].append("The first is the moment of the ask.")
        slide = content_slide("s1", [block], headline=self.HEAD)
        _, report = self.checked("empty_kept", [slide])
        self.assertEqual(self.empties(report), [])

    def test_a_content_slide_with_no_blocks_at_all_is_empty(self):
        _, report = self.checked("empty_bare", [content_slide("s1", [])])
        self.assertEqual(len(self.empties(report)), 1, report["findings"])

    def test_furniture_blocks_do_not_count_as_content(self):
        head = {"id": "b_head", "type": "headline", "col": 2, "row": 3, "cols": 20,
                "rows": 2, "content": {"text": "Referrals stall at the ask"}}
        slide = {"id": "s1", "name": "Furniture only", "kind": "content",
                 "blocks": [head]}
        _, report = self.checked("empty_furniture", [slide])
        self.assertEqual(len(self.empties(report)), 1, report["findings"])

    def test_a_divider_is_never_an_empty_slide(self):
        slides = [{"id": "d", "kind": "divider", "section": "Why referrals",
                   "headline": "Why referrals", "subtitle": "The three moments"},
                  content_slide("s2", [text_block("b1", 2, 7, 12, 4, "A real block.")])]
        _, report = self.checked("empty_divider", slides)
        self.assertEqual(self.empties(report), [])

    def test_strict_keeps_the_block_and_the_slide(self):
        slide = content_slide("s1", [self.only_repeat()], headline=self.HEAD)
        _, report = self.checked("empty_strict", [slide], fidelity="strict")
        self.assertEqual(self.empties(report), [])


class ParallelRowTests(BuildsOnce):
    """Q15: one treatment per row of parallel content."""

    def row(self):
        return [bullets_block("b1", 2, 7, 9, 6, ["One item", "Two items"]),
                bullets_block("b2", 12, 7, 9, 6, ["Three items", "Four items"]),
                text_block("b3", 22, 7, 9, 6, "A rail with no marker", title="Who it is for")]

    def test_a_text_block_in_a_row_of_bullets_takes_the_marker(self):
        result = self.build("par_bullets", [content_slide("s1", self.row())])
        entry = shapes_for(read_snapshot(result["path"]), "b3")[0]
        self.assertEqual(entry["role"], "bullets")

    def test_the_harmonisation_is_recorded(self):
        result = self.build("par_report", [content_slide("s1", self.row())])
        moves = reasons(result, "parallel")
        self.assertEqual([m["blockId"] for m in moves], ["b3"])
        self.assertEqual(moves[0]["to"]["type"], "bullets")

    def test_the_converted_blocks_title_stays_a_title(self):
        result = self.build("par_title", [content_slide("s1", self.row())])
        slide = list(Presentation(result["path"]).slides)[0]
        frames = [s.text_frame for s in slide.shapes
                  if getattr(s, "has_text_frame", False)
                  and s.text_frame.text.startswith("Who it is for")]
        self.assertTrue(frames)
        paras = [p for p in frames[0].paragraphs if p.text.strip()]

        def marker(para):
            pPr = para._p.find(qn("a:pPr"))
            return None if pPr is None else pPr.find(qn("a:buChar"))

        self.assertIsNone(marker(paras[0]))
        self.assertIsNotNone(marker(paras[1]))

    def test_a_row_of_only_text_blocks_is_left_alone(self):
        blocks = [text_block("b1", 2, 7, 12, 5, "Left rail"),
                  text_block("b2", 16, 7, 12, 5, "Right rail")]
        result = self.build("par_text", [content_slide("s1", blocks)])
        self.assertEqual(reasons(result, "parallel"), [])
        self.assertEqual(shapes_for(read_snapshot(result["path"]), "b1")[0]["role"], "body")

    def test_a_text_rail_beside_a_table_is_left_alone(self):
        blocks = [table_block("b1", 2, 7, 16, 6, ["Phase", "Weeks"], [["Build", "2"]]),
                  bullets_block("b2", 19, 7, 6, 6, ["One", "Two"]),
                  text_block("b3", 26, 7, 5, 6, "Not a bullet")]
        result = self.build("par_table", [content_slide("s1", blocks)])
        self.assertEqual(reasons(result, "parallel"), [])
        self.assertEqual(shapes_for(read_snapshot(result["path"]), "b3")[0]["role"], "body")


class ChartDisciplineTests(BuildsOnce):
    """Q15: no legend for one series, horizontal gridlines, labels, ranked bars."""

    def test_one_series_gets_no_legend_and_two_do(self):
        one = chart_block("c1", 2, 7, 14, 8, "column", ["A", "B", "C"],
                          [{"name": "Spend", "values": [3, 5, 4]}])
        result = self.build("chart_one", [content_slide("s1", [one])])
        self.assertFalse(charts_of(result["path"])[0].has_legend)
        two = chart_block("c2", 2, 7, 14, 8, "column", ["A", "B", "C"],
                          [{"name": "Now", "values": [3, 5, 4]},
                           {"name": "Then", "values": [2, 4, 3]}])
        result = self.build("chart_two", [content_slide("s1", [two])])
        self.assertTrue(charts_of(result["path"])[0].has_legend)

    def test_gridlines_run_horizontally_or_not_at_all(self):
        column = chart_block("c1", 2, 7, 14, 8, "column", ["A", "B"],
                             [{"name": "Spend", "values": [3, 5]}])
        result = self.build("chart_grid_col", [content_slide("s1", [column])])
        chart = charts_of(result["path"])[0]
        self.assertTrue(chart.value_axis.has_major_gridlines)
        self.assertFalse(chart.category_axis.has_major_gridlines)
        bar = chart_block("c1", 2, 7, 14, 8, "bar", ["A", "B"],
                          [{"name": "Spend", "values": [3, 5]}])
        result = self.build("chart_grid_bar", [content_slide("s1", [bar])])
        self.assertFalse(charts_of(result["path"])[0].value_axis.has_major_gridlines)

    def test_data_labels_come_on_for_a_chart_the_reader_can_count(self):
        small = chart_block("c1", 2, 7, 14, 8, "column", ["A", "B", "C"],
                            [{"name": "Spend", "values": [3, 5, 4]}])
        result = self.build("chart_labels", [content_slide("s1", [small])])
        self.assertTrue(charts_of(result["path"])[0].plots[0].has_data_labels)

    def test_data_labels_stay_off_past_eight_points(self):
        cats = ["C%d" % i for i in range(9)]
        many = chart_block("c1", 2, 7, 14, 8, "column", cats,
                           [{"name": "Spend", "values": list(range(9))}])
        result = self.build("chart_nolabels", [content_slide("s1", [many])])
        self.assertFalse(charts_of(result["path"])[0].plots[0].has_data_labels)

    def test_a_single_series_bar_chart_is_ranked(self):
        block = chart_block("c1", 2, 7, 14, 8, "column", ["Small", "Big", "Middle"],
                            [{"name": "Spend", "values": [1, 9, 5]}])
        result = self.build("chart_rank", [content_slide("s1", [block])])
        chart = charts_of(result["path"])[0]
        self.assertEqual([str(c) for c in chart.plots[0].categories],
                         ["Big", "Middle", "Small"])

    def test_a_bar_chart_ranks_the_other_way_round(self):
        """PowerPoint plots the first category at the bottom of a bar chart."""
        block = chart_block("c1", 2, 7, 14, 8, "bar", ["Small", "Big", "Middle"],
                            [{"name": "Spend", "values": [1, 9, 5]}])
        result = self.build("chart_rank_bar", [content_slide("s1", [block])])
        chart = charts_of(result["path"])[0]
        self.assertEqual([str(c) for c in chart.plots[0].categories],
                         ["Small", "Middle", "Big"])

    def test_time_categories_keep_their_order(self):
        block = chart_block("c1", 2, 7, 14, 8, "column", ["2024", "2025", "2026"],
                            [{"name": "Spend", "values": [9, 1, 5]}])
        result = self.build("chart_time", [content_slide("s1", [block])])
        chart = charts_of(result["path"])[0]
        self.assertEqual([str(c) for c in chart.plots[0].categories],
                         ["2024", "2025", "2026"])
        self.assertEqual(reasons(result, "chart"), [])

    def test_quarters_count_as_time(self):
        self.assertTrue(build_pptx.looks_like_time(["Q1 2026", "Q2 2026", "Q3 2026"]))
        self.assertFalse(build_pptx.looks_like_time(["Brand", "Product", "Sales"]))

    def test_a_chart_title_that_repeats_the_block_title_is_dropped(self):
        block = chart_block("c1", 2, 7, 14, 8, "column", ["A", "B"],
                            [{"name": "Spend", "values": [3, 5]}],
                            title="Marketing spend by unit")
        block["content"]["title"] = "Marketing spend by unit"
        result = self.build("chart_title", [content_slide("s1", [block])])
        self.assertFalse(charts_of(result["path"])[0].has_title)
        self.assertTrue(any("repeats the title above" in w for w in result["warnings"]))

    def test_a_chart_title_that_says_something_else_is_kept(self):
        block = chart_block("c1", 2, 7, 14, 8, "column", ["A", "B"],
                            [{"name": "Spend", "values": [3, 5]}],
                            title="Marketing spend by unit")
        block["content"]["title"] = "Index, 2019 = 100"
        result = self.build("chart_title2", [content_slide("s1", [block])])
        self.assertTrue(charts_of(result["path"])[0].has_title)

    def test_the_axis_format_comes_from_the_values(self):
        self.assertEqual(build_pptx.parse_value("7.7%"), (7.7, "%"))
        self.assertEqual(build_pptx.parse_value("$60K"), (60000.0, "$"))
        self.assertEqual(build_pptx.parse_value("4.1x"), (4.1, "x"))
        self.assertEqual(build_pptx.axis_format("%", [7.7, 59.0]), '#,##0.0"%"')
        self.assertEqual(build_pptx.axis_format("$", [60000.0]), '"$"#,##0')
        self.assertIsNone(build_pptx.axis_format("", [3.0, 5.0]))

    def test_a_percent_chart_carries_the_percent_format(self):
        block = chart_block("c1", 2, 7, 14, 8, "column", ["A", "B"],
                            [{"name": "Share", "values": ["7.7%", "59%"]}])
        result = self.build("chart_pct", [content_slide("s1", [block])])
        entry = shapes_for(read_snapshot(result["path"]), "c1")[0]
        self.assertEqual(entry["numberFormat"], '#,##0.0"%"')
        self.assertEqual(charts_of(result["path"])[0].value_axis.tick_labels.number_format,
                         '#,##0.0"%"')

    def test_a_named_series_takes_the_accent_without_an_emphasis_flag(self):
        block = chart_block("c1", 2, 7, 14, 8, "column", ["A", "B"],
                            [{"name": "Now", "values": [3, 5]},
                             {"name": "Then", "values": [2, 4]}])
        block["content"]["emphasisSeries"] = "Then"
        result = self.build("chart_emph", [content_slide("s1", [block])],
                            style={"mode": "house"})
        chart = charts_of(result["path"])[0]
        colours = [s.format.fill.fore_color.rgb for s in chart.plots[0].series]
        self.assertNotEqual(str(colours[0]), str(colours[1]))
        self.assertEqual(str(colours[1]), "3B5BDB")


class NotesFloorTests(BuildsOnce):
    """Q15: at the floor the last paragraph goes to the notes, never mid-sentence."""

    PARAS = ["The first paragraph carries the point the slide is making today.",
             "The second paragraph carries the evidence behind that point.",
             "The third paragraph carries the caveat nobody reads on the slide.",
             "The fourth paragraph carries the detail that belongs in a footnote.",
             "The fifth paragraph carries the aside that belongs in the appendix."]

    def tight(self):
        block = text_block("b1", 2, 7, 8, 2, self.PARAS[0])
        block["content"]["paragraphs"] = list(self.PARAS)
        return [content_slide("s1", [block])]

    def drawn(self, path):
        return " ".join(text for text, _, _, _ in runs_of(path))

    def test_the_last_paragraph_moves_to_the_notes(self):
        result = self.build("notes_move", self.tight())
        moves = reasons(result, "moved to notes")
        self.assertTrue(moves)
        self.assertEqual(moves[0]["blockId"], "b1")
        self.assertIn("belongs in the appendix", notes_of(result["path"])[0])

    def test_the_first_paragraph_stays_on_the_slide(self):
        result = self.build("notes_keep", self.tight())
        drawn = self.drawn(result["path"])
        self.assertIn("carries the point", drawn)
        self.assertNotIn("belongs in the appendix", drawn)

    def test_nothing_is_cut_mid_sentence(self):
        result = self.build("notes_whole", self.tight())
        drawn = self.drawn(result["path"])
        notes = notes_of(result["path"])[0]
        for para in self.PARAS:
            self.assertTrue(para in drawn or para in notes, para)

    def test_strict_keeps_the_overflow_and_moves_nothing(self):
        result = self.build("notes_strict", self.tight(), fidelity="strict")
        self.assertEqual(reasons(result, "moved to notes"), [])
        self.assertIn("belongs in the appendix", self.drawn(result["path"]))

    def test_the_move_is_reported_in_the_snapshot(self):
        result = self.build("notes_snap", self.tight())
        slide = read_snapshot(result["path"])["slides"][0]
        self.assertTrue(any(m["reason"] == "moved to notes"
                            for m in slide.get("adjusted") or []))


# ============================================================ Q8 template mode


def make_layout_template(path, width=960, height=540):
    """A .pptx with two layouts of its own, plus a section and a picture layout."""
    prs = Presentation()
    prs.slide_width = build_pptx.emu(width)
    prs.slide_height = build_pptx.emu(height)
    body = prs.slide_layouts[1]
    body.name = "Brand Body"
    places = {0: (48, 40, 864, 80), 1: (48, 150, 864, 300)}
    for shape in body.placeholders:
        box = places.get(shape.placeholder_format.idx)
        if box:
            shape.left, shape.top = build_pptx.emu(box[0]), build_pptx.emu(box[1])
            shape.width, shape.height = build_pptx.emu(box[2]), build_pptx.emu(box[3])
    two = prs.slide_layouts[3]
    two.name = "Brand Two Up"
    places = {0: (48, 40, 864, 80), 1: (48, 150, 420, 300), 2: (492, 150, 420, 300)}
    for shape in two.placeholders:
        box = places.get(shape.placeholder_format.idx)
        if box:
            shape.left, shape.top = build_pptx.emu(box[0]), build_pptx.emu(box[1])
            shape.width, shape.height = build_pptx.emu(box[2]), build_pptx.emu(box[3])
    prs.save(path)
    return path


class TemplateLayoutTests(BuildsOnce):
    """Q8: the template's layouts carry the slide, and its theme wins."""

    @classmethod
    def setUpClass(cls):
        super(TemplateLayoutTests, cls).setUpClass()
        cls.template = make_layout_template(os.path.join(cls.tmp, "brand.pptx"))
        cls.image = os.path.join(cls.tmp, "shot.png")
        make_png(cls.image, 240, 160)

    def templated(self, name, slides, **over):
        return self.build(name, slides,
                          style={"mode": "template", "templatePath": self.template},
                          **over)

    def layout_names(self, path):
        return [slide.slide_layout.name for slide in Presentation(path).slides]

    def test_one_text_block_takes_the_title_and_body_layout(self):
        block = text_block("b1", 2, 6, 14, 6, "One rail of copy for the body slot")
        result = self.templated("tpl_body", [content_slide("s1", [block])])
        self.assertEqual(self.layout_names(result["path"]), ["Brand Body"])

    def test_two_text_blocks_take_the_two_content_layout(self):
        blocks = [text_block("b1", 2, 6, 14, 6, "Left rail"),
                  text_block("b2", 17, 6, 14, 6, "Right rail")]
        result = self.templated("tpl_two", [content_slide("s1", blocks)])
        self.assertEqual(self.layout_names(result["path"]), ["Brand Two Up"])

    def test_a_slide_with_no_blocks_takes_the_title_only_layout(self):
        result = self.templated("tpl_only", [content_slide("s1", [])])
        self.assertEqual(self.layout_names(result["path"]), ["Title Only"])

    def test_a_divider_takes_the_section_layout(self):
        slides = [{"id": "d", "kind": "divider", "section": "Why", "headline": "Why"}]
        result = self.templated("tpl_section", slides)
        self.assertEqual(self.layout_names(result["path"]), ["Section Header"])

    def test_a_cover_takes_the_title_slide_layout(self):
        slides = [{"id": "c", "kind": "title", "headline": "Referral Growth",
                   "subtitle": "Draft for discussion"}]
        result = self.templated("tpl_cover", slides)
        self.assertEqual(self.layout_names(result["path"]), ["Title Slide"])

    def test_the_first_text_block_fills_the_body_placeholder(self):
        block = text_block("b1", 2, 6, 14, 6, "One rail of copy for the body slot")
        result = self.templated("tpl_fill", [content_slide("s1", [block])])
        entry = shapes_for(read_snapshot(result["path"]), "b1")[0]
        self.assertEqual(entry["role"], "templateBody")
        self.assertAlmostEqual(entry["x"], 48.0, delta=1.0)
        self.assertAlmostEqual(entry["y"], 150.0, delta=1.0)
        self.assertEqual(len(shapes_for(read_snapshot(result["path"]), "b1")), 1)

    def test_a_block_the_placeholders_cannot_take_is_drawn_as_usual(self):
        blocks = [text_block("b1", 2, 6, 14, 6, "Into the body slot"),
                  stat_block("b2", 17, 6, 14, 6, [{"value": "59%", "label": "of CMOs"}])]
        result = self.templated("tpl_rest", [content_slide("s1", blocks)])
        snapshot = read_snapshot(result["path"])
        self.assertEqual(shapes_for(snapshot, "b1")[0]["role"], "templateBody")
        self.assertTrue(shapes_for(snapshot, "b2"))
        self.assertEqual(shapes_for(snapshot, "b2")[0]["role"], "stat")

    def test_the_image_block_fills_the_picture_placeholder(self):
        block = {"id": "b1", "type": "image", "role": "visual", "col": 2, "row": 6,
                 "cols": 14, "rows": 6, "content": {"path": self.image}}
        result = self.templated("tpl_pic", [content_slide("s1", [block])])
        self.assertEqual(self.layout_names(result["path"]), ["Picture with Caption"])
        entry = shapes_for(read_snapshot(result["path"]), "b1")[0]
        self.assertEqual(entry["role"], "templatePicture")

    def test_the_headline_lands_in_the_title_placeholder(self):
        result = self.templated("tpl_title", [content_slide("s1", [], headline="The answer")])
        slide = list(Presentation(result["path"]).slides)[0]
        titles = [s for s in slide.placeholders
                  if "TITLE" in str(s.placeholder_format.type)]
        self.assertEqual(len(titles), 1)
        self.assertEqual(titles[0].text_frame.text, "The answer")

    def test_the_body_measure_follows_the_content_layout(self):
        spec = mini_spec(os.path.join(self.tmp, "tpl_measure.pptx"), [],
                         style={"mode": "template", "templatePath": self.template})
        warnings = []
        style, path = build_pptx.make_style(spec, warnings)
        build_pptx.open_presentation(spec, style, path, warnings)
        slot = style.slot("body", "content")
        self.assertAlmostEqual(slot["x"], 48.0, delta=1.0)
        self.assertAlmostEqual(slot["w"], 864.0, delta=1.0)

    def test_the_master_theme_colours_win(self):
        spec = mini_spec(os.path.join(self.tmp, "tpl_colour.pptx"), [],
                         style={"mode": "template", "templatePath": self.template})
        warnings = []
        style, path = build_pptx.make_style(spec, warnings)
        prs, spec = build_pptx.open_presentation(spec, style, path, warnings)
        theme = build_pptx.template_colors(prs)
        self.assertTrue(theme)
        self.assertEqual(style.hex_of("accent", "content"), theme["accent"])
        self.assertEqual(style.background("divider"), theme["bg"])

    def test_the_layout_score_prefers_the_matching_shape(self):
        prs = Presentation(self.template)
        profiles = build_pptx.layout_profiles(prs)
        want = build_pptx.wanted_layout({"kind": "content"},
                                        [{"type": "text"}, {"type": "bullets"}])
        self.assertEqual(want["kind"], "twoContent")
        best = max(profiles, key=lambda p: build_pptx.score_layout(p, want))
        self.assertEqual(best["name"], "Brand Two Up")


if __name__ == "__main__":
    unittest.main(verbosity=2)
